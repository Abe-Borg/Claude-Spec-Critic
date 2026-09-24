"""Tests for preprocessor disposition policy + stale-reference suppression.

Two independent features land in this file because they share the same module
(``preprocessor.py``) and prompt builder (``prompts.py``) surface area:

* **Pre-detected alerts** — feed the deterministic preprocessor's alerts into the LLM
  prompt via a compact ``<pre_detected>`` block so the model knows what was
  already detected locally and does not duplicate those items as new
  findings. The block must be:

    1. compact (count + small example list per rule, no whole-alert dump);
    2. boundary-safe (hostile match payloads cannot close the wrapper);
    3. byte-stable with the legacy message when no alerts are supplied
       (so the prompt-cache breakpoint invariant holds);
    4. toggleable via ``SPEC_CRITIC_PRE_DETECTED_ALERTS=0``;
    5. filtered by filename so a multi-spec project never leaks one spec's
       alerts into another spec's prompt;
    6. wired through the batch path (``submit_review_batch``).

* **Stale-cycle suppression** — add context-aware suppression for the stale-code-cycle
  detector so obvious negated / historical phrasings ("previously per the
  2019 CBC", "shall not follow the 2019 CBC approach", etc.) stop showing
  up as preflight alerts. Active stale references ("Comply with 2019 CBC"
  for a 2025-cycle project) must still be flagged.
"""
from __future__ import annotations


import pytest

from src.core.code_cycles import CALIFORNIA_2025
from src.input.preprocessor import (
    DETERMINISTIC_RULE_STALE_CODE_CYCLE,
    _should_suppress_stale_cycle,
    detect_stale_code_cycle_references,
)
from src.review.prompt_serialization import (
    TAG_PRE_DETECTED,
    render_pre_detected_block,
)
from src.review.prompts import get_single_spec_user_message


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _alert(filename: str, rule: str, match: str, *, atype: str | None = None) -> dict:
    """Build a minimal alert dict shaped like the ones from preprocessor.py."""
    return {
        "filename": filename,
        "type": atype or rule.replace("_", " "),
        "match": match,
        "context": match,
        "position": 0,
        "deterministic_rule": rule,
    }


# ---------------------------------------------------------------------------
# render_pre_detected_block helper
# ---------------------------------------------------------------------------


class TestRenderPreDetectedBlock:
    def test_empty_input_returns_empty_string(self) -> None:
        assert render_pre_detected_block(None) == ""
        assert render_pre_detected_block([]) == ""

    def test_no_matching_filename_returns_empty_string(self) -> None:
        alerts = [_alert("other.docx", "leed_reference", "LEED")]
        assert render_pre_detected_block(alerts, filename="me.docx") == ""

    def test_single_alert_renders_wrapper_count_and_example(self) -> None:
        alerts = [_alert("f.docx", "leed_reference", "LEED Gold")]
        out = render_pre_detected_block(alerts, filename="f.docx")
        assert out.startswith(f"<{TAG_PRE_DETECTED}>")
        assert out.endswith(f"</{TAG_PRE_DETECTED}>")
        assert "leed_reference (count=1)" in out
        assert "LEED Gold" in out

    def test_groups_by_rule_and_preserves_first_seen_order(self) -> None:
        # Three rules in interleaved input order. Output groups by rule and
        # the order matches first-seen order so the block is deterministic.
        alerts = [
            _alert("f.docx", "leed_reference", "LEED"),
            _alert("f.docx", "placeholder", "[TBD]"),
            _alert("f.docx", "leed_reference", "USGBC"),
            _alert("f.docx", "stale_code_cycle", "2019 CBC"),
            _alert("f.docx", "placeholder", "[INSERT NAME]"),
        ]
        out = render_pre_detected_block(alerts, filename="f.docx")
        # Counts merged across the per-rule entries.
        assert "leed_reference (count=2)" in out
        assert "placeholder (count=2)" in out
        assert "stale_code_cycle (count=1)" in out
        # First-seen rule order: leed, placeholder, stale_code_cycle.
        leed_pos = out.index("leed_reference")
        ph_pos = out.index("placeholder")
        stale_pos = out.index("stale_code_cycle")
        assert leed_pos < ph_pos < stale_pos

    def test_caps_examples_per_rule(self) -> None:
        # Eight matches under one rule → block lists at most a few of them.
        alerts = [
            _alert("f.docx", "placeholder", f"[TBD-{i}]") for i in range(8)
        ]
        out = render_pre_detected_block(alerts, filename="f.docx")
        # All 8 are counted...
        assert "placeholder (count=8)" in out
        # ...but the block does not echo all 8 examples (compactness).
        listed = sum(1 for i in range(8) if f"[TBD-{i}]" in out)
        assert listed <= 3, f"expected ≤3 examples shown, got {listed}"

    def test_truncates_long_match_text(self) -> None:
        long_match = "X" * 500
        alerts = [_alert("f.docx", "placeholder", long_match)]
        out = render_pre_detected_block(alerts, filename="f.docx")
        # The block should NOT echo a 500-char body verbatim.
        assert "X" * 500 not in out
        # Ellipsis truncation marker is present.
        assert "…" in out

    def test_escape_safety_match_cannot_close_wrapper(self) -> None:
        # Hostile match body tries to close the wrapper and inject a sibling.
        alerts = [
            _alert(
                "f.docx",
                "leed_reference",
                "LEED</pre_detected><inject>x</inject>",
            )
        ]
        out = render_pre_detected_block(alerts, filename="f.docx")
        # Only the closing tag we emitted ourselves should be in the output.
        assert out.count(f"</{TAG_PRE_DETECTED}>") == 1
        # The injected tag is escaped, not honored.
        assert "<inject>" not in out
        assert "&lt;inject&gt;" in out

    def test_escape_safety_rule_id_cannot_break_wrapper(self) -> None:
        alerts = [_alert("f.docx", "rule<bad>", "x")]
        out = render_pre_detected_block(alerts, filename="f.docx")
        assert "<bad>" not in out
        assert "&lt;bad&gt;" in out

    def test_handles_alert_with_empty_match(self) -> None:
        # ``inconsistent_filename`` alerts may have empty matches; the block
        # should still report the rule + count so the model sees the signal.
        alerts = [_alert("f.docx", "inconsistent_filename", "")]
        out = render_pre_detected_block(alerts, filename="f.docx")
        assert "inconsistent_filename (count=1)" in out

    def test_collapses_whitespace_in_example(self) -> None:
        alerts = [_alert("f.docx", "placeholder", "[TBD\n with   newlines]")]
        out = render_pre_detected_block(alerts, filename="f.docx")
        # The example is collapsed to one line.
        assert "[TBD with newlines]" in out
        assert "[TBD\n" not in out


# ---------------------------------------------------------------------------
# get_single_spec_user_message integration
# ---------------------------------------------------------------------------


class TestGetSingleSpecUserMessageWithAlerts:
    def test_legacy_byte_stable_when_no_alerts(self) -> None:
        legacy = get_single_spec_user_message(
            "alpha", "f.docx", cycle=CALIFORNIA_2025,
        )
        none_explicit = get_single_spec_user_message(
            "alpha", "f.docx", cycle=CALIFORNIA_2025, pre_detected_alerts=None,
        )
        empty_explicit = get_single_spec_user_message(
            "alpha", "f.docx", cycle=CALIFORNIA_2025, pre_detected_alerts=[],
        )
        assert legacy == none_explicit == empty_explicit

    def test_block_appended_when_alerts_provided(self) -> None:
        alerts = [_alert("f.docx", "leed_reference", "LEED")]
        msg = get_single_spec_user_message(
            "alpha", "f.docx", cycle=CALIFORNIA_2025, pre_detected_alerts=alerts,
        )
        assert f"<{TAG_PRE_DETECTED}>" in msg
        assert f"</{TAG_PRE_DETECTED}>" in msg
        # Block sits AFTER the spec body so the cache-prefix invariant holds.
        spec_close = msg.rindex("</spec>")
        block_open = msg.index(f"<{TAG_PRE_DETECTED}>")
        assert block_open > spec_close

    def test_cache_prefix_invariant_holds_with_and_without_alerts(self) -> None:
        # TestPromptCacheBreakpointSafety pins the prefix before
        # ``<spec ``. Adding a pre_detected block at the END must not change
        # that prefix.
        without = get_single_spec_user_message(
            "alpha", "f.docx", cycle=CALIFORNIA_2025,
        )
        with_alerts = get_single_spec_user_message(
            "alpha", "f.docx", cycle=CALIFORNIA_2025,
            pre_detected_alerts=[_alert("f.docx", "leed_reference", "LEED")],
        )
        assert without.split("<spec ")[0] == with_alerts.split("<spec ")[0]

    def test_block_carries_do_not_duplicate_instruction(self) -> None:
        alerts = [_alert("f.docx", "placeholder", "[TBD]")]
        msg = get_single_spec_user_message(
            "alpha", "f.docx", cycle=CALIFORNIA_2025, pre_detected_alerts=alerts,
        )
        # The model is told what to do with the block. We don't pin the
        # exact wording, but the anti-duplication intent must be present.
        lowered = msg.lower()
        assert "do not duplicate" in lowered or "do not report" in lowered

    def test_alerts_for_other_files_filtered_out(self) -> None:
        mixed = [
            _alert("other.docx", "placeholder", "[TBD-other]"),
            _alert("f.docx", "leed_reference", "LEED-mine"),
        ]
        msg = get_single_spec_user_message(
            "alpha", "f.docx", cycle=CALIFORNIA_2025, pre_detected_alerts=mixed,
        )
        assert "LEED-mine" in msg
        assert "TBD-other" not in msg
        # Only the matching spec's rule appears.
        assert "leed_reference" in msg
        assert "placeholder" not in msg

# ---------------------------------------------------------------------------
# pipeline plumbing
# ---------------------------------------------------------------------------


@pytest.fixture
def stub_count_tokens(monkeypatch: pytest.MonkeyPatch) -> None:
    """Replace tiktoken-backed counts with a deterministic word-count proxy.

    Mirrors the helper in ``tests/test_chunk_e_token_budgets.py`` —
    ``_prepare_specs`` calls ``count_tokens`` repeatedly, and the real
    encoder lazily downloads a BPE merge table that fails in fully offline
    environments. The proxy keeps the pipeline path hermetic for tests
    that exercise the per-spec alert map, which doesn't care about exact
    counts.
    """
    def _fake_count(text: str | None) -> int:
        return len((text or "").split()) * 2

    monkeypatch.setattr("src.core.tokenizer.count_tokens", _fake_count)
    monkeypatch.setattr("src.orchestration.pipeline.count_tokens", _fake_count, raising=False)
    # ``src.batch`` no longer imports ``count_tokens`` directly —
    # every batch token count is computed inside the central review
    # request builder. Patch the binding there so the per-spec
    # extended-output gating and the local preflight estimate don't trip
    # the lazy tiktoken download.
    monkeypatch.setattr(
        "src.review.review_request_builder.count_tokens", _fake_count, raising=False
    )
    # Preflight calls the Anthropic API; bypass for hermetic tests.
    monkeypatch.setattr(
        "src.orchestration.pipeline.token_count_preflight_enabled", lambda: False
    )


class TestPipelinePerSpecAlertMap:
    """``_prepare_specs`` populates the per-filename alert map used by the
    reviewer / batch paths to feed each spec's prompt.
    """

    def _make_spec_files(self, tmp_path):
        """Build deterministic .docx files that trip a handful of rules."""
        from docx import Document

        files = []
        for fname in ("23 21 13 - A.docx", "23 22 13 - B.docx"):
            doc = Document()
            doc.add_paragraph("PART 1 - GENERAL")
            doc.add_paragraph("This is a LEED Gold project.")
            doc.add_paragraph("Coordinate with [INSERT PROJECT NAME].")
            doc.add_paragraph("Refer to TODO: confirm capacity later.")
            path = tmp_path / fname
            doc.save(str(path))
            files.append(path)
        return files

    def test_prepare_specs_returns_per_filename_map(
        self, tmp_path, stub_count_tokens
    ) -> None:
        from src.orchestration.pipeline import _prepare_specs

        files = self._make_spec_files(tmp_path)
        prepared = _prepare_specs(
            input_dir=tmp_path,
            files=files,
            project_context="",
            cycle=CALIFORNIA_2025,
        )
        # Every selected spec has its own entry, even if empty.
        assert set(prepared.pre_detected_by_filename) >= {p.name for p in files}
        # Alerts surfaced (LEED + placeholder + template marker at least).
        for f in files:
            spec_alerts = prepared.pre_detected_by_filename[f.name]
            rules = {a["deterministic_rule"] for a in spec_alerts}
            assert "leed_reference" in rules
            assert "placeholder" in rules
            assert "template_marker" in rules
            # Every alert in this spec's bucket carries its own filename.
            for alert in spec_alerts:
                assert alert["filename"] == f.name


class TestPreflightNamingLog:
    """The preflight log line matches the naming notice it summarizes: a
    project with no dominant style is not told its files are "non-dominant"."""

    @staticmethod
    def _log_for(tmp_path, names: list[str]) -> list[str]:
        from docx import Document

        from src.orchestration.pipeline import _prepare_specs

        files = []
        for name in names:
            doc = Document()
            doc.add_paragraph("PART 1 GENERAL")
            doc.add_paragraph("1.01 SUMMARY")
            doc.add_paragraph("A. Provide piping.")
            path = tmp_path / name
            doc.save(str(path))
            files.append(path)
        lines: list[str] = []
        _prepare_specs(
            input_dir=tmp_path,
            files=files,
            project_context="",
            cycle=CALIFORNIA_2025,
            log=lambda message, **_: lines.append(message),
        )
        return [line for line in lines if "naming" in line]

    def test_a_mixture_is_logged_as_a_mixture(self, tmp_path, stub_count_tokens):
        (line,) = self._log_for(tmp_path, ["21 05 00.docx", "211313.docx"])
        assert line == "Preflight: 2 CSI-named file(s) mix naming styles; no single style dominates."

    def test_a_minority_is_logged_as_non_dominant(self, tmp_path, stub_count_tokens):
        (line,) = self._log_for(tmp_path, ["21 05 00.docx", "21 13 13.docx", "211316.docx"])
        assert line == "Preflight: 1 file(s) use a non-dominant CSI naming style."


class TestBatchSubmissionFeedsAlerts:
    """``submit_review_batch`` must pass each spec's alerts into the prompt."""

    def test_per_spec_alerts_land_in_user_message(self, monkeypatch, stub_count_tokens):
        # We capture the kwargs handed to the batch API via a fake client.
        from src.batch import batch as batch_mod
        from src.input.extractor import ExtractedSpec

        captured: list[dict] = []

        class FakeBatches:
            def create(self, **kwargs):
                captured.append(kwargs)
                class _Resp:
                    id = "msgbatch_test"
                return _Resp()

        class FakeBetaBatches(FakeBatches):
            pass

        class FakeBeta:
            class messages:  # noqa: N801 — mimic SDK shape
                batches = FakeBetaBatches()

        class FakeMessages:
            batches = FakeBatches()

        class FakeClient:
            messages = FakeMessages()
            beta = FakeBeta()

        monkeypatch.setattr(batch_mod, "_get_client", lambda **_: FakeClient())

        specs = [
            ExtractedSpec(
                filename="a.docx", content="LEED Gold project body.",
                word_count=4,
            ),
            ExtractedSpec(
                filename="b.docx", content="Coordinate with [INSERT NAME].",
                word_count=3,
            ),
        ]
        alerts = {
            "a.docx": [_alert("a.docx", "leed_reference", "LEED Gold")],
            "b.docx": [_alert("b.docx", "placeholder", "[INSERT NAME]")],
        }
        batch_mod.submit_review_batch(
            specs,
            project_context="",
            cycle=CALIFORNIA_2025,
            pre_detected_alerts=alerts,
        )
        assert captured, "no batch request was issued"
        requests = captured[0]["requests"]
        # Two requests, one per spec. Each carries its own pre_detected block
        # and not the other spec's alerts.
        bodies = {req["custom_id"]: req["params"]["messages"][0]["content"]
                  for req in requests}
        a_body = next(b for cid, b in bodies.items() if "a_docx" in cid or "a__" in cid)
        b_body = next(b for cid, b in bodies.items() if "b_docx" in cid or "b__" in cid)
        assert "LEED Gold" in a_body
        assert "[INSERT NAME]" not in a_body
        assert "[INSERT NAME]" in b_body
        assert "LEED Gold" not in b_body
        # Both bodies carry the block + anti-duplication instruction.
        for body in (a_body, b_body):
            assert f"<{TAG_PRE_DETECTED}>" in body
            assert "do not duplicate" in body.lower()


# ---------------------------------------------------------------------------
# stale-cycle context suppression
# ---------------------------------------------------------------------------


class TestStaleCycleSuppression:
    """Negated / historical phrasings near a stale cycle citation are skipped.

    Plan D4.2: keywords like ``previously``, ``formerly``, ``superseded``,
    ``withdrawn``, ``obsolete``, ``not``, ``no longer``, ``prior``, and
    ``historical`` in a small window before the match indicate the author
    is *describing* an old reference rather than *requiring* it.
    """

    def test_previously_suppresses(self) -> None:
        content = "Previously per the 2019 CBC, now superseded by the current cycle."
        alerts = detect_stale_code_cycle_references(
            content, "s.docx", CALIFORNIA_2025
        )
        assert alerts == []

    def test_superseded_trailing_suppresses(self) -> None:
        # The citation is in the same sentence as ``superseded``; the author
        # is explicitly describing a superseded reference, so the alert is
        # suppressed regardless of whether the keyword sits before or after
        # the cycle citation.
        content = "The 2019 CBC has been superseded for this project."
        alerts = detect_stale_code_cycle_references(
            content, "s.docx", CALIFORNIA_2025
        )
        assert alerts == []

    def test_shall_not_follow_suppresses(self) -> None:
        # "shall not follow the 2019 CBC" — explicit negation.
        content = "The work shall not follow the 2019 CBC approach."
        alerts = detect_stale_code_cycle_references(
            content, "s.docx", CALIFORNIA_2025
        )
        assert alerts == []

    def test_no_longer_suppresses(self) -> None:
        content = "The 2022 CBC is no longer used; comply with the current cycle."
        alerts = detect_stale_code_cycle_references(
            content, "s.docx", CALIFORNIA_2025
        )
        assert alerts == []

    def test_active_requirement_still_flagged(self) -> None:
        # The author actively requires a stale cycle — must still flag.
        content = "Comply with 2019 CBC for piping installations."
        alerts = detect_stale_code_cycle_references(
            content, "s.docx", CALIFORNIA_2025
        )
        assert alerts, "active stale reference should still be flagged"
        assert all(
            a["deterministic_rule"] == DETERMINISTIC_RULE_STALE_CODE_CYCLE
            for a in alerts
        )

    def test_active_requirement_at_start_of_document_still_flagged(self) -> None:
        # No preceding window content. The detector must not over-suppress
        # just because the window is empty.
        content = "2019 CBC governs all piping work."
        alerts = detect_stale_code_cycle_references(
            content, "s.docx", CALIFORNIA_2025
        )
        assert alerts
        assert any(a["found_year"] == "2019" for a in alerts)

    def test_negated_does_not_suppress_unrelated_stale_reference(self) -> None:
        # Two stale citations: one negated, one active. Only the active one
        # should be flagged.
        content = (
            "Previously per the 2019 CBC. Comply with 2022 CBC for all work."
        )
        alerts = detect_stale_code_cycle_references(
            content, "s.docx", CALIFORNIA_2025
        )
        years = {a["found_year"] for a in alerts}
        assert "2019" not in years
        assert "2022" in years


class TestStaleCycleTrailingWindow:
    """The trailing suppression window cuts at the EARLIEST sentence
    terminator by position (B-28 d).

    The old loop tried ``"."`` first and stopped at the first terminator
    *found in tuple order*, so a ``;`` or ``\\n\\n`` that came earlier than
    the ``.`` was ignored and the window kept a whole extra clause — a
    negation belonging to the next clause then suppressed an active stale
    citation.
    """

    def test_semicolon_before_period_bounds_window(self) -> None:
        # Old: cut at "." → window "; the prior edition is no longer
        # referenced" → "prior" / "no longer" suppress the ACTIVE
        # "Comply with 2019 CBC" (wrong). New: cut at ";" (earlier) →
        # empty window → flagged (right).
        content = (
            "Comply with 2019 CBC; the prior edition is no longer referenced. "
            "Provide seismic bracing."
        )
        assert _should_suppress_stale_cycle(
            content, content.index("2019"), content.index("CBC") + 3
        ) is False
        alerts = detect_stale_code_cycle_references(content, "s.docx", CALIFORNIA_2025)
        assert [a["found_year"] for a in alerts] == ["2019"]

    def test_paragraph_break_before_period_bounds_window(self) -> None:
        # Same defect with the "\n\n" terminator: the next paragraph's
        # "no longer" must not bleed back into the active citation.
        content = "Comply with 2019 CBC\n\nThe prior edition is no longer used."
        assert _should_suppress_stale_cycle(
            content, content.index("2019"), content.index("CBC") + 3
        ) is False
        alerts = detect_stale_code_cycle_references(content, "s.docx", CALIFORNIA_2025)
        assert [a["found_year"] for a in alerts] == ["2019"]

    def test_negation_inside_the_clause_still_suppresses(self) -> None:
        # The earliest terminator comes AFTER the negation here, so the
        # descriptive citation is still suppressed — the fix narrows the
        # window, it does not disable trailing suppression.
        content = "The 2019 CBC is no longer referenced; comply with the current cycle."
        assert _should_suppress_stale_cycle(
            content, content.index("2019"), content.index("CBC") + 3
        ) is True
        assert detect_stale_code_cycle_references(content, "s.docx", CALIFORNIA_2025) == []

    def test_period_still_bounds_when_it_is_earliest(self) -> None:
        content = "Comply with 2019 CBC. The prior edition is no longer used; see above."
        assert _should_suppress_stale_cycle(
            content, content.index("2019"), content.index("CBC") + 3
        ) is False
        alerts = detect_stale_code_cycle_references(content, "s.docx", CALIFORNIA_2025)
        assert [a["found_year"] for a in alerts] == ["2019"]

    def test_no_terminator_keeps_full_window(self) -> None:
        content = "Comply with 2019 CBC which is no longer the adopted edition"
        assert _should_suppress_stale_cycle(
            content, content.index("2019"), content.index("CBC") + 3
        ) is True


# ---------------------------------------------------------------------------
# Citation-related suppression cues (plan WP-04B, chunk S03)
# ---------------------------------------------------------------------------


def _flagged(content: str) -> list[str]:
    """The stale citations that are flagged, by their matched text."""
    return [
        alert["match"]
        for alert in detect_stale_code_cycle_references(content, "s.docx", CALIFORNIA_2025)
    ]


class TestCitationRelatedSuppression:
    """A stale citation is suppressed only by text about the CITATION.

    The old keyword list suppressed on unrelated nearby words: "prior" in
    "prior to fabrication", "historical" in "the historical society", and any
    negated modal ("may not deviate from"). A negated requirement to comply is
    still a requirement, so only a closed list of verbs that reject the
    citation (follow, use, apply, reference, cite, ...) counts.
    """

    @pytest.mark.parametrize(
        "sentence",
        [
            # Plan WP-04B's examples and the equivalent forms it names.
            "Submit shop drawings prior to fabrication in accordance with 2022 CBC Section 1704.",
            "Coordinate with the historical society and comply with 2022 CBC.",
            "Contractor may not deviate from 2022 CBC Chapter 17.",
            "Contractor shall not deviate from 2022 CBC Chapter 17.",
            "Work cannot depart from 2022 CBC requirements.",
            # Word autocorrects the apostrophe; a contraction is still a negation.
            "Contractor can\u2019t deviate from 2022 CBC Chapter 17.",
            # Other negated requirements to comply.
            "Anchorage shall not be less than required by 2022 CBC.",
            "Supports shall not exceed the spacing in 2022 CBC Table 1234.",
            "Do not install piping except as permitted by 2022 CBC.",
            "Work not per 2022 CBC shall be removed.",
            "Existing piping that does not comply with 2022 CBC shall be replaced.",
            "Contractor shall not use PVC pipe in accordance with 2022 CBC.",
            # A historical word whose clause is an active requirement.
            "Previously approved submittals shall comply with 2022 CBC.",
            "As previously stated, comply with the 2019 CBC.",
            "Piping no longer in service shall be removed per 2022 CBC.",
            "The Historical Society building shall comply with 2022 CBC.",
        ],
    )
    def test_an_active_citation_is_flagged(self, sentence):
        assert _flagged(sentence), sentence

    @pytest.mark.parametrize(
        "sentence",
        [
            # "Previous edition" contexts.
            "Previously, the 2022 CBC applied to this work.",
            "The building was previously permitted under the 2022 CBC.",
            "Formerly the 2019 CBC governed this work.",
            "The prior edition (2019 CBC) required fewer braces.",
            "Under the previous code cycle, the 2019 CBC applied.",
            "The 2019 CBC was the prior edition.",
            "The 2019 CBC, previously in effect, required fewer braces.",
            "Existing bracing complies with the 2019 CBC (formerly adopted).",
            "The historical 2019 CBC required fewer braces.",
            "Buildings constructed prior to the 2019 CBC are exempt.",
            "Designs no longer follow the 2019 CBC.",
            # "Superseded citation" contexts.
            "The 2019 CBC has been superseded for this project.",
            "The superseded 2019 CBC is not used.",
            "The 2019 CBC (withdrawn) is listed for reference.",
            "The 2025 CBC supersedes the 2022 CBC.",
            "The 2025 CBC replaces the 2022 CBC.",
            "The 2022 CBC is no longer used.",
            "The 2019 CBC edition is no longer used.",
            # Rejections of the citation itself.
            "The work shall not follow the 2019 CBC approach.",
            "Do not use the 2019 CBC.",
            "Don\u2019t reference the 2019 CBC.",
            "Designs shall not be based on the 2019 CBC.",
            "The 2019 CBC shall not be used.",
            "The 2019 CBC is not to be used.",
            "The 2022 CBC does not apply.",
            "The 2019 CBC is not applicable.",
            "The 2019 CBC isn\u2019t in effect.",
            "The 2019 CBC is not the current edition.",
            "Comply with the 2025 CBC instead of the 2022 CBC.",
            "Comply with the 2025 CBC rather than 2022 CBC.",
            "Comply with the 2025 CBC, not the 2022 CBC.",
            "Design per the 2025 CBC, not per the 2022 CBC.",
        ],
    )
    def test_a_historical_or_rejected_citation_is_suppressed(self, sentence):
        assert _flagged(sentence) == [], sentence

    @pytest.mark.parametrize(
        "sentence, flagged",
        [
            # Each citation is judged by its own context: the window of one
            # stops at its neighbors, so a cue about one is never borrowed.
            ("Previously per the 2019 CBC, now per the 2022 CBC.", ["2022 CBC"]),
            (
                "The 2019 CBC was superseded by the 2022 CBC, which governs this work.",
                ["2022 CBC"],
            ),
            ("Comply with the 2022 CBC, not the 2019 CBC.", ["2022 CBC"]),
            ("Comply with 2022 CBC Section 1704 and 2019 CBC Section 1705.", ["2022 CBC", "2019 CBC"]),
            ("The 2019 CBC, not the 2022 CBC, governs anchorage.", ["2019 CBC"]),
            ("Design per ASCE 7-16, not ASCE 7-10.", ["ASCE 7-16"]),
            # The same-year pair in one clause: both active.
            ("Comply with 2019 CBC and 2019 CMC.", ["2019 CBC", "2019 CMC"]),
        ],
    )
    def test_several_citations_in_one_sentence(self, sentence, flagged):
        assert _flagged(sentence) == flagged

    @pytest.mark.parametrize(
        "sentence",
        [
            # One cue governs a coordinated list (found in review): a list
            # is judged as one citation, so every member shares the cue
            # before its first member or after its last.
            "Previously, the 2019 CBC and 2019 CMC applied.",
            "Do not use the 2019 CBC or 2019 CMC.",
            "The 2019 CBC and 2019 CMC were superseded.",
            "The 2019 CBC, the 2019 CMC, and the 2019 CPC have been superseded.",
            "Previously, the 2019 CBC, 2019 CMC and 2019 CPC applied.",
            "Do not use the 2019 CBC and/or 2019 CMC.",
            "Comply with the 2025 CBC instead of the 2019 CBC & 2019 CMC.",
            "ASCE 7-10 and ASCE 7-16 are no longer used.",
        ],
    )
    def test_a_cue_governs_every_citation_in_a_coordinated_list(self, sentence):
        assert _flagged(sentence) == [], sentence

    @pytest.mark.parametrize(
        "sentence, flagged",
        [
            # A list with no cue is active throughout.
            ("Comply with the 2019 CBC, 2019 CMC, and 2019 CPC.", ["2019 CBC", "2019 CMC", "2019 CPC"]),
            # A requirement verb still blocks an in-clause cue for the list.
            (
                "Previously approved submittals shall comply with the 2019 CBC and 2019 CMC.",
                ["2019 CBC", "2019 CMC"],
            ),
            # Only a coordinator joins a list; other words keep citations apart.
            ("Previously per the 2019 CBC, now per the 2019 CMC.", ["2019 CMC"]),
            ("The 2019 CBC was superseded by the 2019 CMC and 2019 CPC.", ["2019 CMC", "2019 CPC"]),
            # A paragraph break ends the list (and the clause).
            ("Previously, the 2019 CBC and\n\n2019 CMC applied.", ["2019 CMC"]),
        ],
    )
    def test_only_a_coordinator_makes_a_list(self, sentence, flagged):
        assert _flagged(sentence) == flagged

    def test_a_cue_about_a_neighboring_citation_is_not_borrowed(self):
        content = "The 2019 CBC was superseded by the 2022 CBC."
        start = content.index("2022")
        end = start + len("2022 CBC")
        # Without the neighbor bound the window still reaches "superseded",
        # but only as "superseded by the", which names the replacement.
        assert _should_suppress_stale_cycle(content, start, end) is False
        historical = content.index("2019")
        assert _should_suppress_stale_cycle(
            content, historical, historical + len("2019 CBC"), window_end=start
        ) is True

    def test_the_window_bounds_are_honored(self):
        content = "Previously per the 2019 CBC, now per the 2022 CBC."
        start = content.index("2022")
        end = start + len("2022 CBC")
        # Unbounded, the in-clause "Previously" reaches the 2022 citation;
        # bounded at the end of the 2019 citation (as the detector passes
        # it), it does not.
        assert _should_suppress_stale_cycle(content, start, end) is True
        bound = content.index("2019") + len("2019 CBC")
        assert _should_suppress_stale_cycle(content, start, end, window_start=bound) is False

    def test_asce_7_citations_use_the_same_cues(self):
        assert _flagged("ASCE 7-10 is no longer used for this work.") == []
        assert _flagged("Designs shall not deviate from ASCE 7-16.") == ["ASCE 7-16"]


class TestLocationAwareModulesStillSuppressStaleCycleChecks:
    """Plan WP-04B: "do not turn a syntax improvement into a new
    governing-edition policy". A location-aware module runs no stale-cycle
    detection at all, and that is unchanged by the new cues and syntax."""

    @staticmethod
    def _location_aware_modules():
        from src.modules.registry import AVAILABLE_MODULES

        modules = [m for m in AVAILABLE_MODULES.values() if m.project_profile_enabled]
        assert modules, "precondition: the registry has location-aware modules"
        return modules

    @pytest.mark.parametrize(
        "sentence",
        [
            "Contractor shall not deviate from the 2018 IBC Chapter 9.",
            "Coordinate with the historical society and comply with 2018 IBC.",
            "Design loads per ASCE/SEI 7-16.",
            "Design loads per ASCE 7\u201310.",
        ],
    )
    def test_no_stale_alert_for_any_location_aware_module(self, sentence):
        from src.input.preprocessor import preprocess_spec

        for module in self._location_aware_modules():
            result = preprocess_spec(sentence, "s.docx", cycle=module.cycle)
            assert result.code_cycle_alerts == [], module.module_id
