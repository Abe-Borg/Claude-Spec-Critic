"""C-13: explain *why* a citation was rejected.

Every search/fetch blocklist entry carries a category; grounding records a
plain-language reason per rejected URL ("blocked domain: <category>" vs
"not among searched or fetched results") on the additive
``VerificationResult.rejected_source_reasons`` field; the cache round-trips
it with no schema bump (legacy rows load ``{}``); and both renderers show it
next to the rejected URL — byte-identical output when nothing was rejected.

Hermetic — no API key, no network.
"""
from __future__ import annotations

import json
import time
from pathlib import Path

import pytest
from docx import Document

from src.core import api_config
from src.core.api_config import (
    BLOCKED_DOMAIN_CATEGORY_AGGREGATOR,
    BLOCKED_DOMAIN_CATEGORY_CONTENT_FARM,
    BLOCKED_DOMAIN_CATEGORY_ENCYCLOPEDIA,
    BLOCKED_DOMAIN_CATEGORY_LLM_OUTPUT,
    BLOCKED_DOMAIN_CATEGORY_MARKETPLACE,
    BLOCKED_DOMAIN_CATEGORY_SOCIAL,
    BLOCKED_DOMAIN_CATEGORY_TRADE_FORUM,
    _WEB_SEARCH_BLOCKED_DOMAIN_ENTRIES,
    _WEB_SEARCH_BLOCKED_DOMAINS,
    blocked_domain_category,
    build_web_fetch_tool,
    build_web_search_tool,
)
from src.core.code_cycles import DEFAULT_CYCLE
from src.output.html_report_exporter import render_html_report
from src.output.report_exporter import export_report
from src.review.reviewer import Finding, ReviewResult
from src.verification.source_grounding import (
    REJECT_EMPTY,
    REJECT_MALFORMED,
    REJECT_UNGROUNDED,
    REJECTION_EXPLANATION_EMPTY,
    REJECTION_EXPLANATION_MALFORMED,
    REJECTION_EXPLANATION_UNGROUNDED,
    SearchedSource,
    describe_rejection,
)
from src.verification.verification_cache import (
    VerificationCache,
    _CACHE_SCHEMA_VERSION,
    _PERSISTED_FIELDS,
    _result_from_dict,
    _result_to_dict,
)
from src.verification.verifier import VerificationResult, _apply_source_grounding

# The flat list as it was before categories were attached. The tool dicts are
# part of the prompt-cache prefix, so content AND order are pinned.
_LEGACY_FLAT_LIST = [
    "reddit.com", "quora.com", "medium.com",
    "stackexchange.com", "stackoverflow.com",
    "answers.yahoo.com", "fixya.com",
    "chatgpt.com", "perplexity.ai", "openai.com", "gemini.google.com",
    "claude.ai", "you.com", "phind.com", "copilot.microsoft.com",
    "poe.com", "character.ai", "jasper.ai", "writesonic.com",
    "diychatroom.com", "forums.jlconline.com", "hvac-talk.com",
    "inspectionnews.net", "inspectorsforum.com", "contractortalk.com",
    "doityourself.com", "homeadvisor.com", "thumbtack.com", "angi.com",
    "ehow.com", "wikihow.com", "about.com", "thespruce.com", "bobvila.com",
    "familyhandyman.com", "hunker.com", "sapling.com", "reference.com",
    "leaf.tv", "sciencing.com", "bizfluent.com", "pocketsense.com",
    "facebook.com", "twitter.com", "x.com", "instagram.com", "tiktok.com",
    "linkedin.com", "pinterest.com", "youtube.com", "threads.net",
    "wikipedia.org", "britannica.com",
]


# ---------------------------------------------------------------------------
# 1. Category lookup + the flat list the tools emit
# ---------------------------------------------------------------------------


class TestBlocklistCategories:
    def test_flat_list_is_byte_identical_to_the_legacy_list(self):
        assert _WEB_SEARCH_BLOCKED_DOMAINS == _LEGACY_FLAT_LIST
        assert build_web_search_tool()["blocked_domains"] == _LEGACY_FLAT_LIST
        assert build_web_fetch_tool()["blocked_domains"] == _LEGACY_FLAT_LIST

    def test_every_entry_has_a_non_empty_category_and_no_duplicates(self):
        domains = [d for d, _ in _WEB_SEARCH_BLOCKED_DOMAIN_ENTRIES]
        assert len(domains) == len(set(domains))
        assert all(isinstance(c, str) and c.strip() for _, c in _WEB_SEARCH_BLOCKED_DOMAIN_ENTRIES)

    @pytest.mark.parametrize(
        "url, category",
        [
            ("https://www.reddit.com/r/hvac/comments/x", BLOCKED_DOMAIN_CATEGORY_AGGREGATOR),
            ("reddit.com/r/hvac", BLOCKED_DOMAIN_CATEGORY_AGGREGATOR),  # bare host
            ("HTTPS://CHATGPT.COM/share/abc", BLOCKED_DOMAIN_CATEGORY_LLM_OUTPUT),
            ("https://hvac-talk.com/vbb/threads/1", BLOCKED_DOMAIN_CATEGORY_TRADE_FORUM),
            ("https://www.homeadvisor.com/r/x", BLOCKED_DOMAIN_CATEGORY_MARKETPLACE),
            ("https://www.thespruce.com/x", BLOCKED_DOMAIN_CATEGORY_CONTENT_FARM),
            ("https://x.com/someone/status/1", BLOCKED_DOMAIN_CATEGORY_SOCIAL),
            ("https://en.wikipedia.org/wiki/NFPA_13", BLOCKED_DOMAIN_CATEGORY_ENCYCLOPEDIA),
            ("https://user:pw@youtube.com:443/watch", BLOCKED_DOMAIN_CATEGORY_SOCIAL),
        ],
    )
    def test_blocked_hosts_resolve_to_their_category(self, url, category):
        assert blocked_domain_category(url) == category

    @pytest.mark.parametrize(
        "url",
        [
            "https://codes.iccsafe.org/content/CBC2025",
            "https://www.nfpa.org/codes-and-standards/nfpa-13",
            "https://yahoo.com/x",          # only answers.yahoo.com is blocked
            "https://jlconline.com/x",      # only forums.jlconline.com is blocked
            "https://notx.com/a",           # suffix match must be on a label boundary
            "https://myreddit.com/a",
            "",
            None,
            "http://:80",                   # no hostname
            "https://[invalid",             # unparseable (urlsplit raises)
        ],
    )
    def test_unblocked_or_unparseable_hosts_are_none(self, url):
        assert blocked_domain_category(url) is None

    def test_every_blocked_domain_matches_itself_and_a_subdomain(self):
        for domain, category in _WEB_SEARCH_BLOCKED_DOMAIN_ENTRIES:
            assert blocked_domain_category(f"https://{domain}/page") == category
            assert blocked_domain_category(f"https://sub.{domain}/page") == category

    def test_todo_is_gone(self):
        import inspect

        src = inspect.getsource(api_config)
        assert "TODO: explore a category-based blocking helper" not in src


# ---------------------------------------------------------------------------
# 2. Reason recording during grounding
# ---------------------------------------------------------------------------


class TestDescribeRejection:
    def test_blocked_domain_wins_over_the_sentinel(self):
        assert (
            describe_rejection("https://reddit.com/r/x", REJECT_UNGROUNDED)
            == "blocked domain: " + BLOCKED_DOMAIN_CATEGORY_AGGREGATOR
        )

    def test_sentinels_map_to_plain_language(self):
        assert describe_rejection("https://a.example/x", REJECT_UNGROUNDED) == REJECTION_EXPLANATION_UNGROUNDED
        assert describe_rejection("https://[invalid", REJECT_MALFORMED) == REJECTION_EXPLANATION_MALFORMED
        assert describe_rejection("", REJECT_EMPTY) == REJECTION_EXPLANATION_EMPTY
        assert describe_rejection("https://a.example/x", "") == REJECTION_EXPLANATION_UNGROUNDED
        assert REJECTION_EXPLANATION_UNGROUNDED == "not among searched or fetched results"


class TestGroundingRecordsReasons:
    def _ground(self, cited: list[str], searched: list[str], fetched: list[str] | None = None):
        r = VerificationResult(verdict="CONFIRMED", sources=list(cited), grounded=True)
        return _apply_source_grounding(
            r,
            searched=[SearchedSource(url=u) for u in searched],
            fetched=[SearchedSource(url=u) for u in (fetched or [])] or None,
        )

    def test_blocked_and_ungrounded_citations_get_distinct_reasons(self):
        out = self._ground(
            cited=[
                "https://codes.iccsafe.org/content/CBC2025",   # accepted
                "https://www.reddit.com/r/hvac/x",             # blocked
                "https://invented.example.com/fake",            # hallucinated
            ],
            searched=["https://codes.iccsafe.org/content/CBC2025"],
        )
        assert out.accepted_sources == ["https://codes.iccsafe.org/content/CBC2025"]
        assert [r["url"] for r in out.rejected_sources] == [
            "https://www.reddit.com/r/hvac/x",
            "https://invented.example.com/fake",
        ]
        # ``rejected_sources[].reason`` keeps the machine sentinel unchanged.
        assert {r["reason"] for r in out.rejected_sources} == {REJECT_UNGROUNDED}
        assert out.rejected_source_reasons == {
            "https://www.reddit.com/r/hvac/x": "blocked domain: " + BLOCKED_DOMAIN_CATEGORY_AGGREGATOR,
            "https://invented.example.com/fake": REJECTION_EXPLANATION_UNGROUNDED,
        }
        assert out.verdict == "CONFIRMED"  # one accepted citation keeps it grounded

    def test_fetched_url_is_accepted_and_carries_no_reason(self):
        out = self._ground(
            cited=["https://www.nfpa.org/13"],
            searched=["https://other.example/x"],
            fetched=["https://www.nfpa.org/13"],
        )
        assert out.accepted_sources == ["https://www.nfpa.org/13"]
        assert out.rejected_source_reasons == {}

    def test_nothing_rejected_leaves_reasons_empty(self):
        out = self._ground(cited=["https://dgs.ca.gov/p"], searched=["https://dgs.ca.gov/p"])
        assert out.rejected_sources == []
        assert out.rejected_source_reasons == {}

    def test_malformed_and_empty_citations_are_explained(self):
        # ``https://[invalid`` is what ``normalize_url`` cannot parse (an
        # unterminated IPv6 literal); a bare ``http://:80`` still normalizes
        # and is merely ungrounded.
        out = self._ground(cited=["", "https://[invalid"], searched=["https://dgs.ca.gov/p"])
        assert out.rejected_source_reasons == {
            "": REJECTION_EXPLANATION_EMPTY,
            "https://[invalid": REJECTION_EXPLANATION_MALFORMED,
        }


# ---------------------------------------------------------------------------
# 3. Cache round-trip (additive telemetry, no schema bump)
# ---------------------------------------------------------------------------


def _finding(issue: str = "NFPA 13 spacing") -> Finding:
    return Finding(
        severity="HIGH",
        fileName="Section_21_1313.docx",
        section="2.1",
        issue=issue,
        actionType="REPORT_ONLY",
        existingText=None,
        replacementText=None,
        codeReference="NFPA 13 §10.2.5",
        confidence=0.8,
    )


def _grounded_with_rejections() -> VerificationResult:
    return VerificationResult(
        verdict="CONFIRMED",
        explanation="NFPA 13 confirms the spacing limit.",
        sources=["https://www.nfpa.org/13"],
        accepted_sources=["https://www.nfpa.org/13"],
        searched_sources=["https://www.nfpa.org/13"],
        cited_sources=["https://www.nfpa.org/13", "https://reddit.com/r/x"],
        rejected_sources=[{"url": "https://reddit.com/r/x", "reason": REJECT_UNGROUNDED}],
        rejected_source_reasons={"https://reddit.com/r/x": "blocked domain: " + BLOCKED_DOMAIN_CATEGORY_AGGREGATOR},
        grounded=True,
        model_used="claude-sonnet-5",
        # ``put`` gates a CONFIRMED on a verbatim source quote as well as an
        # accepted citation, so the fixture must carry one to be cacheable.
        source_quote="The maximum distance between sprinklers shall not exceed 15 ft.",
        web_search_requests=2,
        successful_source_count=1,
    )


class TestCacheRoundTrip:
    def test_field_is_classified_as_persisted(self):
        assert "rejected_source_reasons" in _PERSISTED_FIELDS

    def test_result_to_dict_and_back(self):
        result = _grounded_with_rejections()
        d = _result_to_dict(result)
        assert d["rejected_source_reasons"] == result.rejected_source_reasons
        back = _result_from_dict(d, cache_status="hit")
        assert back.rejected_source_reasons == result.rejected_source_reasons
        assert back.rejected_source_reasons is not result.rejected_source_reasons

    def test_disk_round_trip_preserves_reasons(self, tmp_path: Path):
        cache = VerificationCache()
        finding = _finding()
        cache.put(finding, cycle=DEFAULT_CYCLE, result=_grounded_with_rejections())
        path = tmp_path / "cache.json"
        cache.save_to_disk(path)
        reloaded = VerificationCache()
        reloaded.load_from_disk(path)
        hit = reloaded.get(finding, cycle=DEFAULT_CYCLE)
        assert hit is not None
        assert hit.rejected_source_reasons == {
            "https://reddit.com/r/x": "blocked domain: " + BLOCKED_DOMAIN_CATEGORY_AGGREGATOR
        }
        assert hit.cache_status == "hit"

    def test_legacy_row_without_the_key_loads_empty_dict(self, tmp_path: Path):
        # Same schema version, no ``rejected_source_reasons`` key: legacy rows
        # must load with ``{}`` — no schema bump, no crash.
        path = tmp_path / "cache.json"
        payload = {
            "version": _CACHE_SCHEMA_VERSION,
            "saved_at": time.time(),
            "entries": {
                "legacy-key": {
                    "created_ts": time.time(),
                    "result": {
                        "verdict": "CONFIRMED",
                        "explanation": "legacy",
                        "sources": ["https://www.nfpa.org/13"],
                        "accepted_sources": ["https://www.nfpa.org/13"],
                        "searched_sources": ["https://www.nfpa.org/13"],
                        "cited_sources": ["https://www.nfpa.org/13"],
                        "rejected_sources": [{"url": "https://reddit.com/r/x", "reason": "ungrounded"}],
                        "grounded": True,
                        "model_used": "claude-sonnet-5",
                        "verification_profile": "code_standard",
                        "verification_mode": "standard_reasoning",
                        "source_quote": "legacy quote",
                        "web_search_requests": 1,
                        "successful_source_count": 1,
                        "search_error_count": 0,
                        # ``rejected_source_reasons`` absent.
                    },
                }
            },
        }
        path.write_text(json.dumps(payload), encoding="utf-8")
        cache = VerificationCache()
        assert cache.load_from_disk(path) == 1
        (entry,) = cache._entries.values()
        assert entry.result.rejected_source_reasons == {}
        assert entry.result.rejected_sources == [{"url": "https://reddit.com/r/x", "reason": "ungrounded"}]

    def test_garbage_values_are_coerced_defensively(self):
        d = _result_to_dict(_grounded_with_rejections())
        d["rejected_source_reasons"] = ["not", "a", "dict"]
        assert _result_from_dict(d, cache_status="hit").rejected_source_reasons == {}
        d["rejected_source_reasons"] = {"https://a/x": "", "": "reason", "https://b/y": "kept"}
        assert _result_from_dict(d, cache_status="hit").rejected_source_reasons == {"https://b/y": "kept"}


# ---------------------------------------------------------------------------
# 4. Both renderers show the reason next to the rejected URL
# ---------------------------------------------------------------------------


class _StubPipelineResult:
    def __init__(self, review_result: ReviewResult):
        self.review_result = review_result
        self.cross_check_result = None
        self.compliance_result = None
        self.files_reviewed = [review_result.findings[0].fileName]
        self.leed_alerts = []
        self.placeholder_alerts = []
        self.cycle_label = DEFAULT_CYCLE.label
        self.total_elapsed_seconds = 1.0


def _docx_text(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _result_with(vr: VerificationResult) -> _StubPipelineResult:
    f = _finding()
    f.verification = vr
    return _StubPipelineResult(ReviewResult(findings=[f]))


class TestRenderers:
    def test_docx_renders_reason_beside_each_rejected_url(self, tmp_path: Path):
        vr = _grounded_with_rejections()
        vr.rejected_sources.append({"url": "https://invented.example.com/fake", "reason": REJECT_UNGROUNDED})
        vr.rejected_source_reasons["https://invented.example.com/fake"] = REJECTION_EXPLANATION_UNGROUNDED
        out = tmp_path / "report.docx"
        export_report(_result_with(vr), out)
        text = _docx_text(out)
        assert "https://reddit.com/r/x [ungrounded] — blocked domain: " + BLOCKED_DOMAIN_CATEGORY_AGGREGATOR in text
        assert "https://invented.example.com/fake [ungrounded] — not among searched or fetched results" in text

    def test_html_mirrors_the_docx_panel(self):
        vr = _grounded_with_rejections()
        vr.rejected_sources.append({"url": "https://invented.example.com/fake", "reason": REJECT_UNGROUNDED})
        vr.rejected_source_reasons["https://invented.example.com/fake"] = REJECTION_EXPLANATION_UNGROUNDED
        from datetime import datetime

        html = render_html_report(_result_with(vr), generated_at=datetime(2026, 1, 1))
        assert "Unsupported / rejected sources" in html
        assert "[ungrounded]" in html
        assert "— blocked domain: " + BLOCKED_DOMAIN_CATEGORY_AGGREGATOR.replace("&", "&amp;") in html
        assert "— not among searched or fetched results" in html
        # The machine-readable payload carries the map too (escaped like the
        # rest of the embedded JSON, so search on the value text).
        assert "rejected_source_reasons" in html

    def test_legacy_result_without_reasons_renders_bare_reason_only(self, tmp_path: Path):
        vr = _grounded_with_rejections()
        vr.rejected_source_reasons = {}
        out = tmp_path / "report.docx"
        export_report(_result_with(vr), out)
        text = _docx_text(out)
        assert "https://reddit.com/r/x [ungrounded]" in text
        assert "blocked domain" not in text
        from datetime import datetime

        html = render_html_report(_result_with(vr), generated_at=datetime(2026, 1, 1))
        assert "[ungrounded]" in html and "— blocked domain" not in html

    def test_no_rejected_sources_renders_byte_identically(self, tmp_path: Path):
        # With nothing rejected the new field is inert: a result carrying the
        # default ``{}`` renders exactly like one whose attribute is missing.
        from datetime import datetime

        clean = VerificationResult(
            verdict="CONFIRMED",
            explanation="Confirmed.",
            sources=["https://www.nfpa.org/13"],
            accepted_sources=["https://www.nfpa.org/13"],
            grounded=True,
            model_used="claude-sonnet-5",
            web_search_requests=1,
        )
        stamp = datetime(2026, 1, 1)
        with_field = render_html_report(_result_with(clean), generated_at=stamp)

        class _NoField:
            """Duck-typed legacy result: no ``rejected_source_reasons`` attribute."""

            def __init__(self, src):
                for k, v in vars(src).items():
                    if k != "rejected_source_reasons":
                        setattr(self, k, v)

        without_field = render_html_report(_result_with(_NoField(clean)), generated_at=stamp)
        assert with_field == without_field
        assert "rejected" not in _docx_text(_write(tmp_path, clean)).lower()


def _write(tmp_path: Path, vr: VerificationResult) -> Path:
    out = tmp_path / "clean.docx"
    export_report(_result_with(vr), out)
    return out
