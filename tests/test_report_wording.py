"""Reports say what the code did (plan WP-17, chunk S18).

Two contracts:

* **A withheld edit is explainable wherever it was withheld.** The parser
  demotes a malformed or no-op edit to REPORT_ONLY and records why. A finding
  built any other way now gets the same rule at the step that normalizes it
  before its id is minted (the review dedup, the cross-check and compliance
  id stamping), so the banner's count, the per-finding note, the edit-action
  label, the severity counts, and the sidecar agree — and the banner's row no
  longer claims every demotion happened "at parse time".
* **Four outcomes are never merged.** "Analysis incomplete" (a spec that
  failed review, a repair still outstanding), "verification inconclusive",
  "operational failure", and "no issue found" are worded apart in the
  methodology note, the banner, the Findings section, and the GUI log — in
  both exporters from one wording.
"""
from __future__ import annotations

import re
from pathlib import Path

import pytest
from docx import Document

from src.orchestration import pipeline
from src.orchestration.collection_outcome import CollectionOutcome, RepairOutcome
from src.orchestration.pipeline import PipelineResult
from src.output.edit_sidecar import build_edit_instructions
from src.output.html_report_exporter import render_html_report
from src.output.report_exporter import (
    DEMOTION_ROW_LABEL,
    INCONCLUSIVE_ROW_LABEL,
    _aggregate_run_diagnostics,
    _demotion_row_value,
    _no_findings_notice,
    _report_only_note,
    _summarize_run_diagnostics,
    export_report,
)
from src.output.report_status import (
    EditActionLabel,
    ReportStatus,
    summarize_edit_actions,
    summarize_statuses,
    verification_outcome_counts_line,
    verification_outcome_groups,
    verification_outcome_sentence,
)
from src.review.reviewer import (
    HELD_ADDITION_REASON_PREFIX,
    Finding,
    ReviewResult,
    _parse_findings,
    edit_shape_problem,
    normalize_edit_shapes,
)
from src.verification.verifier import VerificationResult

NO_OP_REASON = "EDIT action is a no-op (existingText equals replacementText)"


def _finding(**overrides) -> Finding:
    fields = dict(
        severity="HIGH",
        fileName="210500.docx",
        section="2.01",
        issue="Wrong valve type.",
        actionType="EDIT",
        existingText="ball valve",
        replacementText="gate valve",
        codeReference="",
        confidence=0.8,
    )
    fields.update(overrides)
    return Finding(**fields)


def _no_op(**overrides) -> Finding:
    return _finding(existingText="same", replacementText="same", **overrides)


def _parsed_no_op() -> Finding:
    [finding] = _parse_findings(
        [
            {
                "severity": "HIGH",
                "fileName": "210500.docx",
                "section": "2.01",
                "issue": "Wrong valve type.",
                "actionType": "EDIT",
                "existingText": "same",
                "replacementText": "same",
                "confidence": 0.8,
            }
        ]
    )
    return finding


def _banner(findings, **kwargs) -> dict:
    return _summarize_run_diagnostics(
        findings=findings,
        status_counts=summarize_statuses(findings),
        edit_action_counts=summarize_edit_actions(findings),
        cross_check_result=None,
        **kwargs,
    )


def _doc_text(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


_PLAINTEXT_RE = re.compile(r'<pre id="sc-plaintext" hidden>(.*?)</pre>', re.DOTALL)


def _html_text(result) -> str:
    html = render_html_report(result, include_chat=False)
    match = _PLAINTEXT_RE.search(html)
    assert match is not None
    return match.group(1)


def _result(findings=(), **overrides) -> PipelineResult:
    fields = dict(
        review_result=ReviewResult(findings=list(findings)),
        files_reviewed=["210500.docx", "211313.docx"],
    )
    fields.update(overrides)
    return PipelineResult(**fields)


def _verified(verdict="CONFIRMED", **fields) -> VerificationResult:
    base = dict(
        verdict=verdict,
        explanation="checked",
        grounded=True,
        sources=["https://example.org/code"],
        source_quote="the code says so",
        outcome="verdict",
    )
    base.update(fields)
    return VerificationResult(**base)


# ---------------------------------------------------------------------------
# 1. A withheld edit is recorded where findings are normalized
# ---------------------------------------------------------------------------


class TestEditShapeNormalization:
    def test_the_reason_is_the_parsers_reason(self):
        assert edit_shape_problem(_no_op()) == NO_OP_REASON
        assert _parsed_no_op().demotion_reason == NO_OP_REASON

    def test_a_valid_edit_and_a_native_report_only_have_no_problem(self):
        assert edit_shape_problem(_finding()) is None
        assert edit_shape_problem(
            _finding(actionType="REPORT_ONLY", existingText=None, replacementText=None)
        ) is None

    def test_an_explicit_proposal_is_checked_like_the_legacy_fields(self):
        finding = _finding()
        proposal = finding.as_edit_proposal()
        finding.edit_proposal = type(proposal)(
            action_type="EDIT", existing_text="x", replacement_text="x"
        )
        assert edit_shape_problem(finding) == NO_OP_REASON
        assert finding.as_edit_proposal() is None

    def test_normalizing_demotes_as_the_parser_does(self):
        finding = _no_op()
        normalize_edit_shapes([finding])
        parsed = _parsed_no_op()
        for field in (
            "actionType",
            "demotion_reason",
            "existingText",
            "replacementText",
            "anchorText",
            "insertPosition",
            "edit_proposal",
        ):
            assert getattr(finding, field) == getattr(parsed, field), field

    def test_normalizing_is_idempotent_and_leaves_recorded_reasons(self):
        finding = _no_op()
        normalize_edit_shapes([finding])
        normalize_edit_shapes([finding])
        assert finding.demotion_reason == NO_OP_REASON
        held = _finding(
            actionType="REPORT_ONLY",
            existingText=None,
            replacementText=None,
            demotion_reason=f"{HELD_ADDITION_REASON_PREFIX}: r-1 not assessed",
        )
        normalize_edit_shapes([held])
        assert held.demotion_reason.startswith(HELD_ADDITION_REASON_PREFIX)

    def test_an_already_merged_groups_members_are_normalized(self):
        # A merged finding normalized on its own (dedup normalizes its inputs,
        # which are the members; this is the path for one built elsewhere).
        members = [_no_op(fileName="a.docx"), _no_op(fileName="b.docx")]
        merged = _finding(occurrence_originals=members)
        normalize_edit_shapes([merged])
        assert merged.demotion_reason is None  # the representative is valid
        assert [m.demotion_reason for m in members] == [NO_OP_REASON, NO_OP_REASON]

    def test_a_recorded_reason_is_never_overwritten(self):
        finding = _no_op(demotion_reason="recorded upstream")
        normalize_edit_shapes([finding])
        assert finding.demotion_reason == "recorded upstream"
        assert finding.as_edit_proposal() is None

    def test_a_valid_edit_is_untouched(self):
        finding = _finding()
        normalize_edit_shapes([finding])
        assert finding.actionType == "EDIT"
        assert finding.demotion_reason is None
        assert finding.as_edit_proposal() is not None

    def test_the_review_dedup_normalizes_before_the_id(self):
        [deduped] = pipeline._deduplicate_findings([_no_op()])
        [parsed] = pipeline._deduplicate_findings([_parsed_no_op()])
        assert deduped.demotion_reason == NO_OP_REASON
        # The id is the one the parser path mints: identity is taken after
        # the edit shape is settled, never before.
        assert deduped.finding_id == parsed.finding_id

    def test_a_merged_groups_members_are_normalized(self):
        first = _no_op(fileName="210500.docx")
        second = _no_op(fileName="211313.docx")
        [merged] = pipeline._deduplicate_findings([first, second])
        assert merged.demotion_reason == NO_OP_REASON
        assert all(m.demotion_reason == NO_OP_REASON for m in merged.occurrence_originals)

    @pytest.mark.parametrize(
        "stamp, prefix",
        [
            (pipeline.assign_cross_check_finding_ids, "cf-"),
            (pipeline.assign_compliance_finding_ids, "lc-"),
        ],
    )
    def test_the_other_passes_normalize_before_stamping(self, stamp, prefix):
        [stamped] = stamp([_no_op()])
        [parsed] = stamp([_parsed_no_op()])
        assert stamped.demotion_reason == NO_OP_REASON
        assert stamped.finding_id.startswith(prefix)
        assert stamped.finding_id == parsed.finding_id


class TestEverySurfaceAgrees:
    """The banner, the edit-action label, the severity count, the note, and
    the sidecar tell one story about a hand-built no-op edit."""

    def _findings(self):
        return pipeline._deduplicate_findings([_no_op(), _finding(fileName="211313.docx")])

    def test_counts(self):
        findings = self._findings()
        summary = _banner(findings)
        assert summary["demotion_count"] == 1
        assert summary["report_only"] == 1
        assert summary["edit_suggested"] == 1
        review = ReviewResult(findings=findings)
        assert review.high_count == 2  # the demoted finding is still a finding
        actions = summarize_edit_actions(findings)
        assert actions[EditActionLabel.REPORT_ONLY] == 1

    def test_the_sidecar_leaves_it_out(self):
        findings = self._findings()
        payload = build_edit_instructions(_result(findings))
        ids = {entry["finding_id"] for entry in payload["edits"]}
        demoted = next(f for f in findings if f.demotion_reason)
        assert demoted.finding_id not in ids
        assert len(payload["edits"]) == 1

    def test_the_note_names_the_reason_and_both_reports_print_it(self, tmp_path):
        findings = self._findings()
        demoted = next(f for f in findings if f.demotion_reason)
        note = _report_only_note(demoted)
        assert note.startswith(f"Edit proposal demoted to REPORT_ONLY: {NO_OP_REASON}.")
        out = export_report(_result(findings), tmp_path / "r.docx")
        text = _doc_text(out)
        assert note in text
        assert note in _html_text(_result(findings))
        for surface in (text, _html_text(_result(findings))):
            assert "at parse time" not in surface
            assert f"{DEMOTION_ROW_LABEL}\n1" in surface or f"{DEMOTION_ROW_LABEL}: 1" in surface


class TestTheDemotionRow:
    def _held(self):
        return _finding(
            actionType="REPORT_ONLY",
            existingText=None,
            replacementText=None,
            demotion_reason=f"{HELD_ADDITION_REASON_PREFIX}: r-1 not assessed.",
        )

    def test_held_additions_are_named_apart(self):
        summary = _banner([self._held(), _parsed_no_op()])
        assert summary["demotion_count"] == 2
        assert summary["held_addition_count"] == 1
        assert _demotion_row_value(summary) == "2 (1 compliance addition held)"
        assert _demotion_row_value({"demotion_count": 3}) == "3"

    def test_an_anchor_validation_demotion_counts(self):
        finding = _finding(existingText="not in the spec")
        from src.review.reviewer import validate_finding_anchors

        validate_finding_anchors([finding], {"210500.docx": "the spec text"})
        assert finding.demotion_reason == "existing text not found in 210500.docx"
        assert _banner([finding])["demotion_count"] == 1
        assert _report_only_note(finding).startswith(
            "Edit proposal demoted to REPORT_ONLY: existing text not found"
        )

    def test_a_program_sums_the_new_keys(self):
        one = _banner([self._held()])
        two = _banner([_parsed_no_op()])
        aggregate = _aggregate_run_diagnostics([("A", one), ("B", two)])
        assert aggregate["demotion_count"] == 2
        assert aggregate["held_addition_count"] == 1
        assert set(aggregate) == set(_banner([]))


# ---------------------------------------------------------------------------
# 2. Verification outcomes are named apart
# ---------------------------------------------------------------------------


def _mixed_findings() -> list[Finding]:
    return [
        _finding(verification=_verified("CONFIRMED")),
        _finding(verification=_verified("UNVERIFIED", grounded=False, sources=[])),
        _finding(
            verification=_verified(
                "UNVERIFIED",
                grounded=False,
                sources=[],
                verification_failed=True,
                outcome="transport_error",
            )
        ),
        _finding(verification=_verified("UNVERIFIED", grounded=False, sources=[], cache_status="local_skip")),
        _finding(),
    ]


class TestVerificationOutcomes:
    def test_each_group_is_counted_apart(self):
        counts = summarize_statuses(_mixed_findings())
        assert counts[ReportStatus.LOCALLY_CLASSIFIED] == 1, counts
        groups = verification_outcome_groups(counts)
        assert groups == {
            "verified": 1,
            "inconclusive": 1,
            "failed": 1,
            "local": 1,
            "not_checked": 1,
        }

    def test_the_sentence_names_every_group_and_merges_none(self):
        sentence = verification_outcome_sentence(summarize_statuses(_mixed_findings()))
        assert sentence.startswith("Verification outcomes for the 5 findings: ")
        for phrase in (
            "1 verified against a retrieved source",
            "1 inconclusive (the verifier ran but could not settle the claim)",
            "1 operational failure (nothing was reliably checked",
            "1 classified locally, without a web search",
            "1 not checked",
        ):
            assert phrase in sentence
        assert "could not be verified" not in sentence

    def test_the_edge_cases(self):
        assert verification_outcome_sentence({}) == "There were no findings to verify."
        assert "No verification outcomes were recorded" in verification_outcome_sentence(
            {ReportStatus.NOT_CHECKED: 3}
        )
        all_failed = verification_outcome_sentence({ReportStatus.VERIFICATION_FAILED: 2})
        assert "failed operationally for every finding" in all_failed
        # A run whose findings were all classified locally was not a failure
        # (the old note said "did not return usable results").
        local = verification_outcome_sentence({ReportStatus.LOCALLY_CLASSIFIED: 2})
        assert "2 classified locally" in local
        assert "usable results" not in local

    def test_the_gui_line(self):
        assert verification_outcome_counts_line([]) == ""
        assert verification_outcome_counts_line(_mixed_findings()) == (
            "Verification: 1 verified, 1 inconclusive, 1 operational failure, "
            "1 classified locally, 1 not checked"
        )

    def test_both_reports_print_the_sentence_and_the_rows(self, tmp_path):
        findings = _mixed_findings()
        sentence = verification_outcome_sentence(summarize_statuses(findings))
        result = _result(findings)
        word = _doc_text(export_report(result, tmp_path / "r.docx"))
        html = _html_text(result)
        for surface in (word, html):
            assert sentence in surface
            assert INCONCLUSIVE_ROW_LABEL in surface
            assert "Verification failures (operational)" in surface

    def test_the_banner_counts_inconclusive_apart_from_failures(self):
        summary = _banner(_mixed_findings())
        assert summary["verification_inconclusive"] == 1
        assert summary["verification_failed"] == 1


# ---------------------------------------------------------------------------
# 3. "No issues found" only when everything was analyzed
# ---------------------------------------------------------------------------


def _provisional_outcome(specs=("211313.docx",)) -> CollectionOutcome:
    return CollectionOutcome(
        batch_id="msgbatch_1",
        repair=RepairOutcome(state="pending", batch_id="msgbatch_2", specs=specs),
        submitted_specs=("210500.docx", "211313.docx"),
        failed_specs=tuple(specs),
    )


class TestTheEmptyFindingsNotice:
    def test_a_clean_run(self):
        assert _no_findings_notice(_result()) == ("No issues found.", True)
        assert _no_findings_notice(None) == ("No issues found.", True)

    def test_some_specs_failed_review(self):
        text, clean = _no_findings_notice(_result(failed_review_specs=["211313.docx"]))
        assert not clean
        assert text.startswith("No issues found in the 1 specification that was reviewed.")
        assert "1 specification failed review and was not reviewed" in text
        assert "211313.docx" in text

    def test_every_spec_failed_review(self):
        text, clean = _no_findings_notice(
            _result(failed_review_specs=["210500.docx", "211313.docx"])
        )
        assert not clean
        assert "no specification was reviewed" in text
        assert "No issues found" not in text

    def test_a_provisional_run(self):
        text, clean = _no_findings_notice(
            _result(
                failed_review_specs=["211313.docx"],
                collection_outcome=_provisional_outcome(),
            )
        )
        assert not clean
        assert "review repair batch is still outstanding" in text

    def test_both_reports_print_it(self, tmp_path):
        result = _result(failed_review_specs=["211313.docx"])
        text, _clean = _no_findings_notice(result)
        word = _doc_text(export_report(result, tmp_path / "r.docx"))
        html = render_html_report(result, include_chat=False)
        assert text in word
        assert "\nNo issues found.\n" not in f"\n{word}\n"
        assert 'class="sc-clean"' not in html
        assert text in _html_text(result)

    def test_a_clean_run_still_prints_green(self, tmp_path):
        result = _result()
        word = Document(str(export_report(result, tmp_path / "r.docx")))
        paragraph = next(p for p in word.paragraphs if p.text == "No issues found.")
        assert paragraph.runs[0].font.color.rgb == (0, 128, 0)
        assert '<p class="sc-clean">No issues found.</p>' in render_html_report(
            result, include_chat=False
        )


# ---------------------------------------------------------------------------
# 4. A merged group's issue suffix says where it was found
# ---------------------------------------------------------------------------


class TestTheMergedIssueSuffix:
    def test_across_files_it_names_the_files(self):
        [merged] = pipeline._deduplicate_findings(
            [_finding(fileName="a.docx"), _finding(fileName="b.docx")]
        )
        assert merged.issue.endswith("(found in 2 specs: a.docx, b.docx)")

    def test_within_one_file_it_counts_the_reports(self):
        # Two places in one file used to read "(found in 1 specs: a.docx)",
        # contradicted by the two edit locations listed under it (S12).
        [merged] = pipeline._deduplicate_findings(
            [
                _finding(fileName="a.docx", evidenceElementId="p4"),
                _finding(fileName="a.docx", evidenceElementId="p8"),
            ]
        )
        assert merged.issue.endswith("(reported 2 times in a.docx)")
        assert "1 specs" not in merged.issue

    def test_the_id_does_not_depend_on_the_suffix(self):
        members = [_finding(fileName="a.docx"), _finding(fileName="a.docx")]
        [merged] = pipeline._deduplicate_findings(members)
        [single] = pipeline._deduplicate_findings([_finding(fileName="a.docx")])
        assert merged.finding_id == single.finding_id
