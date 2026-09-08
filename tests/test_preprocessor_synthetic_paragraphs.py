"""Duplicate-paragraph detector ignores extractor-synthesized entries (B-28 c).

``extract_text_from_docx`` renders section headers / footers, text boxes,
footnotes and endnotes as labeled paragraphs (``[Header] …``, ``[Footer] …``,
``[Text Box] …``, ``[Footnote n] …``, ``[Endnote n] …``) after the body. A
running page header is emitted once per document section by construction,
so on every multi-section spec ``detect_duplicate_paragraphs`` flagged the
spec's own header as a copy-paste duplicate and fed that noise into the
paid review prompt via ``<pre_detected>``.

The fix lives in the detector only: extraction output is byte-unchanged
(the goldens pin it); the detector skips paragraphs that start with one of
the synthetic labels. Genuine repeated body paragraphs still flag.
"""
from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from src.core.code_cycles import CALIFORNIA_2025
from src.input.extractor import extract_text_from_docx
from src.input.preprocessor import (
    DETERMINISTIC_RULE_DUPLICATE_PARAGRAPH,
    detect_duplicate_paragraphs,
    preprocess_spec,
)

_LONG = (
    "SECTION 23 21 13 HYDRONIC PIPING - VALLEY UNIFIED SCHOOL DISTRICT - "
    "NEW CLASSROOM BUILDING - DSA APPLICATION 01-119000 - ISSUED FOR BID"
)
_BODY = (
    "Submittals shall be provided for all piping accessories within 10 days "
    "of award and shall include manufacturer cut sheets."
)

assert len(_LONG) >= 80 and len(_BODY) >= 80


class TestSyntheticPrefixSkip:
    @pytest.mark.parametrize(
        "prefix",
        ["[Header] ", "[Footer] ", "[Text Box] ", "[Footnote 1] ", "[Endnote 3] ", "[Footnote ?] "],
    )
    def test_repeated_synthetic_entry_not_flagged(self, prefix: str) -> None:
        content = f"1.01 SUMMARY\n\n{_BODY}\n\n{prefix}{_LONG}\n\n{prefix}{_LONG}"
        assert detect_duplicate_paragraphs(content, "s.docx") == []

    def test_genuine_body_duplicate_still_flagged_alongside_headers(self) -> None:
        content = (
            f"1.01 SUMMARY\n\n{_BODY}\n\n[Header] {_LONG}\n\n"
            f"2.01 PRODUCTS\n\n{_BODY}\n\n[Header] {_LONG}\n\n[Footer] {_LONG}\n\n[Footer] {_LONG}"
        )
        alerts = detect_duplicate_paragraphs(content, "s.docx")
        assert len(alerts) == 1
        assert alerts[0]["deterministic_rule"] == DETERMINISTIC_RULE_DUPLICATE_PARAGRAPH
        assert alerts[0]["match"].startswith("Submittals shall be provided")
        assert alerts[0]["occurrence_count"] == 2

    def test_label_must_be_at_paragraph_start(self) -> None:
        # A body paragraph that merely mentions the label mid-text is an
        # ordinary paragraph and still flags when repeated.
        para = f"Coordinate the [Header] block with the title sheet. {_BODY}"
        content = f"{para}\n\n{para}"
        assert len(detect_duplicate_paragraphs(content, "s.docx")) == 1

    def test_other_bracketed_body_text_still_flags(self) -> None:
        # Only the extractor's own labels are exempt; an author's bracketed
        # note is body text.
        para = f"[Note] {_BODY}"
        content = f"{para}\n\n{para}"
        assert len(detect_duplicate_paragraphs(content, "s.docx")) == 1

    def test_label_without_trailing_space_is_not_synthetic(self) -> None:
        # The extractor always emits ``[Header] `` + text; ``[Header]x``
        # cannot be one of its entries.
        para = f"[Header]{_LONG}"
        content = f"{para}\n\n{para}"
        assert len(detect_duplicate_paragraphs(content, "s.docx")) == 1


class TestEndToEndMultiSectionSpec:
    def test_linked_header_across_sections_not_flagged_but_body_dup_is(
        self, tmp_path: Path
    ) -> None:
        # The real-world shape: a spec with two sections whose header is
        # linked to the previous section, plus one genuinely duplicated
        # body clause.
        doc = Document()
        doc.add_paragraph("1.01 SUMMARY")
        doc.add_paragraph(_BODY)
        doc.sections[0].header.paragraphs[0].text = _LONG
        doc.add_section()
        doc.add_paragraph("2.01 PRODUCTS")
        doc.add_paragraph(_BODY)
        path = tmp_path / "23 21 13 - Hydronic Piping.docx"
        doc.save(path)

        spec = extract_text_from_docx(path)
        # Extraction output is untouched: the header is still emitted once
        # per section (the goldens pin the extraction shape).
        assert spec.content.count(f"[Header] {_LONG}") == 2

        result = preprocess_spec(spec.content, spec.filename, cycle=CALIFORNIA_2025)
        alerts = result.duplicate_paragraph_alerts
        assert len(alerts) == 1
        assert alerts[0]["match"].startswith("Submittals shall be provided")
        assert not any(a["match"].startswith("[Header]") for a in alerts)

    def test_multi_section_spec_with_clean_body_produces_no_duplicate_alerts(
        self, tmp_path: Path
    ) -> None:
        doc = Document()
        doc.add_paragraph("Body one.")
        doc.sections[0].header.paragraphs[0].text = _LONG
        doc.sections[0].footer.paragraphs[0].text = _LONG[::-1]
        doc.add_section()
        doc.add_paragraph("Body two.")
        doc.add_section()
        doc.add_paragraph("Body three.")
        path = tmp_path / "spec.docx"
        doc.save(path)

        spec = extract_text_from_docx(path)
        assert spec.content.count("[Header] ") == 3
        assert spec.content.count("[Footer] ") == 3
        result = preprocess_spec(spec.content, spec.filename, cycle=CALIFORNIA_2025)
        assert result.duplicate_paragraph_alerts == []
