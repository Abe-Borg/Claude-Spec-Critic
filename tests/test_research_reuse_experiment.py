"""Plan EX-05: reuse a completed requirements-research profile across runs.

Hermetic: every research call is a scripted fake, every clock is injected,
every cache file lives in ``tmp_path``. The decision record is
``plans/experiments/EX-05-research-reuse.md``.

What is pinned, by section:

* the switches (``SPEC_CRITIC_RESEARCH_CACHE``, its age limit) and that off is
  byte-identical — no file read or written, the same runner call, the same
  profile dict, no diagnostics or report change;
* the key: built from the requests the fan-out actually sends, covering every
  request field, split by every materially relevant input and by nothing that
  only changes visibility or cache placement;
* the store: same inputs reuse, changed inputs miss and say why, partial or
  failed research is never stored, age and named dates retire an entry,
  corrupted rows are rejected one by one, the refresh path, failures of the
  cache itself never fail a run, the size and LRU bounds, concurrent stores;
* the pipeline: a hit makes no research call and never replaces the
  operator's Project Context;
* provenance: the reuse record rides the profile through pending state, both
  reports, the banner (single-module and program), and diagnostics;
* the harness: the key matrix, pooled measurement, the profile diff, listing
  and deleting entries, and the protocol (NOT RUN).
"""
from __future__ import annotations

import dataclasses
import json
import threading
import time
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace

import pytest

from src.core import api_config
from src.core.project_profile import ProjectProfile
from src.orchestration.diagnostics import DiagnosticsReport
from src.research import RequirementsProfile, ResearchFanoutError
from src.research import requirements_research as rr
from src.research import research_cache as rc
from src.review.reviewer import ReviewResult
from tests.fixtures.fake_anthropic import research_tool_use_response
from tests.test_requirements_research import (
    FakeResearchClient,
    _LogCollector,
    _complete_profile,
    _dimension,
    _enabled_module,
    _route_by_marker,
)

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

DAY = 86400.0


def _noon(value: str) -> float:
    """Local noon of an ISO date, as an epoch timestamp (never near midnight)."""
    y, m, d = (int(p) for p in value.split("-"))
    return datetime(y, m, d, 12, 0, 0).timestamp()


_CURRENT_CLOCK: dict = {"now": time.time}


class _Clock:
    """An injected clock. The newest one also dates the research it runs.

    The fan-out stamps ``research_date`` from the wall clock and the cache
    checks it against its own clock's creation time, so a test clock must
    drive both (see the autouse fixture below).
    """

    def __init__(self, start: float):
        self.now = start
        _CURRENT_CLOCK["now"] = self

    def __call__(self) -> float:
        return self.now

    def advance(self, days: float) -> None:
        self.now += days * DAY


def _signals(block: str = "Standards cited with edition years:\n- NFPA 13 (2022)"):
    return SimpleNamespace(render_block=lambda: block)


def _two_dimension_module(**overrides):
    return _enabled_module(
        research_dimensions=(_dimension("alpha"), _dimension("beta")), **overrides
    )


def _client(counter: dict | None = None):
    """A fake client that answers every dimension of the test modules."""

    def route(kwargs):
        if counter is not None:
            counter["calls"] = counter.get("calls", 0) + 1
        return research_tool_use_response()

    return FakeResearchClient(route)


def _runner_with(client):
    def runner(*args, **kwargs):
        return rr.run_requirements_research(*args, client=client, **kwargs)

    return runner


def _reuse(
    cache,
    *,
    module=None,
    profile=None,
    signals=None,
    mode="reuse",
    counter=None,
    diag=None,
    log=None,
    max_age_days=30,
    runner=None,
):
    counter = counter if counter is not None else {}
    return rr.run_research_with_reuse(
        module or _enabled_module(),
        profile or _complete_profile(),
        mode=mode,
        corpus_signals=signals if signals is not None else _signals(),
        runner=runner or _runner_with(_client(counter)),
        cache=cache,
        max_age_days=max_age_days,
        diag=diag,
        log=log or _LogCollector(),
    )


def _stored_profile(research_date: str, items: list[dict] | None = None, *, dims=("alpha",), project=None) -> dict:
    project = project if project is not None else _complete_profile().to_dict()
    items = items if items is not None else [
        {
            "item_id": "r-000000000001",
            "dimension_id": dims[0],
            "topic": "Governing code",
            "category": "governing_code",
            "requirement": "The 2024 IBC as adopted governs.",
            "grounded": True,
            "accepted_sources": ["https://codes.example.gov/ibc"],
            "source_urls": ["https://codes.example.gov/ibc"],
            "confidence": 0.9,
            "actionability": "spec_requirement",
        }
    ]
    return {
        "items": items,
        "dimension_statuses": [
            {"dimension_id": d, "status": "completed", "item_count": 1, "grounded_count": 1,
             "web_search_requests": 3, "web_fetch_requests": 0, "error": ""}
            for d in dims
        ],
        "research_date": research_date,
        "project": project,
    }


def _synthetic_key(*, project=None, dims=("alpha",), module_id="m1", **overrides) -> rc.ResearchKey:
    project = project if project is not None else _complete_profile().to_dict()
    components = {name: f"v-{name}" for name in rc.KEY_COMPONENT_ORDER}
    components.update(
        project=rc.digest(project),
        module_id=module_id,
        dimension_ids=json.dumps(list(dims)),
        policy_version=rc.POLICY_VERSION,
        date_basis=rc.DATE_BASIS_CURRENT,
    )
    components.update(overrides)
    return rc.ResearchKey(
        components=components, module_id=module_id, project=project, dimension_ids=tuple(dims)
    )


def _read(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def _write(path: Path, payload: dict) -> None:
    path.write_text(json.dumps(payload), encoding="utf-8")


@pytest.fixture(autouse=True)
def _research_dated_by_the_test_clock(monkeypatch):
    _CURRENT_CLOCK["now"] = time.time
    monkeypatch.setattr(
        rr,
        "time",
        SimpleNamespace(
            strftime=lambda fmt: datetime.fromtimestamp(_CURRENT_CLOCK["now"]()).strftime(fmt)
        ),
    )


@pytest.fixture(autouse=True)
def _no_cache_env(monkeypatch, tmp_path):
    monkeypatch.delenv(api_config.ENV_RESEARCH_CACHE, raising=False)
    monkeypatch.delenv(api_config.ENV_RESEARCH_CACHE_MAX_AGE_DAYS, raising=False)
    # Never the operator's file, even by accident.
    monkeypatch.setenv(rc.ENV_RESEARCH_CACHE_PATH, str(tmp_path / "default_research_cache.json"))


# ---------------------------------------------------------------------------
# Switches
# ---------------------------------------------------------------------------


class TestSwitches:
    @pytest.mark.parametrize("value", [None, "", "  ", "0", "false", "no", "off", "OFF"])
    def test_off_values(self, monkeypatch, value):
        if value is not None:
            monkeypatch.setenv(api_config.ENV_RESEARCH_CACHE, value)
        assert api_config.research_cache_mode() is None

    @pytest.mark.parametrize(
        "value, expected",
        [("reuse", "reuse"), (" Reuse ", "reuse"), ("refresh", "refresh"), ("REFRESH", "refresh")],
    )
    def test_named_values(self, monkeypatch, value, expected):
        monkeypatch.setenv(api_config.ENV_RESEARCH_CACHE, value)
        assert api_config.research_cache_mode() == expected

    @pytest.mark.parametrize("value", ["1", "true", "on", "yes", "reuse!", "cache"])
    def test_no_truthy_shorthand_and_unknown_values_fail_closed(self, monkeypatch, value, caplog):
        monkeypatch.setenv(api_config.ENV_RESEARCH_CACHE, value)
        assert api_config.research_cache_mode() is None

    @pytest.mark.parametrize(
        "value, expected",
        [(None, 30), ("", 30), ("7", 7), ("1", 1), ("90", 90), ("0", 30), ("91", 30),
         ("-3", 30), ("abc", 30), ("7.5", 30)],
    )
    def test_max_age_days(self, monkeypatch, value, expected):
        if value is not None:
            monkeypatch.setenv(api_config.ENV_RESEARCH_CACHE_MAX_AGE_DAYS, value)
        assert api_config.research_cache_max_age_days() == expected

    def test_default_path_and_override(self, monkeypatch, tmp_path):
        monkeypatch.delenv(rc.ENV_RESEARCH_CACHE_PATH, raising=False)
        assert rc.default_research_cache_path() == Path.home() / ".spec_critic" / "research_cache.json"
        monkeypatch.setenv(rc.ENV_RESEARCH_CACHE_PATH, str(tmp_path / "x.json"))
        assert rc.default_research_cache_path() == tmp_path / "x.json"


# ---------------------------------------------------------------------------
# Off is byte-identical
# ---------------------------------------------------------------------------


def _write_spec(tmp_path: Path) -> Path:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Sprinkler systems shall comply with NFPA 13-2022.")
    path = tmp_path / "21 13 13 Wet-Pipe Sprinkler Systems.docx"
    doc.save(str(path))
    return path


def _phase(tmp_path, *, user_context="Operator context.", diag=None, log=None):
    from src.orchestration import pipeline

    spec = tmp_path / "21 13 13 Wet-Pipe Sprinkler Systems.docx"
    if not spec.exists():
        _write_spec(tmp_path)
    return pipeline._run_research_phase(
        module=_enabled_module(),
        profile=_complete_profile(),
        input_dir=tmp_path,
        files=[spec],
        user_context=user_context,
        log=log or _LogCollector(),
        progress=lambda *a, **k: None,
        diagnostics=diag,
    )


@pytest.fixture
def phase_env(monkeypatch):
    """Route the pipeline's research through a counting fake client."""
    counter: dict = {}
    client = _client(counter)
    monkeypatch.setattr(rr, "_get_client", lambda **_: client)
    monkeypatch.setattr(rr, "context_within_token_cap", lambda text: (0, True))
    monkeypatch.setattr(
        "src.research.corpus_signals.count_tokens", lambda text: len(text.split())
    )
    return SimpleNamespace(counter=counter, client=client)


class TestOffIsByteIdentical:
    def test_off_calls_the_runner_exactly_as_before(self, monkeypatch, tmp_path, phase_env):
        seen = {}

        def runner(module, profile, **kwargs):
            seen["kwargs"] = sorted(kwargs)
            return rr.run_requirements_research(module, profile, **kwargs)

        monkeypatch.setattr("src.research.run_requirements_research", runner)

        def forbidden(*a, **k):
            raise AssertionError("the research cache must not be touched when off")

        monkeypatch.setattr(rc.ResearchCache, "__init__", forbidden)
        _effective, profile = _phase(tmp_path)
        assert seen["kwargs"] == ["call_semaphore", "corpus_signals", "diag", "log", "progress"]
        assert "reuse" not in profile
        assert not (tmp_path / "default_research_cache.json").exists()

    def test_fresh_profile_serializes_without_reuse_or_usage(self):
        profile = rr.run_requirements_research(
            _enabled_module(), _complete_profile(), client=_client()
        )
        assert profile.run_usage is not None  # runtime only
        data = profile.to_dict()
        assert set(data) == {"items", "dimension_statuses", "research_date", "project"}
        restored = RequirementsProfile.from_dict(data)
        assert restored.reuse is None
        assert restored.to_dict() == data

    def test_off_run_writes_no_diagnostics_record(self, tmp_path, phase_env):
        diag = DiagnosticsReport()
        _phase(tmp_path, diag=diag)
        summary = diag.summary()
        assert "research_reuse" not in summary
        assert "Research reuse" not in diag.to_text()

    def test_requests_are_unchanged_by_the_builder_refactor(self):
        """The fan-out sends exactly what ``build_dimension_request`` builds."""
        module = _two_dimension_module()
        client = _client()
        rr.run_requirements_research(module, _complete_profile(), client=client)
        for call in client.calls:
            dimension = next(
                d for d in module.research_dimensions
                if d.dimension_id.upper() in call["messages"][0]["content"]
            )
            built = rr.build_dimension_request(module, _complete_profile(), dimension)
            expected = dict(built.request_kwargs)
            expected["messages"] = [{"role": "user", "content": built.user_message}]
            assert call == expected


# ---------------------------------------------------------------------------
# The key
# ---------------------------------------------------------------------------


def _strip_cc(value):
    if isinstance(value, dict):
        return {k: _strip_cc(v) for k, v in value.items() if k != "cache_control"}
    if isinstance(value, list):
        return [_strip_cc(v) for v in value]
    return value


class TestKeyIsTheRequests:
    def test_components_are_digests_of_what_was_sent(self):
        module = _two_dimension_module()
        signals = _signals()
        client = _client()
        rr.run_requirements_research(
            module, _complete_profile(), client=client, corpus_signals=signals
        )
        by_dim = {}
        for call in client.calls:
            for d in module.research_dimensions:
                if f"{d.dimension_id.upper()} research brief" in call["messages"][0]["content"]:
                    by_dim[d.dimension_id] = call
        sent = [by_dim["alpha"], by_dim["beta"]]
        key = rr.research_reuse_key(module, _complete_profile(), corpus_signals=signals)
        c = key.components
        assert c["user_messages"] == rc.digest([call["messages"] for call in sent])
        assert c["system_prompt"] == rc.digest([_strip_cc(call["system"]) for call in sent])
        assert c["tools"] == rc.digest([_strip_cc(call["tools"]) for call in sent])
        assert c["request_settings"] == rc.digest(
            [
                {
                    "model": call["model"],
                    "max_tokens": call["max_tokens"],
                    "thinking": call.get("thinking"),
                    "output_config": call.get("output_config"),
                }
                for call in sent
            ]
        )
        assert c["project"] == rc.digest(_complete_profile().to_dict())
        assert c["corpus_signals"] == rc.digest(signals.render_block())
        assert c["dimension_ids"] == json.dumps(["alpha", "beta"])
        assert c["module_id"] == module.module_id
        assert c["policy_version"] == rc.POLICY_VERSION
        assert c["date_basis"] == rc.DATE_BASIS_CURRENT
        assert key.key == rc.derive_key(c)

    def test_every_request_field_is_keyed(self):
        """A request field outside the key form would be sent without being keyed."""
        client = _client()
        rr.run_requirements_research(_enabled_module(), _complete_profile(), client=client)
        for call in client.calls:
            assert set(call) - {"messages"} <= rr.KEY_FORM_FIELDS
        # And the declared set is exactly what the key form reads.
        module = _enabled_module()
        built = rr.build_dimension_request(
            module, _complete_profile(), module.research_dimensions[0]
        )
        assert set(rr._key_form(built)) - {"messages"} == rr.KEY_FORM_FIELDS

    def test_corpus_signals_rendered_like_the_fan_out(self):
        """No signals and an empty scrape both render nothing, and key alike."""
        module = _enabled_module()
        none = rr.research_reuse_key(module, _complete_profile(), corpus_signals=None)
        empty = rr.research_reuse_key(module, _complete_profile(), corpus_signals=_signals(""))
        assert none.key == empty.key

    def test_the_key_matrix_behaves_as_designed(self):
        from evals import research_reuse as harness

        matrix = harness.key_matrix()
        assert matrix["all_as_expected"], [r for r in matrix["rows"] if not r["as_expected"]]
        rows = {r["case"]: r for r in matrix["rows"]}
        assert rows["same_inputs"]["same_key"] and rows["same_inputs"]["differing"] == []
        assert "project" in rows["city_case"]["differing"]
        assert "project" in rows["client_case"]["differing"]
        assert rows["corpus_edition"]["differing"][0] == "corpus_signals"
        assert rows["module"]["differing"][0] == "module_id"
        assert rows["model"]["differing"][0] == "model"
        assert rows["dimension_budget"]["differing"] == ["tools"]
        assert rows["persona"]["differing"] == ["system_prompt"]
        assert rows["deep_trace_display"]["same_key"]
        assert rows["cache_ttl"]["same_key"]

    def test_deep_trace_display_is_really_sent_but_not_keyed(self, monkeypatch):
        module = _enabled_module()
        base = rr.research_reuse_key(module, _complete_profile())
        monkeypatch.setattr(api_config, "deep_trace_recording", lambda: True)
        built = rr.build_dimension_request(
            module, _complete_profile(), module.research_dimensions[0]
        )
        assert built.request_kwargs["thinking"].get("display") == "summarized"
        assert rr.research_reuse_key(module, _complete_profile()).key == base.key

    @pytest.mark.parametrize(
        "patch, component",
        [
            ("blocklist", "tools"),
            ("strict_off", "tools"),
            ("effort", "request_settings"),
            ("max_tokens", "request_settings"),
            ("policy", "policy_version"),
            ("app_version", "app_version"),
            ("continuations", "max_continuations"),
        ],
    )
    def test_settings_that_change_research_split_the_key(self, monkeypatch, patch, component):
        module = _enabled_module()
        base = rr.research_reuse_key(module, _complete_profile())
        if patch == "blocklist":
            original = api_config.build_web_search_tool

            def tool(**kwargs):
                t = original(**kwargs)
                t["blocked_domains"] = [*t.get("blocked_domains", []), "example.org"]
                return t

            monkeypatch.setattr(rr, "build_web_search_tool", tool)
        elif patch == "strict_off":
            monkeypatch.setenv("SPEC_CRITIC_STRICT_TOOL_USE", "0")
        elif patch == "effort":
            monkeypatch.setattr(
                api_config, "effort_config_for", lambda **_: {"effort": "low"}
            )
        elif patch == "max_tokens":
            monkeypatch.setattr(rr, "research_max_tokens", lambda **_: 1234)
        elif patch == "policy":
            monkeypatch.setattr(rr, "RESEARCH_REUSE_POLICY_VERSION", "rr-test")
        elif patch == "app_version":
            monkeypatch.setattr(rr, "_APP_VERSION", "0.0.0-test")
        elif patch == "continuations":
            monkeypatch.setattr(rr, "RESEARCH_MAX_CONTINUATIONS", 99)
        changed = rr.research_reuse_key(module, _complete_profile())
        assert changed.key != base.key
        assert component in rc.differing_components(base.components, changed.components)

    def test_another_code_basis_splits_the_key(self):
        from src.modules import get_module

        fire = get_module("datacenter_fire")
        other_cycle = dataclasses.replace(fire.cycle, label=fire.cycle.label + "-test")
        base = rr.research_reuse_key(fire, _complete_profile())
        changed = rr.research_reuse_key(
            dataclasses.replace(fire, cycle=other_cycle), _complete_profile()
        )
        assert "cycle_label" in rc.differing_components(base.components, changed.components)

    def test_the_key_takes_no_project_context(self):
        """Research never reads Project Context, so the key cannot depend on it."""
        import inspect

        params = set(inspect.signature(rr.research_reuse_key).parameters)
        assert params == {"module", "profile", "corpus_signals", "model"}

    def test_key_requires_every_component(self):
        with pytest.raises(ValueError, match="missing components"):
            rc.ResearchKey(components={"model": "x"}, module_id="m", project={}, dimension_ids=())


# ---------------------------------------------------------------------------
# Lookup and store
# ---------------------------------------------------------------------------


class TestReuse:
    def test_same_inputs_reuse_without_a_research_call(self, tmp_path):
        clock = _Clock(_noon("2026-09-20"))
        cache = rc.ResearchCache(tmp_path / "rc.json", clock=clock)
        counter = {}
        first = _reuse(cache, counter=counter)
        assert counter["calls"] == 1 and first.reuse is None
        clock.advance(9)
        second_counter = {}
        log = _LogCollector()
        second = _reuse(cache, counter=second_counter, log=log)
        assert second_counter == {}  # no research request at all
        assert second.reuse is not None
        assert second.reuse["source"] == "research_cache"
        assert second.reuse["age_days"] == 9
        assert second.reuse["research_date"] == first.research_date
        assert second.reuse["max_age_days"] == 30
        assert second.reuse["date_basis"] == rc.DATE_BASIS_CURRENT
        # Same content as the research it replaces.
        stripped = second.to_dict()
        stripped.pop("reuse")
        assert stripped == first.to_dict()
        assert second.render_text() == first.render_text()
        # The log shows age and what the research governs.
        warning = log.messages("warning")[0]
        assert "researched" in warning and "9 days ago" in warning
        assert "SPEC_CRITIC_RESEARCH_CACHE=refresh" in warning
        assert any("Governing codes in the reused research" in m for m in log.messages("muted"))

    def test_a_hit_touches_last_used(self, tmp_path):
        clock = _Clock(_noon("2026-09-20"))
        cache = rc.ResearchCache(tmp_path / "rc.json", clock=clock)
        _reuse(cache)
        created = next(iter(_read(cache.path)["entries"].values()))["created_ts"]
        clock.advance(2)
        _reuse(cache)
        row = next(iter(_read(cache.path)["entries"].values()))
        assert row["created_ts"] == created
        assert row["last_used_ts"] == clock.now

    @pytest.mark.parametrize(
        "change, component",
        [
            ({"profile": ProjectProfile("markham", "ON", "Canada", "ExampleCo")}, "project"),
            ({"profile": ProjectProfile("Markham", "ON", "Canada", "OtherCo")}, "project"),
            ({"profile": ProjectProfile("Markham", "QC", "Canada", "ExampleCo")}, "project"),
            ({"signals": _signals("Standards cited with edition years:\n- NFPA 13 (2025)")}, "corpus_signals"),
            ({"module": _two_dimension_module()}, "dimension_ids"),
            ({"module": _enabled_module(research_persona="You are someone else.")}, "system_prompt"),
        ],
    )
    def test_a_changed_input_misses_and_names_why(self, tmp_path, change, component):
        cache = rc.ResearchCache(tmp_path / "rc.json", clock=_Clock(_noon("2026-09-20")))
        _reuse(cache)
        counter = {}
        diag = DiagnosticsReport()
        result = _reuse(cache, counter=counter, diag=diag, **change)
        assert counter["calls"] >= 1 and result.reuse is None
        record = diag.summary()["research_reuse"]
        if "profile" in change:
            # Another project is another subject: nothing stored for it yet.
            assert record["by_outcome"] == {"absent": 1}
        else:
            assert record["by_outcome"] == {"changed": 1}
            assert component in record["changed_components"]

    def test_absent_when_nothing_is_stored(self, tmp_path):
        cache = rc.ResearchCache(tmp_path / "rc.json", clock=_Clock(_noon("2026-09-20")))
        diag = DiagnosticsReport()
        _reuse(cache, diag=diag)
        assert diag.summary()["research_reuse"]["by_outcome"] == {"absent": 1}

    def test_country_alias_is_the_same_project(self, tmp_path):
        cache = rc.ResearchCache(tmp_path / "rc.json", clock=_Clock(_noon("2026-09-20")))
        _reuse(cache, profile=ProjectProfile("Markham", "ON", "Canada", "ExampleCo"))
        again = _reuse(cache, profile=ProjectProfile("Markham", "ON", "CA", "ExampleCo"))
        assert again.reuse is not None

    def test_partial_research_is_never_stored(self, tmp_path):
        module = _two_dimension_module()
        client = FakeResearchClient(
            _route_by_marker(
                {"ALPHA": [research_tool_use_response()], "BETA": [ValueError("boom")] * 3}
            )
        )
        cache = rc.ResearchCache(tmp_path / "rc.json", clock=_Clock(_noon("2026-09-20")))
        log = _LogCollector()
        diag = DiagnosticsReport()
        result = _reuse(cache, module=module, runner=_runner_with(client), log=log, diag=diag)
        assert result.failed_dimensions == 1
        assert not cache.path.exists()
        assert diag.summary()["research_reuse"]["store_outcomes"] == {"partial": 1}
        assert any("A partial profile is never reused" in m for m in log.messages("warning"))
        again = _reuse(cache, module=module)
        assert again.reuse is None

    def test_failed_research_raises_and_stores_nothing(self, tmp_path):
        client = FakeResearchClient(_route_by_marker({"ALPHA": [ValueError("boom")] * 3}))
        cache = rc.ResearchCache(tmp_path / "rc.json", clock=_Clock(_noon("2026-09-20")))
        diag = DiagnosticsReport()
        with pytest.raises(ResearchFanoutError):
            _reuse(cache, runner=_runner_with(client), diag=diag)
        assert not cache.path.exists()
        rollup = diag.summary()["research_reuse"]
        assert rollup["store_outcomes"] == {"not_attempted": 1}

    def test_store_refuses_partial_and_mismatched_profiles_directly(self, tmp_path):
        cache = rc.ResearchCache(tmp_path / "rc.json", clock=_Clock(_noon("2026-09-20")))
        key = _synthetic_key()
        partial = _stored_profile("2026-09-20")
        partial["dimension_statuses"][0]["status"] = "failed"
        assert cache.store(key, partial).outcome == rc.STORE_PARTIAL
        wrong_dims = _stored_profile("2026-09-20", dims=("beta",))
        assert cache.store(key, wrong_dims).outcome == rc.STORE_PARTIAL
        other_project = _stored_profile("2026-09-20", project={"city": "X"})
        assert cache.store(key, other_project).outcome == rc.STORE_PARTIAL
        reused = _stored_profile("2026-09-20")
        reused["reuse"] = {"source": "research_cache"}
        assert cache.store(key, reused).outcome == rc.STORE_PARTIAL
        empty = _stored_profile("2026-09-20")
        empty["dimension_statuses"] = []
        assert cache.store(key, empty).outcome == rc.STORE_PARTIAL
        assert not cache.path.exists()

    def test_a_row_the_reader_would_reject_is_not_written(self, tmp_path):
        """A research date that disagrees with the clock by more than a day."""
        cache = rc.ResearchCache(tmp_path / "rc.json", clock=_Clock(_noon("2026-09-29")))
        result = cache.store(_synthetic_key(), _stored_profile("2026-09-20"))
        assert result.outcome == rc.STORE_INVALID
        assert "research date disagrees" in result.reason
        assert not cache.path.exists()

    def test_too_large_is_not_stored(self, tmp_path, monkeypatch):
        monkeypatch.setattr(rc, "MAX_PROFILE_BYTES", 200)
        cache = rc.ResearchCache(tmp_path / "rc.json", clock=_Clock(_noon("2026-09-20")))
        result = cache.store(_synthetic_key(), _stored_profile("2026-09-20"))
        assert result.outcome == rc.STORE_TOO_LARGE
        assert not cache.path.exists()


class TestFreshness:
    def _stored(self, tmp_path, research_date="2026-09-20", items=None):
        clock = _Clock(_noon(research_date))
        cache = rc.ResearchCache(tmp_path / "rc.json", clock=clock)
        key = _synthetic_key()
        assert cache.store(key, _stored_profile(research_date, items)).outcome == rc.STORE_STORED
        return cache, clock, key

    def test_age_limit_is_inclusive_and_then_stale(self, tmp_path):
        cache, clock, key = self._stored(tmp_path)
        clock.advance(30)
        assert cache.lookup(key, max_age_days=30).outcome == rc.OUTCOME_HIT
        clock.now += 1
        lookup = cache.lookup(key, max_age_days=30)
        assert lookup.outcome == rc.OUTCOME_STALE
        assert "older than 30 day(s)" in lookup.reason
        assert lookup.age_days == 30

    def test_env_age_limit_reaches_the_orchestrator(self, tmp_path, monkeypatch):
        clock = _Clock(_noon("2026-09-20"))
        cache = rc.ResearchCache(tmp_path / "rc.json", clock=clock)
        rr.run_research_with_reuse(
            _enabled_module(), _complete_profile(), mode="reuse",
            corpus_signals=_signals(), runner=_runner_with(_client()), cache=cache,
        )
        monkeypatch.setenv(api_config.ENV_RESEARCH_CACHE_MAX_AGE_DAYS, "7")
        clock.advance(8)
        diag = DiagnosticsReport()
        result = rr.run_research_with_reuse(
            _enabled_module(), _complete_profile(), mode="reuse",
            corpus_signals=_signals(), runner=_runner_with(_client()), cache=cache, diag=diag,
        )
        assert result.reuse is None
        assert diag.summary()["research_reuse"]["stale_reasons"] == {"age": 1}

    @pytest.mark.parametrize(
        "text, researched, today, stale",
        [
            ("The new code takes effect January 18, 2027.", "2026-12-01", "2027-01-17", False),
            ("The new code takes effect January 18, 2027.", "2026-12-01", "2027-01-18", True),
            # A later year named on its own retires the entry when that year
            # begins, even before a later full date in the same text: the
            # conservative direction.
            ("The 2027 code takes effect January 18, 2027.", "2026-12-01", "2027-01-01", True),
            ("Effective 2027-01-05.", "2026-12-01", "2027-01-06", True),
            ("Effective 18 January 2027.", "2026-12-01", "2027-01-19", True),
            ("Adoption expected in March 2027.", "2027-02-01", "2027-02-28", False),
            ("Adoption expected in March 2027.", "2027-02-01", "2027-03-01", True),
            ("The 2027 edition will be adopted.", "2026-12-01", "2026-12-31", False),
            ("The 2027 edition will be adopted.", "2026-12-01", "2027-01-01", True),
            ("NFPA 13-2027 is expected.", "2026-12-01", "2027-01-01", True),
            # Read both ways: April 3 or March 4.
            ("Hearing on 03/04/2027.", "2027-02-01", "2027-03-03", False),
            ("Hearing on 03/04/2027.", "2027-02-01", "2027-03-04", True),
            # Dates on or before the research day are history, not a change.
            ("Adopted January 1, 2026; amended 2026-09-20.", "2026-09-20", "2026-12-15", False),
            ("Adopted September 2026.", "2026-09-20", "2026-10-15", False),
        ],
    )
    def test_named_dates_that_have_begun_retire_an_entry(self, tmp_path, text, researched, today, stale):
        items = [dict(_stored_profile(researched)["items"][0], notes=text)]
        cache, clock, key = self._stored(tmp_path, research_date=researched, items=items)
        # The age limit is not what this test is about.
        lookup = cache.lookup(key, max_age_days=3650, now=_noon(today))
        if stale:
            assert lookup.outcome == rc.OUTCOME_STALE
            assert "a date it names has begun since" in lookup.reason
            assert lookup.passed_dates
        else:
            assert lookup.outcome == rc.OUTCOME_HIT, lookup.reason

    @pytest.mark.parametrize("field", ["requirement", "notes", "topic", "code_reference", "authority"])
    def test_every_text_field_is_read_for_dates(self, field):
        profile = {"items": [{"item_id": "r-1", field: "effective June 1, 2027"}]}
        passed = rc.passed_named_dates(
            profile, researched_on=date(2027, 1, 1), today=date(2027, 6, 1)
        )
        assert [p.item_id for p in passed] == ["r-1"]

    def test_named_periods_parse_the_documented_forms(self):
        text = "Jan. 18, 2027; 18 Feb 2027; 2027-03-04; 5/6/2027; April 2027; 2031"
        starts = sorted(p.start for p in rc.named_periods(text))
        assert starts == [
            date(2027, 1, 18), date(2027, 2, 18), date(2027, 3, 4), date(2027, 4, 1),
            date(2027, 5, 6), date(2027, 6, 5), date(2031, 1, 1),
        ]
        # A year inside a date is not also a bare year.
        assert [p.start for p in rc.named_periods("2027-03-04")] == [date(2027, 3, 4)]
        assert rc.named_periods("February 30, 2027") == []

    def test_a_stale_entry_is_replaced_by_fresh_research(self, tmp_path):
        clock = _Clock(_noon("2026-08-01"))
        cache = rc.ResearchCache(tmp_path / "rc.json", clock=clock)
        _reuse(cache)
        clock.advance(40)
        counter = {}
        result = _reuse(cache, counter=counter)
        assert counter["calls"] == 1 and result.reuse is None
        rows = list(_read(cache.path)["entries"].values())
        assert len(rows) == 1 and rows[0]["created_ts"] == clock.now
        clock.advance(1)
        assert _reuse(cache).reuse is not None


class TestRefresh:
    def test_refresh_researches_again_and_replaces_the_entry(self, tmp_path):
        clock = _Clock(_noon("2026-09-20"))
        cache = rc.ResearchCache(tmp_path / "rc.json", clock=clock)
        _reuse(cache)
        clock.advance(3)
        counter = {}
        diag = DiagnosticsReport()
        log = _LogCollector()
        result = _reuse(cache, mode="refresh", counter=counter, diag=diag, log=log)
        assert counter["calls"] == 1 and result.reuse is None
        rollup = diag.summary()["research_reuse"]
        assert rollup["by_outcome"] == {"refresh": 1}
        assert rollup["lookups"] == 0 and rollup["hit_rate"] is None
        assert rollup["store_outcomes"] == {"stored": 1}
        assert any("refresh requested" in m for m in log.messages("info"))
        row = next(iter(_read(cache.path)["entries"].values()))
        assert row["created_ts"] == clock.now
        clock.advance(1)
        again = _reuse(cache)
        assert again.reuse["age_days"] == 1


# ---------------------------------------------------------------------------
# Corrupted entries and failures of the cache itself
# ---------------------------------------------------------------------------


def _valid_row(tmp_path) -> tuple[rc.ResearchCache, rc.ResearchKey, dict]:
    clock = _Clock(_noon("2026-09-20"))
    cache = rc.ResearchCache(tmp_path / "rc.json", clock=clock)
    key = _synthetic_key()
    cache.store(key, _stored_profile("2026-09-20"))
    return cache, key, _read(cache.path)


def _mutations():
    def edit_component(row):
        row["components"]["model"] = "tampered"

    def edit_key_slot(payload, key):
        payload["entries"]["not-the-key"] = payload["entries"].pop(key)

    def edit_profile(row):
        row["profile"]["items"][0]["requirement"] = "Edited by hand."

    def string_ts(row):
        row["created_ts"] = "yesterday"

    def nan_ts(row):
        row["created_ts"] = float("nan")

    def future_ts(row):
        row["created_ts"] = row["created_ts"] + 5 * DAY

    def future_ts_same_day(row):
        # Five hours ahead of the clock: the research date still agrees, so
        # only the future-time rule can reject it.
        row["created_ts"] = row["created_ts"] + 5 * 3600
        row["last_used_ts"] = row["created_ts"]

    def future_last_used(row):
        row["last_used_ts"] = row["last_used_ts"] + 5 * 3600

    def bad_last_used(row):
        row["last_used_ts"] = -1

    def partial_profile(row):
        row["profile"]["dimension_statuses"][0]["status"] = "failed"
        row["profile_sha256"] = rc.digest(row["profile"])

    def research_date(row):
        row["research_date"] = "2026-09-01"

    def research_date_vs_created(row):
        row["research_date"] = "2026-09-10"
        row["profile"]["research_date"] = "2026-09-10"
        row["profile_sha256"] = rc.digest(row["profile"])

    def other_project(row):
        row["profile"]["project"] = {"city": "Elsewhere"}
        row["profile_sha256"] = rc.digest(row["profile"])

    def subject_module(row):
        row["subject"]["module_id"] = "other"

    def usage_field(row):
        row["usage"] = {"surprise": 1}

    def usage_negative(row):
        row["usage"] = {"input_tokens": -5}

    def no_profile(row):
        del row["profile"]

    return {
        "edited component": edit_component,
        "edited profile": edit_profile,
        "string timestamp": string_ts,
        "nan timestamp": nan_ts,
        "future timestamp": future_ts,
        "future timestamp, same day": future_ts_same_day,
        "future last used": future_last_used,
        "bad last used": bad_last_used,
        "partial profile": partial_profile,
        "research date mismatch": research_date,
        "research date vs created": research_date_vs_created,
        "other project": other_project,
        "subject module": subject_module,
        "unknown usage field": usage_field,
        "negative usage": usage_negative,
        "no profile": no_profile,
        "__slot__": edit_key_slot,
    }


class TestCorruptedEntries:
    @pytest.mark.parametrize("name", list(_mutations()))
    def test_each_corruption_is_rejected_and_counted(self, tmp_path, name):
        cache, key, payload = _valid_row(tmp_path)
        mutate = _mutations()[name]
        if name == "__slot__":
            mutate(payload, key.key)
        else:
            mutate(payload["entries"][key.key])
        _write(cache.path, payload)
        lookup = cache.lookup(key, max_age_days=30)
        assert lookup.outcome == rc.OUTCOME_ABSENT
        assert lookup.rejected == 1

    def test_a_valid_row_beside_a_bad_one_still_hits_and_the_bad_one_is_dropped(self, tmp_path):
        cache, key, payload = _valid_row(tmp_path)
        good_other = _synthetic_key(module_id="m2")
        cache.store(good_other, _stored_profile("2026-09-20"))
        payload = _read(cache.path)
        payload["entries"][key.key]["profile"]["items"][0]["requirement"] = "Edited."
        _write(cache.path, payload)
        lookup = cache.lookup(good_other, max_age_days=30)
        assert lookup.outcome == rc.OUTCOME_HIT and lookup.rejected == 1
        # The hit's save wrote only what loaded: the bad row is gone.
        assert list(_read(cache.path)["entries"]) == [good_other.key]

    @pytest.mark.parametrize(
        "content",
        ["{not json", json.dumps([1, 2]), json.dumps({"version": 99, "entries": {}}),
         json.dumps({"version": 1, "entries": []}), json.dumps({"version": True, "entries": {}})],
    )
    def test_an_unusable_file_is_never_overwritten(self, tmp_path, content):
        path = tmp_path / "rc.json"
        path.write_text(content, encoding="utf-8")
        cache = rc.ResearchCache(path, clock=_Clock(_noon("2026-09-20")))
        diag = DiagnosticsReport()
        counter = {}
        result = _reuse(cache, diag=diag, counter=counter)
        assert counter["calls"] == 1 and result.reuse is None
        rollup = diag.summary()["research_reuse"]
        assert rollup["by_outcome"] == {"unreadable": 1}
        assert rollup["store_outcomes"] == {"unreadable_file": 1}
        assert path.read_text(encoding="utf-8") == content

    def test_a_write_failure_never_fails_the_run(self, tmp_path, monkeypatch):
        cache = rc.ResearchCache(tmp_path / "rc.json", clock=_Clock(_noon("2026-09-20")))

        def refuse(*a, **k):
            raise OSError("disk full")

        monkeypatch.setattr(rc.ResearchCache, "_save", refuse)
        diag = DiagnosticsReport()
        result = _reuse(cache, diag=diag)
        assert result.completed_dimensions == 1
        assert diag.summary()["research_reuse"]["store_outcomes"] == {"write_failed": 1}

    def test_a_lookup_that_raises_falls_back_to_research(self, tmp_path, monkeypatch):
        cache = rc.ResearchCache(tmp_path / "rc.json", clock=_Clock(_noon("2026-09-20")))

        def explode(*a, **k):
            raise RuntimeError("unexpected")

        monkeypatch.setattr(rc.ResearchCache, "lookup", explode)
        diag = DiagnosticsReport()
        counter = {}
        _reuse(cache, diag=diag, counter=counter)
        assert counter["calls"] == 1
        assert diag.summary()["research_reuse"]["by_outcome"] == {"unreadable": 1}

    def test_a_key_error_researches_without_the_cache(self, tmp_path, monkeypatch):
        cache = rc.ResearchCache(tmp_path / "rc.json", clock=_Clock(_noon("2026-09-20")))

        def explode(*a, **k):
            raise RuntimeError("no key")

        monkeypatch.setattr(rr, "research_reuse_key", explode)
        diag = DiagnosticsReport()
        log = _LogCollector()
        counter = {}
        _reuse(cache, diag=diag, log=log, counter=counter)
        assert counter["calls"] == 1
        assert not cache.path.exists()
        rollup = diag.summary()["research_reuse"]
        assert rollup["by_outcome"] == {"key_error": 1}
        assert any("could not build the lookup key" in m for m in log.messages("warning"))


class TestBoundsAndConcurrency:
    def test_lru_keeps_the_most_recently_used(self, tmp_path, monkeypatch):
        monkeypatch.setattr(rc, "MAX_ENTRIES", 3)
        clock = _Clock(_noon("2026-09-20"))
        cache = rc.ResearchCache(tmp_path / "rc.json", clock=clock)
        keys = [_synthetic_key(module_id=f"m{i}") for i in range(4)]
        for i, key in enumerate(keys[:3]):
            clock.now += 60
            cache.store(key, _stored_profile("2026-09-20"))
        clock.now += 60
        assert cache.lookup(keys[0], max_age_days=30).outcome == rc.OUTCOME_HIT
        clock.now += 60
        cache.store(keys[3], _stored_profile("2026-09-20"))
        stored = set(_read(cache.path)["entries"])
        assert stored == {keys[0].key, keys[2].key, keys[3].key}

    def test_concurrent_stores_lose_nothing(self, tmp_path):
        cache = rc.ResearchCache(tmp_path / "rc.json", clock=_Clock(_noon("2026-09-20")))
        keys = [_synthetic_key(module_id=f"m{i}") for i in range(8)]
        barrier = threading.Barrier(len(keys))
        errors = []

        def store(key):
            try:
                barrier.wait()
                assert cache.store(key, _stored_profile("2026-09-20")).outcome == rc.STORE_STORED
            except Exception as exc:  # pragma: no cover - surfaced below
                errors.append(exc)

        threads = [threading.Thread(target=store, args=(k,)) for k in keys]
        for t in threads:
            t.start()
        for t in threads:
            t.join()
        assert errors == []
        assert set(_read(cache.path)["entries"]) == {k.key for k in keys}

    def test_nothing_but_the_profile_and_digests_is_stored(self, tmp_path):
        clock = _Clock(_noon("2026-09-20"))
        cache = rc.ResearchCache(tmp_path / "rc.json", clock=clock)
        secret_block = "Client/owner documents named in the specifications:\n- SECRET-SPEC-EXCERPT"
        _reuse(cache, signals=_signals(secret_block))
        text = cache.path.read_text(encoding="utf-8")
        assert "SECRET-SPEC-EXCERPT" not in text
        row = next(iter(_read(cache.path)["entries"].values()))
        assert set(row) == {
            "key", "components", "subject", "created_ts", "last_used_ts",
            "research_date", "profile", "profile_sha256", "usage",
        }
        assert row["usage"]["dimension_calls"] == 1
        assert row["usage"]["web_search_requests"] >= 1

    def test_delete(self, tmp_path):
        cache, key, _payload = _valid_row(tmp_path)
        other = _synthetic_key(module_id="m2")
        cache.store(other, _stored_profile("2026-09-20"))
        assert cache.delete([key.key]) == 1
        assert list(_read(cache.path)["entries"]) == [other.key]
        assert cache.delete(None) == 1
        assert not cache.path.exists()
        cache.path.write_text("{garbage", encoding="utf-8")
        assert cache.delete(["x"]) == 0 and cache.path.exists()
        assert cache.delete(None) == 0 and not cache.path.exists()


# ---------------------------------------------------------------------------
# Pipeline
# ---------------------------------------------------------------------------


class TestPipeline:
    def test_second_run_reuses_and_keeps_the_operator_context_first(
        self, monkeypatch, tmp_path, phase_env
    ):
        monkeypatch.setenv(api_config.ENV_RESEARCH_CACHE, "reuse")
        effective_1, profile_1 = _phase(tmp_path, user_context="First context.")
        calls_after_first = phase_env.counter["calls"]
        assert "reuse" not in profile_1
        effective_2, profile_2 = _phase(tmp_path, user_context="New project constraint: X.")
        assert phase_env.counter["calls"] == calls_after_first  # no research call
        assert profile_2["reuse"]["source"] == "research_cache"
        assert effective_2.startswith("New project constraint: X.")
        assert "First context." not in effective_2
        # The profile block the reviewers see is the same as the first run's.
        assert effective_2[len("New project constraint: X."):] == effective_1[len("First context."):]

    def test_a_patched_runner_still_intercepts_the_miss(self, monkeypatch, tmp_path, phase_env):
        monkeypatch.setenv(api_config.ENV_RESEARCH_CACHE, "reuse")

        def boom(*a, **k):
            raise AssertionError("intercepted")

        monkeypatch.setattr("src.research.run_requirements_research", boom)
        with pytest.raises(AssertionError, match="intercepted"):
            _phase(tmp_path)

    def test_the_reuse_record_survives_pending_state(self, monkeypatch, tmp_path, phase_env):
        from src.orchestration import batch_resume

        monkeypatch.setenv(api_config.ENV_RESEARCH_CACHE, "reuse")
        _phase(tmp_path)
        _effective, profile = _phase(tmp_path)
        pending = batch_resume._pending_batch_from_mapping(
            {"batch_id": "msgbatch_x", "input_dir": str(tmp_path), "requirements_profile": profile}
        )
        state = tmp_path / "pending.json"
        assert batch_resume.save_pending_batch(pending, path=state)
        loaded = batch_resume.load_pending_batch(path=state)
        assert loaded.requirements_profile["reuse"] == profile["reuse"]
        assert RequirementsProfile.from_dict(loaded.requirements_profile).reuse == profile["reuse"]

    def test_program_modules_share_one_file_safely(self, tmp_path):
        """Two modules researching concurrently both store their profiles."""
        cache = rc.ResearchCache(tmp_path / "rc.json", clock=_Clock(_noon("2026-09-20")))
        modules = [
            _enabled_module(),
            dataclasses.replace(_enabled_module(), module_id="second_module"),
        ]
        barrier = threading.Barrier(2)
        errors = []

        def run(module):
            try:
                barrier.wait()
                _reuse(cache, module=module)
            except Exception as exc:  # pragma: no cover
                errors.append(exc)

        threads = [threading.Thread(target=run, args=(m,)) for m in modules]
        for t in threads:
            t.start()
        for t in threads:
            t.join()
        assert errors == []
        assert len(_read(cache.path)["entries"]) == 2
        for module in modules:
            assert _reuse(cache, module=module).reuse is not None


# ---------------------------------------------------------------------------
# Reports, banner, diagnostics
# ---------------------------------------------------------------------------


def _reused_profile_dict(age_days: int = 9) -> dict:
    data = _stored_profile("2026-09-20")
    data["reuse"] = {
        "source": "research_cache",
        "policy_version": rc.POLICY_VERSION,
        "key": "k" * 64,
        "researched_at": "2026-09-20T12:00:00+00:00",
        "research_date": "2026-09-20",
        "age_days": age_days,
        "max_age_days": 30,
        "reused_at": "2026-09-29T12:00:00+00:00",
        "date_basis": rc.DATE_BASIS_CURRENT,
    }
    return data


def _pipeline_result(requirements_profile):
    from src.orchestration.pipeline import PipelineResult

    module = _enabled_module()
    return PipelineResult(
        review_result=ReviewResult(findings=[]),
        files_reviewed=["21 13 13 Wet-Pipe.docx"],
        cycle_label=module.cycle.label,
        module_id=module.module_id,
        requirements_profile=requirements_profile,
    )


def _summary_for(requirements_profile) -> dict:
    from src.output.report_exporter import _summarize_run_diagnostics

    return _summarize_run_diagnostics(
        findings=[],
        status_counts={},
        edit_action_counts={},
        cross_check_result=None,
        pipeline_result=_pipeline_result(requirements_profile),
    )


def _docx_text(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


class TestReports:
    def test_notice_wording(self):
        notice = rc.reuse_notice(_reused_profile_dict())
        assert "reused from the research cache, not re-run" in notice
        assert "researched 2026-09-20 (9 days before this run)" in notice
        assert "SPEC_CRITIC_RESEARCH_CACHE=refresh" in notice
        assert rc.reuse_notice(_stored_profile("2026-09-20")) is None
        assert rc.age_phrase(0) == "less than a day"
        assert rc.age_phrase(1) == "1 day"

    def test_word_report_says_the_research_was_reused(self, tmp_path):
        from src.output.report_exporter import export_report

        out = tmp_path / "report.docx"
        export_report(_pipeline_result(_reused_profile_dict()), out)
        text = _docx_text(out)
        assert rc.reuse_notice(_reused_profile_dict()) in text
        assert "reused from the research cache (researched 9 days before this run)" in text

    def test_html_report_says_the_research_was_reused(self):
        from src.output.html_report_exporter import render_html_report

        html = render_html_report(_pipeline_result(_reused_profile_dict()), include_chat=False)
        assert "reused from the research cache, not re-run" in html
        assert "reused from the research cache (researched 9 days before this run)" in html

    def test_fresh_research_reports_are_unchanged(self, tmp_path):
        from src.output.html_report_exporter import render_html_report
        from src.output.report_exporter import export_report

        fresh = _stored_profile("2026-09-20")
        out = tmp_path / "report.docx"
        export_report(_pipeline_result(fresh), out)
        assert "research cache" not in _docx_text(out)
        assert "research cache" not in render_html_report(_pipeline_result(fresh), include_chat=False)
        research = _summary_for(fresh)["research"]
        assert set(research) == {
            "dimensions_total", "dimensions_completed", "dimensions_failed",
            "item_count", "ungrounded_count",
        }

    def test_banner_row(self):
        from src.output.report_exporter import _research_banner_row

        base = {"dimensions_total": 4, "dimensions_completed": 4, "dimensions_failed": 0,
                "item_count": 10, "ungrounded_count": 1}
        assert _research_banner_row(base) == ("4 of 4 dimensions completed; 10 items (1 ungrounded)", False)
        one = dict(base, reused_modules=1, reused_oldest_age_days=0)
        assert _research_banner_row(one)[0].endswith(
            "; reused from the research cache (researched less than a day before this run)"
        )
        two = dict(base, reused_modules=2, reused_oldest_age_days=12)
        assert "for 2 modules (oldest researched 12 days before this run)" in _research_banner_row(two)[0]

    def test_program_aggregation_counts_modules_and_keeps_the_oldest_age(self):
        from src.output.report_exporter import _aggregate_run_diagnostics

        reused_old = _summary_for(_reused_profile_dict(12))
        reused_new = _summary_for(_reused_profile_dict(2))
        fresh = _summary_for(_stored_profile("2026-09-20"))
        combined = _aggregate_run_diagnostics(
            [("A", reused_old), ("B", reused_new), ("C", fresh)]
        )
        assert combined["research"]["reused_modules"] == 2
        assert combined["research"]["reused_oldest_age_days"] == 12
        plain = _aggregate_run_diagnostics([("C", fresh)])
        assert "reused_modules" not in plain["research"]

    def test_profile_export_carries_the_reuse_record(self):
        from src.output.edit_sidecar import build_requirements_profile_export

        exported = build_requirements_profile_export(_pipeline_result(_reused_profile_dict()))
        assert exported["requirements_profile"]["reuse"]["age_days"] == 9


class TestDiagnostics:
    def test_reuse_events_are_never_priced(self, tmp_path):
        clock = _Clock(_noon("2026-09-20"))
        cache = rc.ResearchCache(tmp_path / "rc.json", clock=clock)
        diag = DiagnosticsReport()
        _reuse(cache, diag=diag)
        before = diag.summary()["cost_summary"]
        clock.advance(1)
        _reuse(cache, diag=diag)
        after = diag.summary()
        assert after["cost_summary"] == before
        rollup = after["research_reuse"]
        assert rollup["hits"] == 1 and rollup["lookups"] == 2 and rollup["hit_rate"] == 0.5
        assert rollup["saved"]["dimension_calls"] == 1
        assert rollup["saved_by_model"][api_config.RESEARCH_MODEL_DEFAULT]["input_tokens"] > 0
        assert rollup["oldest_hit_age_days"] == 1
        assert "Research reuse (EX-05): 1 absent, 1 hit" in diag.to_text()

    def test_summarize_ignores_garbage(self):
        rollup = rc.summarize_research_reuse(
            [{"outcome": "hit", "saved": {"input_tokens": "x"}, "age_days": "y", "rejected": "z"}, "junk"]
        )
        assert rollup["hits"] == 1 and rollup["saved"]["input_tokens"] == 0
        assert rc.summarize_research_reuse([]) is None
        assert rc.summary_line(None) is None


# ---------------------------------------------------------------------------
# The harness
# ---------------------------------------------------------------------------


class TestHarness:
    def test_measure_pools_runs(self):
        from evals import research_reuse as harness

        def summary(hits, lookups, model="claude-sonnet-5-5"):
            return {
                "research_reuse": {
                    "lookups": lookups,
                    "hits": hits,
                    "by_outcome": {"hit": hits, "absent": lookups - hits},
                    "by_mode": {"reuse": lookups},
                    "changed_components": {"corpus_signals": 1},
                    "stale_reasons": {},
                    "store_outcomes": {"stored": lookups - hits},
                    "rejected_entries": 0,
                    "saved": {"dimension_calls": 4 * hits, "web_search_requests": 40 * hits},
                    "saved_by_model": {model: {"input_tokens": 1_000_000 * hits, "web_search_requests": 40 * hits}},
                    "oldest_hit_age_days": 5,
                }
            }

        pooled = harness.measure([summary(1, 2), summary(2, 3), {"other": 1}])
        assert pooled["runs_with_research_reuse"] == 2
        assert pooled["lookups"] == 5 and pooled["hits"] == 3 and pooled["hit_rate"] == 0.6
        assert pooled["hit_rate_wilson95"] is not None
        assert pooled["saved"]["dimension_calls"] == 12
        assert pooled["miss_components"] == {"corpus_signals": 2}
        # 3M input tokens at Sonnet 5.5's $2/MTok plus 120 searches at $10/1k.
        assert pooled["saved_cost"]["usd"] == pytest.approx(6.0 + 1.2)
        unknown = harness.saved_cost_lower_bound({"not-a-model": {"input_tokens": 5}})
        assert unknown["unpriced_models"] == ["not-a-model"]

    def test_saved_cost_prices_cache_writes_at_the_lower_rate(self):
        from evals import research_reuse as harness

        cost = harness.saved_cost_lower_bound(
            {"claude-sonnet-5-5": {"cache_creation_input_tokens": 1_000_000}}
        )
        assert cost["usd"] == pytest.approx(2.0 * 1.25)

    def test_diff_profiles(self, tmp_path):
        from evals import research_reuse as harness

        a = _stored_profile("2026-09-01")
        b = _stored_profile("2026-09-29")
        b["items"] = [
            dict(a["items"][0]),
            dict(a["items"][0], item_id="r-000000000002", requirement="New amendment.",
                 category="local_amendment", code_reference="Ord. 12"),
            # Neither of these is a controlling requirement.
            dict(a["items"][0], item_id="r-000000000003", requirement="Unverified claim.",
                 category="governing_code", code_reference="IBC 2027", grounded=False),
            dict(a["items"][0], item_id="r-000000000004", requirement="Permit fee due.",
                 category="ahj_requirement", actionability="process_advisory"),
        ]
        diff = harness.diff_profiles(a, b)
        assert diff["shared_items"] == 1 and diff["jaccard"] == 0.25
        assert [r["item_id"] for r in diff["only_in_fresh"]] == [
            "r-000000000002", "r-000000000003", "r-000000000004",
        ]
        assert [r["item_id"] for r in diff["controlling_only_in_fresh"]] == ["r-000000000002"]
        assert diff["governing_references"]["only_in_fresh"] == [["local_amendment", "ord. 12"]]
        flipped = dict(b, items=[dict(a["items"][0], grounded=False)])
        assert [r["item_id"] for r in harness.diff_profiles(a, flipped)["grounding_changed"]] == [
            a["items"][0]["item_id"]
        ]
        path = tmp_path / "program.profile.json"
        path.write_text(json.dumps({"module_profiles": {"a": {"requirements_profile": a},
                                                        "b": {"requirements_profile": b}}}))
        with pytest.raises(ValueError, match="--module"):
            harness.load_profile(path)
        assert harness.load_profile(path, module_id="b") == b

    def test_eval_arms_never_touch_the_operators_research_cache(self, tmp_path):
        """EX-03's arm environment points the research cache into the arm's directory."""
        from evals import model_effort

        assert "SPEC_CRITIC_RESEARCH_CACHE_PATH" in model_effort.ARM_STATE_ENV
        arm = model_effort.ARMS[model_effort.BASELINE_ARM]
        env = model_effort.arm_environment(
            arm, state_dir=tmp_path, base_env={"SPEC_CRITIC_RESEARCH_CACHE": "reuse"}
        )
        assert env["SPEC_CRITIC_RESEARCH_CACHE_PATH"] == str(tmp_path / "research_cache.json")
        assert "SPEC_CRITIC_RESEARCH_CACHE" not in env

    def test_cli(self, tmp_path, capsys):
        from evals import research_reuse as harness

        cache, key, _payload = _valid_row(tmp_path)
        assert harness.main(["entries", "--path", str(cache.path)]) == 0
        listing = json.loads(capsys.readouterr().out)
        assert [e["key"] for e in listing["entries"]] == [key.key]
        assert "profile" not in listing["entries"][0]
        assert harness.main(["protocol"]) == 0
        protocol = json.loads(capsys.readouterr().out)
        assert protocol["evaluation"]["status"] == "NOT RUN"
        export = tmp_path / "diag.json"
        export.write_text(json.dumps({"summary": {"research_reuse": {"lookups": 1, "hits": 1}}}))
        assert harness.main(["measure", str(export)]) == 0
        assert json.loads(capsys.readouterr().out)["hits"] == 1
        assert harness.main(["delete", "--all", "--path", str(cache.path)]) == 0
        assert json.loads(capsys.readouterr().out)["deleted"] == 1
        assert not cache.path.exists()
