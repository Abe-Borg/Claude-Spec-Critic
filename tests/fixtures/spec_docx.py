"""Deterministic DOCX builders for extraction, detector, routing, and applier tests.

Implementation plan WP-01 (chunk S01). Later chunks test Word structures the
extractor does not read yet (content controls, field results, smart tags,
automatic numbering), so every test needs the same small, anonymized
documents. They are built here in code rather than committed as binaries:
the XML a test depends on is visible where it is written, and a builder can
be varied one element at a time.

Three rules keep these fixtures honest:

* **Clean and defective are separate builders.** A clean document is the
  control for its detector category; a mutation changes exactly one thing
  about a clean document and names the single alert it should cause. A
  document that is clean *and* carries a TBD cannot prove the TBD detector
  works (plan WP-01 item 4). :func:`changed_blocks` lets a test prove a
  mutation differs from its clean source by one block.
* **Everything is deterministic.** No wall clock, no random ids, no counters
  shared between documents: two calls to a builder produce identical XML,
  so an extraction result can be pinned exactly.
* **Nothing is written except where the caller says.** Builders return an
  in-memory ``Document``; :func:`save_docx` writes it into a directory the
  test supplies (always a pytest ``tmp_path``). No real project document is
  ever read or altered.

The Word structures are built from raw WordprocessingML fragments, because
python-docx has no API for most of them. Each fragment helper returns an
XML string using the ``w:`` / ``r:`` prefixes; :class:`SpecDocBuilder`
parses fragments inside a namespaced wrapper and moves the resulting
elements into the document body, so the elements are the ordinary
python-docx ``CT_*`` classes the extractor and the applier expect.

**Checked against a real word processor.** When these builders were written
(2026-09-23, chunk S01) every ready-made document was opened in LibreOffice
Writer 24.2 and exported to text. The automatically numbered fixture
displayed exactly ``PART 1 GENERAL`` / ``1.01 SUMMARY`` / ``A. Provide the
specified piping system.`` … ``3.01 INSTALLATION``, and every wrapped
sentinel (block and inline controls, both drop-downs, the ``23 05 00``
field result, the smart tag, the insertion inside a hyperlink) was visible
text — the text the extractor did not read at the time. That check is not
repeated in CI (no word processor there); ``tests/test_spec_docx_fixtures.py``
pins the XML that produced it.
"""
from __future__ import annotations

from dataclasses import dataclass, replace
from pathlib import Path
from typing import Iterable, Sequence
from xml.sax.saxutils import escape, quoteattr

from docx import Document
from docx.document import Document as DocxDocument
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

#: Author and date stamped on every tracked revision these builders write.
#: Fixed values keep the XML identical from run to run.
REVISION_AUTHOR = "Fixture Author"
REVISION_DATE = "2026-01-01T00:00:00Z"

#: The two namespaces fragments may use. ``r`` is needed only by hyperlinks.
_FRAGMENT_NAMESPACES = ("w", "r")


# ===========================================================================
# 1. The clean three-PART specification (plan Appendix A) and its variants
# ===========================================================================


@dataclass(frozen=True)
class Para:
    """One body paragraph of a structured spec fixture.

    ``role`` says what the paragraph *is* in CSI terms — ``"part"``,
    ``"article"``, or ``"body"`` — independent of how it is numbered, so a
    test can ask "which paragraphs are headings?" without re-deriving it
    from the text. ``text`` is the literal text typed in the document.

    ``auto_label`` is set only on automatically numbered paragraphs: the
    label Word *displays* (``"PART 1"``, ``"1.01"``, ``"A."``) but that is
    not stored in any run. ``level`` is the list level the numbering
    definition assigns it (0 = PART, 1 = article, 2 = lettered paragraph).
    """

    text: str
    role: str
    auto_label: str | None = None
    level: int | None = None

    @property
    def displayed_text(self) -> str:
        """The paragraph as a reader sees it in Word."""
        if self.auto_label is None:
            return self.text
        return f"{self.auto_label} {self.text}"

    @property
    def is_heading(self) -> bool:
        return self.role in ("part", "article")


@dataclass(frozen=True)
class TableBlock:
    """A body-level table: a tuple of rows, each a tuple of cell texts."""

    rows: tuple[tuple[str, ...], ...]
    role: str = "body"

    @property
    def displayed_text(self) -> str:
        return "\n\n".join(" | ".join(row) for row in self.rows)

    @property
    def is_heading(self) -> bool:
        return False


Block = Para | TableBlock

#: Plan Appendix A, verbatim: manual numbering typed as text, one ordinary
#: body paragraph under every article. It must produce no empty-heading and
#: no duplicate-heading alerts (before chunk S03 it produced three false
#: "Empty section" alerts, one per PART — the P1-1 defect).
CLEAN_THREE_PART: tuple[Block, ...] = (
    Para("PART 1 GENERAL", "part"),
    Para("1.01 SUMMARY", "article"),
    Para("A. Provide the specified piping system.", "body"),
    Para("1.02 SUBMITTALS", "article"),
    Para("A. Submit product data before fabrication.", "body"),
    Para("PART 2 PRODUCTS", "part"),
    Para("2.01 MATERIALS", "article"),
    Para("A. Provide materials meeting the scheduled requirements.", "body"),
    Para("PART 3 EXECUTION", "part"),
    Para("3.01 INSTALLATION", "article"),
    Para("A. Install in accordance with the approved product instructions.", "body"),
)

#: The body of 2.01 MATERIALS in the table-only variant. No cell starts with
#: a number, so the variant exercises "an article whose only body is a
#: table" and nothing else (a leading quantity would also exercise the
#: separate rule that a quantity line is body text, not a heading).
MATERIALS_TABLE_ROWS: tuple[tuple[str, ...], ...] = (
    ("Component", "Material"),
    ("Piping", "Copper tube, Type L"),
    ("Fittings", "Wrought copper, solder joint"),
)


def clean_three_part_blocks() -> tuple[Block, ...]:
    """The clean control (plan Appendix A)."""
    return CLEAN_THREE_PART


def table_only_article_blocks() -> tuple[Block, ...]:
    """Clean variant: 2.01 MATERIALS's only body is a table.

    Still clean — an article whose content is a table is not empty.
    """
    blocks = list(CLEAN_THREE_PART)
    index = _index_of(blocks, "A. Provide materials meeting the scheduled requirements.")
    blocks[index] = TableBlock(MATERIALS_TABLE_ROWS)
    return tuple(blocks)


#: Displayed labels for the automatic-numbering variant, by level. Level 0
#: renders ``PART %1``, level 1 ``%1.%2`` with a two-digit article number
#: (``decimalZero``), level 2 ``%3.`` as an upper-case letter.
_AUTO_LEVEL_BY_ROLE = {"part": 0, "article": 1, "body": 2}


def auto_numbered_blocks() -> tuple[Block, ...]:
    """Clean variant: the same structure, numbered by Word, not typed.

    Every paragraph's ``text`` loses its typed number and carries it as
    ``auto_label`` instead, so ``displayed_text`` equals the clean fixture's
    text exactly. Today the extractor reads only the literal text (the
    WP-03 defect S14 fixes), so the labels never reach review.
    """
    out: list[Block] = []
    for block in CLEAN_THREE_PART:
        assert isinstance(block, Para)
        label, _, rest = block.text.partition(" ")
        if block.role == "part":
            # "PART 1 GENERAL" -> label "PART 1", text "GENERAL"
            number, _, rest = rest.partition(" ")
            label = f"{label} {number}"
        out.append(
            Para(rest, block.role, auto_label=label, level=_AUTO_LEVEL_BY_ROLE[block.role])
        )
    return tuple(out)


def empty_article_blocks() -> tuple[Block, ...]:
    """Mutation: 1.02 SUBMITTALS loses its only body paragraph.

    Exactly one defect: a truly empty leaf article. Expected alert: one
    ``empty_section`` naming ``1.02 SUBMITTALS``, and nothing else.
    """
    blocks = list(CLEAN_THREE_PART)
    del blocks[_index_of(blocks, "A. Submit product data before fabrication.")]
    return tuple(blocks)


def empty_part_blocks() -> tuple[Block, ...]:
    """Mutation: PART 2's only article, 2.01 MATERIALS, loses its only body.

    Exactly one defect, and it empties a whole PART: nothing under
    ``PART 2 PRODUCTS`` is text any more. Alerts are not repeated down a tree
    (chunk S03 defined the policy: an empty heading is reported only when its
    parent is not empty), so the expected alert is one ``empty_section``
    naming ``PART 2 PRODUCTS`` — not 2.01 as well.
    """
    blocks = list(CLEAN_THREE_PART)
    del blocks[_index_of(blocks, "A. Provide materials meeting the scheduled requirements.")]
    return tuple(blocks)


def duplicate_heading_blocks() -> tuple[Block, ...]:
    """Mutation: the second article of PART 1 repeats ``1.01 SUMMARY``.

    Exactly one defect: a heading duplicated verbatim (number *and* title),
    the copy/paste slip the detector exists for. The repeated article keeps
    its own body, so the mutation cannot also read as an empty section.
    Expected alert: one ``duplicate_heading`` for ``1.01``, and nothing else.
    """
    blocks = list(CLEAN_THREE_PART)
    blocks[_index_of(blocks, "1.02 SUBMITTALS")] = Para("1.01 SUMMARY", "article")
    return tuple(blocks)


#: Every structured variant, keyed by a stable name for parametrized tests.
#: ``clean`` says whether the variant is a control (no structural defect) or
#: a single-defect mutation.
@dataclass(frozen=True)
class SpecVariant:
    name: str
    blocks: tuple[Block, ...]
    clean: bool
    #: For a mutation: the ``deterministic_rule`` it must trigger and the
    #: ``match`` text of the one alert it should produce.
    expected_rule: str | None = None
    expected_match: str | None = None


def three_part_variants() -> tuple[SpecVariant, ...]:
    return (
        SpecVariant("clean", clean_three_part_blocks(), clean=True),
        SpecVariant("table_only_article", table_only_article_blocks(), clean=True),
        SpecVariant("auto_numbered", auto_numbered_blocks(), clean=True),
        SpecVariant(
            "empty_article",
            empty_article_blocks(),
            clean=False,
            expected_rule="empty_section",
            expected_match="1.02 SUBMITTALS",
        ),
        SpecVariant(
            "empty_part",
            empty_part_blocks(),
            clean=False,
            expected_rule="empty_section",
            expected_match="PART 2 PRODUCTS",
        ),
        SpecVariant(
            "duplicate_heading",
            duplicate_heading_blocks(),
            clean=False,
            expected_rule="duplicate_heading",
            expected_match="1.01 SUMMARY",
        ),
    )


def with_section_heading(
    blocks: Sequence[Block], number: str, title: str
) -> tuple[Block, ...]:
    """Prefix a block list with a two-paragraph CSI SECTION heading.

    ``with_section_heading(clean_three_part_blocks(), "21 05 00", "COMMON
    WORK RESULTS FOR FIRE SUPPRESSION")`` opens the document the way a
    real Division 21 spec does, which is what routing reads (plan WP-05).
    """
    return (
        Para(f"SECTION {number}", "section"),
        Para(title, "section"),
        *blocks,
    )


def blocks_text(blocks: Iterable[Block]) -> str:
    """The text the extractor produces for ``blocks`` as it reads today.

    Paragraphs contribute their *literal* text (an automatic label is not
    in any run, so it is absent); tables contribute one entry per row, cells
    joined with ``" | "``; entries are joined with a blank line. This is the
    extractor's own rendering, stated independently so a test can pin it.
    """
    parts: list[str] = []
    for block in blocks:
        if isinstance(block, TableBlock):
            parts.extend(" | ".join(row) for row in block.rows)
        elif block.text:
            parts.append(block.text)
    return "\n\n".join(parts)


def displayed_text(blocks: Iterable[Block]) -> str:
    """The text a reader sees in Word, automatic labels included."""
    parts: list[str] = []
    for block in blocks:
        if isinstance(block, TableBlock):
            parts.extend(" | ".join(row) for row in block.rows)
        else:
            parts.append(block.displayed_text)
    return "\n\n".join(parts)


def changed_blocks(
    clean: Sequence[Block], mutated: Sequence[Block]
) -> list[tuple[str, Block | None, Block | None]]:
    """The differences between two block lists, as ``(kind, old, new)``.

    ``kind`` is ``"replaced"``, ``"removed"``, or ``"added"``. Aligns the two
    lists on their longest common prefix and suffix, which is exact for the
    single-block edits mutations make and deliberately reports anything
    larger as several changes.
    """
    prefix = 0
    while (
        prefix < len(clean)
        and prefix < len(mutated)
        and clean[prefix] == mutated[prefix]
    ):
        prefix += 1
    suffix = 0
    while (
        suffix < len(clean) - prefix
        and suffix < len(mutated) - prefix
        and clean[len(clean) - 1 - suffix] == mutated[len(mutated) - 1 - suffix]
    ):
        suffix += 1
    old = list(clean[prefix : len(clean) - suffix])
    new = list(mutated[prefix : len(mutated) - suffix])
    changes: list[tuple[str, Block | None, Block | None]] = []
    for index in range(max(len(old), len(new))):
        before = old[index] if index < len(old) else None
        after = new[index] if index < len(new) else None
        if before is not None and after is not None:
            changes.append(("replaced", before, after))
        elif before is not None:
            changes.append(("removed", before, None))
        else:
            changes.append(("added", None, after))
    return changes


def _index_of(blocks: Sequence[Block], text: str) -> int:
    for index, block in enumerate(blocks):
        if isinstance(block, Para) and block.text == text:
            return index
    raise ValueError(f"no paragraph {text!r} in the fixture")


# ===========================================================================
# 2. Inline WordprocessingML fragments (strings; see SpecDocBuilder)
# ===========================================================================
#
# Inline fragments go inside a ``<w:p>``; block fragments are direct
# children of ``<w:body>`` (or of a table cell / a block content control).
# Text is XML-escaped here, so callers pass plain strings.


def run(text: str, *, bold: bool = False) -> str:
    """A run with one text node. ``bold`` adds run properties so a test can
    check that formatting is carried (a run's ``rPr`` is not text)."""
    properties = "<w:rPr><w:b/></w:rPr>" if bold else ""
    return (
        f'<w:r>{properties}<w:t xml:space="preserve">{escape(text)}</w:t></w:r>'
    )


def deleted_run(text: str) -> str:
    """A run as it appears inside ``<w:del>``: its text is ``w:delText``."""
    return f'<w:r><w:delText xml:space="preserve">{escape(text)}</w:delText></w:r>'


def tab_run() -> str:
    """A run carrying only a tab character."""
    return "<w:r><w:tab/></w:r>"


def _revision_attrs(revision_id: int) -> str:
    return (
        f'w:id="{revision_id}" w:author={quoteattr(REVISION_AUTHOR)} '
        f'w:date="{REVISION_DATE}"'
    )


def inserted(*inline: str, revision_id: int) -> str:
    """A tracked insertion wrapping ``inline`` runs (kept by Accept All)."""
    return f"<w:ins {_revision_attrs(revision_id)}>{''.join(inline)}</w:ins>"


def deleted(*deleted_runs: str, revision_id: int) -> str:
    """A tracked deletion wrapping :func:`deleted_run` runs (dropped by Accept All)."""
    return f"<w:del {_revision_attrs(revision_id)}>{''.join(deleted_runs)}</w:del>"


def moved_from(*deleted_runs: str, revision_id: int) -> str:
    """A move source (dropped by Accept All)."""
    return f"<w:moveFrom {_revision_attrs(revision_id)}>{''.join(deleted_runs)}</w:moveFrom>"


def moved_to(*inline: str, revision_id: int) -> str:
    """A move destination (kept by Accept All)."""
    return f"<w:moveTo {_revision_attrs(revision_id)}>{''.join(inline)}</w:moveTo>"


def _sdt_properties(*, tag: str, control_id: int, alias: str | None, extra: str = "") -> str:
    alias_xml = f"<w:alias w:val={quoteattr(alias)}/>" if alias else ""
    return (
        f"<w:sdtPr>{alias_xml}<w:tag w:val={quoteattr(tag)}/>"
        f'<w:id w:val="{control_id}"/>{extra}</w:sdtPr>'
    )


def inline_control(
    *inline: str, tag: str, control_id: int, alias: str | None = None
) -> str:
    """A rich-text content control *inside* a paragraph, wrapping runs."""
    return (
        f"<w:sdt>{_sdt_properties(tag=tag, control_id=control_id, alias=alias)}"
        f"<w:sdtContent>{''.join(inline)}</w:sdtContent></w:sdt>"
    )


def dropdown_control(
    displayed: str,
    items: Sequence[tuple[str, str]],
    *,
    tag: str,
    control_id: int,
    alias: str | None = None,
    showing_placeholder: bool = False,
) -> str:
    """An inline drop-down content control.

    ``items`` are ``(display_text, value)`` pairs; ``displayed`` is the text
    Word stores in the control's content and shows on the page. With
    ``showing_placeholder`` the control is unresolved: Word shows its
    placeholder text (``displayed``) instead of a chosen item, and marks the
    control ``w:showingPlcHdr``. An unresolved choice is exactly what the
    placeholder detector should see once the text is readable (WP-02).
    """
    list_items = "".join(
        f"<w:listItem w:displayText={quoteattr(text)} w:value={quoteattr(value)}/>"
        for text, value in items
    )
    extra = ("<w:showingPlcHdr/>" if showing_placeholder else "") + (
        f"<w:dropDownList>{list_items}</w:dropDownList>"
    )
    return (
        f"<w:sdt>{_sdt_properties(tag=tag, control_id=control_id, alias=alias, extra=extra)}"
        f"<w:sdtContent>{run(displayed)}</w:sdtContent></w:sdt>"
    )


def simple_field(instruction: str, result: str) -> str:
    """A simple field (``w:fldSimple``) with its stored, displayed result.

    Only the stored result is visible text. The instruction (for example
    ``REF sec_210500 \\h``) is field code and must never be read as prose.
    """
    return (
        f"<w:fldSimple w:instr={quoteattr(instruction)}>{run(result)}</w:fldSimple>"
    )


def complex_field(instruction: str, result: str) -> str:
    """A complex field: begin / instruction / separate / result / end runs.

    The result run is an ordinary direct run, so it is extracted today; the
    ``w:instrText`` run is field code and is not (python-docx run text
    excludes it). The pair pins "results yes, instructions never".
    """
    return (
        '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        f'<w:r><w:instrText xml:space="preserve">{escape(instruction)}</w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        f"{run(result)}"
        '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
    )


def smart_tag(*inline: str, element: str = "place") -> str:
    """A smart-tag container (``w:smartTag``) wrapping runs."""
    return (
        '<w:smartTag w:uri="urn:schemas-microsoft-com:office:smarttags" '
        f"w:element={quoteattr(element)}>{''.join(inline)}</w:smartTag>"
    )


# ===========================================================================
# 3. Block-level fragments
# ===========================================================================


def paragraph(*inline: str, properties: str = "") -> str:
    """A ``<w:p>`` holding ``inline`` fragments.

    A string starting with ``<w:`` is a WordprocessingML fragment; anything
    else is plain text and becomes one run — so literal text that looks like
    markup (``"<VERIFY>"``, a real placeholder form) stays text.
    """
    body = "".join(_as_inline(item) for item in inline)
    return f"<w:p>{properties}{body}</w:p>"


def table(rows: Sequence[Sequence[str]]) -> str:
    """A minimal valid table; each cell holds one paragraph of its text."""
    columns = max((len(row) for row in rows), default=1)
    grid = "".join('<w:gridCol w:w="2000"/>' for _ in range(columns))
    body = "".join(
        "<w:tr>"
        + "".join(
            f'<w:tc><w:tcPr><w:tcW w:w="2000" w:type="dxa"/></w:tcPr>{paragraph(cell)}</w:tc>'
            for cell in row
        )
        + "</w:tr>"
        for row in rows
    )
    return (
        '<w:tbl><w:tblPr><w:tblW w:w="0" w:type="auto"/></w:tblPr>'
        f"<w:tblGrid>{grid}</w:tblGrid>{body}</w:tbl>"
    )


def block_control(
    *blocks: str, tag: str, control_id: int, alias: str | None = None
) -> str:
    """A block-level content control wrapping paragraphs and/or tables."""
    return (
        f"<w:sdt>{_sdt_properties(tag=tag, control_id=control_id, alias=alias)}"
        f"<w:sdtContent>{''.join(blocks)}</w:sdtContent></w:sdt>"
    )


def _as_inline(item: str) -> str:
    """Pass a ``<w:`` fragment through; wrap anything else as run text."""
    return item if item.startswith("<w:") else run(item)


# ===========================================================================
# 4. The document builder
# ===========================================================================


class SpecDocBuilder:
    """Assemble a document body element by element, in order.

    Starts from python-docx's blank template, whose body holds only
    ``w:sectPr`` — so the first element added is body child 0 and its
    extracted element id is ``p0``. Every ``add_*`` method appends before
    the section properties, exactly as Word orders a body.

    Revision and control ids are allocated by :meth:`next_id`, per builder,
    so a document's XML never depends on what other tests built first.
    """

    def __init__(self) -> None:
        self.document: DocxDocument = Document()
        self._next_id = 1
        self._numbering_id: int | None = None

    # -- ids ---------------------------------------------------------------
    def next_id(self) -> int:
        value = self._next_id
        self._next_id += 1
        return value

    # -- body elements -----------------------------------------------------
    @property
    def body(self):
        return self.document.element.body

    def add_xml(self, *fragments: str) -> "SpecDocBuilder":
        """Append body-level fragments (``<w:p>``, ``<w:tbl>``, ``<w:sdt>``)."""
        for element in parse_fragments(*fragments):
            self._append_body_element(element)
        return self

    def add_paragraph(self, *inline: str, properties: str = "") -> "SpecDocBuilder":
        """Append a paragraph built from inline fragments or plain text."""
        return self.add_xml(paragraph(*inline, properties=properties))

    def add_text(self, text: str) -> "SpecDocBuilder":
        """Append a plain one-run paragraph (python-docx's own API)."""
        self.document.add_paragraph(text)
        return self

    def add_table(self, rows: Sequence[Sequence[str]]):
        """Append a table through python-docx and return it, so a caller can
        merge cells or nest tables with the library's own operations."""
        columns = max((len(row) for row in rows), default=1)
        docx_table = self.document.add_table(rows=len(rows), cols=columns)
        for r, row in enumerate(rows):
            for c, text in enumerate(row):
                docx_table.cell(r, c).text = text
        return docx_table

    def add_blocks(self, blocks: Iterable[Block]) -> "SpecDocBuilder":
        """Append a structured block list (see :data:`CLEAN_THREE_PART`)."""
        for block in blocks:
            if isinstance(block, TableBlock):
                self.add_table(block.rows)
            elif block.auto_label is not None:
                self.add_paragraph(
                    block.text,
                    properties=self._numbering_properties(block.level or 0),
                )
            else:
                self.add_text(block.text)
        return self

    def hyperlink(self, url: str, *inline: str) -> str:
        """An inline ``<w:hyperlink>`` fragment with a real external
        relationship in this document (so the XML is valid on reopen)."""
        rel_id = self.document.part.relate_to(url, RT.HYPERLINK, is_external=True)
        body = "".join(_as_inline(item) for item in inline)
        return f'<w:hyperlink r:id="{rel_id}">{body}</w:hyperlink>'

    # -- automatic numbering ----------------------------------------------
    def _numbering_properties(self, level: int) -> str:
        num_id = self._ensure_csi_numbering()
        return (
            f'<w:pPr><w:numPr><w:ilvl w:val="{level}"/>'
            f'<w:numId w:val="{num_id}"/></w:numPr></w:pPr>'
        )

    def _ensure_csi_numbering(self) -> int:
        """Register the CSI list definition once and return its ``numId``.

        Levels: 0 = ``PART %1``; 1 = ``%1.%2`` with a zero-padded article
        number (``1.01``); 2 = ``%3.`` as an upper-case letter (``A.``). The
        suffix after each label is a space, so the displayed paragraph reads
        exactly like the manually numbered fixture. Lower levels restart
        when a higher level advances (Word's default), which is what makes
        PART 2's first article ``2.01`` and every article's first paragraph
        ``A.``. Ids are one past the template's own definitions, so the
        list cannot collide with a built-in one.
        """
        if self._numbering_id is not None:
            return self._numbering_id
        numbering = self.document.part.numbering_part.element
        abstract_ids = [
            int(el.get(qn("w:abstractNumId")))
            for el in numbering.findall(qn("w:abstractNum"))
        ]
        num_ids = [int(el.get(qn("w:numId"))) for el in numbering.findall(qn("w:num"))]
        abstract_id = max(abstract_ids, default=-1) + 1
        num_id = max(num_ids, default=0) + 1
        levels = (
            (0, "decimal", "PART %1"),
            (1, "decimalZero", "%1.%2"),
            (2, "upperLetter", "%3."),
        )
        level_xml = "".join(
            f'<w:lvl w:ilvl="{ilvl}"><w:start w:val="1"/>'
            f'<w:numFmt w:val="{fmt}"/><w:suff w:val="space"/>'
            f'<w:lvlText w:val="{text}"/><w:lvlJc w:val="left"/></w:lvl>'
            for ilvl, fmt, text in levels
        )
        abstract = parse_xml(
            f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{abstract_id}">'
            f'<w:multiLevelType w:val="multilevel"/>{level_xml}</w:abstractNum>'
        )
        # Schema order: every w:abstractNum precedes every w:num.
        first_num = numbering.find(qn("w:num"))
        if first_num is not None:
            first_num.addprevious(abstract)
        else:
            numbering.append(abstract)
        numbering.append(
            parse_xml(
                f'<w:num {nsdecls("w")} w:numId="{num_id}">'
                f'<w:abstractNumId w:val="{abstract_id}"/></w:num>'
            )
        )
        self._numbering_id = num_id
        return num_id

    # -- output ----------------------------------------------------------
    def save(self, path: Path) -> Path:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.document.save(path)
        return path

    def _append_body_element(self, element) -> None:
        sect_pr = self.body.find(qn("w:sectPr"))
        if sect_pr is not None:
            sect_pr.addprevious(element)
        else:
            self.body.append(element)


def parse_fragments(*fragments: str) -> list:
    """Parse block-level fragments into python-docx ``CT_*`` elements.

    Fragments are parsed inside a ``<w:body>`` wrapper that declares the
    ``w`` and ``r`` namespaces, then detached from it. Parsing goes through
    python-docx's ``parse_xml`` so the elements get the library's custom
    classes (``CT_P``, ``CT_Tbl``) — the extractor wraps body tables in
    ``docx.table.Table``, which needs a real ``CT_Tbl``.
    """
    wrapper = parse_xml(
        f"<w:body {nsdecls(*_FRAGMENT_NAMESPACES)}>{''.join(fragments)}</w:body>"
    )
    children = list(wrapper)
    for child in children:
        wrapper.remove(child)
    return children


def save_docx(builder_or_document, directory: Path, filename: str) -> Path:
    """Write a built document to ``directory / filename`` and return the path."""
    document = getattr(builder_or_document, "document", builder_or_document)
    path = Path(directory) / filename
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    return path


# ===========================================================================
# 5. Ready-made documents
# ===========================================================================


def build_blocks(blocks: Iterable[Block]) -> SpecDocBuilder:
    """A document holding exactly ``blocks``, in order."""
    return SpecDocBuilder().add_blocks(blocks)


def build_clean_three_part() -> SpecDocBuilder:
    return build_blocks(clean_three_part_blocks())


def build_table_only_article() -> SpecDocBuilder:
    return build_blocks(table_only_article_blocks())


def build_auto_numbered_three_part() -> SpecDocBuilder:
    return build_blocks(auto_numbered_blocks())


def build_empty_article_mutation() -> SpecDocBuilder:
    return build_blocks(empty_article_blocks())


def build_empty_part_mutation() -> SpecDocBuilder:
    return build_blocks(empty_part_blocks())


def build_duplicate_heading_mutation() -> SpecDocBuilder:
    return build_blocks(duplicate_heading_blocks())


# Sentinel texts. Each appears exactly once in its document, inside exactly
# one kind of container, so a test can ask "did the text of *this*
# container reach extraction?" with a substring check and nothing else can
# satisfy it by accident.

BEFORE_CONTROL = "Provide fire protection piping as specified."
BLOCK_CONTROL_PARAGRAPH = "Provide a listed backflow preventer at each service."
BLOCK_CONTROL_TABLE_ROWS: tuple[tuple[str, ...], ...] = (
    ("Assembly", "Listing"),
    ("Backflow preventer", "Listed double check"),
)
AFTER_CONTROL = "Install valves where accessible for maintenance."
ORDINARY_TABLE_ROWS: tuple[tuple[str, ...], ...] = (
    ("Service", "Pipe size"),
    ("Riser", "Four inch"),
)


def build_block_control_spec(*, wrapped_table: bool = True) -> SpecDocBuilder:
    """A block content control between two ordinary paragraphs, then an
    ordinary table.

    Body children, in order::

        0  <w:p>   BEFORE_CONTROL                       -> p0
        1  <w:sdt> BLOCK_CONTROL_PARAGRAPH [+ a table]  (not read today)
        2  <w:p>   AFTER_CONTROL                        -> p2
        3  <w:tbl> ORDINARY_TABLE_ROWS                  -> t0r0, t0r1

    The legacy id contract this pins: ``p2`` is the *physical* body child 2
    (the control occupies index 1), and the ordinary table is ``t0`` even
    though a wrapped table precedes it — only direct body tables count.
    """
    builder = SpecDocBuilder()
    wrapped = [paragraph(BLOCK_CONTROL_PARAGRAPH)]
    if wrapped_table:
        wrapped.append(table(BLOCK_CONTROL_TABLE_ROWS))
    builder.add_paragraph(BEFORE_CONTROL)
    builder.add_xml(
        block_control(
            *wrapped, tag="backflow_block", control_id=builder.next_id(), alias="Backflow"
        )
    )
    builder.add_paragraph(AFTER_CONTROL)
    builder.add_xml(table(ORDINARY_TABLE_ROWS))
    return builder


INLINE_CONTROL_TEXT = "schedule 40 black steel"
INLINE_CONTROL_BEFORE = "Provide "
INLINE_CONTROL_AFTER = " pipe for sprinkler mains."

DROPDOWN_ITEMS: tuple[tuple[str, str], ...] = (
    ("Copper Type L", "copper_l"),
    ("Steel Schedule 40", "steel_40"),
)
DROPDOWN_CHOSEN = "Copper Type L"
DROPDOWN_LEAD = "Pipe material: "
UNRESOLVED_DROPDOWN_PLACEHOLDER = "[SELECT HANGER FINISH]"
UNRESOLVED_DROPDOWN_LEAD = "Hanger finish: "


def build_inline_controls_spec() -> SpecDocBuilder:
    """Inline rich-text and drop-down controls between ordinary runs.

    ``p0``: ``"Provide " [control: schedule 40 black steel] " pipe for
    sprinkler mains."`` — Word shows one sentence; the extractor today reads
    ``"Provide  pipe for sprinkler mains."`` (two spaces, the control gone).

    ``p1``: ``"Pipe material: " [drop-down showing "Copper Type L"]``.

    ``p2``: ``"Hanger finish: " [unresolved drop-down showing the
    placeholder "[SELECT HANGER FINISH]"]`` — once controls are read, the
    placeholder detector must see that marker (WP-02 acceptance).
    """
    builder = SpecDocBuilder()
    builder.add_paragraph(
        INLINE_CONTROL_BEFORE,
        inline_control(
            run(INLINE_CONTROL_TEXT), tag="pipe_spec", control_id=builder.next_id()
        ),
        INLINE_CONTROL_AFTER,
    )
    builder.add_paragraph(
        DROPDOWN_LEAD,
        dropdown_control(
            DROPDOWN_CHOSEN, DROPDOWN_ITEMS, tag="pipe_material", control_id=builder.next_id()
        ),
    )
    builder.add_paragraph(
        UNRESOLVED_DROPDOWN_LEAD,
        dropdown_control(
            UNRESOLVED_DROPDOWN_PLACEHOLDER,
            (("Galvanized", "galv"), ("Painted", "paint")),
            tag="hanger_finish",
            control_id=builder.next_id(),
            showing_placeholder=True,
        ),
    )
    return builder


REF_FIELD_INSTRUCTION = " REF sec_230500 \\h "
REF_FIELD_RESULT = "23 05 00"
COMPLEX_FIELD_INSTRUCTION = " REF sec_211313 \\h "
COMPLEX_FIELD_RESULT = "21 13 13"


def build_fields_spec() -> SpecDocBuilder:
    """A simple field and a complex field, each a stored REF result.

    ``p0``: ``"Refer to Section " [fldSimple REF -> "23 05 00"] " for common
    work results."`` — the stored result is lost today (WP-02).

    ``p1``: ``"Coordinate with Section " [complex REF -> "21 13 13"] "."`` —
    the result is an ordinary run and is read today; the instruction text
    must never be (it is field code, not prose).
    """
    builder = SpecDocBuilder()
    builder.add_paragraph(
        "Refer to Section ",
        simple_field(REF_FIELD_INSTRUCTION, REF_FIELD_RESULT),
        " for common work results.",
    )
    builder.add_paragraph(
        "Coordinate with Section ",
        complex_field(COMPLEX_FIELD_INSTRUCTION, COMPLEX_FIELD_RESULT),
        ".",
    )
    return builder


SMART_TAG_TEXT = "City of Oakland"


def build_smart_tag_spec() -> SpecDocBuilder:
    """``p0``: ``"Obtain approval from the " [smartTag: City of Oakland] "
    fire marshal."`` — the tagged words are lost today (WP-02)."""
    builder = SpecDocBuilder()
    builder.add_paragraph(
        "Obtain approval from the ",
        smart_tag(run(SMART_TAG_TEXT)),
        " fire marshal.",
    )
    return builder


HYPERLINK_URL = "https://example.org/listing-guide"
HYPERLINK_TEXT = "the listing directory"
HYPERLINK_INSERTED_TEXT = "manufacturer "


def build_hyperlink_spec() -> SpecDocBuilder:
    """Hyperlinks, one plain and one carrying a tracked revision.

    ``p0``: ``"Verify each listing in " [link: the listing directory] "."``
    — read today (a direct hyperlink child).

    ``p1``: ``"Submit " [link: ins("manufacturer ") + "data sheets"] "."``
    — Word's Accept-All view reads "Submit manufacturer data sheets."; the
    insertion sits *inside* the hyperlink, a depth the revision walk does
    not reach today (WP-02 item 2 asks for that audit).
    """
    builder = SpecDocBuilder()
    builder.add_paragraph(
        "Verify each listing in ",
        builder.hyperlink(HYPERLINK_URL, run(HYPERLINK_TEXT)),
        ".",
    )
    builder.add_paragraph(
        "Submit ",
        builder.hyperlink(
            HYPERLINK_URL,
            inserted(run(HYPERLINK_INSERTED_TEXT), revision_id=builder.next_id()),
            run("data sheets"),
        ),
        ".",
    )
    return builder


def build_tracked_changes_spec() -> SpecDocBuilder:
    """Pending revisions in the body and in a table cell.

    ``p0``: ``"Comply with " del("2019") ins("2025") " CBC."`` — Accept All
    reads ``"Comply with 2025 CBC."``.

    ``p1``: ``"Pipe shall be " del("copper ") "steel."`` -> ``"Pipe shall be
    steel."``

    ``p2``: ``moveFrom("Test before concealment. ")`` ``"Flush all piping."``
    -> ``"Flush all piping."``; ``p3``: ``moveTo("Test before concealment.")``.

    ``p4``: a paragraph whose only content is deleted — empty on accept, so
    it produces no element at all.

    ``t0``: one row whose second cell reads ``del("Type M")
    ins("Type L")`` -> ``"Copper | Type L"``.
    """
    builder = SpecDocBuilder()
    builder.add_paragraph(
        "Comply with ",
        deleted(deleted_run("2019"), revision_id=builder.next_id()),
        inserted(run("2025"), revision_id=builder.next_id()),
        " CBC.",
    )
    builder.add_paragraph(
        "Pipe shall be ",
        deleted(deleted_run("copper "), revision_id=builder.next_id()),
        "steel.",
    )
    builder.add_paragraph(
        moved_from(deleted_run("Test before concealment. "), revision_id=builder.next_id()),
        "Flush all piping.",
    )
    builder.add_paragraph(
        moved_to(run("Test before concealment."), revision_id=builder.next_id()),
    )
    builder.add_paragraph(
        deleted(deleted_run("This whole paragraph was struck."), revision_id=builder.next_id()),
    )
    cell_revision = (
        '<w:tc><w:tcPr><w:tcW w:w="2000" w:type="dxa"/></w:tcPr><w:p>'
        + deleted(deleted_run("Type M"), revision_id=builder.next_id())
        + inserted(run("Type L"), revision_id=builder.next_id())
        + "</w:p></w:tc>"
    )
    builder.add_xml(
        '<w:tbl><w:tblPr><w:tblW w:w="0" w:type="auto"/></w:tblPr>'
        '<w:tblGrid><w:gridCol w:w="2000"/><w:gridCol w:w="2000"/></w:tblGrid>'
        f"<w:tr>{_cell_xml('Copper')}{cell_revision}</w:tr></w:tbl>"
    )
    return builder


def _cell_xml(text: str) -> str:
    return f'<w:tc><w:tcPr><w:tcW w:w="2000" w:type="dxa"/></w:tcPr>{paragraph(text)}</w:tc>'


def build_merged_and_nested_tables_spec() -> SpecDocBuilder:
    """Merged cells and a nested table, through python-docx's own operations.

    ``t0`` (3 x 3): row 0 is one heading merged across all three columns;
    rows 1-2 have their first column merged vertically. Extracted rows::

        t0r0  "PIPING SCHEDULE"
        t0r1  "Sprinkler | Black steel | Sch 40"
        t0r2  "Galvanized | Sch 10"

    ``t1`` (1 x 2): a layout table whose second cell holds a nested 2 x 2
    schedule — its rows come out after ``t1r0`` as ``t1r0c1t0r0`` and
    ``t1r0c1t0r1``.
    """
    builder = SpecDocBuilder()
    schedule = builder.add_table(
        [
            ("", "", ""),
            ("", "Black steel", "Sch 40"),
            ("", "Galvanized", "Sch 10"),
        ]
    )
    schedule.cell(0, 0).merge(schedule.cell(0, 2)).text = "PIPING SCHEDULE"
    schedule.cell(1, 0).merge(schedule.cell(2, 0)).text = "Sprinkler"
    layout = builder.add_table([("Layout note", "")])
    nested = layout.cell(0, 1).add_table(rows=2, cols=2)
    for r, row in enumerate((("Hanger", "Rod size"), ("Clevis", "3/8 inch"))):
        for c, text in enumerate(row):
            nested.cell(r, c).text = text
    return builder


# ===========================================================================
# 6. Representative specification file names
# ===========================================================================


@dataclass(frozen=True)
class FilenameExample:
    """A file name as a user might supply it, with what it *means*.

    ``style`` is the naming convention: ``"separated"`` (``21 05 00``),
    ``"dashed"`` (``21-05-00``), ``"compact"`` (``210500``),
    ``"section_prefixed"`` (``SECTION 21 13 16``), or ``"unrecognized"``
    when the name carries no CSI section number at all. ``section`` is the
    six-digit section number the name encodes, normalized to ``"NN NN NN"``,
    or ``None`` — including for the guard cases whose digits are a date, a
    project number, or a standard's number rather than a section.
    """

    name: str
    style: str
    section: str | None
    note: str = ""


FILENAME_EXAMPLES: tuple[FilenameExample, ...] = (
    FilenameExample("21 05 00.docx", "separated", "21 05 00"),
    FilenameExample(
        "21 05 00 - Common Work Results for Fire Suppression.docx",
        "separated",
        "21 05 00",
    ),
    FilenameExample("21-13-13.docx", "dashed", "21 13 13"),
    FilenameExample("210500.docx", "compact", "21 05 00"),
    FilenameExample("211313.docx", "compact", "21 13 13"),
    FilenameExample(
        "211313 - Wet-Pipe Sprinkler Systems.docx", "compact", "21 13 13"
    ),
    FilenameExample(
        "SECTION 21 13 16.DOCX",
        "section_prefixed",
        "21 13 16",
        "upper-case extension",
    ),
    FilenameExample("Section 211316.docx", "section_prefixed", "21 13 16"),
    FilenameExample(
        "2024-05-01 Addendum 2.docx", "unrecognized", None, "a date, not a section"
    ),
    FilenameExample(
        "Project 230415 Fire Pump.docx",
        "unrecognized",
        None,
        "an embedded project number, not a leading section number",
    ),
    FilenameExample(
        "NFPA 13 Checklist.docx", "unrecognized", None, "a standard's number"
    ),
    FilenameExample("Fire Protection Narrative.docx", "unrecognized", None),
)


def filename_examples(style: str) -> tuple[FilenameExample, ...]:
    """The examples of one naming ``style``, in declaration order."""
    found = tuple(example for example in FILENAME_EXAMPLES if example.style == style)
    if not found:
        raise ValueError(f"no filename examples of style {style!r}")
    return found


__all__ = [
    "AFTER_CONTROL",
    "BEFORE_CONTROL",
    "BLOCK_CONTROL_PARAGRAPH",
    "BLOCK_CONTROL_TABLE_ROWS",
    "Block",
    "CLEAN_THREE_PART",
    "COMPLEX_FIELD_INSTRUCTION",
    "COMPLEX_FIELD_RESULT",
    "DROPDOWN_CHOSEN",
    "DROPDOWN_ITEMS",
    "DROPDOWN_LEAD",
    "FILENAME_EXAMPLES",
    "FilenameExample",
    "HYPERLINK_INSERTED_TEXT",
    "HYPERLINK_TEXT",
    "HYPERLINK_URL",
    "INLINE_CONTROL_AFTER",
    "INLINE_CONTROL_BEFORE",
    "INLINE_CONTROL_TEXT",
    "MATERIALS_TABLE_ROWS",
    "ORDINARY_TABLE_ROWS",
    "Para",
    "REF_FIELD_INSTRUCTION",
    "REF_FIELD_RESULT",
    "REVISION_AUTHOR",
    "REVISION_DATE",
    "SMART_TAG_TEXT",
    "SpecDocBuilder",
    "SpecVariant",
    "TableBlock",
    "UNRESOLVED_DROPDOWN_LEAD",
    "UNRESOLVED_DROPDOWN_PLACEHOLDER",
    "auto_numbered_blocks",
    "block_control",
    "blocks_text",
    "build_auto_numbered_three_part",
    "build_block_control_spec",
    "build_blocks",
    "build_clean_three_part",
    "build_duplicate_heading_mutation",
    "build_empty_article_mutation",
    "build_empty_part_mutation",
    "build_fields_spec",
    "build_hyperlink_spec",
    "build_inline_controls_spec",
    "build_merged_and_nested_tables_spec",
    "build_smart_tag_spec",
    "build_table_only_article",
    "build_tracked_changes_spec",
    "changed_blocks",
    "clean_three_part_blocks",
    "complex_field",
    "deleted",
    "deleted_run",
    "displayed_text",
    "dropdown_control",
    "duplicate_heading_blocks",
    "empty_article_blocks",
    "empty_part_blocks",
    "filename_examples",
    "inline_control",
    "inserted",
    "moved_from",
    "moved_to",
    "paragraph",
    "parse_fragments",
    "run",
    "save_docx",
    "simple_field",
    "smart_tag",
    "tab_run",
    "table",
    "table_only_article_blocks",
    "three_part_variants",
    "with_section_heading",
]
