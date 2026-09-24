"""A fake Message Batches service for repair-recovery tests (plan WP-14).

Replaces every remote boundary a batch collection crosses and **counts the
paid ones**, so a test can assert that resuming a run submitted no second
repair and ran each dependent stage exactly once:

* ``repair_submits`` — review repair batches created (billed);
* ``verification_rounds`` — verification rounds that had findings to verify;
* ``cross_checks`` / ``compliance_passes`` / ``drawing_impacts`` — package
  passes run.

Polling and result retrieval are recorded too (``polls`` / ``retrieved``),
but they are not paid: re-reading a finished batch's results costs nothing,
so tests count submissions, never reads.

Primary batches return the scripted ``primary`` results; a repair batch's
items all succeed with one finding naming their file, once its status (set in
``status``) is ``"ended"``. A repair's status can be ``"processing"``
(polling detaches), ``"unreachable"`` (polling fails), ``"ended"``, or a
non-``ended`` terminal status such as ``"expired"``.
"""
from __future__ import annotations

import copy
from pathlib import Path

from docx import Document

from src.batch.batch import BatchJob, _review_custom_id
from src.batch.batch_runtime import PollOutcome
from src.input.extractor import ExtractedSpec
from src.orchestration import pipeline as pl
from src.orchestration.pipeline import BatchSubmission
from src.review.reviewer import Finding, ReviewResult

PRIMARY_ID = "msgbatch_PRIMARY"


def finding(file_name: str, issue: str = "") -> Finding:
    # The default issue names something specific to the file besides its
    # name: finding identity ignores the run's own file names (plan WP-06A),
    # so "Issue in A.docx" and "Issue in B.docx" would rightly merge.
    return Finding(
        severity="MEDIUM",
        fileName=file_name,
        section="1.01",
        issue=issue or f"Requirement {Path(file_name).stem} cites a superseded edition",
        actionType="REPORT_ONLY",
        existingText=None,
        replacementText=None,
        codeReference="",
    )


def review_ok(file_name: str) -> ReviewResult:
    return ReviewResult(
        findings=[finding(file_name)],
        parse_status="ok",
        input_tokens=1_000,
        output_tokens=400,
    )


def review_truncated() -> ReviewResult:
    return ReviewResult(
        findings=[],
        parse_status="incomplete",
        stop_reason="max_tokens",
        input_tokens=1_000,
        output_tokens=128_000,
    )


def review_refused() -> ReviewResult:
    return ReviewResult(
        findings=[],
        parse_status="refusal",
        stop_reason="refusal",
        error="Review refused by the model (stop_reason: refusal)",
    )


def request_id(index: int) -> str:
    return f"review__{index}__{index}"


def spec(name: str) -> ExtractedSpec:
    body = f"PART 1 - GENERAL. Body of {name}. Provide the specified equipment."
    return ExtractedSpec(filename=name, content=body, word_count=len(body.split()))


def submission(
    names: list[str],
    *,
    batch_id: str = PRIMARY_ID,
    module_id: str = "california_k12_mep",
    prepared: bool = True,
    cross_check: bool = True,
    project_context: str = "",
    transport: str = "batch",
) -> BatchSubmission:
    request_map = {
        request_id(i): {"filename": name, "index": i, "type": "review"}
        for i, name in enumerate(names)
    }
    return BatchSubmission(
        job=BatchJob(
            batch_id=batch_id,
            job_type="review",
            request_map=request_map,
            created_at=0.0,
        ),
        files_reviewed=list(names),
        review_request_ids=list(request_map),
        model="claude-opus-5",
        prepared_specs=[spec(n) for n in names] if prepared else None,
        cycle_label=pl.get_module(module_id).cycle.label,
        module_id=module_id,
        cross_check_enabled=cross_check,
        project_context=project_context,
        review_transport=transport,
    )


def write_docx_specs(directory: Path, names: list[str]) -> list[Path]:
    """Real .docx files, so a resumed run re-extracts its specs from disk."""
    paths = []
    for name in names:
        document = Document()
        document.add_paragraph("PART 1 - GENERAL")
        document.add_paragraph(f"1.01 SUMMARY of {name}")
        document.add_paragraph("A. Provide the specified equipment per the drawings.")
        path = directory / name
        document.save(str(path))
        paths.append(path)
    return paths


class FakeBatchService:
    """Every remote boundary of a collection, faked; the paid ones counted."""

    def __init__(self, monkeypatch, *, primary: dict[str, dict[str, ReviewResult]]):
        self.primary = primary
        self.status: dict[str, str] = {}
        self.default_repair_status = "ended"
        self.repair_submits: list[list[str]] = []
        self.polls: list[str] = []
        self.retrieved: list[str] = []
        self.verification_rounds: list[int] = []
        self.cross_checks = 0
        self.compliance_passes = 0
        self.drawing_impacts = 0
        monkeypatch.setattr(pl, "submit_review_batch", self._submit)
        monkeypatch.setattr(pl, "poll_batch_bounded", self._poll)
        monkeypatch.setattr(pl, "retrieve_review_results", self._retrieve)
        monkeypatch.setattr(pl, "verify_findings_for_run", self._verify)
        monkeypatch.setattr(pl, "run_chunked_cross_check", self._cross_check)
        monkeypatch.setattr(
            "src.compliance.run_chunked_compliance_check", self._compliance
        )
        monkeypatch.setattr(
            "src.drawing_impact.run_drawing_impact", self._drawing_impact
        )

    # -- review batches ------------------------------------------------------

    def _submit(self, repair_specs, **_kwargs) -> BatchJob:
        repair_id = f"msgbatch_REPAIR_{len(self.repair_submits) + 1}"
        self.repair_submits.append([s.filename for s in repair_specs])
        self.status.setdefault(repair_id, self.default_repair_status)
        return BatchJob(
            batch_id=repair_id,
            job_type="review",
            request_map={
                _review_custom_id(s.filename, i): {
                    "filename": s.filename,
                    "index": i,
                    "type": "review",
                }
                for i, s in enumerate(repair_specs)
            },
            created_at=0.0,
        )

    def _poll(self, batch_id, **_kwargs) -> PollOutcome:
        self.polls.append(batch_id)
        status = self.status.get(batch_id, "ended")
        if status == "processing":
            return PollOutcome(detached=True, detach_reason="max_elapsed")
        if status == "unreachable":
            return PollOutcome(poll_failed=True, poll_error="poll_error_threshold: 503")
        return PollOutcome(terminal=True, terminal_status=status)

    def _retrieve(self, job, *, model) -> dict[str, ReviewResult]:
        self.retrieved.append(job.batch_id)
        if job.batch_id in self.primary:
            return {
                custom_id: copy.deepcopy(result)
                for custom_id, result in self.primary[job.batch_id].items()
            }
        if self.status.get(job.batch_id, "ended") != "ended":
            # A repair that ended without results (expired / canceled items).
            return {
                custom_id: ReviewResult(findings=[], error="Batch request expired")
                for custom_id in job.request_map
            }
        return {
            custom_id: review_ok(meta["filename"])
            for custom_id, meta in job.request_map.items()
        }

    # -- dependent paid stages -------------------------------------------------

    def _verify(self, findings, **_kwargs) -> None:
        if findings:
            self.verification_rounds.append(len(findings))

    def _cross_check(self, specs, findings, **_kwargs) -> ReviewResult:
        self.cross_checks += 1
        return ReviewResult(findings=[], cross_check_status="completed")

    def _compliance(self, specs, profile, existing, **_kwargs) -> ReviewResult:
        self.compliance_passes += 1
        return ReviewResult(findings=[], cross_check_status="completed", coverage=[])

    def _drawing_impact(self, **_kwargs):
        from src.drawing_impact import DrawingImpactResult

        self.drawing_impacts += 1
        return DrawingImpactResult(status="completed", impact_level="minimal")

    @property
    def paid_downstream(self) -> int:
        """Dependent paid stages run so far (verification rounds included)."""
        return (
            len(self.verification_rounds)
            + self.cross_checks
            + self.compliance_passes
            + self.drawing_impacts
        )


# ---------------------------------------------------------------------------
# One outcome table for every entry point (plan WP-14 parity)
# ---------------------------------------------------------------------------

#: ``name -> (primary results for A.docx / B.docx, repair status, record kept)``.
#: The GUI, the recovery CLI and the headless decision are each driven through
#: every row and must agree on whether the saved record survives.
PARITY_NAMES = ["A.docx", "B.docx"]


def parity_primary(kind: str) -> dict[str, dict[str, ReviewResult]]:
    a, b = request_id(0), request_id(1)
    if kind == "truncated_b":
        results = {a: review_ok("A.docx"), b: review_truncated()}
    elif kind == "refused_b":
        results = {a: review_ok("A.docx"), b: review_refused()}
    elif kind == "refused_both":
        results = {a: review_refused(), b: review_refused()}
    elif kind == "clean_empty":
        results = {
            a: ReviewResult(findings=[], parse_status="ok"),
            b: ReviewResult(findings=[], parse_status="ok"),
        }
    else:  # pragma: no cover - table typo
        raise ValueError(kind)
    return {PRIMARY_ID: results}


PARITY_SCENARIOS: dict[str, tuple[str, str, bool]] = {
    "repair pending": ("truncated_b", "processing", True),
    "repair unreachable": ("truncated_b", "unreachable", True),
    "repair consumed": ("truncated_b", "ended", False),
    "repair ended without results": ("truncated_b", "expired", False),
    "partial, not repairable": ("refused_b", "ended", False),
    "all failed": ("refused_both", "ended", True),
    "valid, zero findings": ("clean_empty", "ended", False),
}
