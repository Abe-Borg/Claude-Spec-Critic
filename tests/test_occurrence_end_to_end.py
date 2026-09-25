"""Every edit location, end to end (plan WP-06B, chunk S12).

The report shows a finding once; each place its edit applies is an
*occurrence*, which reaches the edit sidecar as its own entry (schemas 6 and
7) and the applier as its own instruction. These tests drive real documents
through the whole chain — findings deduplicated as a run deduplicates them,
both report exporters, the sidecar writer, the applier, and its receipt — and
check that every executable occurrence survives each step under one
occurrence id:

* the same fix at p4 and p8 of one file is two entries and two tracked
  changes, and a duplicate emission at p4 is one;
* two files with two places each give four instructions, each with its own
  anchor;
* one finding in two modules of a program is two keys, and neither collides
  with the other;
* instructions that disagree about one place are held, not raced;
* a place with no location, or no instruction, of its own is accounted for,
  never lent another place's.

Before S12 the first case reached the sidecar as one entry (p8 was never
edited) and the second as two.
"""
from __future__ import annotations

import json
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document

from applier.models import OutcomeStatus
from applier.receipt import build_receipt, render_summary
from applier.run import RunSettings, apply_sidecar
from applier.sidecar import load_sidecar
from src.core.code_cycles import CALIFORNIA_2025
from src.input.extractor import ExtractedSpec, _accept_all_paragraph_text, extract_text_from_docx
from src.orchestration.pipeline import (
    FindingIdentityContext,
    PipelineResult,
    _deduplicate_findings,
    assign_cross_check_finding_ids,
)
from src.output.edit_sidecar import build_edit_instructions, write_edit_instructions_sidecar
from src.output.html_report_exporter import render_html_report
from src.output.report_exporter import export_report
from src.review.reviewer import Finding, ReviewResult
from src.verification.verifier import VerificationResult

_OCCURRENCE_ID = re.compile(r"\boc-[0-9a-f]{12}\b")

_SPEC_A = [
    "SECTION 21 05 00",
    "PART 2 PRODUCTS",
    "2.01 VALVES",
    "A. Provide ball valves at each riser.",
    "B. Provide a gate valve at each branch line.",  # p4
    "C. Provide check valves at each pump discharge.",
    "PART 3 EXECUTION",
    "3.01 INSTALLATION",
    "A. Install a gate valve at the connection to each riser.",  # p8
]
_SPEC_B = [
    "SECTION 21 13 13",
    "PART 2 PRODUCTS",
    "2.01 VALVES",
    "A. Provide a gate valve at each floor control assembly.",  # p3
    "B. Provide drain valves at each low point.",
    "PART 3 EXECUTION",
    "3.01 INSTALLATION",
    "A. Install drain valves where shown.",
    "B. Install a gate valve upstream of each alarm check valve.",  # p8
]


def _write_spec(directory: Path, name: str, paragraphs: list[str]) -> tuple[Path, ExtractedSpec]:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    path = directory / name
    document.save(path)
    return path, extract_text_from_docx(path)


def _edit(file_name: str, element: str | None, **overrides) -> Finding:
    fields = dict(
        severity="HIGH",
        fileName=file_name,
        section="2.01 VALVES",
        issue="Gate valves are specified where ball valves are required.",
        actionType="EDIT",
        existingText="gate valve",
        replacementText="ball valve",
        codeReference=None,
        confidence=0.9,
        evidenceElementId=element,
    )
    fields.update(overrides)
    return Finding(**fields)


def _addition(file_name: str, element: str | None, anchor: str | None, **overrides) -> Finding:
    fields = dict(
        issue="Supervisory switches are missing from control valves.",
        actionType="ADD",
        existingText=None,
        replacementText="Provide a supervisory tamper switch on each control valve.",
        anchorText=anchor,
        insertPosition="after",
    )
    fields.update(overrides)
    return _edit(file_name, element, **fields)


def _confirmed(finding: Finding, verdict: str = "CONFIRMED") -> Finding:
    """Verification runs on each finding the report shows, after dedup."""
    finding.verification = VerificationResult(
        verdict=verdict,
        grounded=True,
        sources=["https://example.gov/nfpa-13"],
        source_quote="Control valves shall be supervised.",
    )
    return finding


def _result(
    specs: list[ExtractedSpec],
    review: list[Finding],
    *,
    cross_check: list[Finding] = (),
    module_id: str = "california_k12_mep",
) -> PipelineResult:
    names = [spec.filename for spec in specs]
    merged = _deduplicate_findings(review, context=FindingIdentityContext.from_filenames(names))
    for finding in merged:
        if finding.verification is None:
            _confirmed(finding)
    return PipelineResult(
        review_result=ReviewResult(findings=merged, model="claude-opus-5"),
        cross_check_result=(
            ReviewResult(findings=list(cross_check), cross_check_status="completed")
            if cross_check
            else None
        ),
        files_reviewed=names,
        cycle_label=CALIFORNIA_2025.label,
        module_id=module_id,
        extracted_specs=list(specs),
    )


@dataclass
class _Chain:
    report_text: str
    html: str
    sidecar: dict
    receipt: dict
    results: list

    def report_ids(self) -> list[str]:
        """Occurrence ids the Word report prints, in report order."""
        return _OCCURRENCE_ID.findall(self.report_text)

    def html_ids(self) -> list[str]:
        return [
            location["occurrence_id"]
            for finding in json.loads(_payload_json(self.html))["findings"]
            for location in finding["locations"]
        ]

    def sidecar_ids(self) -> list[str]:
        return [entry["occurrence_id"] for entry in self.sidecar["edits"]]

    def outcomes(self) -> dict[str, list[str]]:
        """``{occurrence id: [outcome, ...]}`` over every receipt entry."""
        found: dict[str, list[str]] = {}
        for file_entry in self.receipt["files"]:
            for outcome in file_entry["outcomes"]:
                found.setdefault(outcome["occurrence_id"], []).append(outcome["outcome"])
        return found


def _payload_json(html: str) -> str:
    import html as _html

    match = re.search(r'<script type="application/json" id="sc-report-data">(.*?)</script>', html, re.S)
    assert match, "the report carries no data payload"
    return _html.unescape(match.group(1))


def _run_chain(tmp_path: Path, result, spec_paths: list[Path], settings: RunSettings | None = None) -> _Chain:
    settings = settings or RunSettings()
    report = tmp_path / "report.docx"
    export_report(result, report)
    report_text = "\n".join(p.text for p in Document(report).paragraphs)
    html = render_html_report(result, generated_at=datetime(2026, 9, 25, 12))
    sidecar_path = write_edit_instructions_sidecar(result, report)
    sidecar = json.loads(sidecar_path.read_text(encoding="utf-8"))
    loaded = load_sidecar(sidecar_path)
    results = apply_sidecar(loaded, spec_paths, settings)
    receipt = build_receipt(
        sidecar=loaded,
        file_results=results,
        settings=settings.to_dict(),
        entries_in=len(loaded.entries) + len(loaded.malformed),
    )
    receipt["summary"] = render_summary(receipt, results)
    return _Chain(report_text=report_text, html=html, sidecar=sidecar, receipt=receipt, results=results)


def _paragraph(path: Path, index: int):
    return Document(path).paragraphs[index]


def _tracked(paragraph) -> bool:
    xml = paragraph._p.xml
    return "<w:ins " in xml or "<w:ins>" in xml


def _accepted(path: Path) -> list[str]:
    return [_accept_all_paragraph_text(p._p) for p in Document(path).paragraphs]


def _document_after(source: Path) -> list[str]:
    """The document's text after the run: its edited copy, or, when nothing
    was written (the applier makes no copy then), the untouched source."""
    copy = source.with_name(source.stem + ".applied.docx")
    return _accepted(copy if copy.exists() else source)


def _assert_one_outcome_per_entry(chain: _Chain) -> None:
    """Receipts account for every occurrence: each sidecar entry, exactly once."""
    assert chain.receipt["accounting"]["balanced"] is True
    assert chain.receipt["accounting"]["entries_in_sidecar"] == len(chain.sidecar["edits"])
    outcomes = chain.outcomes()
    assert sorted(outcomes) == sorted(chain.sidecar_ids())
    assert all(len(statuses) == 1 for statuses in outcomes.values()), outcomes


# ---------------------------------------------------------------------------
# One file, repeated places
# ---------------------------------------------------------------------------


class TestTheSameFixAtTwoPlacesInOneFile:
    def test_p4_and_p8_are_two_entries_and_two_tracked_changes(self, tmp_path):
        path, spec = _write_spec(tmp_path, "210500.docx", _SPEC_A)
        result = _result([spec], [_edit("210500.docx", "p4"), _edit("210500.docx", "p8")])
        assert len(result.review_result.findings) == 1  # one finding in the report

        chain = _run_chain(tmp_path, result, [path])

        entries = chain.sidecar["edits"]
        assert chain.sidecar["schema_version"] == 6
        assert [(e["evidenceElementId"], e["location_basis"]) for e in entries] == [
            ("p4", "validated"),
            ("p8", "validated"),
        ]
        assert len({e["finding_id"] for e in entries}) == 1
        assert all(v == ["APPLIED"] for v in chain.outcomes().values())
        _assert_one_outcome_per_entry(chain)
        edited = tmp_path / "210500.applied.docx"
        texts = _accepted(edited)
        assert texts[4] == "B. Provide a ball valve at each branch line."
        assert texts[8] == "A. Install a ball valve at the connection to each riser."
        assert texts[3] == _SPEC_A[3]
        assert _tracked(_paragraph(edited, 4)) and _tracked(_paragraph(edited, 8))
        assert not _tracked(_paragraph(edited, 3))

    def test_both_reports_name_the_places_the_sidecar_lists(self, tmp_path):
        path, spec = _write_spec(tmp_path, "210500.docx", _SPEC_A)
        result = _result([spec], [_edit("210500.docx", "p8"), _edit("210500.docx", "p4")])
        chain = _run_chain(tmp_path, result, [path])
        assert chain.report_ids() == chain.sidecar_ids() == chain.html_ids()
        assert "Edit locations (2):" in chain.report_text
        assert "210500.docx, element p4 — confirmed in the reviewed text" in chain.report_text
        assert "210500.docx, element p8 — confirmed in the reviewed text" in chain.report_text
        assert "210500.docx, element p8 — confirmed in the reviewed text" in chain.html

    def test_a_duplicate_emission_at_p4_is_one_instruction(self, tmp_path):
        path, spec = _write_spec(tmp_path, "210500.docx", _SPEC_A)
        result = _result([spec], [_edit("210500.docx", "p4"), _edit("210500.docx", "p4")])
        chain = _run_chain(tmp_path, result, [path])
        assert [e["evidenceElementId"] for e in chain.sidecar["edits"]] == ["p4"]
        assert chain.outcomes() == {chain.sidecar_ids()[0]: ["APPLIED"]}
        assert "Edit location: 210500.docx, element p4" in chain.report_text
        texts = _accepted(tmp_path / "210500.applied.docx")
        assert texts[4] == "B. Provide a ball valve at each branch line."
        assert texts[8] == _SPEC_A[8]


class TestTwoFilesWithTwoPlacesEach:
    def test_four_instructions_each_with_its_own_anchor(self, tmp_path):
        path_a, spec_a = _write_spec(tmp_path, "210500.docx", _SPEC_A)
        path_b, spec_b = _write_spec(tmp_path, "211313.docx", _SPEC_B)
        result = _result(
            [spec_a, spec_b],
            [
                _addition("210500.docx", "p4", "B. Provide a gate valve at each branch line."),
                _addition("210500.docx", "p8", "A. Install a gate valve at the connection"),
                _addition("211313.docx", "p3", "A. Provide a gate valve at each floor control assembly."),
                _addition("211313.docx", "p8", "B. Install a gate valve upstream"),
            ],
        )
        assert len(result.review_result.findings) == 1

        chain = _run_chain(tmp_path, result, [path_a, path_b])

        entries = chain.sidecar["edits"]
        assert [(e["fileName"], e["evidenceElementId"], e["edit_proposal"]["anchor_text"]) for e in entries] == [
            ("210500.docx", "p4", "B. Provide a gate valve at each branch line."),
            ("210500.docx", "p8", "A. Install a gate valve at the connection"),
            ("211313.docx", "p3", "A. Provide a gate valve at each floor control assembly."),
            ("211313.docx", "p8", "B. Install a gate valve upstream"),
        ]
        assert len(set(chain.sidecar_ids())) == 4
        assert all(v == ["APPLIED"] for v in chain.outcomes().values())
        _assert_one_outcome_per_entry(chain)
        added = "Provide a supervisory tamper switch on each control valve."
        texts_a = _accepted(tmp_path / "210500.applied.docx")
        texts_b = _accepted(tmp_path / "211313.applied.docx")
        assert texts_a[5] == added and texts_a[10] == added  # after p4, and after p8 (shifted by one)
        assert texts_b[4] == added and texts_b[10] == added
        assert texts_a.count(added) == 2 and texts_b.count(added) == 2
        assert chain.report_ids() == chain.sidecar_ids()
        assert 'insert after “A. Install a gate valve at the connection”' in chain.report_text
        (finding,) = json.loads(_payload_json(chain.html))["findings"]
        assert [
            (loc["fileName"], loc["element_id"], loc["insert_position"], loc["anchor_text"], loc["has_instruction"])
            for loc in finding["locations"]
        ] == [
            (e["fileName"], e["evidenceElementId"], "after", e["edit_proposal"]["anchor_text"], True)
            for e in entries
        ]


# ---------------------------------------------------------------------------
# A program: one finding in two modules
# ---------------------------------------------------------------------------


def _program(children: dict[str, PipelineResult]):
    return SimpleNamespace(
        module_results=children,
        program_id="hyperscale_datacenter",
        assignments=[],
        files_reviewed=["210500.docx"],
        expected_files_reviewed=["210500.docx"],
        routed_request_count=len(children),
        expected_routed_request_count=len(children),
    )


class TestModulesDoNotCollide:
    def test_one_finding_in_two_modules_is_two_keys_written_once(self, tmp_path):
        path, spec = _write_spec(tmp_path, "210500.docx", _SPEC_A)
        children = {
            module_id: _result([spec], [_edit("210500.docx", "p4")], module_id=module_id)
            for module_id in ("datacenter_fire", "datacenter_electrical")
        }
        payload = build_edit_instructions(_program(children))
        assert payload["schema_version"] == 7
        keys = [(e["module_id"], e["occurrence_id"]) for e in payload["edits"]]
        assert [module for module, _ in keys] == ["datacenter_fire", "datacenter_electrical"]
        assert len({occurrence for _, occurrence in keys}) == 2
        assert len({e["finding_id"] for e in payload["edits"]}) == 1  # one finding id, two keys

        sidecar_path = tmp_path / "program.edits.json"
        sidecar_path.write_text(json.dumps(payload), encoding="utf-8")
        loaded = load_sidecar(sidecar_path)
        assert not loaded.malformed  # neither copy is refused as a repeated key
        (result,) = apply_sidecar(loaded, [path], RunSettings())
        statuses = {o.entry.module_id: o.status for o in result.outcomes}
        assert sorted(statuses.values(), key=lambda s: s.value) == [
            OutcomeStatus.APPLIED,
            OutcomeStatus.DUPLICATE,
        ]  # the same change at one place is written once
        assert _accepted(tmp_path / "210500.applied.docx")[4] == "B. Provide a ball valve at each branch line."

    def test_each_modules_report_section_names_its_own_keys(self, tmp_path):
        """The Word and HTML program reports print the ids the program
        sidecar keys each module's places by."""
        import sys

        sys.path.insert(0, str(Path(__file__).parent))
        from test_html_report_exporter import build_program_result

        program = build_program_result()
        sidecar = build_edit_instructions(program)
        by_module: dict[str, set[str]] = {}
        for entry in sidecar["edits"]:
            by_module.setdefault(entry["module_id"], set()).add(entry["occurrence_id"])
        assert by_module, "the program fixture no longer carries an edit"

        html = render_html_report(program, generated_at=datetime(2026, 9, 25, 12))
        payload = json.loads(_payload_json(html))
        for module_id, module_payload in payload["modules"].items():
            shown = {loc["occurrence_id"] for f in module_payload["findings"] for loc in f["locations"]}
            assert shown == by_module.get(module_id, set()), module_id

        report = tmp_path / "program.docx"
        export_report(program, report)
        printed = set(_OCCURRENCE_ID.findall("\n".join(p.text for p in Document(report).paragraphs)))
        assert printed == set().union(*by_module.values())


# ---------------------------------------------------------------------------
# Instructions that disagree, and places with nothing of their own
# ---------------------------------------------------------------------------


class TestConflictsAndGaps:
    def test_two_findings_that_disagree_about_one_place_are_both_held(self, tmp_path):
        path, spec = _write_spec(tmp_path, "210500.docx", _SPEC_A)
        result = _result(
            [spec],
            [
                _edit("210500.docx", "p4"),
                _edit(
                    "210500.docx",
                    "p4",
                    issue="Gate valves at branch lines must be butterfly valves.",
                    existingText="gate valve at each branch line",
                    replacementText="butterfly valve at each branch line",
                ),
            ],
        )
        assert len(result.review_result.findings) == 2
        chain = _run_chain(tmp_path, result, [path])
        assert all(v == ["EDIT_CONFLICT"] for v in chain.outcomes().values())
        _assert_one_outcome_per_entry(chain)
        assert _document_after(path)[4] == _SPEC_A[4]
        assert _accepted(path) == _SPEC_A  # the source is never written

    def test_a_file_with_no_original_is_found_by_text_or_refused_for_want_of_a_place(self, tmp_path):
        path_a, spec_a = _write_spec(tmp_path, "210500.docx", _SPEC_A)
        path_b, spec_b = _write_spec(tmp_path, "211313.docx", _SPEC_B)
        legacy_edit = _edit("210500.docx", "p4")
        legacy_edit.affected_files = ["210500.docx", "211313.docx"]  # no originals recorded
        legacy_edit.finding_id = "rf-0000000000aa"
        legacy_add = _addition("210500.docx", "p4", "B. Provide a gate valve at each branch line.")
        legacy_add.affected_files = ["210500.docx", "211313.docx"]
        legacy_add.finding_id = "rf-0000000000bb"
        result = _result([spec_a, spec_b], [])
        result.review_result = ReviewResult(findings=[_confirmed(legacy_edit), _confirmed(legacy_add)], model="m")

        chain = _run_chain(tmp_path, result, [path_a, path_b])

        missing = [e for e in chain.sidecar["edits"] if e["location_basis"] == "missing_original"]
        assert {(e["finding_id"], e["fileName"]) for e in missing} == {
            ("rf-0000000000aa", "211313.docx"),
            ("rf-0000000000bb", "211313.docx"),
        }
        for entry in missing:  # nothing lent from 210500.docx
            assert entry["evidenceElementId"] is None
            assert entry["edit_proposal"]["target_element_id"] is None
            assert entry["edit_proposal"]["anchor_text"] is None
        outcomes = {
            (o["finding_id"], o["file_name"]): o for f in chain.receipt["files"] for o in f["outcomes"]
        }
        # The EDIT is found by its own text: it occurs at p3 and p8 there, and
        # the finding's section ("2.01 VALVES") singles out p3.
        located = outcomes[("rf-0000000000aa", "211313.docx")]
        assert (located["outcome"], located["location_status"], located["element_id"]) == (
            "APPLIED",
            "RESOLVED_BY_SECTION",
            "p3",
        )
        # The ADD has no anchor of its own there, and none is lent.
        held = outcomes[("rf-0000000000bb", "211313.docx")]
        assert held["outcome"] == "MALFORMED" and "recorded no original here" in held["reason"]
        _assert_one_outcome_per_entry(chain)
        assert "211313.docx — no location recorded for this file" in chain.report_text
        assert "no anchor recorded here, so the addition cannot be placed" in chain.report_text

    def test_a_place_whose_finding_proposed_no_edit_is_refused_with_that_reason(self, tmp_path):
        path_a, spec_a = _write_spec(tmp_path, "210500.docx", _SPEC_A)
        path_b, spec_b = _write_spec(tmp_path, "211313.docx", _SPEC_B)
        result = _result([spec_a, spec_b], [_edit("210500.docx", "p4"), _edit("211313.docx", "p3")])
        (group,) = result.review_result.findings
        demoted = next(o for o in group.occurrence_originals if o.fileName == "211313.docx")
        demoted.actionType = "REPORT_ONLY"

        chain = _run_chain(tmp_path, result, [path_a, path_b])

        by_file = {e["fileName"]: e for e in chain.sidecar["edits"]}
        assert by_file["211313.docx"]["edit_proposal"] is None
        (outcome,) = [o for f in chain.receipt["files"] for o in f["outcomes"] if o["file_name"] == "211313.docx"]
        assert outcome["outcome"] == "MALFORMED"
        assert "no edit instruction for this place" in outcome["reason"]
        assert "no edit instruction for this place" in chain.report_text
        (finding,) = json.loads(_payload_json(chain.html))["findings"]
        assert {loc["fileName"]: loc["has_instruction"] for loc in finding["locations"]} == {
            "210500.docx": True,
            "211313.docx": False,
        }
        assert all(loc["insert_position"] is None and loc["anchor_text"] is None for loc in finding["locations"])
        assert _accepted(tmp_path / "210500.applied.docx")[4] == "B. Provide a ball valve at each branch line."
        _assert_one_outcome_per_entry(chain)

    def test_content_twins_are_one_instruction_under_the_more_cautious_verdict(self, tmp_path):
        path, spec = _write_spec(tmp_path, "210500.docx", _SPEC_A)
        first = _confirmed(_edit("210500.docx", "p4", issue="Coordination: the valve schedules disagree."))
        second = _confirmed(
            _edit("210500.docx", "p4", issue="Coordination: the valve schedules disagree."), "DISPUTED"
        )
        assign_cross_check_finding_ids([first, second])
        result = _result([spec], [], cross_check=[first, second])

        chain = _run_chain(tmp_path, result, [path])

        (entry,) = chain.sidecar["edits"]
        assert entry["report_status"] == "DISPUTED"
        assert chain.outcomes() == {entry["occurrence_id"]: ["HELD_BY_POLICY"]}
        # Both findings stay in the report, each naming the one place.
        assert chain.report_ids() == [entry["occurrence_id"]] * 2
        assert _document_after(path)[4] == _SPEC_A[4]


# ---------------------------------------------------------------------------
# Every executable occurrence survives the whole chain
# ---------------------------------------------------------------------------


class TestEveryExecutableOccurrenceSurvives:
    def test_report_sidecar_applier_and_receipt_agree_place_by_place(self, tmp_path):
        path_a, spec_a = _write_spec(tmp_path, "210500.docx", _SPEC_A)
        path_b, spec_b = _write_spec(tmp_path, "211313.docx", _SPEC_B)
        result = _result(
            [spec_a, spec_b],
            [
                _edit("210500.docx", "p4"),
                _edit("210500.docx", "p8"),
                _edit("211313.docx", "p3"),
                _edit("211313.docx", "p8"),
                _addition("210500.docx", "p3", "A. Provide ball valves at each riser."),
                _addition("211313.docx", "p4", "B. Provide drain valves at each low point."),
            ],
        )
        chain = _run_chain(tmp_path, result, [path_a, path_b])

        assert chain.report_ids() == chain.sidecar_ids() == chain.html_ids()
        assert len(chain.sidecar_ids()) == 6
        assert all(v == ["APPLIED"] for v in chain.outcomes().values())
        _assert_one_outcome_per_entry(chain)
        for occurrence_id in chain.sidecar_ids():
            assert occurrence_id in chain.receipt["summary"]
        texts_a = _accepted(tmp_path / "210500.applied.docx")
        texts_b = _accepted(tmp_path / "211313.applied.docx")
        assert sum("ball valve at" in t for t in texts_a) == 2
        assert sum("ball valve " in t for t in texts_b) == 2
        assert sum("gate valve" in t for t in texts_a + texts_b) == 0
        added = "Provide a supervisory tamper switch on each control valve."
        assert texts_a.count(added) == 1 and texts_b.count(added) == 1


@pytest.fixture(autouse=True)
def _hermetic(monkeypatch):
    """The chain writes reports; nothing here may reach a cache or the network."""
    monkeypatch.setenv("SPEC_CRITIC_VERIFICATION_CACHE_PERSIST", "0")
    yield


class TestBothReportsWordEveryPlaceAlike:
    """Both exporters render locations from one wording
    (``report_exporter._edit_location_text``); the Word report's lines and
    the HTML report's text lines are the same, place by place."""

    def test_word_and_html_location_lines_match(self, tmp_path):
        import html as _html

        path_a, spec_a = _write_spec(tmp_path, "210500.docx", _SPEC_A)
        path_b, spec_b = _write_spec(tmp_path, "211313.docx", _SPEC_B)
        result = _result(
            [spec_a, spec_b],
            [
                _edit("210500.docx", "p4"),
                _edit("210500.docx", "p8"),
                _edit("211313.docx", "p5"),  # names an element that lacks the text
                _addition("210500.docx", "p3", "A. Provide ball valves at each riser."),
            ],
        )
        chain = _run_chain(tmp_path, result, [path_a, path_b])

        word = [line.strip() for line in chain.report_text.splitlines() if _OCCURRENCE_ID.search(line)]
        text = re.search(r'<pre id="sc-plaintext" hidden>(.*?)</pre>', chain.html, re.S).group(1)
        html_lines = [
            re.sub(r"^- ", "", line.strip())
            for line in _html.unescape(text).splitlines()
            if _OCCURRENCE_ID.search(line)
        ]
        assert word == html_lines
        assert len(word) == len(chain.sidecar_ids()) == 4
        assert any(
            "211313.docx — place not identified (the review named p5, which does not contain "
            "the text this edit targets)" in line
            for line in word
        )


class TestTheOutputLayerNeverLoadsThePipeline:
    """The occurrence model moved to ``orchestration/occurrences.py`` (stdlib
    only) so the sidecar writer and both exporters could use it without
    importing the pipeline, which the HTML exporter documents it never does.
    Checked in a fresh interpreter: this test session has long since
    imported the pipeline."""

    def test_importing_the_writer_and_both_exporters_loads_no_pipeline(self):
        import subprocess
        import sys

        probe = (
            "import sys\n"
            "import src.output.edit_sidecar, src.output.report_exporter, src.output.html_report_exporter\n"
            "loaded = sorted(m for m in sys.modules if m in "
            "('src.orchestration.pipeline', 'src.orchestration.program_pipeline'))\n"
            "print(loaded)\n"
        )
        repo = Path(__file__).resolve().parent.parent
        completed = subprocess.run(
            [sys.executable, "-c", probe], cwd=repo, capture_output=True, text=True, timeout=120
        )
        assert completed.returncode == 0, completed.stderr
        assert completed.stdout.strip() == "[]"
