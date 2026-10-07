"""Pin shipped trust facts to code; the saved dossier is the reviewed snapshot."""
from __future__ import annotations

import ast
from dataclasses import fields
from pathlib import Path
import re
from string import Formatter
from xml.etree import ElementTree

import pytest

from src.gui import trust_content as trust

ROOT = Path(__file__).resolve().parents[1]
LEDGER = (ROOT / "docs/TRUST_CLAIMS.md").read_text(encoding="utf-8")


def strings(value):
    if isinstance(value, str):
        yield value
    elif isinstance(value, (tuple, list)):
        for item in value:
            yield from strings(item)


def templates():
    yield trust.LEAD
    yield trust.CLOSING
    for block in trust.SHORT_POINTS:
        yield block.title
        yield from strings(block.body)
    for topic in trust.TOPICS:
        yield topic.title
        for block in topic.blocks:
            yield block.title
            yield from strings(block.body)
    for action in trust.ACTIONS:
        for field in fields(action):
            if field.name not in ("identity", "claims"):
                yield getattr(action, field.name)


def test_trust_inventory_complete():
    inventoried = re.findall(r"^\| ([AB]\d+) \|", LEDGER, re.M)
    assert [a.identity for a in trust.ACTIONS] == inventoried
    assert len(set(inventoried)) == len(inventoried)
    assert sum(i.startswith("A") for i in inventoried) == 40
    assert sum(i.startswith("B") for i in inventoried) == 13
    assert len(re.findall(r"^\| N\d+ \|", LEDGER, re.M)) == 8
    assert len(trust.TOPICS) == 12
    assert [t.anchor for t in trust.TOPICS] == [
        "answer", "origins", "engine", "boundary", "runtime", "tools",
        "terms", "local", "security", "money", "limits", "audit",
    ]


def test_every_runtime_card_has_the_same_five_rows():
    facts = trust.fact_values()
    for action in trust.ACTIONS:
        rows = action.rows(facts)
        assert tuple(label for label, _ in rows) == trust.RUNTIME_LABELS
        assert all(value.strip() for _, value in rows)
        assert action.claims
        if action.ai == "None.":
            assert rows[3] == ("AI involved", "None.")
    assert sum(a.ai == "None." for a in trust.ACTIONS) >= 25


def test_all_blocks_reference_registered_source_claims():
    registered = set(re.findall(r"^\| (C\d+) \|", LEDGER, re.M))
    used = set()
    blocks = list(trust.SHORT_POINTS) + [b for t in trust.TOPICS for b in t.blocks] + list(trust.ACTIONS)
    for block in blocks:
        assert block.claims
        assert set(block.claims) <= registered
        used.update(block.claims)
    assert used == registered
    # Files named as sources must exist. Imported dependency source is pinned
    # separately below; module glob references intentionally cover the registry.
    for path in re.findall(r"`((?:src|applier|scripts|packaging)/[^` :]+\.py)(?:[:`])", LEDGER):
        assert list(ROOT.glob(path)), path
    # Pin ledger references to real definitions/imports, without importing the
    # GUI or developer commands merely to inspect their names.
    for path, references in re.findall(r"`([^`]+\.py)` — ([^;|\n]+)", LEDGER):
        if "*" in path:
            continue
        tree = ast.parse((ROOT / path).read_text(encoding="utf-8"))
        names = {n.name for n in ast.walk(tree) if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef))}
        names.update(n.id for n in ast.walk(tree) if isinstance(n, ast.Name))
        names.update(n.attr for n in ast.walk(tree) if isinstance(n, ast.Attribute))
        names.update(n.asname or n.name for n in ast.walk(tree) if isinstance(n, ast.alias))
        for symbol in re.findall(r"`(\w+)`", references):
            assert symbol in names, f"{path}::{symbol}"


def test_quoted_numbers_models_and_hosts_use_ledger_bindings():
    facts = trust.fact_values()
    bindings = set()
    for template in templates():
        assert not re.search(r"\d", template), f"Literal numeric claim: {template}"
        assert "https://" not in template and "http://" not in template
        assert not re.search(r"\b\w+\.(com|net|org)\b", template)
        for _, name, _, _ in Formatter().parse(template):
            if name:
                assert name in facts
                bindings.add(name)
                assert re.search(r"\b" + re.escape(name) + r"\b", LEDGER), name
        template.format_map(facts)
    assert "engine_rows" in LEDGER and "price_rows" in LEDGER
    assert bindings


def test_trust_models_and_settings():
    from src.core import api_config as cfg
    from src.review.review_request_builder import ReviewRequestSpec, build_review_request
    from src.verification.verification_modes import mode_policy
    from src.output import html_report_exporter as html

    facts = trust.fact_values()
    for key, model, phase in (
        ("review_ai", cfg.REVIEW_MODEL_DEFAULT, cfg.PHASE_REVIEW),
        ("research_ai", cfg.RESEARCH_MODEL_DEFAULT, cfg.PHASE_RESEARCH),
        ("cross_ai", cfg.CROSS_CHECK_MODEL_DEFAULT, cfg.PHASE_CROSS_CHECK),
        ("compliance_ai", cfg.COMPLIANCE_MODEL_DEFAULT, cfg.PHASE_COMPLIANCE),
        ("impact_ai", cfg.DRAWING_IMPACT_MODEL_DEFAULT, cfg.PHASE_DRAWING_IMPACT),
        ("coordination_ai", cfg.COORDINATION_MODEL_DEFAULT, cfg.PHASE_COORDINATION),
    ):
        assert facts[key].startswith(model + ";")
        assert f"{cfg.phase_output_cap(phase, model=model):,} tokens" in facts[key]
        assert cfg.effort_config_for(model=model, phase=phase)["effort"] in facts[key]
    request = build_review_request(ReviewRequestSpec("A small spec", "spec.docx", cfg.REVIEW_MODEL_DEFAULT, force_allow_extended_output=False)).params
    assert request["model"] in facts["review_ai"]
    assert request["output_config"]["effort"] in facts["review_ai"]
    assert f"{request['max_tokens']:,}" in facts["review_ai"]
    assert "temperature" not in request
    assert mode_policy("strict_structured").effort == "low"
    assert mode_policy("strict_structured").model in facts["verifier_ai"]
    assert html.CHAT_DEFAULT_MODEL in facts["chat_ai"]
    assert cfg.MODEL_HAIKU_55 in facts["chat_ai"]
    assert str(html.CHAT_MAX_TOKENS).replace("000", ",000") in facts["chat_ai"]
    assert facts["triage_ai"].startswith(cfg.MODEL_HAIKU_55 + ";")
    assert "effort medium; adaptive thinking; output cap 16,000 tokens" in facts["triage_ai"]
    assert len(set(facts["models"].split(", "))) == 3
    assert len(cfg._MODEL_CAPABILITIES) == 8


def test_trust_fact_sources(monkeypatch):
    from anthropic import Anthropic
    from anthropic._constants import DEFAULT_MAX_RETRIES, DEFAULT_TIMEOUT
    from src.core import api_config as cfg, tokenizer, updates, pricing
    from src.verification import retry_policy as retry, verification_cache as cache
    from src.tracing import config as trace
    from src.batch import batch_runtime as poll
    from src.input import drawing_analysis
    from src.research.requirements_research import RESEARCH_MAX_CONTINUATIONS
    from applier import assist
    from src.output import html_report_exporter as html
    from src.programs.routing import _RESTRICTED_UNIMPLEMENTED_DIVISIONS

    facts = trust.fact_values()
    from platformdirs import user_state_dir
    monkeypatch.delenv(trace.ENV_TRACE_DIR, raising=False)
    assert trace.default_trace_root() == Path(user_state_dir("SpecCritic", appauthor=False)) / "traces"
    assert trust.TRACE_LOCATIONS_PIN == "%LOCALAPPDATA%/SpecCritic/traces (Windows), ~/Library/Application Support/SpecCritic/traces (macOS), or $XDG_STATE_HOME/SpecCritic/traces (Linux; usually ~/.local/state/SpecCritic/traces)"
    expected = {
        "context_cap": tokenizer.PROJECT_CONTEXT_MAX_TOKENS,
        "review_extended": cfg.REVIEW_OUTPUT_CAP_BATCH_EXTENDED,
        "review_threshold": cfg.LARGE_REVIEW_INPUT_THRESHOLD,
        "retry_attempts": retry.DEFAULT_REALTIME_RETRY_POLICY.max_attempts,
        "retry_wait": retry.DEFAULT_REALTIME_RETRY_POLICY.max_retry_wait_seconds,
        "continuations": retry.DEFAULT_MAX_CONTINUATIONS,
        "deep_continuations": retry.DEEP_MAX_CONTINUATIONS,
        "cache_days": cache._DEFAULT_CACHE_TTL_DAYS,
        "cache_entries": cache._DEFAULT_CACHE_MAX_ENTRIES,
        "flight_wait": cache._DEFAULT_SINGLEFLIGHT_WAIT_SECONDS,
        "trace_days": trace.DEFAULT_TRACE_RETENTION_DAYS,
        "trace_runs": trace.DEFAULT_TRACE_MAX_RUNS,
        "analysis_mib": drawing_analysis.MAX_DRAWING_ANALYSIS_BYTES // (1024 * 1024),
        "sdk_retries": DEFAULT_MAX_RETRIES,
        "sdk_timeout": DEFAULT_TIMEOUT.read,
        "sdk_connect": DEFAULT_TIMEOUT.connect,
        "research_continuations": RESEARCH_MAX_CONTINUATIONS,
        "poll_errors": poll.DEFAULT_REVIEW_POLL_POLICY.max_consecutive_errors,
        "update_days": updates.DEFAULT_MIN_INTERVAL_DAYS,
        "manifest_timeout": updates.DEFAULT_MANIFEST_TIMEOUT,
        "download_timeout": updates.DEFAULT_DOWNLOAD_TIMEOUT,
        "manifest_kib": updates.MAX_MANIFEST_BYTES // 1024,
        "chat_tokens": html.CHAT_MAX_TOKENS,
        "assist_rounds": assist.MAX_TOOL_ROUNDS,
        "assist_tokens": assist.ASSIST_MAX_TOKENS,
    }
    for name, value in expected.items():
        assert facts[name] == (f"{value:,}" if isinstance(value, int) else f"{value:g}")
    assert trust.ASSIST_ROUNDS_PIN == assist.MAX_TOOL_ROUNDS
    assert trust.ASSIST_TOKENS_PIN == assist.ASSIST_MAX_TOKENS
    assert assist.AssistConfig().model in facts["assist_ai"]
    assert _RESTRICTED_UNIMPLEMENTED_DIVISIONS == frozenset({"27", "28"})
    assert all(value in facts["unsupported_divisions"] for value in _RESTRICTED_UNIMPLEMENTED_DIVISIONS)
    client = Anthropic(api_key="test-key", base_url="https://api.anthropic.com")
    assert client.base_url.host == facts["api_host"]
    client.close()
    # Check the dependency's actual default rather than just the report URL.
    import anthropic._client as sdk
    assert 'base_url = f"https://' + facts["api_host"] + '"' in Path(sdk.__file__).read_text()
    assert facts["update_url"] == updates._DEFAULT_MANIFEST_URL
    assert facts["tokenizer_url"] == tokenizer.CL100K_BASE_BLOB_URL
    from urllib.parse import urlsplit
    from src.gui import about_usage_dialogs as help_topics
    assert set(trust.ABOUT_REFERENCE_HOSTS_PIN) == {urlsplit(url).hostname for url in (help_topics._LICENSE_URL, help_topics._LINKEDIN_URL, help_topics._GITHUB_PROFILE_URL)}
    assert set(facts["browser_hosts"].split(", ")) == set(trust.ABOUT_REFERENCE_HOSTS_PIN) | {urlsplit(url).hostname for _, url in trust.FURTHER_READING}
    assert facts["update_hash"] == "SHA-256"
    import inspect
    assert "hashlib.sha256()" in inspect.getsource(updates.verify_sha256)
    assert "hashlib.sha256()" in inspect.getsource(updates.download_installer)
    assert facts["batch_discount"] == f"{(1-pricing.BATCH_DISCOUNT)*100:g}"
    for model, rates, read in trust.price_rows(facts):
        tier = re.fullmatch(r"(.+) \(prompt ([≤>]) ([\d,]+) tokens\)", model)
        if tier:
            model, comparison, threshold = tier.groups()
            prompt_tokens = int(threshold.replace(",", "")) + (comparison == ">")
        else:
            prompt_tokens = 0
        price = pricing.price_for(model, prompt_tokens=prompt_tokens)
        assert rates == f"${price.input_per_mtok:g} / ${price.output_per_mtok:g}"
        assert read == f"${price.cache_read_rate_per_mtok:g}"
    # JS-literal facts are read from the same declarations the browser executes.
    assert facts["chat_tools"] == re.search(r"MAX_TOOL_ROUNDS = (\d+)", html._CHAT_JS).group(1)
    assert facts["chat_continuations"] == re.search(r"MAX_CONTINUATIONS = (\d+)", html._CHAT_JS).group(1)
    assert facts["chat_history"] == re.search(r"MAX_HISTORY_MESSAGES = (\d+)", html._CHAT_JS).group(1)


def test_reviewed_dossier_pins_all_rendered_facts(monkeypatch):
    # Clear path/experiment switches; this snapshot describes fresh-install
    # defaults. The GUI still displays the running process's configured paths.
    import os
    for name in list(os.environ):
        if name.startswith("SPEC_CRITIC_"):
            monkeypatch.delenv(name)
    assert (ROOT / "docs/TRUST.md").read_text(encoding="utf-8") == trust.markdown_dossier()


def test_changed_limit_requires_reviewed_copy_update(monkeypatch):
    from src.core import tokenizer
    before = trust.markdown_dossier()
    monkeypatch.setattr(tokenizer, "PROJECT_CONTEXT_MAX_TOKENS", tokenizer.PROJECT_CONTEXT_MAX_TOKENS + 1)
    assert trust.markdown_dossier() != before


def test_trust_prices_disclose_both_haiku_tiers_and_cache_reads():
    rows = trust.price_rows(trust.fact_values())
    assert ("claude-haiku-5-5 (prompt ≤ 100,000 tokens)", "$0.1 / $0.5", "$0.01") in rows
    assert ("claude-haiku-5-5 (prompt > 100,000 tokens)", "$0.5 / $2.5", "$0.05") in rows
    dossier = trust.markdown_dossier()
    assert "total prompt size selects the rate for the whole request" in dossier
    assert "including output and cache tokens" in dossier


def test_report_quote_label_does_not_claim_the_gate_compared_words():
    from types import SimpleNamespace
    from docx import Document
    from src.output.html_report_exporter import _render_evidence_panel
    from src.output.report_exporter import _write_evidence_panel
    from src.verification.verifier import VerificationResult
    result = VerificationResult(verdict="CONFIRMED", source_quote="A claim the model supplied")
    finding = SimpleNamespace(code_reference="", issue="", severity="HIGH")
    html, lines = _render_evidence_panel(finding, result)
    doc = Document()
    _write_evidence_panel(doc, finding, result)
    rendered = html + "\n".join(lines) + "\n".join(p.text for p in doc.paragraphs)
    assert rendered.count("Source quote (supplied by verifier):") == 3
    assert "verbatim from search result" not in rendered


def test_export_can_overwrite_a_source_path_as_disclosed(tmp_path):
    from docx import Document
    from src.gui.report_controller import _write_report_and_sidecars
    from src.orchestration.pipeline import PipelineResult
    from src.review.reviewer import ReviewResult
    path = tmp_path / "spec.docx"
    original = Document()
    original.add_paragraph("Original specification sentinel")
    original.save(path)
    result = PipelineResult(review_result=ReviewResult(), files_reviewed=[path.name])
    _write_report_and_sidecars(result, path)
    assert "Original specification sentinel" not in "\n".join(p.text for p in Document(path).paragraphs)
    assert path.with_suffix(".edits.json").is_file()


def test_trust_no_external_assets():
    svg = ElementTree.fromstring(trust.FLOW_SVG)
    assert svg.attrib["role"] == "img"
    assert svg.attrib["aria-label"] == trust.FLOW_DESCRIPTION
    assert "prefers-color-scheme:dark" in trust.FLOW_SVG
    assert "stroke-dasharray" in trust.FLOW_SVG
    assert not any(e.tag.rsplit("}", 1)[-1] in {"script", "image", "use", "foreignObject"} for e in svg.iter())
    assert not re.search(r"(?:href|src)=|url\(", trust.FLOW_SVG)
    tree = ast.parse((ROOT / "src/gui/trust_dialogs.py").read_text())
    imported = {n.module for n in ast.walk(tree) if isinstance(n, ast.ImportFrom)} | {
        alias.name for n in ast.walk(tree) if isinstance(n, ast.Import) for alias in n.names
    }
    assert imported <= {"__future__", "tkinter", "webbrowser", "xml.etree", "customtkinter", "widgets", None}
    assert not any(isinstance(n, ast.Call) and isinstance(n.func, ast.Attribute) and n.func.attr in {"urlopen", "fetch", "get", "load_font", "PhotoImage"} for n in ast.walk(tree))
    external_urls = re.findall(r'https://[^"\s)]+', (ROOT / "src/gui/trust_dialogs.py").read_text())
    assert external_urls == []
    assert all(url.startswith("https://") for _, url in trust.FURTHER_READING)


def test_short_topic_is_short_and_has_mechanisms():
    assert 6 <= len(trust.SHORT_POINTS) <= 8
    assert all(5 <= len(point.title.split()) <= 8 for point in trust.SHORT_POINTS)
    words = len((trust.LEAD + " ".join(p.title + " " + str(p.body) for p in trust.SHORT_POINTS) + trust.CLOSING).split())
    assert words < 450
    assert trust.DETAIL_BUTTON == "I'm not convinced — show me exactly what runs →"
    assert trust.TOPICS[-1].blocks[-1].body.startswith("Good.")
