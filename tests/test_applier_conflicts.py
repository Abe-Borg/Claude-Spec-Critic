"""Instructions that disagree about one place are held, not raced (plan WP-06B).

Before S11 the applier resolved every *element* before writing, but found each
edit's *text* again after the edits before it had changed the document. So
the sidecar's order decided outcomes: of two edits to overlapping text, the
first was written and the second failed to find its text; an identical edit
listed twice read as a failure; two additions at one anchor came out in list
order; and even two compatible edits could block each other, when one's
replacement repeated the other's target, or changed an addition's anchor.

Now every edit is planned against the unmutated document — the paragraph and
the characters it changes, or the gap a new paragraph goes into — and the
plans are settled before anything is written (``applier.conflicts``):

* different instructions for one place are an ``EDIT_CONFLICT``, every one;
* identical ones are written once, the rest ``DUPLICATE``;
* everything else is applied from the planned positions.

The property all of it serves is asserted directly: permuting the sidecar's
entries changes nothing — not the document, not any entry's outcome.
"""
from __future__ import annotations

import itertools
import json
import random
from types import SimpleNamespace

import pytest
from docx import Document
from docx.oxml.ns import qn

from applier.cli import EXIT_OK, EXIT_UNAPPLIED, main
from applier.conflicts import Settlement, application_order, settle
from applier.docx_edit import DIRECT, DocumentEditor, EditError, PlannedEdit
from applier.models import EditEntry, ElementKind, Location, LocationStatus, OutcomeStatus
from applier.receipt import build_receipt, render_summary
from applier.run import RunSettings, apply_sidecar
from applier.sidecar import load_sidecar
from src.input.extractor import _accept_all_paragraph_text

_PARAS = [
    "SECTION 21 05 00",
    "Provide gate valves at each branch line.",
    "Provide gate valves and check valves.",
    "Install hangers per the manufacturer.",
]


def edit(finding_id, existing, replacement, *, element="p1", **overrides):
    proposal = {
        "action_type": "EDIT",
        "existing_text": existing,
        "replacement_text": replacement,
        "anchor_text": None,
        "insert_position": None,
        "target_element_id": element,
        "edit_confidence": 0.9,
    }
    proposal.update(overrides.pop("proposal", {}))
    base = {
        "finding_id": finding_id,
        "fileName": "spec.docx",
        "affected_files": ["spec.docx"],
        "has_per_file_original": True,
        "section": "21 05 00",
        "severity": "HIGH",
        "issue": f"Issue {finding_id}.",
        "codeReference": None,
        "evidenceElementId": element,
        "verification_verdict": "CONFIRMED",
        "report_status": "VERIFIED_SUPPORTED",
        "edit_proposal": proposal,
    }
    base.update(overrides)
    return base


def delete(finding_id, existing, *, element="p1", **overrides):
    proposal = {"action_type": "DELETE", **overrides.pop("proposal", {})}
    return edit(finding_id, existing, None, element=element, proposal=proposal, **overrides)


def add(finding_id, text, *, anchor, position="after", element="p1", **overrides):
    return edit(
        finding_id,
        None,
        text,
        element=element,
        proposal={"action_type": "ADD", "anchor_text": anchor, "insert_position": position},
        **overrides,
    )


def build_spec(path, paragraphs=_PARAS):
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    document.save(path)
    return path


def write_sidecar(tmp_path, edits, *, version=4):
    path = tmp_path / "report.edits.json"
    path.write_text(json.dumps({"schema_version": version, "edits": edits}), encoding="utf-8")
    return path


def _inserted(p_el) -> bool:
    properties = p_el.find(qn("w:pPr"))
    run_properties = properties.find(qn("w:rPr")) if properties is not None else None
    return run_properties is not None and run_properties.find(qn("w:ins")) is not None


def _reject_text(p_el) -> str:
    parts = []
    for node in p_el.iter():
        if node.tag == qn("w:delText"):
            parts.append(node.text or "")
        elif node.tag == qn("w:t") and not any(a.tag == qn("w:ins") for a in node.iterancestors()):
            parts.append(node.text or "")
    return "".join(parts)


def views(path):
    """``(accept-all paragraphs, reject-all paragraphs)`` of an edited copy."""
    document = Document(path)
    accept = [_accept_all_paragraph_text(p._p) for p in document.paragraphs]
    reject = [_reject_text(p._p) for p in document.paragraphs if not _inserted(p._p)]
    return accept, reject


def run(tmp_path, edits, *, paragraphs=_PARAS, version=4, settings=None):
    build_spec(tmp_path / "spec.docx", paragraphs)
    sidecar = load_sidecar(write_sidecar(tmp_path, edits, version=version))
    (result,) = apply_sidecar(sidecar, [tmp_path / "spec.docx"], settings or RunSettings())
    by_id = {o.entry.finding_id: o for o in result.outcomes}
    return sidecar, result, by_id


def statuses(by_id) -> dict:
    return {finding_id: outcome.status for finding_id, outcome in by_id.items()}


# ---------------------------------------------------------------------------
# Different instructions for one place are held
# ---------------------------------------------------------------------------


class TestOverlappingEditsAreHeld:
    EDITS = [
        edit("rf-A", "gate valves", "ball valves"),
        edit("rf-B", "gate valves at each branch", "butterfly valves at each branch"),
    ]

    @pytest.mark.parametrize("order", [(0, 1), (1, 0)])
    def test_neither_is_applied_in_either_order(self, tmp_path, order):
        _, result, by_id = run(tmp_path, [self.EDITS[i] for i in order])
        assert statuses(by_id) == {
            "rf-A": OutcomeStatus.EDIT_CONFLICT,
            "rf-B": OutcomeStatus.EDIT_CONFLICT,
        }
        assert result.applied == 0
        assert not (tmp_path / "spec.applied.docx").exists()

    def test_each_names_the_other_and_says_why(self, tmp_path):
        _, _, by_id = run(tmp_path, self.EDITS)
        assert by_id["rf-A"].related == ("rf-B",)
        assert by_id["rf-B"].related == ("rf-A",)
        reason = by_id["rf-A"].reason
        assert "rf-B" in reason and "overlapping text in p1" in reason
        assert "--only" in reason

    def test_an_independent_edit_beside_them_still_lands(self, tmp_path):
        _, result, by_id = run(
            tmp_path, self.EDITS + [edit("rf-C", "hangers", "seismic hangers", element="p3")]
        )
        assert by_id["rf-C"].status is OutcomeStatus.APPLIED
        assert result.applied == 1
        accept, _ = views(tmp_path / "spec.applied.docx")
        assert accept[1] == _PARAS[1]
        assert accept[3] == "Install seismic hangers per the manufacturer."

    def test_a_chain_of_overlaps_holds_every_link(self, tmp_path):
        """A overlaps B and B overlaps C, though A and C do not touch."""
        _, _, by_id = run(
            tmp_path,
            [
                edit("rf-A", "Provide gate", "Install gate"),
                edit("rf-B", "gate valves at", "ball valves at"),
                edit("rf-C", "at each branch", "at every branch"),
            ],
        )
        assert set(statuses(by_id).values()) == {OutcomeStatus.EDIT_CONFLICT}
        assert by_id["rf-A"].related == ("rf-B",)
        assert by_id["rf-B"].related == ("rf-A", "rf-C")

    def test_an_edit_and_a_deletion_of_the_same_text_conflict(self, tmp_path):
        _, _, by_id = run(tmp_path, [edit("rf-A", "gate valves", "ball valves"), delete("rf-D", "gate valves")])
        assert set(statuses(by_id).values()) == {OutcomeStatus.EDIT_CONFLICT}

    def test_a_deletion_carrying_stray_replacement_text_is_still_a_deletion(self, tmp_path):
        """A DELETE may carry a ``replacement_text`` the writer ignores; the
        action, not the text alone, says what the change is."""
        deletion = delete("rf-D", "gate valves", proposal={"replacement_text": "ball valves"})
        _, _, by_id = run(tmp_path, [edit("rf-A", "gate valves", "ball valves"), deletion])
        assert set(statuses(by_id).values()) == {OutcomeStatus.EDIT_CONFLICT}

    def test_two_replacements_of_the_same_text_conflict(self, tmp_path):
        _, _, by_id = run(tmp_path, [edit("rf-A", "gate valves", "ball valves"), edit("rf-B", "gate valves", "check valves")])
        assert set(statuses(by_id).values()) == {OutcomeStatus.EDIT_CONFLICT}

    def test_edits_that_only_touch_at_a_boundary_are_compatible(self, tmp_path):
        _, result, by_id = run(
            tmp_path,
            [edit("rf-A", "Provide gate", "Install ball"), edit("rf-B", " valves at", " valves near")],
        )
        assert set(statuses(by_id).values()) == {OutcomeStatus.APPLIED}
        accept, reject = views(tmp_path / "spec.applied.docx")
        assert accept[1] == "Install ball valves near each branch line."
        assert reject == _PARAS

    def test_a_policy_hold_takes_no_part(self, tmp_path):
        """Only what would be written can conflict: a withheld instruction
        never blocks an allowed one."""
        _, _, by_id = run(
            tmp_path,
            [
                edit("rf-A", "gate valves", "ball valves"),
                edit("rf-B", "gate valves at each", "check valves at each", report_status="DISPUTED"),
            ],
        )
        assert statuses(by_id) == {
            "rf-A": OutcomeStatus.APPLIED,
            "rf-B": OutcomeStatus.HELD_BY_POLICY,
        }

    def test_conflicts_are_found_in_the_occurrence_schemas_too(self, tmp_path):
        entries = []
        for number, entry in enumerate(self.EDITS, start=1):
            entry = dict(entry, occurrence_id=f"oc-00000000000{number}", location_basis="validated")
            del entry["has_per_file_original"]
            entries.append(entry)
        _, _, by_id = run(tmp_path, entries, version=6)
        assert set(statuses(by_id).values()) == {OutcomeStatus.EDIT_CONFLICT}
        assert by_id["rf-A"].related == ("rf-B [oc-000000000002]",)


# ---------------------------------------------------------------------------
# Identical instructions are written once
# ---------------------------------------------------------------------------


class TestIdenticalEditsAreWrittenOnce:
    @pytest.mark.parametrize("order", [(0, 1), (1, 0)])
    def test_one_is_applied_and_the_other_is_its_duplicate(self, tmp_path, order):
        twins = [edit("rf-D1", "gate valves", "ball valves"), edit("rf-D2", "gate valves", "ball valves")]
        _, result, by_id = run(tmp_path, [twins[i] for i in order])
        # The written one is chosen by content, never by position.
        assert statuses(by_id) == {
            "rf-D1": OutcomeStatus.APPLIED,
            "rf-D2": OutcomeStatus.DUPLICATE,
        }
        assert by_id["rf-D2"].related == ("rf-D1",)
        assert "identical to rf-D1" in by_id["rf-D2"].reason
        assert result.applied == 1
        document = Document(tmp_path / "spec.applied.docx")
        assert len(document.element.body.findall(".//" + qn("w:del"))) == 1
        accept, reject = views(tmp_path / "spec.applied.docx")
        assert accept[1] == "Provide ball valves at each branch line."
        assert reject == _PARAS

    def test_a_review_and_a_coordination_finding_with_one_fix(self, tmp_path):
        _, _, by_id = run(
            tmp_path,
            [edit("rf-000000000001", "gate valves", "ball valves"), edit("cf-000000000001", "gate valves", "ball valves")],
        )
        assert statuses(by_id) == {
            "cf-000000000001": OutcomeStatus.APPLIED,
            "rf-000000000001": OutcomeStatus.DUPLICATE,
        }

    def test_the_same_edit_located_two_ways_is_one_change(self, tmp_path):
        """One names p1; the other names nothing and is found by its text."""
        _, _, by_id = run(
            tmp_path,
            [
                edit("rf-D1", "gate valves at each", "ball valves at each"),
                edit("rf-D2", "gate valves at each", "ball valves at each", element=None),
            ],
        )
        assert sorted(o.value for o in statuses(by_id).values()) == ["APPLIED", "DUPLICATE"]

    def test_a_dry_run_reports_the_duplicate_too(self, tmp_path):
        twins = [edit("rf-D1", "gate valves", "ball valves"), edit("rf-D2", "gate valves", "ball valves")]
        _, _, by_id = run(tmp_path, twins, settings=RunSettings(dry_run=True))
        assert statuses(by_id) == {
            "rf-D1": OutcomeStatus.WOULD_APPLY,
            "rf-D2": OutcomeStatus.DUPLICATE,
        }
        assert "would be written" in by_id["rf-D2"].reason
        assert not (tmp_path / "spec.applied.docx").exists()

    def test_an_identical_addition_is_one_new_paragraph(self, tmp_path):
        text = "Provide tamper switches on each valve."
        _, _, by_id = run(tmp_path, [add("rf-X1", text, anchor="gate valves at"), add("rf-X2", text, anchor="Provide gate")])
        assert sorted(o.value for o in statuses(by_id).values()) == ["APPLIED", "DUPLICATE"]
        accept, reject = views(tmp_path / "spec.applied.docx")
        assert accept.count(text) == 1
        assert reject == _PARAS

    def test_an_addition_after_one_paragraph_and_before_the_next_is_one_place(self, tmp_path):
        text = "Provide tamper switches on each valve."
        _, _, by_id = run(
            tmp_path,
            [
                add("rf-X1", text, anchor="at each branch", position="after", element="p1"),
                add("rf-X2", text, anchor="check valves", position="before", element="p2"),
            ],
        )
        assert sorted(o.value for o in statuses(by_id).values()) == ["APPLIED", "DUPLICATE"]
        accept, _ = views(tmp_path / "spec.applied.docx")
        assert accept[:4] == [_PARAS[0], _PARAS[1], text, _PARAS[2]]

    def test_when_the_written_one_fails_its_duplicate_says_so(self, tmp_path, monkeypatch):
        def boom(self, planned):
            raise RuntimeError("disk on fire")

        monkeypatch.setattr(DocumentEditor, "apply_planned", boom)
        twins = [edit("rf-D1", "gate valves", "ball valves"), edit("rf-D2", "gate valves", "ball valves")]
        _, _, by_id = run(tmp_path, twins)
        assert statuses(by_id) == {"rf-D1": OutcomeStatus.FAILED, "rf-D2": OutcomeStatus.FAILED}
        assert "which was not applied" in by_id["rf-D2"].reason

    def test_when_the_copy_cannot_be_saved_nothing_counts_as_written(self, tmp_path, monkeypatch):
        import docx.document

        def refuse(self, path):
            raise OSError("read-only volume")

        build_spec(tmp_path / "spec.docx")
        twins = [edit("rf-D1", "gate valves", "ball valves"), edit("rf-D2", "gate valves", "ball valves")]
        sidecar = load_sidecar(write_sidecar(tmp_path, twins))
        monkeypatch.setattr(docx.document.Document, "save", refuse)
        (result,) = apply_sidecar(sidecar, [tmp_path / "spec.docx"], RunSettings())
        by_id = {o.entry.finding_id: o for o in result.outcomes}
        assert statuses(by_id) == {"rf-D1": OutcomeStatus.FAILED, "rf-D2": OutcomeStatus.FAILED}
        assert all("read-only volume" in o.reason for o in by_id.values())
        assert result.applied == 0


# ---------------------------------------------------------------------------
# Additions at one place
# ---------------------------------------------------------------------------


class TestAdditionsAtOnePlace:
    @pytest.mark.parametrize(
        "second",
        [
            dict(anchor="Provide gate", position="after", element="p1"),
            dict(anchor="check valves", position="before", element="p2"),
        ],
        ids=["same-anchor", "next-paragraph"],
    )
    def test_two_different_new_paragraphs_at_one_place_are_held(self, tmp_path, second):
        _, _, by_id = run(
            tmp_path,
            [
                add("rf-X1", "First addition.", anchor="at each branch", position="after", element="p1"),
                add("rf-X2", "Second addition.", **second),
            ],
        )
        assert set(statuses(by_id).values()) == {OutcomeStatus.EDIT_CONFLICT}
        assert "insert different paragraphs at the same place" in by_id["rf-X1"].reason

    def test_the_same_addition_at_two_places_is_two_new_paragraphs(self, tmp_path):
        """The WP-06B case for additions: one fix needed at two places. Same
        text, different gaps — neither is a duplicate of the other."""
        text = "Brace every valve."
        _, _, by_id = run(
            tmp_path,
            [
                add("rf-X1", text, anchor="gate valves at", element="p1"),
                add("rf-X2", text, anchor="Install hangers", element="p3"),
            ],
        )
        assert set(statuses(by_id).values()) == {OutcomeStatus.APPLIED}
        accept, _ = views(tmp_path / "spec.applied.docx")
        assert accept == [_PARAS[0], _PARAS[1], text, _PARAS[2], _PARAS[3], text]

    def test_additions_on_either_side_of_one_paragraph_are_independent(self, tmp_path):
        _, _, by_id = run(
            tmp_path,
            [
                add("rf-X1", "Before it.", anchor="gate valves at", position="before"),
                add("rf-X2", "After it.", anchor="gate valves at", position="after"),
            ],
        )
        assert set(statuses(by_id).values()) == {OutcomeStatus.APPLIED}
        accept, _ = views(tmp_path / "spec.applied.docx")
        assert accept[:4] == [_PARAS[0], "Before it.", _PARAS[1], "After it."]

    def test_a_new_paragraph_copies_its_anchor_as_reviewed(self, tmp_path):
        """Additions go first, so a new paragraph takes the formatting of the
        anchor's first run as the review saw it — even when another edit
        deletes that run — and so neither order of the sidecar changes it."""
        document = Document()
        document.add_paragraph("SECTION 21 05 00")
        anchor = document.add_paragraph()
        anchor.add_run("Provide ").bold = True
        anchor.add_run("gate valves at each branch line.")
        document.save(tmp_path / "spec.docx")
        edits = [
            delete("rf-E", "Provide ", element="p1"),
            add("rf-X", "Brace every valve.", anchor="gate valves", element="p1"),
        ]
        for order in (edits, list(reversed(edits))):
            sidecar = load_sidecar(write_sidecar(tmp_path, order))
            (result,) = apply_sidecar(sidecar, [tmp_path / "spec.docx"], RunSettings())
            assert result.applied == 2
            added = Document(tmp_path / "spec.applied.docx").paragraphs[2]._p
            assert _accept_all_paragraph_text(added) == "Brace every valve."
            (inserted_run,) = added.find(qn("w:ins")).findall(qn("w:r"))
            properties = inserted_run.find(qn("w:rPr"))
            assert properties is not None and properties.find(qn("w:b")) is not None

    def test_an_addition_and_an_edit_to_its_anchor_both_land(self, tmp_path):
        """Before S11, whichever came second failed: the edit changed the
        anchor's text, and the addition could no longer find it."""
        for edits in (
            [edit("rf-E", "Provide gate", "Install ball"), add("rf-X", "Added.", anchor="Provide gate")],
            [add("rf-X", "Added.", anchor="Provide gate"), edit("rf-E", "Provide gate", "Install ball")],
        ):
            _, _, by_id = run(tmp_path, edits)
            assert set(statuses(by_id).values()) == {OutcomeStatus.APPLIED}
            accept, reject = views(tmp_path / "spec.applied.docx")
            assert accept[1:3] == ["Install ball valves at each branch line.", "Added."]
            assert reject == _PARAS


# ---------------------------------------------------------------------------
# The sidecar's order decides nothing
# ---------------------------------------------------------------------------


_MIXED = [
    # Compatible, though one's replacement repeats the other's target text.
    edit("rf-E1", "gate valves", "ball valves or approved check valves", element="p2"),
    edit("rf-E2", "check valves", "swing check valves", element="p2"),
    # An addition whose anchor another edit changes.
    edit("rf-F1", "Install hangers", "Install seismic hangers", element="p3"),
    add("rf-F2", "Brace every hanger.", anchor="Install hangers", element="p3"),
    # A conflict and a duplicate.
    edit("rf-C1", "at each branch", "at every branch", element="p1"),
    edit("rf-C2", "each branch line", "each branch main", element="p1"),
    edit("rf-D1", "SECTION", "Section", element="p0"),
    edit("rf-D2", "SECTION", "Section", element="p0"),
]


class TestOrderDecidesNothing:
    def _outcome(self, tmp_path, edits, name):
        folder = tmp_path / name
        folder.mkdir()
        _, _, by_id = run(folder, edits)
        return statuses(by_id), views(folder / "spec.applied.docx")

    def test_every_order_of_a_mixed_sidecar_gives_one_result(self, tmp_path):
        baseline_status, baseline_views = self._outcome(tmp_path, _MIXED, "baseline")
        assert baseline_status == {
            "rf-E1": OutcomeStatus.APPLIED,
            "rf-E2": OutcomeStatus.APPLIED,
            "rf-F1": OutcomeStatus.APPLIED,
            "rf-F2": OutcomeStatus.APPLIED,
            "rf-C1": OutcomeStatus.EDIT_CONFLICT,
            "rf-C2": OutcomeStatus.EDIT_CONFLICT,
            "rf-D1": OutcomeStatus.APPLIED,
            "rf-D2": OutcomeStatus.DUPLICATE,
        }
        accept, reject = baseline_views
        assert accept == [
            "Section 21 05 00",
            _PARAS[1],
            "Provide ball valves or approved check valves and swing check valves.",
            "Install seismic hangers per the manufacturer.",
            "Brace every hanger.",
        ]
        assert reject == _PARAS
        orders = [list(reversed(_MIXED))]
        for seed in range(12):
            shuffled = list(_MIXED)
            random.Random(seed).shuffle(shuffled)
            orders.append(shuffled)
        for number, order in enumerate(orders):
            assert self._outcome(tmp_path, order, f"order{number}") == (baseline_status, baseline_views)

    @pytest.mark.parametrize("order", [(0, 1), (1, 0)])
    def test_a_replacement_that_repeats_an_earlier_target_hides_nothing(self, tmp_path, order):
        """Edits in one paragraph are written last span first, so the right
        edit's new text lands before the left edit is written. It repeats the
        left edit's target, which a writer re-finding its text would now see
        twice and refuse; the planned span is unaffected."""
        edits = [
            edit("rf-L", "gate valves", "ball valves", element="p2"),
            edit("rf-R", "check valves", "gate valves or check valves", element="p2"),
        ]
        _, _, by_id = run(tmp_path, [edits[i] for i in order])
        assert set(statuses(by_id).values()) == {OutcomeStatus.APPLIED}
        accept, reject = views(tmp_path / "spec.applied.docx")
        assert accept[2] == "Provide ball valves and gate valves or check valves."
        assert reject == _PARAS

    @pytest.mark.parametrize("mode", ["tracked", "direct"])
    def test_several_edits_in_one_paragraph_land_from_their_planned_spans(self, tmp_path, mode):
        settings = RunSettings(mode=DIRECT) if mode == "direct" else RunSettings()
        edits = [
            edit("rf-1", "Provide", "Furnish", element="p2"),
            edit("rf-2", "gate valves", "ball valves", element="p2"),
            edit("rf-3", "check valves", "swing check valves", element="p2"),
        ]
        for order in itertools.permutations(edits):
            folder = tmp_path / f"{mode}-{'-'.join(e['finding_id'] for e in order)}"
            folder.mkdir()
            _, _, by_id = run(folder, list(order), settings=settings)
            assert set(statuses(by_id).values()) == {OutcomeStatus.APPLIED}
            accept, reject = views(folder / "spec.applied.docx")
            assert accept[2] == "Furnish ball valves and swing check valves."
            if mode == "tracked":
                assert reject == _PARAS


# ---------------------------------------------------------------------------
# Everything is decided before the first write
# ---------------------------------------------------------------------------


def _document_with_a_tab_run():
    document = Document()
    document.add_paragraph("SECTION 21 05 00")
    paragraph = document.add_paragraph("Provide gate")
    run = paragraph.add_run(" valves")
    run._r.append(run._r.makeelement(qn("w:tab"), {}))
    paragraph.add_run(" at each branch.")
    return document


def _entry(existing, replacement="X"):
    return EditEntry(
        finding_id="rf-1", file_name="spec.docx", action_type="EDIT",
        existing_text=existing, replacement_text=replacement, anchor_text=None,
        insert_position=None, target_element_id="p1", evidence_element_id="p1",
        edit_confidence=0.9, report_status="VERIFIED_SUPPORTED", verification_verdict="CONFIRMED",
        severity="HIGH", section="", issue="", code_reference=None, has_per_file_original=True,
    )


class TestDecidedBeforeTheFirstWrite:
    def test_an_unsplittable_boundary_is_refused_while_planning(self):
        """Before S11 the start boundary was split first and the end refused
        afterwards, leaving a split run behind a refused edit."""
        document = _document_with_a_tab_run()
        before = document.element.body.xml
        editor = DocumentEditor(document)
        paragraphs = editor.resolve(
            Location(status=LocationStatus.RESOLVED_BY_ID, element_id="p1", kind=ElementKind.BODY_PARAGRAPH)
        )
        with pytest.raises(EditError, match="tab, a line"):
            editor.plan(_entry("gate val"), paragraphs)
        assert document.element.body.xml == before

    def test_planning_never_changes_the_document(self, tmp_path):
        build_spec(tmp_path / "spec.docx")
        document = Document(tmp_path / "spec.docx")
        before = document.element.body.xml
        editor = DocumentEditor(document)
        location = Location(status=LocationStatus.RESOLVED_BY_ID, element_id="p2", kind=ElementKind.BODY_PARAGRAPH)
        for existing in ("gate valves", "check valves", "Provide"):
            editor.plan(_entry(existing), editor.resolve(location))
        assert document.element.body.xml == before

    def test_a_writer_refusal_reaches_the_receipt_as_before(self, tmp_path):
        paragraphs = ["Intro.", "Install per NFPA 13 and test per NFPA 13 before acceptance."]
        _, _, by_id = run(tmp_path, [edit("rf-dup", "per NFPA 13", "per NFPA 13 (2025)")], paragraphs=paragraphs)
        assert by_id["rf-dup"].status is OutcomeStatus.UNLOCATED
        assert "occurs 2 times" in by_id["rf-dup"].reason


# ---------------------------------------------------------------------------
# The settlement on its own
# ---------------------------------------------------------------------------


def _plan(finding_id, paragraph, span=None, *, action="EDIT", text="X", gap=None):
    entry = SimpleNamespace(
        action_type=action,
        replacement_text=text,
        sort_key=(finding_id,),
    )
    return PlannedEdit(entry=entry, paragraph=paragraph, span=span, direct_span=span, gap=gap)


class TestSettle:
    def test_nothing_to_settle(self):
        assert settle([]) == Settlement()

    def test_independent_plans_all_apply(self):
        p1, p2 = object(), object()
        plans = [_plan("b", p1, (0, 4)), _plan("a", p2, (2, 6)), _plan("c", p1, (5, 9))]
        settled = settle(plans)
        assert not settled.conflicts and not settled.duplicates
        # Paragraphs follow their smallest key (p2's "a" before p1's "b"), and
        # within p1 the later span goes first.
        assert settled.order == (1, 2, 0)

    def test_additions_go_first(self):
        p1 = object()
        plans = [_plan("a", p1, (0, 4)), _plan("z", p1, action="ADD", gap=(p1, None))]
        assert settle(plans).order == (1, 0)

    def test_identity_and_overlap(self):
        p1 = object()
        plans = [
            _plan("a", p1, (0, 5)),
            _plan("b", p1, (0, 5)),
            _plan("c", p1, (3, 8), text="Y"),
            _plan("d", p1, (10, 12)),
        ]
        settled = settle(plans)
        assert settled.conflicts == {0: (2,), 1: (2,), 2: (0, 1)}
        assert settled.duplicates == {}
        assert settled.order == (3,)

    def test_the_written_duplicate_is_the_first_by_content(self):
        p1 = object()
        settled = settle([_plan("z", p1, (0, 5)), _plan("m", p1, (0, 5)), _plan("q", p1, (0, 5))])
        assert settled.order == (1,)
        assert settled.duplicates == {0: 1, 2: 1}

    def test_application_order_is_last_span_first_within_a_paragraph(self):
        p1 = object()
        plans = [_plan("a", p1, (0, 2)), _plan("b", p1, (8, 9)), _plan("c", p1, (4, 6))]
        assert application_order(plans, [0, 1, 2]) == (1, 2, 0)


# ---------------------------------------------------------------------------
# The receipt, the summary, and the exit status
# ---------------------------------------------------------------------------


class TestReporting:
    def _run_all(self, tmp_path):
        edits = [
            edit("rf-A", "gate valves", "ball valves"),
            edit("rf-B", "gate valves at each", "check valves at each"),
            edit("rf-D1", "hangers", "seismic hangers", element="p3"),
            edit("rf-D2", "hangers", "seismic hangers", element="p3"),
        ]
        sidecar, result, by_id = run(tmp_path, edits)
        receipt = build_receipt(
            sidecar=sidecar, file_results=[result], settings=RunSettings().to_dict(), entries_in=len(edits)
        )
        return receipt, result, by_id

    def test_the_receipt_balances_and_counts_both_new_outcomes(self, tmp_path):
        receipt, _, _ = self._run_all(tmp_path)
        accounting = receipt["accounting"]
        assert accounting["balanced"] is True
        assert accounting["by_outcome"]["EDIT_CONFLICT"] == 2
        assert accounting["by_outcome"]["DUPLICATE"] == 1
        assert accounting["by_outcome"]["APPLIED"] == 1
        outcomes = {o["finding_id"]: o for o in receipt["files"][0]["outcomes"]}
        assert outcomes["rf-A"]["related_entries"] == ["rf-B"]
        assert outcomes["rf-D2"]["related_entries"] == ["rf-D1"]

    def test_the_summary_puts_conflicts_in_front_of_a_person(self, tmp_path):
        receipt, result, _ = self._run_all(tmp_path)
        text = render_summary(receipt, [result])
        assert "2 edit conflict" in text
        assert "1 duplicate(s) of an instruction applied" in text
        not_applied = text.split("Not applied:")[1].split("Applied:")[0]
        assert "rf-A" in not_applied and "rf-B" in not_applied
        assert "Duplicates (the change is made once):" in text

    def test_strict_counts_a_conflict_but_not_a_duplicate(self, tmp_path):
        build_spec(tmp_path / "spec.docx")
        twins = write_sidecar(
            tmp_path, [edit("rf-D1", "hangers", "seismic hangers", element="p3"), edit("rf-D2", "hangers", "seismic hangers", element="p3")]
        )
        assert main([str(twins), "--specs", str(tmp_path), "--quiet", "--strict"]) == EXIT_OK
        (tmp_path / "spec.applied.docx").unlink()
        clash = write_sidecar(tmp_path, TestOverlappingEditsAreHeld.EDITS)
        assert main([str(clash), "--specs", str(tmp_path), "--quiet", "--strict"]) == EXIT_UNAPPLIED
