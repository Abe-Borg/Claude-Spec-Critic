"""The resolution ladder, and the refusals at the bottom of it.

The locator's value is not that it finds targets — an index lookup finds
targets. It is that it declines to find one when the evidence does not single
one out. Most of these tests are about the declining.
"""
from __future__ import annotations

import pytest
from docx import Document

from applier.locator import (
    build_candidates,
    classify_element_id,
    locate,
)
from applier.models import (
    Candidate,
    EditEntry,
    ElementKind,
    LocationStatus,
)
from src.input.extractor import extract_text_from_docx


def candidate(element_id, text, section_id="", element_type="paragraph"):
    return Candidate(
        element_id=element_id,
        text=text,
        section_id=section_id,
        element_type=element_type,
        kind=classify_element_id(element_id),
    )


def entry(**overrides):
    base = dict(
        finding_id="rf-1",
        file_name="215000.docx",
        action_type="EDIT",
        existing_text="NFPA 13, 2019 edition",
        replacement_text="NFPA 13, 2025 edition",
        anchor_text=None,
        insert_position=None,
        target_element_id=None,
        evidence_element_id=None,
        edit_confidence=0.9,
        report_status="VERIFIED_SUPPORTED",
        verification_verdict="CONFIRMED",
        severity="HIGH",
        section="21 13 13",
        issue="…",
        code_reference=None,
        has_per_file_original=True,
    )
    base.update(overrides)
    return EditEntry(**base)


class TestElementKindClassification:
    @pytest.mark.parametrize(
        "element_id, kind",
        [
            ("p0", ElementKind.BODY_PARAGRAPH),
            ("p137", ElementKind.BODY_PARAGRAPH),
            ("t0r2", ElementKind.TABLE_ROW),
            ("t1r0c0t0r3", ElementKind.TABLE_ROW),
            ("s0h1", ElementKind.HEADER_FOOTER),
            ("s2f0", ElementKind.HEADER_FOOTER),
            # Text inside a block content control (plan WP-02): readable,
            # never written, and never mistaken for a legacy id.
            ("cc1p0", ElementKind.CONTENT_CONTROL),
            ("cc1cc2p0", ElementKind.CONTENT_CONTROL),
            ("cc1t1r0", ElementKind.CONTENT_CONTROL),
            ("cc1t1r0c0t0r1", ElementKind.CONTENT_CONTROL),
            ("t0cc5r0", ElementKind.CONTENT_CONTROL),
            ("t0r1c0t0cc3r0", ElementKind.CONTENT_CONTROL),
            ("s0hcc2p0", ElementKind.CONTENT_CONTROL),
            ("s1fcc0cc1p2", ElementKind.CONTENT_CONTROL),
            ("tb0cc1p0", ElementKind.CONTENT_CONTROL),
            ("fn3cc0p0", ElementKind.CONTENT_CONTROL),
            ("en2cc1p0", ElementKind.CONTENT_CONTROL),
            ("cc1", ElementKind.UNSUPPORTED),
            ("cc1t1", ElementKind.UNSUPPORTED),
            ("p1cc0p0", ElementKind.UNSUPPORTED),
            ("tb0p1", ElementKind.UNSUPPORTED),
            ("fn3p0", ElementKind.UNSUPPORTED),
            ("en1p0", ElementKind.UNSUPPORTED),
            ("meta:hf", ElementKind.UNSUPPORTED),
            ("", ElementKind.UNSUPPORTED),
            (None, ElementKind.UNSUPPORTED),
            ("p", ElementKind.UNSUPPORTED),
            ("px1", ElementKind.UNSUPPORTED),
        ],
    )
    def test_ids_classify_by_the_extractors_scheme(self, element_id, kind):
        assert classify_element_id(element_id) is kind


class TestCandidatesMatchTheExtractor:
    """The applier must agree with the code that minted the ids — if the two
    ever disagree, edits land on the wrong paragraph."""

    def test_ids_come_straight_from_the_paragraph_map(self, tmp_path):
        document = Document()
        document.add_paragraph("SECTION 21 13 13")
        document.add_paragraph("Sprinklers shall comply with NFPA 13.")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Hazard"
        table.cell(0, 1).text = "Density"
        path = tmp_path / "215000.docx"
        document.save(path)

        extracted = extract_text_from_docx(path)
        candidates = build_candidates(extracted)
        assert [c.element_id for c in candidates] == [
            mapping.element_id
            for mapping in extracted.paragraph_map
            if not mapping.element_id.startswith("meta:")
        ]
        assert candidates[0].element_id == "p0"
        assert any(c.kind is ElementKind.TABLE_ROW for c in candidates)

    def test_synthetic_block_delimiters_are_not_targets(self):
        class FakeMapping:
            def __init__(self, element_id, text):
                self.element_id = element_id
                self.text = text
                self.section_id = ""
                self.element_type = "paragraph"

        class FakeSpec:
            paragraph_map = [FakeMapping("meta:hf", "===== HEADER ====="), FakeMapping("p1", "x")]

        assert [c.element_id for c in build_candidates(FakeSpec())] == ["p1"]

    def test_a_spec_with_no_map_yields_no_candidates(self):
        class Empty:
            paragraph_map = None

        assert build_candidates(Empty()) == []


class TestResolutionLadder:
    def test_by_element_id_when_the_text_confirms(self):
        candidates = [
            candidate("p1", "Sprinklers shall comply with NFPA 13, 2019 edition."),
            candidate("p5", "Hangers shall comply with NFPA 13, 2019 edition."),
        ]
        location = locate(entry(target_element_id="p5"), candidates)
        assert location.status is LocationStatus.RESOLVED_BY_ID
        assert location.element_id == "p5"
        assert location.is_applicable

    def test_the_proposals_own_id_outranks_the_evidence_id(self):
        candidates = [
            candidate("p1", "A: NFPA 13, 2019 edition"),
            candidate("p5", "B: NFPA 13, 2019 edition"),
        ]
        location = locate(
            entry(target_element_id="p5", evidence_element_id="p1"), candidates
        )
        assert location.element_id == "p5"

    def test_by_unique_text_when_no_id_is_given(self):
        candidates = [
            candidate("p1", "Sprinklers shall comply with NFPA 13, 2019 edition."),
            candidate("p2", "Unrelated requirement."),
        ]
        location = locate(entry(), candidates)
        assert location.status is LocationStatus.RESOLVED_BY_UNIQUE_TEXT
        assert location.element_id == "p1"

    def test_whitespace_differences_do_not_defeat_the_match(self):
        candidates = [candidate("p1", "Comply with NFPA 13,\n  2019   edition today.")]
        assert locate(entry(), candidates).element_id == "p1"

    def test_by_section_when_the_text_is_not_unique(self):
        candidates = [
            candidate("p1", "Comply with NFPA 13, 2019 edition.", "21 13 13 WET-PIPE"),
            candidate("p9", "Comply with NFPA 13, 2019 edition.", "21 22 00 STANDPIPES"),
        ]
        location = locate(entry(section="21 13 13"), candidates)
        assert location.status is LocationStatus.RESOLVED_BY_SECTION
        assert location.element_id == "p1"

    def test_add_resolves_on_its_anchor_text(self):
        candidates = [
            candidate("p1", "Provide hangers per manufacturer instructions."),
            candidate("p2", "Something else."),
        ]
        location = locate(
            entry(
                action_type="ADD",
                existing_text=None,
                anchor_text="Provide hangers",
                insert_position="after",
            ),
            candidates,
        )
        assert location.status is LocationStatus.RESOLVED_BY_UNIQUE_TEXT
        assert location.element_id == "p1"

    def test_add_may_resolve_on_an_element_id_with_no_anchor(self):
        candidates = [candidate("p3", "Anything at all.")]
        location = locate(
            entry(
                action_type="ADD",
                existing_text=None,
                anchor_text=None,
                target_element_id="p3",
                insert_position="before",
            ),
            candidates,
        )
        assert location.status is LocationStatus.RESOLVED_BY_ID


class TestRefusals:
    def test_ambiguous_text_is_refused_and_the_candidates_are_named(self):
        candidates = [
            candidate("p1", "Comply with NFPA 13, 2019 edition.", "PART 1"),
            candidate("p9", "Comply with NFPA 13, 2019 edition.", "PART 2"),
        ]
        location = locate(entry(section="99 99 99"), candidates)
        assert location.status is LocationStatus.AMBIGUOUS
        assert not location.is_applicable
        assert {c.element_id for c in location.candidates} == {"p1", "p9"}

    def test_a_section_matching_several_candidates_stays_ambiguous(self):
        candidates = [
            candidate("p1", "NFPA 13, 2019 edition", "21 13 13"),
            candidate("p9", "NFPA 13, 2019 edition", "21 13 13"),
        ]
        assert locate(entry(section="21 13 13"), candidates).status is (
            LocationStatus.AMBIGUOUS
        )

    def test_an_empty_section_never_disambiguates(self):
        candidates = [
            candidate("p1", "NFPA 13, 2019 edition", ""),
            candidate("p9", "NFPA 13, 2019 edition", ""),
        ]
        assert locate(entry(section=""), candidates).status is LocationStatus.AMBIGUOUS

    def test_missing_text_is_not_found(self):
        location = locate(entry(), [candidate("p1", "Nothing relevant here.")])
        assert location.status is LocationStatus.NOT_FOUND
        assert "may have been edited since the review" in location.detail

    def test_drift_is_reported_with_what_the_element_now_says(self):
        candidates = [candidate("p4", "Comply with NFPA 25, 2013 California Edition.")]
        location = locate(entry(target_element_id="p4"), candidates)
        assert location.status is LocationStatus.DRIFTED
        assert "NFPA 25" in location.detail

    def test_drift_falls_through_when_the_clause_merely_moved(self):
        candidates = [
            candidate("p4", "Something else entirely."),
            candidate("p7", "Comply with NFPA 13, 2019 edition."),
        ]
        location = locate(entry(target_element_id="p4"), candidates)
        assert location.status is LocationStatus.RESOLVED_BY_UNIQUE_TEXT
        assert location.element_id == "p7"
        assert "no longer contains" in location.detail

    @pytest.mark.parametrize("element_id", ["tb0p1", "fn2p0", "en1p0"])
    def test_unsupported_containers_are_reported_not_written(self, element_id):
        candidates = [candidate(element_id, "NFPA 13, 2019 edition")]
        location = locate(entry(target_element_id=element_id), candidates)
        assert location.status is LocationStatus.UNSUPPORTED_ELEMENT
        assert not location.is_applicable
        assert "by hand" in location.detail

    def test_text_found_only_in_an_unsupported_container_is_refused(self):
        candidates = [candidate("fn1p0", "Comply with NFPA 13, 2019 edition.")]
        assert locate(entry(), candidates).status is (
            LocationStatus.UNSUPPORTED_ELEMENT
        )

    def test_an_unwritable_copy_makes_an_id_less_match_ambiguous(self):
        """The same text in a note and in the body, and no id to say which
        the finding meant. Before S10 the writable copy won by default, which
        is how an edit meant for text inside a content control could land on
        a plain copy elsewhere (plan WP-02). Neither wins now: the note is
        never written, and the body copy is not assumed."""
        candidates = [
            candidate("fn1p0", "Comply with NFPA 13, 2019 edition."),
            candidate("p3", "Comply with NFPA 13, 2019 edition."),
        ]
        location = locate(entry(), candidates)
        assert location.status is LocationStatus.AMBIGUOUS
        assert not location.is_applicable
        assert {c.element_id for c in location.candidates} == {"fn1p0", "p3"}
        assert "fn1p0" in location.detail

    def test_the_findings_section_can_still_single_out_the_writable_copy(self):
        candidates = [
            candidate("cc4p0", "Comply with NFPA 13, 2019 edition.", section_id="1.01 SUMMARY"),
            candidate("p9", "Comply with NFPA 13, 2019 edition.", section_id="21 13 13 PIPING"),
        ]
        location = locate(entry(section="21 13 13"), candidates)
        assert location.status is LocationStatus.RESOLVED_BY_SECTION
        assert location.element_id == "p9"

    def test_a_section_that_points_at_the_unwritable_copy_is_refused(self):
        candidates = [
            candidate("cc4p0", "Comply with NFPA 13, 2019 edition.", section_id="21 13 13 PIPING"),
            candidate("p9", "Comply with NFPA 13, 2019 edition.", section_id="1.01 SUMMARY"),
        ]
        location = locate(entry(section="21 13 13"), candidates)
        assert location.status is LocationStatus.UNSUPPORTED_ELEMENT
        assert location.element_id == "cc4p0"
        assert "content control" in location.detail

    def test_no_id_and_no_text_cannot_be_located(self):
        location = locate(
            entry(action_type="DELETE", existing_text=None), [candidate("p1", "x")]
        )
        assert location.status is LocationStatus.NOT_FOUND


class TestBorrowedLocator:
    """``has_per_file_original=False`` means the id points into a *different*
    document, so a mismatch there is expected rather than evidence of drift."""

    def test_a_borrowed_id_that_misses_falls_through_to_text(self):
        candidates = [
            candidate("p4", "A completely different clause."),
            candidate("p8", "Comply with NFPA 13, 2019 edition."),
        ]
        location = locate(
            entry(target_element_id="p4", has_per_file_original=False), candidates
        )
        assert location.status is LocationStatus.RESOLVED_BY_UNIQUE_TEXT
        assert location.element_id == "p8"
        assert "no longer contains" not in location.detail

    def test_the_same_miss_is_drift_when_the_locator_is_this_files_own(self):
        candidates = [candidate("p4", "A completely different clause.")]
        assert locate(
            entry(target_element_id="p4", has_per_file_original=True), candidates
        ).status is LocationStatus.DRIFTED

    def test_a_borrowed_id_that_hits_is_still_used(self):
        candidates = [candidate("p4", "Comply with NFPA 13, 2019 edition.")]
        location = locate(
            entry(target_element_id="p4", has_per_file_original=False), candidates
        )
        assert location.status is LocationStatus.RESOLVED_BY_ID
