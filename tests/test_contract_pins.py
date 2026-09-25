"""Contract pins for behavior that is already right (implementation plan WP-01, S01).

Later chunks rewrite extraction (S10, S14), finding identity (S02, S11, S12),
and the applier. These tests pin what those chunks must *keep*, so a change
that breaks one reads as a regression rather than slipping through:

* **Extraction reconstruction.** ``"\\n\\n".join(m.text for m in
  paragraph_map) == content`` for every fixture, and re-extraction is
  deterministic.
* **Unique element ids** in a known grammar. S10 (plan WP-02) added one
  namespace to it: text read from inside a block content control carries a
  ``cc<n>`` step (``cc1p0``, ``cc1t1r0``, ``t0cc5r0``).
* **The meaning of legacy ids.** ``pN`` is *physical* body child ``N``
  (anything else at the body level, such as a block content control, still
  occupies an index); ``tN`` is the ``N``-th *direct* body table (a table
  wrapped in a control does not count). Pinned against the XML and through
  the applier's own resolver, because the extractor and the applier must
  agree or an edit lands in the wrong paragraph. Every ``cc`` id is checked
  against an independent reading of the XML and must be one the applier
  classifies as a content control and refuses to write.
* **Established text keeps its location.** Text the extractor reads today
  keeps its element id and text when new structures become readable.
* **Group-versus-occurrence identity**: what a merged finding's identity is.
  S11 replaced the old ``grp-0000`` presentation counters with
  content-derived occurrence ids and one occurrence per place, not per file
  (pinned in ``tests/test_edit_occurrences.py``).
* **One extracted input, two requests.** The review repair request
  re-sends exactly the primary request's input plus the retry suffix.
* **The applier boundary for every new container**: an edit aimed at text
  inside a content control, field, smart tag, or hyperlink is either
  applied exactly or refused, and never changes anything else. "Exactly"
  includes structure: every wrapper survives with its identity (link
  target, field instruction, control), and the new text stays inside the
  wrapper that held the old text. (S10 made the text readable and chose
  refusal: the writer never edits through a wrapper; see
  ``tests/test_applier_wrapped_content.py``.)

The automatically numbered fixture is held to numbering-neutral pins only
(same ids, literal text still present): WP-03 allows S14 to show displayed
labels in the extracted text, so an exact-text pin there would pin the
numbering defect instead of protecting a correct behavior.

Behavior that is *wrong* today is not pinned here; it is tracked as strict
xfails in ``tests/test_plan_open_defects.py``.
"""
from __future__ import annotations

import json
import re
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from applier.docx_edit import DocumentEditor, EditError
from applier.locator import classify_element_id
from applier.models import ElementKind, OutcomeStatus
from applier.run import RunSettings, apply_sidecar
from applier.sidecar import load_sidecar
from src.core.code_cycles import CALIFORNIA_2025
from src.input.extractor import _accept_all_paragraph_text, extract_text_from_docx
from src.orchestration import pipeline
from src.output import edit_sidecar
from src.review.review_request_builder import RETRY_TRUNCATED_REVIEW_INSTRUCTION
from src.review.reviewer import Finding, ReviewResult
from tests.fixtures import spec_docx as fx

# ---------------------------------------------------------------------------
# Fixture catalogue
# ---------------------------------------------------------------------------

_FIXTURES = {
    **{
        f"three_part:{variant.name}": (lambda blocks=variant.blocks: fx.build_blocks(blocks))
        for variant in fx.three_part_variants()
    },
    "block_control": fx.build_block_control_spec,
    "inline_controls": fx.build_inline_controls_spec,
    "fields": fx.build_fields_spec,
    "smart_tag": fx.build_smart_tag_spec,
    "hyperlinks": fx.build_hyperlink_spec,
    "tracked_changes": fx.build_tracked_changes_spec,
    "merged_and_nested_tables": fx.build_merged_and_nested_tables_spec,
    "table_controls": fx.build_table_controls_spec,
    "nested_controls": fx.build_nested_controls_spec,
    "table_of_contents": fx.build_table_of_contents_spec,
}

#: The three-PART variants whose numbers are typed text. The automatically
#: numbered variant is excluded from the exact-text pins: WP-03 (S14) may put
#: its displayed labels into the extracted text, so pinning today's label-free
#: text would pin the numbering defect itself. It gets numbering-neutral pins.
_TYPED_VARIANTS = [
    variant
    for variant in fx.three_part_variants()
    if not any(getattr(block, "auto_label", None) for block in variant.blocks)
]

#: Every element-id shape the extractor mints (ParagraphMapping docstring).
#: The ``cc<n>`` shapes are S10's (plan WP-02): text inside a block content
#: control, a row-level control within a table, or a control in a header,
#: footer, text box, or note.
_ROW_PATH = r"(?:cc\d+)*r\d+(?:c\d+t\d+(?:cc\d+)*r\d+)*"
_ELEMENT_ID_RE = re.compile(
    rf"^(?:p\d+|t\d+{_ROW_PATH}|s\d+[hf]\d+|tb\d+p\d+|fn\d+p\d+|en\d+p\d+"
    rf"|(?:cc\d+)+(?:p\d+|t\d+{_ROW_PATH})"
    r"|(?:s\d+[hf]|tb\d+|(?:fn|en)\d+)(?:cc\d+)+p\d+"
    r"|meta:(?:tb|fn|en|hf))$"
)


def _extract(name: str, tmp_path: Path):
    path = fx.save_docx(_FIXTURES[name](), tmp_path, f"{name.replace(':', '_')}.docx")
    return path, extract_text_from_docx(path)


# ---------------------------------------------------------------------------
# Reconstruction, determinism, and id uniqueness
# ---------------------------------------------------------------------------


class TestReconstruction:
    @pytest.mark.parametrize("name", sorted(_FIXTURES))
    def test_the_paragraph_map_reconstructs_the_content(self, name, tmp_path):
        _, spec = _extract(name, tmp_path)
        assert "\n\n".join(m.text for m in spec.paragraph_map) == spec.content
        assert spec.word_count == len(spec.content.split())

    @pytest.mark.parametrize("name", sorted(_FIXTURES))
    def test_re_extraction_is_identical(self, name, tmp_path):
        """The resume path re-extracts specs from disk, so a second
        extraction must equal the first — content, ids, and sections."""
        path, first = _extract(name, tmp_path)
        second = extract_text_from_docx(path)
        assert second.content == first.content
        assert second.paragraph_map == first.paragraph_map
        assert second.tracked_changes_detected == first.tracked_changes_detected

    @pytest.mark.parametrize("variant", _TYPED_VARIANTS, ids=lambda v: v.name)
    def test_typed_fixtures_extract_to_their_declared_text(self, variant, tmp_path):
        path = fx.save_docx(fx.build_blocks(variant.blocks), tmp_path, "spec.docx")
        assert extract_text_from_docx(path).content == fx.blocks_text(variant.blocks)


class TestElementIds:
    @pytest.mark.parametrize("name", sorted(_FIXTURES))
    def test_ids_are_unique_and_well_formed(self, name, tmp_path):
        _, spec = _extract(name, tmp_path)
        ids = [m.element_id for m in spec.paragraph_map]
        assert len(ids) == len(set(ids)), ids
        assert all(_ELEMENT_ID_RE.match(i) for i in ids), ids


# ---------------------------------------------------------------------------
# The meaning of legacy pN / tN ids
# ---------------------------------------------------------------------------


def _row_pieces(text: str) -> list[str]:
    return [piece for cell in text.split(" | ") for piece in cell.split("\n") if piece]


def _reads_as(p_el, mapping_text: str) -> bool:
    """Whether a mapping's text is what paragraph ``p_el`` holds.

    Equal to the paragraph's Accept-All text — except that a paragraph
    carrying automatic numbering (``w:numPr``) may also show its displayed
    label in front of that text, which WP-03 (S14) is allowed to add.
    """
    literal = _accept_all_paragraph_text(p_el).strip()
    if mapping_text == literal:
        return True
    numbered = p_el.find(f"{qn('w:pPr')}/{qn('w:numPr')}") is not None
    return numbered and bool(literal) and mapping_text.endswith(literal)


_BLOCK_WRAPPER_TAGS = (qn("w:sdt"), qn("w:customXml"))
_CC_STEP_RE = re.compile(r"(cc|p|t|r)(\d+)")


def _shown_children(wrapper) -> list:
    """A block control's shown children, read straight from the XML."""
    if wrapper.tag == qn("w:sdt"):
        container = wrapper.find(qn("w:sdtContent"))
        children = list(container) if container is not None else []
    else:
        children = [child for child in wrapper if child.tag != qn("w:customXmlPr")]
    return [child for child in children if isinstance(child.tag, str)]


def _block_texts(children) -> list[str]:
    """Non-empty Accept-All texts of the paragraphs among ``children`` and
    inside any block control among them (a cell's content, or a control's)."""
    texts = []
    for child in children:
        if child.tag == qn("w:p"):
            text = _accept_all_paragraph_text(child).strip()
            if text:
                texts.append(text)
        elif child.tag in _BLOCK_WRAPPER_TAGS:
            texts.extend(_block_texts(_shown_children(child)))
    return texts


def _row_texts(tr) -> list[str]:
    """Every non-empty paragraph text of a row, cells (wrapped or not) in order."""
    texts = []
    for child in tr:
        if child.tag == qn("w:tc"):
            texts.extend(_block_texts(child))
        elif child.tag in _BLOCK_WRAPPER_TAGS:
            for wrapped in _shown_children(child):
                if wrapped.tag == qn("w:tc"):
                    texts.extend(_block_texts(wrapped))
    return texts


def _read_cc_id(body, element_id: str) -> list[str]:
    """What a ``cc`` id names, found by walking its steps through the XML.

    Deliberately independent of the extractor's own helpers, so the two
    cannot share a mistake: ``cc<n>`` at the body is physical body child
    ``n``; inside a control every index is the position among the control's
    shown children; inside a table ``cc<k>`` is the table's physical child
    ``k`` and a bare ``r<n>`` its ``n``-th own row. Returns the non-empty
    paragraph texts the id's element holds. Covers every shape the fixture
    catalogue produces; any other shape fails loudly rather than passing.
    """
    steps = [(kind, int(number)) for kind, number in _CC_STEP_RE.findall(element_id)]
    assert "".join(f"{kind}{number}" for kind, number in steps) == element_id, element_id
    kind, index = steps[0]
    if kind == "cc":
        node = list(body)[index]
    elif kind == "t":
        node = [child for child in body if child.tag == qn("w:tbl")][index]
    else:
        pytest.fail(f"unexpected first step in {element_id}")
    for kind, index in steps[1:]:
        if node.tag in _BLOCK_WRAPPER_TAGS:
            node = _shown_children(node)[index]
            expected = {"cc": _BLOCK_WRAPPER_TAGS, "p": (qn("w:p"),), "t": (qn("w:tbl"),),
                        "r": (qn("w:tr"),)}[kind]
            assert node.tag in expected, (element_id, kind, node.tag)
        elif node.tag == qn("w:tbl"):
            if kind == "cc":
                node = list(node)[index]
                assert node.tag in _BLOCK_WRAPPER_TAGS, element_id
            elif kind == "r":
                node = [child for child in node if child.tag == qn("w:tr")][index]
            else:
                pytest.fail(f"unexpected step {kind} after a table in {element_id}")
        else:
            pytest.fail(f"step {kind}{index} not covered by this reader: {element_id}")
    if node.tag == qn("w:p"):
        return [_accept_all_paragraph_text(node).strip()]
    assert node.tag == qn("w:tr"), element_id
    return _row_texts(node)


class TestLegacyIdMeaning:
    @pytest.mark.parametrize("name", sorted(_FIXTURES))
    def test_every_pN_is_the_physical_body_child_it_names(self, name, tmp_path):
        path, spec = _extract(name, tmp_path)
        children = list(Document(path).element.body)
        for mapping in spec.paragraph_map:
            match = re.fullmatch(r"p(\d+)", mapping.element_id)
            if not match:
                continue
            element = children[int(match.group(1))]
            assert element.tag == qn("w:p"), mapping.element_id
            assert _reads_as(element, mapping.text), (mapping.element_id, mapping.text)

    @pytest.mark.parametrize("name", sorted(_FIXTURES))
    def test_the_applier_resolves_every_id_to_the_element_that_was_read(self, name, tmp_path):
        """The extractor mints ids and the applier resolves them; they must
        agree on every body paragraph and table row (nested rows included),
        or an edit lands in a different paragraph than the one reviewed.

        Every ``cc`` id (text inside a content control, plan WP-02) must name
        the element that was read — checked by an independent reading of the
        XML — and be one the applier classifies as a content control and
        refuses to resolve for writing: explicitly unsupported, never
        mistaken for a legacy id."""
        path, spec = _extract(name, tmp_path)
        document = Document(path)
        editor = DocumentEditor(document)
        checked = 0
        for mapping in spec.paragraph_map:
            kind = classify_element_id(mapping.element_id)
            if kind is ElementKind.BODY_PARAGRAPH:
                (p_el,) = editor.resolve_paragraphs(mapping.element_id, kind)
                assert _reads_as(p_el, mapping.text), (mapping.element_id, mapping.text)
                checked += 1
            elif kind is ElementKind.TABLE_ROW:
                paragraphs = editor.resolve_paragraphs(mapping.element_id, kind)
                texts = [
                    t for t in (_accept_all_paragraph_text(p).strip() for p in paragraphs) if t
                ]
                assert texts == _row_pieces(mapping.text), mapping.element_id
                checked += 1
            elif kind is ElementKind.CONTENT_CONTROL:
                assert mapping.container_type == "content_control", mapping.element_id
                texts = [t for t in _read_cc_id(document.element.body, mapping.element_id) if t]
                assert texts == _row_pieces(mapping.text), mapping.element_id
                with pytest.raises(EditError):
                    editor.resolve_paragraphs(mapping.element_id, kind)
                checked += 1
        assert checked == len(spec.paragraph_map)

    def test_a_block_control_still_occupies_its_body_index(self, tmp_path):
        """``p2`` is the paragraph *after* the control: body child 1 is the
        ``w:sdt``. Reading the control's text later must not renumber it."""
        path, spec = _extract("block_control", tmp_path)
        by_text = {m.text: m.element_id for m in spec.paragraph_map}
        assert by_text[fx.BEFORE_CONTROL] == "p0"
        assert by_text[fx.AFTER_CONTROL] == "p2"
        assert Document(path).element.body[1].tag == qn("w:sdt")

    def test_a_wrapped_table_does_not_renumber_ordinary_tables(self, tmp_path):
        """``tN`` counts direct body tables only — python-docx's
        ``Document.tables`` — so the ordinary table after a control that
        wraps a table is still ``t0``, for the extractor and the applier.
        The wrapped table's rows (read since S10) carry ``cc`` ids."""
        path, spec = _extract("block_control", tmp_path)
        rows = {m.element_id: m.text for m in spec.paragraph_map if m.element_type == "table_cell"}
        assert {k: v for k, v in rows.items() if k.startswith("t")} == {
            "t0r0": "Service | Pipe size",
            "t0r1": "Riser | Four inch",
        }
        assert {k for k in rows if not k.startswith("t")} == {"cc1t1r0", "cc1t1r1"}
        document = Document(path)
        assert len(document.tables) == 1
        assert document.tables[0].cell(1, 0).text == "Riser"

    def test_a_control_wrapping_rows_does_not_renumber_the_tables_own_rows(self, tmp_path):
        """``r<n>`` counts the table's own ``<w:tr>`` children, so rows a
        control wraps (a repeating section) get ``t0cc5r0`` and leave the
        others' numbers — and cells a control wraps join their own row."""
        _, spec = _extract("table_controls", tmp_path)
        assert [(m.element_id, m.text) for m in spec.paragraph_map] == [
            ("t0r0", "Service | Pipe size"),
            ("t0r1", f"Riser | {fx.CELL_CONTROL_TEXT}\n{fx.CELL_CONTROL_REVISED}"),
            ("t0r2", f"Branch | {fx.WRAPPED_CELL_TEXT}"),
            ("t0cc5r0", " | ".join(fx.WRAPPED_ROW_CELLS)),
        ]


class TestEstablishedTextKeepsItsLocation:
    """Everything the extractor reads today, with the id it reads it at.

    Exact for fixtures with no unsupported structure; a subset for the
    wrapper fixtures, where S10 *added* text (a paragraph holding an inline
    control reads differently now) but did not move or change what was
    already read.
    """

    @pytest.mark.parametrize("variant", _TYPED_VARIANTS, ids=lambda v: v.name)
    def test_typed_three_part_fixtures_are_read_exactly(self, variant, tmp_path):
        path = fx.save_docx(fx.build_blocks(variant.blocks), tmp_path, "spec.docx")
        spec = extract_text_from_docx(path)
        expected: list[tuple[str, str]] = []
        for index, block in enumerate(variant.blocks):
            if isinstance(block, fx.TableBlock):
                expected.extend((f"t0r{r}", " | ".join(row)) for r, row in enumerate(block.rows))
            else:
                expected.append((f"p{index}", block.text))
        assert [(m.element_id, m.text) for m in spec.paragraph_map] == expected

    def test_automatic_numbering_keeps_ids_and_literal_text(self, tmp_path):
        """Numbering-neutral on purpose. WP-03 lets S14 show the labels in the
        extracted text, so this pins only what must survive either way:
        every paragraph keeps its physical id, in order, and still ends with
        its literal source text (a displayed label may precede it)."""
        path = fx.save_docx(fx.build_auto_numbered_three_part(), tmp_path, "spec.docx")
        spec = extract_text_from_docx(path)
        read = [(m.element_id, m.text) for m in spec.paragraph_map if m.element_id.startswith("p")]
        blocks = fx.auto_numbered_blocks()
        assert [element_id for element_id, _ in read] == [f"p{i}" for i in range(len(blocks))]
        for (element_id, text), block in zip(read, blocks):
            assert text.endswith(block.text), (element_id, text)

    def test_tracked_changes_read_as_accept_all(self, tmp_path):
        _, spec = _extract("tracked_changes", tmp_path)
        assert [(m.element_id, m.text) for m in spec.paragraph_map] == [
            ("p0", "Comply with 2025 CBC."),
            ("p1", "Pipe shall be steel."),
            ("p2", "Flush all piping."),
            ("p3", "Test before concealment."),
            # p4 is a wholly deleted paragraph: empty on accept, no element.
            ("t0r0", "Copper | Type L"),
        ]
        assert spec.tracked_changes_detected is True

    def test_merged_and_nested_tables_read_once_each(self, tmp_path):
        _, spec = _extract("merged_and_nested_tables", tmp_path)
        assert [(m.element_id, m.text) for m in spec.paragraph_map] == [
            ("t0r0", "PIPING SCHEDULE"),
            ("t0r1", "Sprinkler | Black steel | Sch 40"),
            ("t0r2", "Galvanized | Sch 10"),
            ("t1r0", "Layout note"),
            ("t1r0c1t0r0", "Hanger | Rod size"),
            ("t1r0c1t0r1", "Clevis | 3/8 inch"),
        ]

    @pytest.mark.parametrize(
        "name,expected",
        [
            (
                "block_control",
                {
                    "p0": fx.BEFORE_CONTROL,
                    "p2": fx.AFTER_CONTROL,
                    "t0r0": "Service | Pipe size",
                    "t0r1": "Riser | Four inch",
                },
            ),
            ("fields", {"p1": "Coordinate with Section 21 13 13."}),
            ("hyperlinks", {"p0": "Verify each listing in the listing directory."}),
        ],
    )
    def test_supported_text_beside_wrappers_keeps_its_id(self, name, expected, tmp_path):
        _, spec = _extract(name, tmp_path)
        read = {m.element_id: m.text for m in spec.paragraph_map}
        for element_id, text in expected.items():
            assert read.get(element_id) == text, element_id

    @pytest.mark.parametrize(
        "name,ids",
        [("inline_controls", {"p0", "p1", "p2"}), ("smart_tag", {"p0"})],
    )
    def test_paragraphs_holding_inline_wrappers_keep_their_ids(self, name, ids, tmp_path):
        _, spec = _extract(name, tmp_path)
        assert ids <= {m.element_id for m in spec.paragraph_map}

    def test_field_instructions_are_never_read_as_prose(self, tmp_path):
        """A complex field's stored result is text; its instruction is code.
        (A simple field's stored result is read too since S10; its
        instruction is an attribute and never is.)"""
        _, spec = _extract("fields", tmp_path)
        assert fx.COMPLEX_FIELD_RESULT in spec.content
        assert "REF" not in spec.content
        assert "sec_211313" not in spec.content and "sec_230500" not in spec.content


# ---------------------------------------------------------------------------
# Group-versus-occurrence identity as it stands today
# ---------------------------------------------------------------------------


def _finding(file_name: str, *, element_id: str, **overrides) -> Finding:
    fields = dict(
        severity="HIGH",
        fileName=file_name,
        section="2.01",
        issue="Gate valves are specified where ball valves are required.",
        actionType="EDIT",
        existingText="gate valve",
        replacementText="ball valve",
        codeReference=None,
        evidenceElementId=element_id,
    )
    fields.update(overrides)
    return Finding(**fields)


def _three_file_group() -> list[Finding]:
    return [
        _finding("a.docx", element_id="p4"),
        _finding("b.docx", element_id="p7"),
        _finding("c.docx", element_id="p2"),
    ]


class TestGroupVersusOccurrenceIdentity:
    """Per-file original binding is pinned in ``test_dedup_edit_identity``
    and per-file sidecar fan-out in ``test_edit_sidecar``; these pin the
    *identity* those rest on. Deliberately not pinned here: representative
    choice, and the occurrence model S11 introduced (content-derived ids, one
    occurrence per place), which ``test_edit_occurrences`` pins. Since S12 the
    sidecar has one entry per occurrence (schema 6)."""

    def test_a_cross_file_group_has_one_occurrence_per_file(self):
        (merged,) = pipeline._deduplicate_findings(_three_file_group())
        (group,) = pipeline.group_findings([merged])
        assert sorted(group.file_names) == ["a.docx", "b.docx", "c.docx"]
        own_ids = {"a.docx": "p4", "b.docx": "p7", "c.docx": "p2"}
        for occurrence in group.occurrences:
            assert occurrence.has_original()
            assert occurrence.executable_finding().evidenceElementId == own_ids[occurrence.file_name]
        occurrence_ids = [o.occurrence_id for o in group.occurrences]
        assert len(occurrence_ids) == len(set(occurrence_ids))

    def test_group_identity_is_content_derived_and_order_free(self):
        members = _three_file_group()
        (forward,) = pipeline._deduplicate_findings(_three_file_group())
        (backward,) = pipeline._deduplicate_findings(list(reversed(_three_file_group())))
        assert forward.finding_id == backward.finding_id
        assert forward.finding_id == pipeline.compute_finding_id(members[0])
        assert all(pipeline.compute_finding_id(m) == forward.finding_id for m in members)
        assert forward.finding_id.startswith("rf-")

    def test_a_singleton_is_its_own_occurrence(self):
        (only,) = pipeline._deduplicate_findings([_finding("a.docx", element_id="p4")])
        assert only.finding_id == pipeline.compute_finding_id(only)
        (group,) = pipeline.group_findings([only])
        (occurrence,) = group.occurrences
        assert occurrence.executable_finding() is only

    def test_the_sidecar_key_module_and_occurrence_is_unique(self):
        merged = pipeline._deduplicate_findings(
            _three_file_group() + [_finding("a.docx", element_id="p9", existingText="globe valve")]
        )
        payload = edit_sidecar.build_edit_instructions(
            SimpleNamespace(review_result=ReviewResult(findings=merged), module_id="datacenter_fire")
        )
        keys = [(e["module_id"], e["occurrence_id"]) for e in payload["edits"]]
        assert len(keys) == 4
        assert len(keys) == len(set(keys))
        assert {module for module, _ in keys} == {"datacenter_fire"}
        by_file = {e["fileName"]: e for e in payload["edits"] if e["edit_proposal"]["existing_text"] == "gate valve"}
        assert {name: e["evidenceElementId"] for name, e in by_file.items()} == {
            "a.docx": "p4",
            "b.docx": "p7",
            "c.docx": "p2",
        }

    def test_origin_prefixes_keep_identical_content_apart(self):
        review = _finding("a.docx", element_id="p4")
        coordination = _finding("a.docx", element_id="p4")
        compliance = _finding("a.docx", element_id="p4")
        pipeline._deduplicate_findings([review])
        pipeline.assign_cross_check_finding_ids([coordination])
        pipeline.assign_compliance_finding_ids([compliance])
        ids = {review.finding_id, coordination.finding_id, compliance.finding_id}
        assert len(ids) == 3
        assert {i.split("-", 1)[0] for i in ids} == {"rf", "cf", "lc"}
        assert len({i.split("-", 1)[1] for i in ids}) == 1  # same content digest


# ---------------------------------------------------------------------------
# One extracted input reaches the primary and the repair request alike
# ---------------------------------------------------------------------------


class TestPrimaryAndRepairRequestsShareTheirInput:
    def test_the_repair_request_is_the_primary_plus_the_retry_suffix(self, monkeypatch, tmp_path):
        from src.batch import batch as batch_mod
        from src.review import review_request_builder as builder
        from src.review.prompt_serialization import render_spec_with_ids

        # The local cl100k count only decides the 300k-output path; a word
        # count keeps the test offline (the rank file is not downloadable).
        monkeypatch.setattr(builder, "count_tokens", lambda text: len(text.split()))
        submitted: list[list[dict]] = []

        def fake_create(client, batch_requests, *, use_beta, model):
            submitted.append(batch_requests)
            return SimpleNamespace(id=f"msgbatch_{len(submitted)}"), False

        monkeypatch.setattr(batch_mod, "_create_review_batch", fake_create)
        monkeypatch.setattr(batch_mod, "_get_client", lambda **_: SimpleNamespace())

        path = fx.save_docx(fx.build_clean_three_part(), tmp_path, "230500.docx")
        alerts = {
            "230500.docx": [
                {
                    "filename": "230500.docx",
                    "type": "Empty section",
                    "match": "PART 1 GENERAL",
                    "context": "PART 1 GENERAL",
                    "position": 0,
                    "deterministic_rule": "empty_section",
                }
            ]
        }
        primary_spec = extract_text_from_docx(path)
        batch_mod.submit_review_batch(
            [primary_spec], project_context="Project context.", cycle=CALIFORNIA_2025,
            pre_detected_alerts=alerts,
        )
        # A resumed collection re-extracts the spec from disk for its repair.
        repair_spec = extract_text_from_docx(path)
        batch_mod.submit_review_batch(
            [repair_spec], project_context="Project context.", cycle=CALIFORNIA_2025,
            retry_instruction=RETRY_TRUNCATED_REVIEW_INSTRUCTION, pre_detected_alerts=alerts,
        )

        (primary,), (repair,) = submitted
        p_params, r_params = primary["params"], repair["params"]
        p_message = p_params["messages"][0]["content"]
        r_message = r_params["messages"][0]["content"]
        assert r_message == p_message + "\n\n" + RETRY_TRUNCATED_REVIEW_INSTRUCTION
        assert {k: v for k, v in r_params.items() if k != "messages"} == {
            k: v for k, v in p_params.items() if k != "messages"
        }
        spec_block = render_spec_with_ids(
            primary_spec.content, primary_spec.paragraph_map, filename="230500.docx"
        )
        assert spec_block in p_message and spec_block in r_message
        assert primary["custom_id"] == repair["custom_id"]


# ---------------------------------------------------------------------------
# The applier boundary for every new container
# ---------------------------------------------------------------------------

_TEXT_SKIP = {
    qn("w:del"),
    qn("w:moveFrom"),
    qn("w:instrText"),
    qn("w:delText"),
    qn("w:pPr"),
    qn("w:rPr"),
    qn("w:sdtPr"),
}


def _visible_paragraph_text(p_el) -> str:
    """Accept-All text of a paragraph *including* every wrapper's runs.

    Independent of the extractor on purpose: it reads inside content
    controls, smart tags, simple fields, and hyperlinks by its own simple
    rule, so it can see damage the extractor's walk would miss.
    """
    parts: list[str] = []

    def walk(element) -> None:
        for child in element:
            tag = child.tag
            if not isinstance(tag, str) or tag in _TEXT_SKIP or tag == qn("w:p"):
                continue
            if tag == qn("w:t"):
                parts.append(child.text or "")
            elif tag == qn("w:tab"):
                parts.append("\t")
            elif tag in (qn("w:br"), qn("w:cr")):
                parts.append("\n")
            else:
                walk(child)

    walk(p_el)
    return "".join(parts)


def _visible_text(path: Path) -> str:
    body = Document(path).element.body
    return "\n".join(_visible_paragraph_text(p) for p in body.iter(qn("w:p")))


def _write_sidecar(directory: Path, file_name: str, entries: list[dict]) -> Path:
    payload = {
        "schema_version": 4,
        "generated_at": "2026-09-23T00:00:00Z",
        "report_file": "report.docx",
        "edit_count": len(entries),
        "edits": [
            {
                "finding_id": f"rf-{index:012d}",
                "fileName": file_name,
                "affected_files": [file_name],
                "has_per_file_original": True,
                "section": "",
                "severity": "HIGH",
                "issue": "Fixture edit.",
                "codeReference": None,
                "evidenceElementId": entry.get("element_id"),
                "verification_verdict": "CONFIRMED",
                "report_status": "VERIFIED_SUPPORTED",
                "edit_proposal": {
                    "action_type": "EDIT",
                    "existing_text": entry["existing"],
                    "replacement_text": entry["replacement"],
                    "anchor_text": None,
                    "insert_position": None,
                    "target_element_id": entry.get("element_id"),
                    "edit_confidence": 0.9,
                },
            }
            for index, entry in enumerate(entries)
        ],
    }
    path = directory / "report.edits.json"
    path.write_text(json.dumps(payload), encoding="utf-8")
    return path


#: Containers whose *identity* an edit must never change: the link target,
#: the field instruction, the control, the smart tag. Tracked-change markup
#: (``w:ins`` / ``w:del``) is deliberately absent; an edit adds it.
_WRAPPER_TAGS = (qn("w:sdt"), qn("w:fldSimple"), qn("w:hyperlink"), qn("w:smartTag"))
_R_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"


def _wrapper_key(element, rels) -> tuple:
    """A wrapper's identity: what would be lost if it were replaced by a run."""
    tag = element.tag
    if tag == qn("w:sdt"):
        properties = element.find(qn("w:sdtPr"))

        def value(child: str) -> str | None:
            found = properties.find(qn(child)) if properties is not None else None
            return found.get(qn("w:val")) if found is not None else None

        return ("sdt", value("w:tag"), value("w:id"))
    if tag == qn("w:fldSimple"):
        return ("fldSimple", element.get(qn("w:instr")))
    if tag == qn("w:hyperlink"):
        rel = rels.get(element.get(_R_ID))
        return ("hyperlink", getattr(rel, "target_ref", None), element.get(qn("w:anchor")))
    return ("smartTag", element.get(qn("w:uri")), element.get(qn("w:element")))


def _wrapper_signature(path: Path) -> list[tuple]:
    """Every wrapper and complex-field element in the body, in order, by identity."""
    document = Document(path)
    rels = document.part.rels
    body = document.element.body
    signature = [_wrapper_key(el, rels) for el in body.iter(*_WRAPPER_TAGS)]
    signature += [("fldChar", el.get(qn("w:fldCharType"))) for el in body.iter(qn("w:fldChar"))]
    signature += [("instrText", el.text) for el in body.iter(qn("w:instrText"))]
    return signature


def _visible_text_of(element) -> str:
    """Accept-All text under a wrapper: a block wrapper holds paragraphs, an
    inline one holds runs."""
    paragraphs = list(element.iter(qn("w:p")))
    if paragraphs:
        return "\n".join(_visible_paragraph_text(p) for p in paragraphs)
    return _visible_paragraph_text(element)


def _innermost_wrapper_holding(path: Path, needle: str) -> tuple | None:
    """Identity of the innermost wrapper whose visible text contains
    ``needle``; ``None`` when the text sits in no wrapper at all."""
    document = Document(path)
    rels = document.part.rels
    best, best_depth = None, -1
    for element in document.element.body.iter(*_WRAPPER_TAGS):
        if needle in _visible_text_of(element):
            depth = sum(1 for _ in element.iterancestors())
            if depth > best_depth:
                best, best_depth = _wrapper_key(element, rels), depth
    return best


def _assert_edit_is_exact(source: Path, output: Path, existing: str, replacement: str) -> None:
    """An applied edit changed its target text and nothing else.

    Comparing visible text alone is not enough: an editor could replace a
    hyperlink, field, control, or smart tag with a plain run holding the
    same words, and the document would silently lose the link, the field,
    or the control. So the wrappers must survive with the same identity,
    and the new text must sit inside the wrapper that held the old text.
    """
    assert _visible_text(output) == _visible_text(source).replace(existing, replacement, 1)
    assert _wrapper_signature(output) == _wrapper_signature(source), (
        "a hyperlink, field, content control, or smart tag was added, removed, or rewritten"
    )
    assert _innermost_wrapper_holding(output, replacement) == _innermost_wrapper_holding(
        source, existing
    ), "the replacement is not inside the wrapper that held the original text"


def _apply_one(tmp_path: Path, name: str, *, existing: str, replacement: str,
               element_id: str | None, allow_tracked_source: bool = False) -> SimpleNamespace:
    source_dir = tmp_path / "source"
    file_name = f"{name}.docx"
    source = fx.save_docx(_FIXTURES[name](), source_dir, file_name)
    before_bytes = source.read_bytes()
    assert _visible_text(source).count(existing) == 1, "the target must be unique in the fixture"
    sidecar = load_sidecar(
        _write_sidecar(
            tmp_path,
            file_name,
            [{"existing": existing, "replacement": replacement, "element_id": element_id}],
        )
    )
    out_dir = tmp_path / "out"
    (result,) = apply_sidecar(
        sidecar,
        [source],
        RunSettings(output_dir=out_dir, allow_tracked_source=allow_tracked_source),
    )
    assert source.read_bytes() == before_bytes, "the applier must never write the source"
    (outcome,) = result.outcomes
    return SimpleNamespace(
        source=source, outcome=outcome, result=result, out_dir=out_dir,
        existing=existing, replacement=replacement,
    )


def _assert_applied_exactly_or_refused(run) -> None:
    if run.outcome.status is OutcomeStatus.APPLIED:
        _assert_edit_is_exact(
            run.source, Path(run.result.output_path), run.existing, run.replacement
        )
    else:
        assert run.outcome.status is OutcomeStatus.UNLOCATED, run.outcome.status
        assert run.outcome.reason
        assert not list(run.out_dir.glob("*.docx")), "a refused edit must write nothing"


class TestApplierBoundaryForWrappedContent:
    """Aimed at text inside a wrapper, an edit is applied exactly or refused.

    Since S10 the extractor reads all of this text, and the writer still
    refuses every one of these edits (plan WP-02: readable is not writable;
    the specific reasons are pinned in ``test_applier_wrapped_content.py``).
    Either way nothing outside the target may change, and no wrapper may be
    lost or rewritten.
    """

    @pytest.mark.parametrize(
        "name,existing,replacement,element_id,allow_tracked",
        [
            ("block_control", "listed backflow preventer", "listed reduced-pressure assembly", None, False),
            ("inline_controls", fx.INLINE_CONTROL_TEXT, "schedule 10 galvanized steel", "p0", False),
            ("inline_controls", fx.DROPDOWN_CHOSEN, "Steel Schedule 40", "p1", False),
            ("fields", fx.REF_FIELD_RESULT, "23 05 29", "p0", False),
            ("smart_tag", fx.SMART_TAG_TEXT, "City of Fremont", "p0", False),
            ("hyperlinks", fx.HYPERLINK_TEXT, "the manufacturer's listing", "p0", True),
            ("hyperlinks", "manufacturer data sheets", "manufacturer product data", "p1", True),
        ],
        ids=[
            "block_control",
            "inline_control",
            "dropdown",
            "simple_field",
            "smart_tag",
            "hyperlink",
            "hyperlink_with_insertion",
        ],
    )
    def test_an_edit_inside_a_wrapper_is_exact_or_refused(
        self, tmp_path, name, existing, replacement, element_id, allow_tracked
    ):
        _assert_applied_exactly_or_refused(
            _apply_one(
                tmp_path, name, existing=existing, replacement=replacement,
                element_id=element_id, allow_tracked_source=allow_tracked,
            )
        )

    def test_ordinary_text_beside_a_block_control_is_still_editable(self, tmp_path):
        """The positive control for the check above: a direct-run edit next
        to a wrapper applies, and changes exactly its target."""
        run = _apply_one(
            tmp_path, "block_control", existing="fire protection piping",
            replacement="fire sprinkler piping", element_id="p0",
        )
        assert run.outcome.status is OutcomeStatus.APPLIED
        _assert_applied_exactly_or_refused(run)

    def test_an_edit_beside_a_complex_field_leaves_the_field_intact(self, tmp_path):
        run = _apply_one(
            tmp_path, "fields", existing="Coordinate with", replacement="Coordinate work with",
            element_id="p1",
        )
        assert run.outcome.status is OutcomeStatus.APPLIED
        _assert_applied_exactly_or_refused(run)
        field_paragraph = Document(run.result.output_path).element.body.findall(qn("w:p"))[1]
        kinds = [el.get(qn("w:fldCharType")) for el in field_paragraph.iter(qn("w:fldChar"))]
        assert kinds == ["begin", "separate", "end"]
        (instruction,) = list(field_paragraph.iter(qn("w:instrText")))
        assert instruction.text == fx.COMPLEX_FIELD_INSTRUCTION


class TestTheExactnessCheckItself:
    """The applied branch above is unreachable for wrapped text (the writer
    refuses to edit through a wrapper), so the check is proven here against
    hand-made "edited" documents: it must accept an edit made inside the
    wrapper and reject one that destroys it."""

    _REPLACEMENT = "the manufacturer's listing"

    def _edited_hyperlink(self, tmp_path: Path, how: str) -> tuple[Path, Path]:
        source = fx.save_docx(fx.build_hyperlink_spec(), tmp_path / "source", "hyperlinks.docx")
        document = Document(source)
        p0 = document.element.body.findall(qn("w:p"))[0]
        link = p0.find(qn("w:hyperlink"))
        (link_text,) = list(link.iter(qn("w:t")))
        if how == "inside_the_link":
            link_text.text = self._REPLACEMENT
        else:
            plain = parse_xml(
                f'<w:r {nsdecls("w")}><w:t xml:space="preserve">{self._REPLACEMENT}</w:t></w:r>'
            )
            link.addnext(plain)
            if how == "link_removed":
                p0.remove(link)
            else:  # "text_moved_out": the link survives, emptied
                link_text.text = ""
        output = tmp_path / "out" / "hyperlinks.applied.docx"
        output.parent.mkdir(parents=True, exist_ok=True)
        document.save(output)
        return source, output

    def test_an_edit_inside_the_kept_link_is_exact(self, tmp_path):
        source, output = self._edited_hyperlink(tmp_path, "inside_the_link")
        _assert_edit_is_exact(source, output, fx.HYPERLINK_TEXT, self._REPLACEMENT)

    @pytest.mark.parametrize("how", ["link_removed", "text_moved_out"])
    def test_an_edit_that_breaks_the_link_is_not_exact(self, tmp_path, how):
        source, output = self._edited_hyperlink(tmp_path, how)
        # Same visible text either way: only the structure checks can see it.
        assert _visible_text(output) == _visible_text(source).replace(
            fx.HYPERLINK_TEXT, self._REPLACEMENT, 1
        )
        with pytest.raises(AssertionError):
            _assert_edit_is_exact(source, output, fx.HYPERLINK_TEXT, self._REPLACEMENT)
