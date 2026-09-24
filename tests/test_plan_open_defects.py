"""The implementation plan's open defects, as strict expected failures.

Chunk S01 (plan WP-01) converted ``plans/check_plan_status.py`` — the
review's offline reproductions — into this module and deleted the script.
Together with ``plans/PROGRESS.md`` it is the plan's progress tracker, and
unlike the script it is enforced by CI.

How to read a result
--------------------
Every reproduction carries a literal marker of this shape::

    @pytest.mark.xfail(strict=True, raises=AssertionError,
                       reason="open: fixed by S0N (WP-xx)")

* **xfailed** — the defect is still present (an assertion about the fixed
  behavior failed). This is the expected state until chunk S0N lands.
* **FAILED ... [XPASS(strict)] open: fixed by S0N** — the code now behaves
  as the plan requires. The session doing chunk S0N removes the marker in
  the same pull request, turning the reproduction into a regression test.
* **FAILED with any other exception** — the *check* is broken (a function it
  uses was renamed or reshaped). ``raises=AssertionError`` is what keeps
  that from reading as "still open": only a failed assertion counts as the
  defect, never an ``AttributeError``. The same goes for the probes'
  preconditions, which fail through ``pytest.fail`` — a probe that no longer
  reaches the behavior it tests must say so, not pass for "still broken".

Controls
--------
A fix that makes something *disappear* (an alert, a merge, a cache hit)
could pass by breaking the detector, the dedup, or the cache outright. Each
such reproduction therefore has an ordinary test beside it — a control —
that passes today and must keep passing: the truly empty article is still
flagged, identical findings still merge, a grounded CONFIRMED still caches.
Controls are separate tests, not assertions inside the xfail, because an
assertion failing inside a strict xfail would read as "still open".

The reproductions that only inspected source text ("source hints") are
behavioral here: they drive the real retry loop, the exact JavaScript the
HTML exporter ships, and the real GUI controllers. The one exception is the
Haiku cache-minimum check (WP-17), which is about documentation and has no
behavior to exercise; it parses the specific claims instead of searching
loosely, and requires them to exist.

Everything is hermetic: no API key, no network, temporary files only.
"""
from __future__ import annotations

import html
import json
import os
import re
import shutil
import subprocess
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace

import pytest

from src.core.code_cycles import CALIFORNIA_2025
from src.input.extractor import ExtractedSpec, extract_text_from_docx
from src.input.preprocessor import (
    detect_duplicate_headings,
    detect_empty_sections,
    detect_inconsistent_file_naming,
    detect_placeholders,
    detect_stale_code_cycle_references,
    preprocess_spec,
)
from src.review.reviewer import Finding
from tests.fixtures import spec_docx as fx

#: A credential-shaped string that is obviously not a real key.
FAKE_KEY = "sk-ant-api03-FIXTURE-not-a-real-key-0123456789"

_REPO_ROOT = Path(__file__).resolve().parents[1]


def _finding(issue: str = "Wrong valve type.", **overrides) -> Finding:
    fields = dict(
        severity="HIGH",
        fileName="210500.docx",
        section="2.01",
        issue=issue,
        actionType="REPORT_ONLY",
        existingText=None,
        replacementText=None,
        codeReference=None,
    )
    fields.update(overrides)
    return Finding(**fields)


def _paragraphs(*items: str) -> str:
    return "\n\n".join(items)


def _extract(builder, tmp_path: Path, name: str = "spec.docx") -> ExtractedSpec:
    return extract_text_from_docx(fx.save_docx(builder, tmp_path, name))


# ===========================================================================
# WP-04 — deterministic structure and text detectors (chunk S03)
# ===========================================================================


def _structural(spec: ExtractedSpec) -> list[tuple[str, str]]:
    result = preprocess_spec(spec.content, spec.filename, cycle=CALIFORNIA_2025)
    return [(a["deterministic_rule"], a["match"]) for a in result.structural_alerts]


def _variant(name: str) -> fx.SpecVariant:
    (variant,) = [v for v in fx.three_part_variants() if v.name == name]
    return variant


class TestHeadingStructure:
    """Fixed by S03 (WP-04A); kept as the regression tests. Headings were
    read flat, so a PART followed by its first article looked "empty", and
    integer-led prose read as a heading. Focused cases live in
    ``test_heading_structure.py``."""

    @pytest.mark.parametrize("name", ["clean", "table_only_article"])
    def test_clean_three_part_spec_has_no_structural_alerts(self, name, tmp_path):
        spec = _extract(fx.build_blocks(_variant(name).blocks), tmp_path)
        assert _structural(spec) == []

    @pytest.mark.parametrize("name", ["empty_article", "duplicate_heading"])
    def test_a_mutation_produces_only_its_own_alert(self, name, tmp_path):
        variant = _variant(name)
        spec = _extract(fx.build_blocks(variant.blocks), tmp_path)
        assert _structural(spec) == [(variant.expected_rule, variant.expected_match)]

    @pytest.mark.parametrize(
        "line",
        [
            "2 coats of primer shall be applied.",
            "12 inches minimum clearance shall be maintained.",
            "1.5 inches minimum cover is required.",
            "1 year from Substantial Completion.",
        ],
    )
    def test_quantity_lines_are_not_headings(self, line):
        empty = detect_empty_sections(
            _paragraphs("3.01 PAINTING", line, "3.02 CLEANING", "A. Clean."), "q.docx"
        )
        duplicates = detect_duplicate_headings(
            _paragraphs("1.01 SUMMARY", line, "A. Text.", line, "B. Text."), "q.docx"
        )
        assert [a["match"] for a in empty] == []
        assert [a["match"] for a in duplicates] == []

    # -- controls ------------------------------------------------------------

    def test_control_a_truly_empty_article_is_flagged(self, tmp_path):
        spec = _extract(fx.build_empty_article_mutation(), tmp_path)
        assert ("empty_section", "1.02 SUBMITTALS") in _structural(spec)
        text_only = detect_empty_sections(
            _paragraphs("1.01 SUMMARY", "1.02 SUBMITTALS", "A. Submit product data."), "c.docx"
        )
        assert any(a["match"].startswith("1.01") for a in text_only)

    def test_control_a_truly_duplicated_heading_is_flagged(self, tmp_path):
        spec = _extract(fx.build_duplicate_heading_mutation(), tmp_path)
        assert ("duplicate_heading", "1.01 SUMMARY") in _structural(spec)
        text_only = detect_duplicate_headings(
            _paragraphs("1.01 SUMMARY", "A. Text.", "1.01 SUMMARY", "B. Text."), "c.docx"
        )
        assert text_only

    def test_control_the_clean_spec_is_clean_for_every_text_detector(self, tmp_path):
        spec = _extract(fx.build_clean_three_part(), tmp_path)
        result = preprocess_spec(spec.content, spec.filename, cycle=CALIFORNIA_2025)
        for category in (
            "leed_alerts",
            "placeholder_alerts",
            "code_cycle_alerts",
            "template_marker_alerts",
            "invalid_code_cycle_alerts",
            "duplicate_paragraph_alerts",
            "polity_alerts",
        ):
            assert getattr(result, category) == [], category

    def test_control_the_auto_numbered_spec_raises_no_structural_alert(self, tmp_path):
        """Vacuous today — its numbers are not read, so no headings are seen
        (WP-03). Once S14 shows the numbers to the detectors, this is the
        test that proves the same clean structure stays clean."""
        spec = _extract(fx.build_auto_numbered_three_part(), tmp_path)
        assert _structural(spec) == []


def _stale(sentence: str) -> int:
    return len(detect_stale_code_cycle_references(sentence, "s.docx", CALIFORNIA_2025))


class TestStaleCitationSuppression:
    """Fixed by S03 (WP-04B); kept as the regression tests. Nearby words such
    as "prior", "historical", and "may not" suppressed an active citation of
    a stale code year. Focused cases live in ``test_preprocessor_policy.py``."""

    @pytest.mark.parametrize(
        "sentence",
        [
            "Submit shop drawings prior to fabrication in accordance with 2022 CBC Section 1704.",
            "Coordinate with the historical society and comply with 2022 CBC.",
            "Contractor may not deviate from 2022 CBC Chapter 17.",
            "Contractor shall not deviate from 2022 CBC Chapter 17.",
            "Work cannot depart from 2022 CBC requirements.",
        ],
    )
    def test_an_active_stale_citation_is_flagged(self, sentence):
        assert _stale(sentence) > 0

    def test_control_a_genuinely_historical_reference_stays_suppressed(self):
        assert _stale("Previously, the 2022 CBC applied to this work.") == 0

    def test_control_a_plain_active_stale_citation_is_flagged(self):
        assert _stale("Comply with 2022 CBC Section 1704.") == 1


class TestCitationSyntax:
    """Fixed by S03 (WP-04C); kept as the regression tests. ASCE 7 editions
    written as ASCE/SEI, with the word Standard, with a Unicode dash, or with
    a four-digit year were not recognized. Focused cases live in
    ``test_asce7_stale_editions.py``."""

    @pytest.mark.parametrize(
        "token",
        [
            "ASCE/SEI 7-16",
            "ASCE 7–16",  # en dash
            "ASCE 7—16",  # em dash
            "ASCE 7-2016",
            "ASCE Standard 7-16",
        ],
    )
    def test_a_stale_asce_7_edition_is_recognized(self, token):
        assert _stale(f"Design loads per {token}.") > 0

    def test_control_the_plain_form_is_flagged(self):
        assert _stale("Design loads per ASCE 7-16.") == 1

    @pytest.mark.parametrize("token", ["ASCE 7-22", "ASCE 7-42"])
    def test_control_current_and_implausible_editions_are_not(self, token):
        # 7-22 is the cycle's own edition; 7-42 is not a real edition.
        assert _stale(f"Design loads per {token}.") == 0


class TestPlaceholders:
    """Fixed by S03 (WP-04D); kept as the regression tests. Bare TBD was
    missed, and the bracket patterns had no word boundary, so [EDITION ...]
    read as an EDIT placeholder. Focused cases live in
    ``test_deterministic_checks.py``."""

    def test_bare_tbd_is_detected_once(self):
        assert len(detect_placeholders("Pipe size: TBD by engineer.", "p.docx")) == 1

    def test_edition_and_selected_are_not_edit_or_select_placeholders(self):
        alerts = detect_placeholders("See [EDITION 2024] and [SELECTED ITEMS].", "p.docx")
        assert [a["type"] for a in alerts] == []

    def test_control_a_real_select_placeholder_is_flagged(self):
        assert detect_placeholders("Provide [SELECT ONE] finish.", "p.docx")

    def test_control_bracketed_tbd_is_counted_once(self):
        assert len(detect_placeholders("Pipe size: [TBD].", "p.docx")) == 1

    def test_control_a_tbd_product_identifier_stays_clean(self):
        assert detect_placeholders("Provide model TBDF-200 fitting.", "p.docx") == []


class TestFileNaming:
    """Fixed by S03 (WP-04E); kept as the regression tests. Compact and
    SECTION-prefixed names were "unrecognized", and an unrecognized majority
    silenced the mixture notice. Focused cases live in
    ``test_deterministic_checks.py``."""

    def test_a_mix_of_naming_styles_is_reported(self):
        # One separated, one compact, one SECTION-prefixed (upper-case
        # extension) name — each in spec_docx.FILENAME_EXAMPLES.
        names = ["21 05 00.docx", "211313.docx", "SECTION 21 13 16.DOCX"]
        known = {example.name for example in fx.FILENAME_EXAMPLES}
        if not set(names) <= known:
            pytest.fail("precondition: these names are no longer fixture examples")
        assert detect_inconsistent_file_naming(names)

    def test_control_uniform_names_raise_no_notice(self):
        assert detect_inconsistent_file_naming(
            ["21 05 00.docx", "21 13 13.docx", "21 13 16.docx"]
        ) == []

    def test_control_a_recognized_minority_is_still_reported(self):
        alerts = detect_inconsistent_file_naming(
            ["21 05 00.docx", "21 13 13.docx", "21-13-16.docx"]
        )
        assert [a["filename"] for a in alerts] == ["21-13-16.docx"]


# ===========================================================================
# WP-05 — route from the document's own SECTION heading (chunk S13)
# ===========================================================================


def _assignment_for(tmp_path: Path, file_name: str, number: str, title: str):
    from src.programs import assignments_for_specs, get_program

    blocks = fx.with_section_heading(fx.clean_three_part_blocks(), number, title)
    path = fx.save_docx(fx.build_blocks(blocks), tmp_path, file_name)
    spec = extract_text_from_docx(path)
    (assignment,) = assignments_for_specs(
        [spec], [path], program=get_program("hyperscale_datacenter")
    )
    return assignment


class TestSectionHeadingRouting:
    """Converted at the assignment seam (extracted spec + source path ->
    routing), not at ``route_spec``: S13 carries the heading on the
    extracted spec, so a check that bypassed extraction could stay red after
    a correct fix."""

    @pytest.mark.xfail(strict=True, raises=AssertionError, reason="open: fixed by S13 (WP-05)")
    @pytest.mark.parametrize(
        "file_name,number,title",
        [
            ("210500.docx", "21 05 00", "COMMON WORK RESULTS FOR FIRE SUPPRESSION"),
            ("211313.docx", "21 13 13", "WET-PIPE SPRINKLER SYSTEMS"),
        ],
    )
    def test_a_compact_name_with_its_section_heading_routes_to_fire(
        self, tmp_path, file_name, number, title
    ):
        decision = _assignment_for(tmp_path, file_name, number, title).decision
        assert decision.automatic_state.value == "supported"
        assert decision.automatic_module_ids == ("datacenter_fire",)

    def test_control_a_separated_name_routes_to_fire(self, tmp_path):
        decision = _assignment_for(
            tmp_path,
            "21 05 00 - Common Work Results for Fire Suppression.docx",
            "21 05 00",
            "COMMON WORK RESULTS FOR FIRE SUPPRESSION",
        ).decision
        assert decision.automatic_state.value == "supported"
        assert decision.automatic_module_ids == ("datacenter_fire",)


# ===========================================================================
# WP-06A — distinct findings must not merge (chunk S02)
# ===========================================================================


class TestFindingNormalization:
    """Fixed by S02 (WP-06A); kept as the regression test. Plan Appendix A
    puts 210500.docx in the known corpus, so both the corpus and the
    no-corpus contexts are exercised. Focused cases live in
    ``test_finding_identity_normalization.py``."""

    @pytest.mark.parametrize("corpus", [(), ("210500.docx",)], ids=["no-corpus", "known-corpus"])
    def test_copper_and_pvc_findings_stay_distinct(self, corpus):
        from src.orchestration import pipeline

        context = pipeline.FindingIdentityContext.from_filenames(corpus)
        copper = _finding("Section 21 05 00 requires copper pipe in 210500.docx.")
        pvc = _finding("Section 21 05 00 requires PVC pipe in 210500.docx.")
        assert len(pipeline._deduplicate_findings([copper, pvc], context=context)) == 2

    def test_control_identical_findings_still_merge(self):
        from src.orchestration import pipeline

        twins = [_finding("Exact same issue."), _finding("Exact same issue.")]
        assert len(pipeline._deduplicate_findings(twins)) == 1


# ===========================================================================
# WP-06B — every edit location survives to the sidecar (chunk S12)
# ===========================================================================


def _sidecar_entries(findings) -> list:
    from src.orchestration import pipeline
    from src.output import edit_sidecar
    from src.review.reviewer import ReviewResult

    merged = pipeline._deduplicate_findings(findings)
    payload = edit_sidecar.build_edit_instructions(
        SimpleNamespace(review_result=ReviewResult(findings=merged), module_id="datacenter_fire")
    )
    return list(payload.get("edits") or [])


_SAME_EDIT = dict(actionType="EDIT", existingText="gate valve", replacementText="ball valve")


class TestRepeatedLocations:
    @pytest.mark.xfail(strict=True, raises=AssertionError, reason="open: fixed by S12 (WP-06B)")
    def test_the_same_edit_at_p4_and_p8_gives_two_entries(self):
        entries = _sidecar_entries(
            [
                _finding(evidenceElementId="p4", **_SAME_EDIT),
                _finding(evidenceElementId="p8", **_SAME_EDIT),
            ]
        )
        targets = sorted(
            e.get("evidenceElementId") or (e.get("edit_proposal") or {}).get("target_element_id")
            for e in entries
        )
        assert targets == ["p4", "p8"]

    def test_control_a_duplicate_emission_at_p4_gives_one_entry(self):
        entries = _sidecar_entries(
            [
                _finding(evidenceElementId="p4", **_SAME_EDIT),
                _finding(evidenceElementId="p4", **_SAME_EDIT),
            ]
        )
        assert len(entries) == 1


# ===========================================================================
# WP-07 — the applier refuses ambiguous file bindings (chunk S02)
# ===========================================================================


def _two_same_named_specs(tmp_path: Path) -> tuple[Path, Path]:
    first = fx.save_docx(fx.build_clean_three_part(), tmp_path / "projA", "spec.docx")
    second = fx.save_docx(fx.build_clean_three_part(), tmp_path / "projB", "spec.docx")
    return first, second


def _spec_edit_sidecar(tmp_path: Path):
    from applier.sidecar import load_sidecar

    payload = {
        "schema_version": 4,
        "generated_at": "2026-09-23T00:00:00Z",
        "report_file": "report.docx",
        "edit_count": 1,
        "edits": [
            {
                "finding_id": "rf-000000000001",
                "fileName": "spec.docx",
                "affected_files": ["spec.docx"],
                "has_per_file_original": True,
                "section": "1.02 SUBMITTALS",
                "severity": "HIGH",
                "issue": "Submittal timing.",
                "codeReference": None,
                "evidenceElementId": "p4",
                "verification_verdict": "CONFIRMED",
                "report_status": "VERIFIED_SUPPORTED",
                "edit_proposal": {
                    "action_type": "EDIT",
                    "existing_text": "before fabrication",
                    "replacement_text": "before ordering materials",
                    "anchor_text": None,
                    "insert_position": None,
                    "target_element_id": "p4",
                    "edit_confidence": 0.9,
                },
            }
        ],
    }
    path = tmp_path / "report.edits.json"
    path.write_text(json.dumps(payload), encoding="utf-8")
    return load_sidecar(path)


class TestAmbiguousFileBindings:
    """Fixed by S02 (WP-07); kept as the regression tests. Focused cases
    (destinations, receipt, exit status, assist, dry run) live in
    ``test_applier_bindings.py``."""

    def test_a_name_maps_to_every_distinct_input(self, tmp_path):
        from applier import run as applier_run

        first, second = _two_same_named_specs(tmp_path)
        forward = applier_run._index_specs([first, second])
        backward = applier_run._index_specs([second, first])
        for index in (forward, backward):
            bound = index.get("spec.docx")
            assert not isinstance(bound, Path), f"first input wins: {bound}"
            assert {Path(p).resolve() for p in bound} == {first.resolve(), second.resolve()}

    @pytest.mark.parametrize("order", ["A_then_B", "B_then_A"])
    def test_two_same_named_inputs_are_not_edited(self, tmp_path, order):
        from applier.models import OutcomeStatus
        from applier.run import RunSettings, apply_sidecar

        first, second = _two_same_named_specs(tmp_path)
        inputs = [first, second] if order == "A_then_B" else [second, first]
        results = apply_sidecar(_spec_edit_sidecar(tmp_path), inputs, RunSettings())
        outcomes = [o for r in results for o in r.outcomes]
        if len(outcomes) != 1:
            pytest.fail(f"precondition: expected one accounted instruction, got {outcomes}")
        written = sorted(str(p.relative_to(tmp_path)) for p in tmp_path.rglob("*.applied.docx"))
        assert written == [], f"edited a guessed file: {written}"
        assert outcomes[0].status not in (OutcomeStatus.APPLIED, OutcomeStatus.WOULD_APPLY)
        assert outcomes[0].reason

    def test_control_repeating_one_input_is_not_ambiguous(self, tmp_path):
        from applier.models import OutcomeStatus
        from applier.run import RunSettings, apply_sidecar

        first, _ = _two_same_named_specs(tmp_path)
        results = apply_sidecar(_spec_edit_sidecar(tmp_path), [first, first], RunSettings())
        (outcome,) = [o for r in results for o in r.outcomes]
        assert outcome.status is OutcomeStatus.APPLIED
        assert (first.parent / "spec.applied.docx").exists()


# ===========================================================================
# WP-10 — uncertainty is not a reusable verdict (chunk S05)
# ===========================================================================


class TestVerificationCacheAndCitations:
    @pytest.mark.xfail(strict=True, raises=AssertionError, reason="open: fixed by S05 (WP-10)")
    def test_a_grounded_unverified_is_not_replayed_from_the_cache(self):
        from src.verification.verification_cache import VerificationCache
        from src.verification.verifier import VerificationResult

        cache = VerificationCache()
        uncertain = _finding("ASME B31.9 requires X.", codeReference="ASME B31.9")
        cache.put(
            uncertain,
            cycle=CALIFORNIA_2025,
            result=VerificationResult(
                verdict="UNVERIFIED",
                explanation="could not settle",
                grounded=True,
                searched_sources=["https://example.org/a"],
                successful_source_count=1,
            ),
        )
        assert cache.get(uncertain, cycle=CALIFORNIA_2025) is None

    def test_control_a_grounded_confirmed_with_a_quote_is_cached(self):
        from src.verification.verification_cache import VerificationCache
        from src.verification.verifier import VerificationResult

        source = "https://example.org/a"
        cache = VerificationCache()
        settled = _finding("ASME B31.9 requires Y.", codeReference="ASME B31.9")
        cache.put(
            settled,
            cycle=CALIFORNIA_2025,
            result=VerificationResult(
                verdict="CONFIRMED",
                explanation="settled",
                grounded=True,
                sources=[source],
                searched_sources=[source],
                accepted_sources=[source],
                source_quote="the requirement text",
                successful_source_count=1,
            ),
        )
        assert cache.get(settled, cycle=CALIFORNIA_2025) is not None

    @pytest.mark.xfail(strict=True, raises=AssertionError, reason="open: fixed by S05 (WP-10)")
    @pytest.mark.parametrize("blank", ["", "   "], ids=["empty", "whitespace"])
    def test_a_blank_source_is_not_a_citation(self, blank):
        from src.output.report_status import ReportStatus, classify_status
        from src.verification.verifier import VerificationResult

        finding = _finding()
        finding.verification = VerificationResult(
            verdict="DISPUTED", grounded=True, sources=[blank], accepted_sources=[blank]
        )
        assert classify_status(finding) != ReportStatus.DISPUTED

    def test_control_a_disputed_with_a_real_source_is_disputed(self):
        from src.output.report_status import ReportStatus, classify_status
        from src.verification.verifier import VerificationResult

        finding = _finding()
        finding.verification = VerificationResult(
            verdict="DISPUTED",
            grounded=True,
            sources=["https://example.org/a"],
            accepted_sources=["https://example.org/a"],
        )
        assert classify_status(finding) == ReportStatus.DISPUTED


# ===========================================================================
# WP-02 — content controls, field results, smart tags (chunk S10)
# ===========================================================================


class TestWrappedWordContent:
    @pytest.mark.xfail(strict=True, raises=AssertionError, reason="open: fixed by S10 (WP-02)")
    def test_a_block_control_paragraph_is_extracted(self, tmp_path):
        spec = _extract(fx.build_block_control_spec(), tmp_path)
        assert fx.BLOCK_CONTROL_PARAGRAPH in spec.content

    @pytest.mark.xfail(strict=True, raises=AssertionError, reason="open: fixed by S10 (WP-02)")
    def test_a_block_control_yields_its_paragraph_and_table_in_order(self, tmp_path):
        content = _extract(fx.build_block_control_spec(), tmp_path).content
        order = [
            fx.BEFORE_CONTROL,
            fx.BLOCK_CONTROL_PARAGRAPH,
            "Backflow preventer",
            "Listed double check",
            fx.AFTER_CONTROL,
        ]
        positions = [content.find(text) for text in order]
        assert -1 not in positions and positions == sorted(positions)

    @pytest.mark.xfail(strict=True, raises=AssertionError, reason="open: fixed by S10 (WP-02)")
    def test_a_simple_fields_stored_result_is_extracted(self, tmp_path):
        spec = _extract(fx.build_fields_spec(), tmp_path)
        assert "Refer to Section 23 05 00 for common work results." in spec.content

    @pytest.mark.xfail(strict=True, raises=AssertionError, reason="open: fixed by S10 (WP-02)")
    def test_inline_control_text_is_extracted_in_place(self, tmp_path):
        spec = _extract(fx.build_inline_controls_spec(), tmp_path)
        assert "Provide schedule 40 black steel pipe for sprinkler mains." in spec.content

    @pytest.mark.xfail(strict=True, raises=AssertionError, reason="open: fixed by S10 (WP-02)")
    def test_a_dropdowns_shown_value_is_extracted(self, tmp_path):
        spec = _extract(fx.build_inline_controls_spec(), tmp_path)
        assert "Pipe material: Copper Type L" in spec.content

    @pytest.mark.xfail(strict=True, raises=AssertionError, reason="open: fixed by S10 (WP-02)")
    def test_an_unresolved_dropdown_reaches_the_placeholder_detector(self, tmp_path):
        spec = _extract(fx.build_inline_controls_spec(), tmp_path)
        matches = [a["match"] for a in detect_placeholders(spec.content, spec.filename)]
        assert fx.UNRESOLVED_DROPDOWN_PLACEHOLDER in matches

    @pytest.mark.xfail(strict=True, raises=AssertionError, reason="open: fixed by S10 (WP-02)")
    def test_smart_tag_text_is_extracted(self, tmp_path):
        spec = _extract(fx.build_smart_tag_spec(), tmp_path)
        assert "Obtain approval from the City of Oakland fire marshal." in spec.content

    @pytest.mark.xfail(strict=True, raises=AssertionError, reason="open: fixed by S10 (WP-02)")
    def test_an_insertion_inside_a_hyperlink_is_accepted(self, tmp_path):
        spec = _extract(fx.build_hyperlink_spec(), tmp_path)
        assert "Submit manufacturer data sheets." in spec.content

    @pytest.mark.xfail(strict=True, raises=AssertionError, reason="open: fixed by S10 (WP-02)")
    def test_text_inside_a_control_never_redirects_an_edit_elsewhere(self, tmp_path):
        """The same words inside a control (p0) and in plain text (p1): an
        edit the review aimed at p0 must not land in p1. Today p0's control
        text is invisible, so the locator treats p0 as drifted and edits the
        unique plain-text copy in p1 instead."""
        from applier.models import OutcomeStatus
        from applier.run import RunSettings, apply_sidecar
        from applier.sidecar import load_sidecar
        from docx import Document
        from docx.oxml.ns import qn

        from src.input.extractor import _accept_all_paragraph_text

        builder = fx.SpecDocBuilder()
        builder.add_paragraph(
            "Provide ",
            fx.inline_control(fx.run(fx.INLINE_CONTROL_TEXT), tag="pipe", control_id=builder.next_id()),
            " pipe for mains.",
        )
        builder.add_paragraph(f"Provide {fx.INLINE_CONTROL_TEXT} pipe for branch lines.")
        source = fx.save_docx(builder, tmp_path, "ctl.docx")
        sidecar_path = tmp_path / "report.edits.json"
        sidecar_path.write_text(
            json.dumps(
                {
                    "schema_version": 4,
                    "generated_at": "2026-09-23T00:00:00Z",
                    "report_file": "report.docx",
                    "edit_count": 1,
                    "edits": [
                        {
                            "finding_id": "rf-000000000002",
                            "fileName": "ctl.docx",
                            "affected_files": ["ctl.docx"],
                            "has_per_file_original": True,
                            "section": "",
                            "severity": "HIGH",
                            "issue": "Main pipe material.",
                            "codeReference": None,
                            "evidenceElementId": "p0",
                            "verification_verdict": "CONFIRMED",
                            "report_status": "VERIFIED_SUPPORTED",
                            "edit_proposal": {
                                "action_type": "EDIT",
                                "existing_text": fx.INLINE_CONTROL_TEXT,
                                "replacement_text": "schedule 10 galvanized steel",
                                "anchor_text": None,
                                "insert_position": None,
                                "target_element_id": "p0",
                                "edit_confidence": 0.9,
                            },
                        }
                    ],
                }
            ),
            encoding="utf-8",
        )
        (result,) = apply_sidecar(load_sidecar(sidecar_path), [source], RunSettings())
        (outcome,) = result.outcomes
        if outcome.status is not OutcomeStatus.APPLIED:
            return  # refused: the plain-text copy in p1 cannot have been touched
        # Read the Accept-All view: a tracked edit keeps the old words as
        # w:delText, so a raw text read would still "find" them.
        edited = Document(result.output_path).element.body.findall(qn("w:p"))[1]
        assert fx.INLINE_CONTROL_TEXT in _accept_all_paragraph_text(edited), "p1 was edited"

    # -- controls ------------------------------------------------------------

    def test_control_text_around_a_block_control_is_read_once_in_order(self, tmp_path):
        spec = _extract(fx.build_block_control_spec(), tmp_path)
        for text in (fx.BEFORE_CONTROL, fx.AFTER_CONTROL, "Riser | Four inch"):
            assert spec.content.count(text) == 1, text
        assert spec.content.index(fx.BEFORE_CONTROL) < spec.content.index(fx.AFTER_CONTROL)

    def test_control_a_complex_fields_result_is_read_and_its_code_is_not(self, tmp_path):
        spec = _extract(fx.build_fields_spec(), tmp_path)
        assert "Coordinate with Section 21 13 13." in spec.content
        assert "REF" not in spec.content


# ===========================================================================
# WP-03 — automatic numbering reaches review (chunk S14)
# ===========================================================================


def _review_user_message(spec: ExtractedSpec) -> str:
    from src.review.review_request_builder import ReviewRequestSpec, build_user_message

    return build_user_message(
        ReviewRequestSpec(
            spec_content=spec.content,
            filename=spec.filename,
            model="claude-opus-5",
            paragraph_map=spec.paragraph_map,
        )
    )


class TestAutomaticNumbering:
    """Checked in the review prompt itself, so it holds whether S14 puts the
    labels into ``content`` or renders them from paragraph metadata."""

    @pytest.mark.xfail(strict=True, raises=AssertionError, reason="open: fixed by S14 (WP-03)")
    def test_every_displayed_label_reaches_the_review_prompt(self, tmp_path):
        message = _review_user_message(_extract(fx.build_auto_numbered_three_part(), tmp_path))
        missing = [
            f"{block.auto_label} {block.text}"
            for block in fx.auto_numbered_blocks()
            if not re.search(rf"{re.escape(block.auto_label)}\s+{re.escape(block.text)}", message)
        ]
        assert missing == []

    @pytest.mark.xfail(strict=True, raises=AssertionError, reason="open: fixed by S14 (WP-03)")
    def test_a_numbered_article_is_the_section_of_its_body(self, tmp_path):
        spec = _extract(fx.build_auto_numbered_three_part(), tmp_path)
        (body,) = [m for m in spec.paragraph_map if m.text == "Provide the specified piping system."]
        assert "1.01" in body.section_id

    def test_control_literal_body_text_is_still_read(self, tmp_path):
        spec = _extract(fx.build_auto_numbered_three_part(), tmp_path)
        assert "Provide the specified piping system." in spec.content

    def test_control_typed_numbers_are_not_repeated(self, tmp_path):
        message = _review_user_message(_extract(fx.build_clean_three_part(), tmp_path))
        assert "1.01 SUMMARY" in message
        assert not re.search(r"1\.01\s+1\.01", message)
        assert not re.search(r"\bA\.\s+A\.", message)


# ===========================================================================
# WP-11 — rate-limit timing (chunk S15)
# ===========================================================================


def _rate_limit_error(headers: dict):
    import anthropic
    import httpx2

    request = httpx2.Request("GET", "https://api.anthropic.com/v1/messages/batches/x/results")
    return anthropic.RateLimitError(
        "rate limited", response=httpx2.Response(429, headers=headers, request=request), body=None
    )


class _FlakyResults:
    """``client.messages.batches`` whose first ``results()`` raises."""

    def __init__(self, exc):
        self.exc = exc
        self.calls = 0

    def results(self, batch_id):
        self.calls += 1
        if self.calls == 1:
            raise self.exc
        return iter([SimpleNamespace(custom_id="a")])


def _drive_results_retry(monkeypatch, exc) -> tuple[list[float], _FlakyResults]:
    from src.batch import batch as batch_mod

    batches = _FlakyResults(exc)
    client = SimpleNamespace(messages=SimpleNamespace(batches=batches))
    monkeypatch.setattr(batch_mod, "_get_client", lambda **_: client)
    sleeps: list[float] = []
    monkeypatch.setattr(batch_mod.time, "sleep", lambda seconds: sleeps.append(seconds))
    out = batch_mod._collect_batch_results_with_retry("msgbatch_probe")
    if set(out) != {"a"} or batches.calls != 2:
        pytest.fail(f"precondition: one retry then success (calls={batches.calls}, out={out})")
    return sleeps, batches


class TestRetryAfter:
    """Driven through the batch-results download, an app-owned retry loop
    with SDK retries off (so the app, not the SDK, owns the wait)."""

    @pytest.mark.xfail(strict=True, raises=AssertionError, reason="open: fixed by S15 (WP-11)")
    def test_a_retry_waits_at_least_the_server_retry_after(self, monkeypatch):
        # 12 s exceeds today's first rate-limit backoff (5 s) and fits any
        # sane elapsed budget; the fake sleep means nothing actually waits.
        sleeps, _ = _drive_results_retry(monkeypatch, _rate_limit_error({"retry-after": "12"}))
        assert sleeps and sleeps[0] >= 12

    def test_control_without_a_header_it_still_retries_after_a_local_backoff(self, monkeypatch):
        sleeps, _ = _drive_results_retry(monkeypatch, _rate_limit_error({}))
        assert len(sleeps) == 1 and sleeps[0] > 0


# ===========================================================================
# WP-12 — the chat keeps the API key out of web storage (chunk S04)
# ===========================================================================

_REQUIRE_TOOLS_ENV = "SPEC_CRITIC_REQUIRE_HTML_TEST_TOOLS"


def _node_or_skip() -> str:
    node = shutil.which("node")
    if node:
        return node
    if os.environ.get(_REQUIRE_TOOLS_ENV, "").strip().lower() not in ("", "0", "false", "no", "off"):
        pytest.fail(f"node is required when {_REQUIRE_TOOLS_ENV} is set")
    pytest.skip(f"node not installed; set {_REQUIRE_TOOLS_ENV}=1 to make this a failure")


def _run_chat_key_probe(tmp_path: Path) -> dict:
    """Write a real report, then run its exact script under the Node probe."""
    node = _node_or_skip()
    from test_html_report_exporter import _EXEC_SCRIPT_RE, build_full_pipeline_result

    from src.output.html_report_exporter import write_html_report

    report = tmp_path / "report.html"
    write_html_report(
        build_full_pipeline_result(), report, generated_at=datetime(2026, 1, 1, 12), include_chat=True
    )
    text = report.read_bytes().decode("utf-8")
    scripts = _EXEC_SCRIPT_RE.findall(text)
    if len(scripts) != 1:
        pytest.fail(f"precondition: expected one executable script, found {len(scripts)}")
    elements = {
        match.group(1): match.group(2)
        for match in re.finditer(
            r'<script type="application/json" id="([^"]+)">(.*?)</script>', text, re.S
        )
    }
    plaintext = re.search(r'<pre id="sc-plaintext" hidden>(.*?)</pre>', text, re.S)
    if plaintext is None or "sc-chat-config" not in elements:
        pytest.fail("precondition: the report no longer carries the chat's data blocks")
    elements["sc-plaintext"] = html.unescape(plaintext.group(1))
    script_path = tmp_path / "report-script.js"
    script_path.write_bytes(scripts[0].encode("utf-8"))
    elements_path = tmp_path / "elements.json"
    elements_path.write_text(json.dumps(elements), encoding="utf-8")
    proc = subprocess.run(
        [
            node,
            str(_REPO_ROOT / "tests" / "fixtures" / "chat_key_probe.js"),
            str(script_path),
            str(elements_path),
            FAKE_KEY,
        ],
        capture_output=True,
        text=True,
        timeout=60,
    )
    if proc.returncode != 0 or not proc.stdout.strip():
        pytest.fail(f"precondition: the probe did not run: {proc.stderr[-2000:]}")
    result = json.loads(proc.stdout)
    if not result.get("ok"):
        pytest.fail(f"precondition: the shipped script did not load under the probe: {result}")
    return result


class TestChatKeyStorage:
    @pytest.mark.xfail(strict=True, raises=AssertionError, reason="open: fixed by S04 (WP-12)")
    def test_a_saved_key_is_never_written_to_web_storage(self, tmp_path):
        result = _run_chat_key_probe(tmp_path)
        leaked = [w for w in result["storageWrites"] if FAKE_KEY in w["value"]]
        assert leaked == []
        stored = {**result["sessionStorage"], **result["localStorage"]}
        assert FAKE_KEY not in stored.values()

    def test_control_a_saved_key_makes_the_chat_ready(self, tmp_path):
        result = _run_chat_key_probe(tmp_path)
        assert result["chatReady"] is True
        assert result["keyFieldCleared"] is True
        assert result["fetchCalls"] == []  # saving a key sends nothing


# ===========================================================================
# WP-13 — a key typed into the GUI stays out of os.environ (chunk S16)
# ===========================================================================


@pytest.fixture
def restored_environ(monkeypatch, tmp_path):
    """Snapshot the whole environment and restore it, however the code
    under test mutated it, and keep the GUI's state file out of $HOME."""
    saved = dict(os.environ)
    monkeypatch.setenv("SPEC_CRITIC_UI_STATE_PATH", str(tmp_path / "ui_state.json"))
    yield
    os.environ.clear()
    os.environ.update(saved)


class _RecordedThread:
    """Stands in for ``threading.Thread`` inside one controller module: the
    worker is captured, never started."""

    started: list = []

    def __init__(self, target=None, args=(), kwargs=None, daemon=None, **_):
        self.target, self.args = target, args

    def start(self):
        _RecordedThread.started.append(self)


def _record_threads(monkeypatch, module) -> list:
    import threading

    _RecordedThread.started = []
    monkeypatch.setattr(
        module, "threading", SimpleNamespace(Thread=_RecordedThread, Lock=threading.Lock)
    )
    return _RecordedThread.started


def _key_leaked() -> bool:
    return any(FAKE_KEY in value for value in os.environ.values())


class TestGuiKeyStaysOutOfTheEnvironment:
    """Each flow runs the real controller against a fake app until the
    point where the typed key has been consumed (a worker is handed off or
    the file picker opens); only then is the environment inspected."""

    @pytest.mark.xfail(strict=True, raises=AssertionError, reason="open: fixed by S16 (WP-13)")
    def test_the_drawing_digest_flow(self, monkeypatch, restored_environ):
        pytest.importorskip("tkinter")
        from src.gui import context_controller as cc

        picker_calls = []
        monkeypatch.setattr(
            cc.filedialog, "askopenfilenames", lambda **kw: picker_calls.append(kw) or ()
        )
        monkeypatch.setattr(
            cc.messagebox,
            "showerror",
            lambda *a, **k: pytest.fail(f"precondition: the key gate refused a key: {a}"),
        )
        app = SimpleNamespace(
            _drawing_digest_running=False,
            is_processing=False,
            api_key_entry=SimpleNamespace(get=lambda: FAKE_KEY),
        )
        cc.attach_drawing_files(app)
        if not picker_calls:
            pytest.fail("precondition: the flow never reached the file picker")
        assert not _key_leaked(), "the key typed into the GUI is in os.environ"

    @pytest.mark.xfail(strict=True, raises=AssertionError, reason="open: fixed by S16 (WP-13)")
    def test_starting_a_review(self, monkeypatch, restored_environ, tmp_path):
        pytest.importorskip("tkinter")
        from unittest.mock import MagicMock

        from src.gui import review_run_controller as rrc

        started = _record_threads(monkeypatch, rrc)
        spec_path = fx.save_docx(fx.build_clean_three_part(), tmp_path, "230500.docx")
        app = MagicMock()
        app.is_processing = False
        app.api_key_entry.get.return_value = FAKE_KEY
        app._selected_files = [spec_path]
        app.file_list_panel.get_selected_files.return_value = [spec_path]
        app._get_project_context.return_value = ""
        app._gather_project_profile.return_value = None
        app._cross_check_var.get.return_value = False
        app._realtime_var = None
        app._realtime_workers_var = None
        app._selected_program_id = None
        app._project_context_tokens = 0
        app._extracted_specs = [extract_text_from_docx(spec_path)]
        rrc.start_review(app)
        if len(started) != 1:
            pytest.fail(f"precondition: no review worker was handed off ({len(started)})")
        assert not _key_leaked(), "the key typed into the GUI is in os.environ"

    @pytest.mark.xfail(strict=True, raises=AssertionError, reason="open: fixed by S16 (WP-13)")
    def test_reconnecting_to_a_batch(self, monkeypatch, restored_environ):
        pytest.importorskip("tkinter")
        from unittest.mock import MagicMock

        from src.gui import batch_controller as bc

        started = _record_threads(monkeypatch, bc)
        app = MagicMock()
        app.is_processing = False
        app.api_key_entry.get.return_value = FAKE_KEY
        bc._begin_reconnect_run(
            app,
            reconstruct_fn=lambda log, progress: None,
            model="claude-opus-5",
            cycle_label=CALIFORNIA_2025.label,
            project_context="",
            cross_check_enabled=False,
            files_for_review=[],
            module_id="california_k12_mep",
            batch_label="msgbatch_probe",
        )
        if len(started) != 1:
            pytest.fail(f"precondition: no reconnect worker was handed off ({len(started)})")
        assert not _key_leaked(), "the key typed into the GUI is in os.environ"

    def test_control_the_digest_flow_still_requires_a_key(self, monkeypatch, restored_environ):
        pytest.importorskip("tkinter")
        from src.gui import context_controller as cc

        errors, picker_calls = [], []
        monkeypatch.setattr(cc.messagebox, "showerror", lambda *a, **k: errors.append(a))
        monkeypatch.setattr(
            cc.filedialog, "askopenfilenames", lambda **kw: picker_calls.append(kw) or ()
        )
        app = SimpleNamespace(
            _drawing_digest_running=False,
            is_processing=False,
            api_key_entry=SimpleNamespace(get=lambda: "   "),
        )
        cc.attach_drawing_files(app)
        assert errors and not picker_calls


# ===========================================================================
# WP-17 — documentation and banners match the code (chunk S18)
# ===========================================================================

#: Files that state the Haiku prompt-cache minimum. The handbook line is not
#: in the plan's original "Known-wrong statements" table; S01 found it.
_HAIKU_CLAIM_FILES = (
    "src/core/api_config.py",
    "CLAUDE.md",
    "handbook/12_configuration_and_models.md",
)

# A token count written as "2048" or "4,096", immediately followed by
# "-token" / " tokens" — the form every claim uses ("2048-token Haiku cache
# minimum", "(2048 tokens for Haiku)").
_TOKEN_COUNT_RE = re.compile(r"\b(\d{1,3}(?:,\d{3})+|\d{4,})(?=[- ]tokens?\b)")


def _haiku_cache_minimum_claims(relative: str) -> list[int]:
    claims: list[int] = []
    for line in (_REPO_ROOT / relative).read_text(encoding="utf-8").splitlines():
        lowered = line.lower()
        if "haiku" in lowered and "cache minimum" in lowered:
            claims.extend(int(n.replace(",", "")) for n in _TOKEN_COUNT_RE.findall(line))
    return claims


class TestDocumentationMatchesTheCode:
    @pytest.mark.xfail(strict=True, raises=AssertionError, reason="open: fixed by S18 (WP-17)")
    @pytest.mark.parametrize("relative", _HAIKU_CLAIM_FILES)
    def test_the_haiku_cache_minimum_is_stated_as_4096(self, relative):
        claims = _haiku_cache_minimum_claims(relative)
        if not claims:
            pytest.fail(
                f"precondition: {relative} no longer states the Haiku cache minimum; "
                "correct the statement rather than deleting it, or update this check"
            )
        assert set(claims) == {4096}

    def test_control_triage_is_still_not_cached(self):
        from src.core.api_config import PHASE_TRIAGE, cache_policy_for

        assert cache_policy_for(PHASE_TRIAGE).caches_anything is False

    @pytest.mark.xfail(strict=True, raises=AssertionError, reason="open: fixed by S18 (WP-17)")
    def test_the_banner_counts_a_hand_built_no_op_edit_as_a_demotion(self):
        from src.orchestration import pipeline
        from src.output.report_exporter import _summarize_run_diagnostics
        from src.output.report_status import summarize_edit_actions

        # Built directly (not parsed), then sent through dedup, the review
        # path's normalization step. The parser stamps demotion_reason; this
        # path does not.
        noop = _finding(actionType="EDIT", existingText="same", replacementText="same")
        findings = pipeline._deduplicate_findings([noop])
        summary = _summarize_run_diagnostics(
            findings=findings,
            status_counts={},
            edit_action_counts=summarize_edit_actions(findings),
            cross_check_result=None,
        )
        assert summary.get("demotion_count") == 1

    def test_control_a_parsed_no_op_edit_counts_as_a_demotion(self):
        from src.output.report_exporter import _summarize_run_diagnostics
        from src.output.report_status import summarize_edit_actions
        from src.review.reviewer import _parse_findings

        findings = _parse_findings(
            [
                {
                    "severity": "HIGH",
                    "fileName": "210500.docx",
                    "section": "2.01",
                    "issue": "Wrong valve type.",
                    "actionType": "EDIT",
                    "existingText": "same",
                    "replacementText": "same",
                    "confidence": 0.8,
                }
            ]
        )
        summary = _summarize_run_diagnostics(
            findings=findings,
            status_counts={},
            edit_action_counts=summarize_edit_actions(findings),
            cross_check_result=None,
        )
        assert summary.get("demotion_count") == 1
