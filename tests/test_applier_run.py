"""End to end: a sidecar and some specifications in, edited copies out.

The property these tests exist for is **conservation**: the number of
instructions accounted for in the receipt equals the number the sidecar
listed, in every scenario — including the unhappy ones, where it matters
most. An applier that drops four of twenty-three instructions makes four
defective specs look clean.
"""
from __future__ import annotations

import json

import pytest
from docx import Document

from applier.cli import EXIT_ERROR, EXIT_OK, EXIT_UNAPPLIED, _collect_specs, main
from applier.docx_edit import DIRECT, TRACKED
from applier.models import OutcomeStatus
from applier.policy import Policy, PolicyConfig
from applier.receipt import build_receipt, render_summary
from applier.run import RunSettings, apply_sidecar
from applier.sidecar import load_sidecar
from src.input.extractor import extract_text_from_docx


def build_spec(path, *, tracked=False, repeated=False):
    document = Document()
    document.add_paragraph("SECTION 21 13 13 - WET-PIPE SPRINKLER SYSTEMS")
    document.add_paragraph("PART 1 - GENERAL")
    document.add_paragraph("Systems shall comply with NFPA 13, 2019 edition.")
    document.add_paragraph("PART 2 - PRODUCTS")
    document.add_paragraph("Provide [SELECT] type sprinklers throughout.")
    document.add_paragraph("Hangers shall comply with NFPA 13, 2019 edition.")
    if repeated:
        # One paragraph, two occurrences of the same target text: an element
        # id names the paragraph, not which occurrence.
        document.add_paragraph(
            "Install per NFPA 13 and test per NFPA 13 before acceptance."
        )
    if tracked:
        # A pending revision on the source: the review read its accept-all view.
        from docx.oxml.ns import qn

        paragraph = document.paragraphs[2]._p
        run = paragraph.findall(qn("w:r"))[0]
        insertion = paragraph.makeelement(
            qn("w:ins"),
            {qn("w:id"): "1", qn("w:author"): "Someone", qn("w:date"): "2026-01-01T00:00:00Z"},
        )
        paragraph.replace(run, insertion)
        insertion.append(run)
    document.save(path)
    return path


def edit(finding_id, **overrides):
    proposal = {
        "action_type": "EDIT",
        "existing_text": "NFPA 13, 2019 edition",
        "replacement_text": "NFPA 13, 2025 edition",
        "anchor_text": None,
        "insert_position": None,
        "target_element_id": "p2",
        "edit_confidence": 0.9,
    }
    proposal.update(overrides.pop("proposal", {}))
    base = {
        "finding_id": finding_id,
        "fileName": "215000.docx",
        "affected_files": ["215000.docx"],
        "has_per_file_original": True,
        "section": "21 13 13",
        "severity": "HIGH",
        "issue": "Stale edition reference.",
        "codeReference": "NFPA 13",
        "evidenceElementId": "p2",
        "verification_verdict": "CONFIRMED",
        "report_status": "VERIFIED_SUPPORTED",
        "edit_proposal": proposal,
    }
    base.update(overrides)
    return base


def write_sidecar(tmp_path, edits, *, version=4, **top):
    payload = {
        "schema_version": version,
        "generated_at": "2026-09-21T00:00:00Z",
        "report_file": "report.docx",
        "edit_count": len(edits),
        "edits": edits,
    }
    payload.update(top)
    path = tmp_path / "report.edits.json"
    path.write_text(json.dumps(payload), encoding="utf-8")
    return path


def run(tmp_path, edits, settings=None, **kwargs):
    build_spec(
        tmp_path / "215000.docx",
        tracked=kwargs.pop("tracked_source", False),
        repeated=kwargs.pop("repeated", False),
    )
    sidecar = load_sidecar(write_sidecar(tmp_path, edits, **kwargs))
    results = apply_sidecar(
        sidecar,
        [tmp_path / "215000.docx"],
        settings or RunSettings(),
    )
    return sidecar, results


def outcomes(results):
    return {
        outcome.entry.finding_id: outcome
        for result in results
        for outcome in result.outcomes
    }


class TestHappyPath:
    def test_an_edit_is_applied_to_a_new_file(self, tmp_path):
        _, results = run(tmp_path, [edit("rf-1")])
        (result,) = results
        assert result.applied == 1
        assert result.output_path.endswith("215000.applied.docx")
        assert (tmp_path / "215000.applied.docx").exists()

    def test_the_source_is_left_byte_identical(self, tmp_path):
        build_spec(tmp_path / "215000.docx")
        before = (tmp_path / "215000.docx").read_bytes()
        sidecar = load_sidecar(write_sidecar(tmp_path, [edit("rf-1")]))
        apply_sidecar(sidecar, [tmp_path / "215000.docx"], RunSettings())
        assert (tmp_path / "215000.docx").read_bytes() == before

    def test_the_edited_copy_reads_as_intended(self, tmp_path):
        run(tmp_path, [edit("rf-1")])
        extracted = extract_text_from_docx(tmp_path / "215000.applied.docx")
        assert extracted.paragraph_map[2].text == (
            "Systems shall comply with NFPA 13, 2025 edition."
        )
        assert extracted.tracked_changes_detected is True

    def test_direct_mode_writes_no_revisions(self, tmp_path):
        run(tmp_path, [edit("rf-1")], RunSettings(mode=DIRECT))
        extracted = extract_text_from_docx(tmp_path / "215000.applied.docx")
        assert extracted.tracked_changes_detected is False
        assert "2025 edition" in extracted.paragraph_map[2].text

    def test_several_edits_to_one_document_all_land(self, tmp_path):
        _, results = run(
            tmp_path,
            [
                edit("rf-1"),
                edit(
                    "rf-2",
                    evidenceElementId="p4",
                    proposal={
                        "existing_text": "[SELECT] type",
                        "replacement_text": "quick-response",
                        "target_element_id": "p4",
                    },
                ),
                edit(
                    "rf-3",
                    evidenceElementId="p4",
                    proposal={
                        "action_type": "ADD",
                        "existing_text": None,
                        "replacement_text": "Sprinklers shall be listed for the hazard.",
                        "anchor_text": "Provide",
                        "insert_position": "after",
                        "target_element_id": "p4",
                    },
                ),
            ],
        )
        assert results[0].applied == 3
        texts = [
            mapping.text
            for mapping in extract_text_from_docx(
                tmp_path / "215000.applied.docx"
            ).paragraph_map
        ]
        assert "Systems shall comply with NFPA 13, 2025 edition." in texts
        assert "Provide quick-response sprinklers throughout." in texts
        assert "Sprinklers shall be listed for the hazard." in texts

    def test_a_program_sidecar_runs_the_same_way(self, tmp_path):
        _, results = run(
            tmp_path,
            [edit("rf-1", module_id="datacenter_fire")],
            version=5,
            program_id="hyperscale_datacenter",
        )
        assert results[0].applied == 1
        assert outcomes(results)["rf-1"].entry.module_id == "datacenter_fire"


class TestNothingIsLost:
    @pytest.mark.parametrize("dry_run", [False, True])
    def test_every_instruction_is_accounted_for(self, tmp_path, dry_run):
        edits = [
            edit("rf-applied"),
            edit("rf-held", report_status="DISPUTED"),
            edit("rf-corrected", report_status="VERIFIED_CONTRADICTED"),
            edit(
                "rf-missing",
                evidenceElementId=None,
                proposal={
                    "existing_text": "ASCE 7-16 bracing",
                    "target_element_id": None,
                },
            ),
            edit("rf-bad", proposal={"action_type": "REPORT_ONLY"}),
            edit("rf-elsewhere", fileName="230000.docx"),
        ]
        sidecar, results = run(tmp_path, edits, RunSettings(dry_run=dry_run))
        receipt = build_receipt(
            sidecar=sidecar,
            file_results=results,
            settings=RunSettings(dry_run=dry_run).to_dict(),
            entries_in=len(sidecar.entries) + len(sidecar.malformed),
        )
        assert receipt["accounting"]["entries_in_sidecar"] == len(edits)
        assert receipt["accounting"]["balanced"] is True

        by_id = outcomes(results)
        applied = (
            OutcomeStatus.WOULD_APPLY if dry_run else OutcomeStatus.APPLIED
        )
        assert by_id["rf-applied"].status is applied
        assert by_id["rf-held"].status is OutcomeStatus.HELD_BY_POLICY
        assert by_id["rf-corrected"].status is OutcomeStatus.HELD_BY_POLICY
        assert by_id["rf-missing"].status is OutcomeStatus.UNLOCATED
        assert by_id["rf-bad"].status is OutcomeStatus.MALFORMED
        assert by_id["rf-elsewhere"].status is OutcomeStatus.FILE_MISSING

    def test_every_outcome_carries_a_reason_or_a_note(self, tmp_path):
        _, results = run(
            tmp_path,
            [
                edit("rf-1"),
                edit("rf-2", report_status="DISPUTED"),
                edit("rf-3", proposal={"existing_text": "absent text", "target_element_id": None}),
            ],
        )
        for outcome in outcomes(results).values():
            assert outcome.reason or outcome.change_note


class TestDryRun:
    def test_nothing_is_written(self, tmp_path):
        _, results = run(tmp_path, [edit("rf-1")], RunSettings(dry_run=True))
        assert not (tmp_path / "215000.applied.docx").exists()
        assert results[0].output_path is None
        assert results[0].applied == 0

    def test_the_outcome_says_what_would_happen(self, tmp_path):
        _, results = run(tmp_path, [edit("rf-1")], RunSettings(dry_run=True))
        outcome = outcomes(results)["rf-1"]
        assert outcome.status is OutcomeStatus.WOULD_APPLY
        # The note is the writer's own, not a guess made before it ran.
        assert "tracked replacement" in outcome.change_note
        assert outcome.location.element_id == "p2"

    def test_a_dry_run_refuses_exactly_what_a_real_run_refuses(self, tmp_path):
        """The preview must not be rosier than the thing it previews: a target
        repeated inside its element is a writer-level refusal, so a dry run
        has to reach the writer to see it."""
        build_spec(tmp_path / "215000.docx", repeated=True)
        instruction = [
            edit(
                "rf-dup",
                evidenceElementId="p6",
                proposal={
                    "existing_text": "per NFPA 13",
                    "replacement_text": "per NFPA 13 (2025)",
                    "target_element_id": "p6",
                },
            )
        ]
        sidecar = load_sidecar(write_sidecar(tmp_path, instruction))
        dry = apply_sidecar(
            sidecar, [tmp_path / "215000.docx"], RunSettings(dry_run=True)
        )
        wet = apply_sidecar(sidecar, [tmp_path / "215000.docx"], RunSettings())
        assert outcomes(dry)["rf-dup"].status is OutcomeStatus.UNLOCATED
        assert outcomes(wet)["rf-dup"].status is OutcomeStatus.UNLOCATED
        assert "occurs 2 times" in outcomes(dry)["rf-dup"].reason

    def test_strict_is_not_fooled_by_a_dry_run(self, tmp_path):
        build_spec(tmp_path / "215000.docx", repeated=True)
        sidecar = write_sidecar(
            tmp_path,
            [
                edit(
                    "rf-dup",
                    evidenceElementId="p6",
                    proposal={
                        "existing_text": "per NFPA 13",
                        "replacement_text": "per NFPA 13 (2025)",
                        "target_element_id": "p6",
                    },
                )
            ],
        )
        from applier.cli import EXIT_UNAPPLIED, main

        assert (
            main([str(sidecar), "--specs", str(tmp_path), "--dry-run", "--strict"])
            == EXIT_UNAPPLIED
        )


class TestDocumentLevelRefusals:
    def test_a_source_with_pending_revisions_is_skipped(self, tmp_path):
        _, results = run(tmp_path, [edit("rf-1")], tracked_source=True)
        assert results[0].applied == 0
        assert not (tmp_path / "215000.applied.docx").exists()
        assert any("pending tracked changes" in e for e in results[0].errors)
        assert outcomes(results)["rf-1"].status is OutcomeStatus.FAILED

    def test_the_override_admits_the_document_but_not_the_redlined_clause(
        self, tmp_path
    ):
        """--allow-tracked-source is not a blanket permission: the document is
        processed, but an edit landing inside someone else's undecided
        revision is still refused, with a message that says so."""
        _, results = run(
            tmp_path,
            [edit("rf-1")],
            RunSettings(allow_tracked_source=True),
            tracked_source=True,
        )
        assert not any("pending tracked changes" in e for e in results[0].errors)
        outcome = outcomes(results)["rf-1"]
        assert outcome.status is OutcomeStatus.UNLOCATED
        assert "inside an existing tracked revision" in outcome.reason

    def test_the_override_lets_a_clean_clause_through(self, tmp_path):
        _, results = run(
            tmp_path,
            [
                edit(
                    "rf-clean",
                    evidenceElementId="p5",
                    proposal={
                        "existing_text": "Hangers shall comply",
                        "replacement_text": "Hangers and bracing shall comply",
                        "target_element_id": "p5",
                    },
                )
            ],
            RunSettings(allow_tracked_source=True),
            tracked_source=True,
        )
        assert results[0].applied == 1

    def test_the_source_is_never_the_destination(self, tmp_path):
        build_spec(tmp_path / "215000.docx")
        before = (tmp_path / "215000.docx").read_bytes()
        sidecar = load_sidecar(write_sidecar(tmp_path, [edit("rf-1")]))
        results = apply_sidecar(
            sidecar,
            [tmp_path / "215000.docx"],
            RunSettings(output_suffix=""),
        )
        assert (tmp_path / "215000.docx").read_bytes() == before
        assert results[0].applied == 0
        assert any("refusing to write over" in e for e in results[0].errors)
        assert outcomes(results)["rf-1"].status is OutcomeStatus.FAILED

    def test_an_unreadable_document_is_reported_not_raised(self, tmp_path):
        (tmp_path / "215000.docx").write_text("not a docx", encoding="utf-8")
        sidecar = load_sidecar(write_sidecar(tmp_path, [edit("rf-1")]))
        results = apply_sidecar(sidecar, [tmp_path / "215000.docx"], RunSettings())
        assert outcomes(results)["rf-1"].status is OutcomeStatus.FAILED
        assert results[0].errors

    def test_output_dir_is_honoured(self, tmp_path):
        out = tmp_path / "edited"
        run(tmp_path, [edit("rf-1")], RunSettings(output_dir=out))
        assert (out / "215000.applied.docx").exists()


class TestPolicyIntegration:
    def test_a_strict_run_withholds_a_locally_classified_finding(self, tmp_path):
        _, results = run(
            tmp_path,
            [edit("rf-1", report_status="LOCALLY_CLASSIFIED")],
            RunSettings(policy=PolicyConfig(policy=Policy.STRICT)),
        )
        assert outcomes(results)["rf-1"].status is OutcomeStatus.HELD_BY_POLICY

    def test_force_status_reaches_the_document(self, tmp_path):
        _, results = run(
            tmp_path,
            [edit("rf-1", report_status="DISPUTED")],
            RunSettings(
                policy=PolicyConfig(force_statuses=frozenset({"DISPUTED"}))
            ),
        )
        assert outcomes(results)["rf-1"].status is OutcomeStatus.APPLIED


class TestSpecDiscovery:
    def test_directories_expand_and_lock_files_are_skipped(self, tmp_path):
        build_spec(tmp_path / "215000.docx")
        build_spec(tmp_path / "230000.DOCX")
        (tmp_path / "~$215000.docx").write_text("lock", encoding="utf-8")
        (tmp_path / "notes.txt").write_text("x", encoding="utf-8")
        found = _collect_specs([tmp_path], tmp_path / "report.edits.json")
        assert [p.name for p in found] == ["215000.docx", "230000.DOCX"]

    def test_the_same_file_twice_is_listed_once(self, tmp_path):
        build_spec(tmp_path / "215000.docx")
        found = _collect_specs(
            [tmp_path, tmp_path / "215000.docx"], tmp_path / "report.edits.json"
        )
        assert len(found) == 1

    def test_the_sidecar_directory_is_the_default(self, tmp_path):
        build_spec(tmp_path / "215000.docx")
        assert _collect_specs([], tmp_path / "report.edits.json")

    def test_file_names_match_case_insensitively(self, tmp_path):
        build_spec(tmp_path / "215000.DOCX")
        sidecar = load_sidecar(write_sidecar(tmp_path, [edit("rf-1")]))
        results = apply_sidecar(sidecar, [tmp_path / "215000.DOCX"], RunSettings())
        assert results[0].applied == 1


class TestSummary:
    def test_the_summary_leads_with_what_needs_a_person(self, tmp_path):
        sidecar, results = run(
            tmp_path,
            [edit("rf-1"), edit("rf-2", report_status="DISPUTED")],
        )
        receipt = build_receipt(
            sidecar=sidecar,
            file_results=results,
            settings=RunSettings().to_dict(),
            entries_in=2,
        )
        text = render_summary(receipt, results)
        assert text.index("Not applied:") < text.index("Applied:")
        assert "1 of 2 instructions applied." in text
        assert "Review > Accept/Reject" in text

    def test_integrity_warnings_are_surfaced(self, tmp_path):
        sidecar, results = run(
            tmp_path,
            [edit("rf-1", module_id="datacenter_fire")],
            version=5,
            program_id="p",
            integrity_warnings=["submitted_files entry dropped"],
        )
        receipt = build_receipt(
            sidecar=sidecar,
            file_results=results,
            settings=RunSettings().to_dict(),
            entries_in=1,
        )
        assert "submitted_files entry dropped" in render_summary(receipt, results)


class TestCli:
    def _setup(self, tmp_path, edits, **kwargs):
        build_spec(tmp_path / "215000.docx")
        return write_sidecar(tmp_path, edits, **kwargs)

    def test_a_clean_run_exits_zero_and_writes_a_receipt(self, tmp_path, capsys):
        sidecar = self._setup(tmp_path, [edit("rf-1")])
        assert main([str(sidecar), "--specs", str(tmp_path)]) == EXIT_OK
        assert (tmp_path / "report.edits.applied.json").exists()
        assert "1 of 1 instructions applied." in capsys.readouterr().out

    def test_dry_run_writes_no_document(self, tmp_path, capsys):
        sidecar = self._setup(tmp_path, [edit("rf-1")])
        assert main([str(sidecar), "--specs", str(tmp_path), "--dry-run"]) == EXIT_OK
        assert not (tmp_path / "215000.applied.docx").exists()
        assert "DRY RUN" in capsys.readouterr().out

    def test_strict_flags_an_unapplied_instruction(self, tmp_path):
        sidecar = self._setup(
            tmp_path,
            [edit("rf-1", proposal={"existing_text": "absent", "target_element_id": None})],
        )
        argv = [str(sidecar), "--specs", str(tmp_path)]
        assert main(argv) == EXIT_OK
        assert main(argv + ["--strict"]) == EXIT_UNAPPLIED

    def test_a_policy_hold_alone_does_not_fail_strict(self, tmp_path):
        sidecar = self._setup(tmp_path, [edit("rf-1", report_status="DISPUTED")])
        assert main([str(sidecar), "--specs", str(tmp_path), "--strict"]) == EXIT_OK

    def test_only_restricts_the_run(self, tmp_path):
        sidecar = self._setup(tmp_path, [edit("rf-1"), edit("rf-2")])
        main([str(sidecar), "--specs", str(tmp_path), "--only", "rf-1"])
        receipt = json.loads((tmp_path / "report.edits.applied.json").read_text())
        assert receipt["accounting"]["by_outcome"]["APPLIED"] == 1
        assert receipt["accounting"]["by_outcome"]["HELD_BY_POLICY"] == 1

    def test_an_unknown_schema_exits_with_an_error(self, tmp_path, capsys):
        sidecar = self._setup(tmp_path, [edit("rf-1")], version=99)
        assert main([str(sidecar), "--specs", str(tmp_path)]) == EXIT_ERROR
        assert "schema_version" in capsys.readouterr().err

    def test_an_unknown_force_status_exits_with_an_error(self, tmp_path, capsys):
        sidecar = self._setup(tmp_path, [edit("rf-1")])
        code = main(
            [str(sidecar), "--specs", str(tmp_path), "--force-status", "DISPUTEED"]
        )
        assert code == EXIT_ERROR
        assert "Unknown report status" in capsys.readouterr().err

    def test_an_out_of_range_confidence_exits_with_an_error(self, tmp_path):
        sidecar = self._setup(tmp_path, [edit("rf-1")])
        assert (
            main(
                [str(sidecar), "--specs", str(tmp_path), "--min-edit-confidence", "1.5"]
            )
            == EXIT_ERROR
        )

    def test_no_specifications_found_exits_with_an_error(self, tmp_path, capsys):
        sidecar = write_sidecar(tmp_path, [edit("rf-1")])
        empty = tmp_path / "empty"
        empty.mkdir()
        assert main([str(sidecar), "--specs", str(empty)]) == EXIT_ERROR
        assert "no .docx specifications found" in capsys.readouterr().err

    def test_assist_is_off_unless_asked_for(self, tmp_path):
        sidecar = self._setup(tmp_path, [edit("rf-1")])
        main([str(sidecar), "--specs", str(tmp_path)])
        receipt = json.loads((tmp_path / "report.edits.applied.json").read_text())
        assert receipt["settings"]["assist"] is False
        assert receipt["settings"]["assist_model"] is None

    def test_the_default_mode_is_tracked_and_the_policy_conservative(self, tmp_path):
        sidecar = self._setup(tmp_path, [edit("rf-1")])
        main([str(sidecar), "--specs", str(tmp_path)])
        receipt = json.loads((tmp_path / "report.edits.applied.json").read_text())
        assert receipt["settings"]["mode"] == "tracked"
        assert receipt["settings"]["policy"] == "conservative"
