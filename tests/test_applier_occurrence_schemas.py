"""The applier reads schemas 4, 5, 6, and 7 (plan WP-06B, chunk S11).

Schemas 4 and 5 list one entry per affected file, so one file's second
location of a fix never reached them; they are read exactly as before. Schemas
6 and 7 are their occurrence-aware successors — one entry per occurrence, with
an ``occurrence_id`` and a ``location_basis`` — and the reader lands before the
writer does (S12), so the contract is pinned here from the reader's side:
required fields, what each location basis permits, module provenance in a
program sidecar, and the unique key ``(module_id, occurrence_id)``.
"""
from __future__ import annotations

import json

import pytest
from docx import Document

from applier.cli import EXIT_OK, main
from applier.models import EditEntry, OutcomeStatus
from applier.receipt import build_receipt, render_summary
from applier.run import RunSettings, apply_sidecar
from applier.sidecar import (
    LEGACY_SCHEMA_VERSIONS,
    OCCURRENCE_SCHEMA_VERSIONS,
    PROGRAM_SCHEMA_VERSIONS,
    SidecarSchemaError,
    load_sidecar,
)
from src.input.extractor import _accept_all_paragraph_text, extract_text_from_docx


def proposal(**overrides):
    base = {
        "action_type": "EDIT",
        "existing_text": "gate valve",
        "replacement_text": "ball valve",
        "anchor_text": None,
        "insert_position": None,
        "target_element_id": "p1",
        "edit_confidence": 0.9,
    }
    base.update(overrides)
    return base


def legacy_entry(**overrides):
    base = {
        "finding_id": "rf-abc123",
        "fileName": "spec.docx",
        "affected_files": ["spec.docx"],
        "has_per_file_original": True,
        "section": "2.01",
        "severity": "HIGH",
        "issue": "Gate valves are specified where ball valves are required.",
        "codeReference": None,
        "evidenceElementId": "p1",
        "verification_verdict": "CONFIRMED",
        "report_status": "VERIFIED_SUPPORTED",
        "edit_proposal": proposal(),
    }
    base.update(overrides)
    return base


def occurrence_entry(occurrence_id="oc-000000000001", element="p1", basis="validated", **overrides):
    base = legacy_entry()
    del base["has_per_file_original"]
    base.update(
        occurrence_id=occurrence_id,
        location_basis=basis,
        evidenceElementId=element,
        edit_proposal=proposal(target_element_id=element),
    )
    base.update(overrides)
    return base


def write(tmp_path, edits, *, version, name="report.edits.json", **top):
    payload = {"schema_version": version, "generated_at": "2026-09-24T00:00:00Z", "edits": edits}
    payload.update(top)
    path = tmp_path / name
    path.write_text(json.dumps(payload), encoding="utf-8")
    return path


def load(tmp_path, edits, *, version, **top):
    return load_sidecar(write(tmp_path, edits, version=version, **top))


# ---------------------------------------------------------------------------
# Which schemas, and which are programs
# ---------------------------------------------------------------------------


class TestSchemaFamilies:
    def test_the_families(self):
        assert LEGACY_SCHEMA_VERSIONS == {4, 5}
        assert OCCURRENCE_SCHEMA_VERSIONS == {6, 7}
        assert PROGRAM_SCHEMA_VERSIONS == {5, 7}

    @pytest.mark.parametrize(
        "version,program,occurrences",
        [(4, False, False), (5, True, False), (6, False, True), (7, True, True)],
    )
    def test_is_program_and_is_occurrence_aware(self, tmp_path, version, program, occurrences):
        entry = occurrence_entry(module_id="datacenter_fire") if occurrences else legacy_entry()
        loaded = load(tmp_path, [entry], version=version)
        assert loaded.is_program is program
        assert loaded.is_occurrence_aware is occurrences
        assert len(loaded.entries) == 1 and not loaded.malformed

    @pytest.mark.parametrize("version", [3, 8, 0, -1, "6", "7", 6.0, True, None])
    def test_anything_else_is_still_refused(self, tmp_path, version):
        with pytest.raises(SidecarSchemaError):
            load(tmp_path, [occurrence_entry()], version=version)

    def test_the_writers_future_numbers_are_the_ones_read(self):
        """6 and 7 were verified unused before they were reserved; the writer
        must keep emitting 4 and 5 until it emits the whole occurrence shape
        (plan WP-06B: never half a contract migration)."""
        from src.output.edit_sidecar import PROGRAM_SIDECAR_SCHEMA_VERSION, SIDECAR_SCHEMA_VERSION

        assert (SIDECAR_SCHEMA_VERSION, PROGRAM_SIDECAR_SCHEMA_VERSION) == (4, 5)


# ---------------------------------------------------------------------------
# Legacy sidecars read exactly as before
# ---------------------------------------------------------------------------


class TestLegacyReadExactlyAsBefore:
    @pytest.mark.parametrize("version", [4, 5])
    def test_occurrence_keys_in_a_legacy_sidecar_mean_nothing(self, tmp_path, version):
        entry = legacy_entry(occurrence_id="oc-ffffffffffff", location_basis="unresolved", module_id="m")
        (read,) = load(tmp_path, [entry], version=version).entries
        assert read.occurrence_id is None and read.location_basis is None
        assert read.key == ("rf-abc123", "spec.docx")
        assert read.evidence_element_id == "p1"

    def test_every_field_reads_as_it_always_has(self, tmp_path):
        (read,) = load(tmp_path, [legacy_entry(has_per_file_original=False, module_id="m")], version=5).entries
        assert read == EditEntry(
            finding_id="rf-abc123",
            file_name="spec.docx",
            action_type="EDIT",
            existing_text="gate valve",
            replacement_text="ball valve",
            anchor_text=None,
            insert_position=None,
            target_element_id="p1",
            evidence_element_id="p1",
            edit_confidence=0.9,
            report_status="VERIFIED_SUPPORTED",
            verification_verdict="CONFIRMED",
            severity="HIGH",
            section="2.01",
            issue="Gate valves are specified where ball valves are required.",
            code_reference=None,
            has_per_file_original=False,
            affected_files=("spec.docx",),
            module_id="m",
        )

    def test_a_repeated_legacy_key_is_still_read(self, tmp_path):
        """Schema 4 never promised a unique key across content twins (two
        identical coordination findings share a ``cf-`` id), so the reader
        keeps both; the run decides what to write from where they land."""
        loaded = load(tmp_path, [legacy_entry(), legacy_entry()], version=4)
        assert len(loaded.entries) == 2 and not loaded.malformed

    def test_a_legacy_program_entry_may_name_no_module(self, tmp_path):
        loaded = load(tmp_path, [legacy_entry()], version=5)
        assert loaded.entries and loaded.entries[0].module_id is None


# ---------------------------------------------------------------------------
# The occurrence contract
# ---------------------------------------------------------------------------


def _problem(tmp_path, entry, version=6):
    loaded = load(tmp_path, [entry], version=version)
    assert not loaded.entries, "expected the entry to be refused"
    (_, why), = loaded.malformed
    return why


class TestOccurrenceContract:
    def test_a_well_formed_entry(self, tmp_path):
        (read,) = load(tmp_path, [occurrence_entry()], version=6).entries
        assert (read.occurrence_id, read.location_basis) == ("oc-000000000001", "validated")
        assert read.key == ("", "oc-000000000001")
        assert read.has_per_file_original is True

    @pytest.mark.parametrize("missing", [None, "", "   "])
    def test_an_occurrence_id_is_required(self, tmp_path, missing):
        assert "occurrence_id" in _problem(tmp_path, occurrence_entry(occurrence_id=missing))

    @pytest.mark.parametrize("basis", [None, "", "guessed", "VALIDATEDX"])
    @pytest.mark.parametrize("element", ["p1", None])
    def test_an_unknown_basis_is_refused(self, tmp_path, basis, element):
        why = _problem(tmp_path, occurrence_entry(basis=basis, element=element))
        assert why.startswith("unknown location_basis")

    def test_the_basis_is_read_case_insensitively(self, tmp_path):
        (read,) = load(tmp_path, [occurrence_entry(basis=" Validated ")], version=6).entries
        assert read.location_basis == "validated"

    @pytest.mark.parametrize("basis", ["validated", "claimed"])
    def test_an_element_basis_needs_an_element(self, tmp_path, basis):
        entry = occurrence_entry(basis=basis, element=None)
        assert "names no element" in _problem(tmp_path, entry)

    @pytest.mark.parametrize("basis", ["unresolved", "missing_original"])
    def test_a_basis_without_a_place_may_not_carry_an_element(self, tmp_path, basis):
        entry = occurrence_entry(basis=basis, element="p1")
        assert "may not carry one" in _problem(tmp_path, entry)

    @pytest.mark.parametrize("basis", ["unresolved", "missing_original"])
    def test_a_basis_without_a_place_reads_with_no_element(self, tmp_path, basis):
        (read,) = load(tmp_path, [occurrence_entry(basis=basis, element=None)], version=6).entries
        assert read.target_element_id is None and read.evidence_element_id is None
        assert read.has_per_file_original is (basis != "missing_original")

    def test_the_legacy_per_file_flag_is_derived_not_read(self, tmp_path):
        entry = occurrence_entry(has_per_file_original=False)
        (read,) = load(tmp_path, [entry], version=6).entries
        assert read.has_per_file_original is True

    def test_a_program_entry_must_name_its_module(self, tmp_path):
        assert "module_id" in _problem(tmp_path, occurrence_entry(), version=7)

    def test_a_single_module_entry_may_name_its_module(self, tmp_path):
        (read,) = load(tmp_path, [occurrence_entry(module_id="datacenter_fire")], version=6).entries
        assert read.key == ("datacenter_fire", "oc-000000000001")

    def test_an_addition_with_no_original_has_its_own_reason(self, tmp_path):
        entry = occurrence_entry(
            basis="missing_original",
            element=None,
            edit_proposal=proposal(
                action_type="ADD", existing_text=None, replacement_text="New clause.",
                anchor_text=None, insert_position="after", target_element_id=None,
            ),
        )
        assert "recorded no original here" in _problem(tmp_path, entry)

    def test_the_action_shape_rules_still_apply(self, tmp_path):
        entry = occurrence_entry(edit_proposal=proposal(action_type="REPORT_ONLY"))
        assert "unsupported action_type" in _problem(tmp_path, entry)


class TestUniqueKey:
    def test_every_copy_of_a_repeated_key_is_refused(self, tmp_path):
        loaded = load(
            tmp_path,
            [occurrence_entry(), occurrence_entry("oc-000000000002", element="p2"), occurrence_entry()],
            version=6,
        )
        assert [e.occurrence_id for e in loaded.entries] == ["oc-000000000002"]
        assert len(loaded.malformed) == 2
        assert all("is listed 2 times" in why for _, why in loaded.malformed)

    def test_the_refusal_does_not_depend_on_order(self, tmp_path):
        """Two different instructions under one key: neither is chosen."""
        a = occurrence_entry()
        b = occurrence_entry(edit_proposal=proposal(replacement_text="butterfly valve"))
        first = load(tmp_path, [a, b], version=6)
        second = load(tmp_path, [b, a], version=6)
        assert not first.entries and not second.entries
        assert len(first.malformed) == len(second.malformed) == 2

    def test_one_occurrence_id_in_two_modules_is_two_keys(self, tmp_path):
        loaded = load(
            tmp_path,
            [occurrence_entry(module_id="datacenter_fire"), occurrence_entry(module_id="datacenter_electrical")],
            version=7,
        )
        assert len(loaded.entries) == 2 and not loaded.malformed

    def test_the_repeated_key_names_its_module(self, tmp_path):
        loaded = load(
            tmp_path,
            [occurrence_entry(module_id="datacenter_fire"), occurrence_entry(module_id="datacenter_fire")],
            version=7,
        )
        assert all("in module datacenter_fire" in why for _, why in loaded.malformed)

    def test_every_entry_is_still_accounted_for(self, tmp_path):
        edits = [occurrence_entry(), occurrence_entry(), occurrence_entry("oc-3", element=None, basis="validated")]
        loaded = load(tmp_path, edits, version=6)
        assert len(loaded.entries) + len(loaded.malformed) == len(edits)

    @pytest.mark.parametrize("broken_first", [True, False], ids=["broken-first", "broken-last"])
    def test_a_copy_refused_for_its_own_defect_still_repeats_the_key(self, tmp_path, broken_first):
        """A broken copy beside an executable one is still a key listed twice.
        Counting only the executable copies let it through, so the sidecar's
        corruption chose which instruction was written."""
        valid = occurrence_entry()
        broken = occurrence_entry(edit_proposal=proposal(action_type="EDTI"))
        edits = [broken, valid] if broken_first else [valid, broken]
        loaded = load(tmp_path, edits, version=6)
        assert loaded.entries == []
        reasons = {entry.action_type: why for entry, why in loaded.malformed}
        assert reasons == {
            "EDTI": "unsupported action_type 'EDTI'",
            "EDIT": (
                "occurrence oc-000000000001 is listed 2 times; each occurrence "
                "must appear once, so none of these copies is applied"
            ),
        }

    def test_a_broken_copy_repeats_a_program_key_too(self, tmp_path):
        valid = occurrence_entry(module_id="datacenter_fire")
        broken = occurrence_entry(module_id="datacenter_fire", basis="guessed")
        loaded = load(tmp_path, [valid, broken], version=7)
        assert loaded.entries == []
        assert any(
            "occurrence oc-000000000001 in module datacenter_fire is listed 2 times" in why
            for _, why in loaded.malformed
        )

    def test_a_copy_that_states_no_key_repeats_nothing(self, tmp_path):
        """A copy with no occurrence id has no key, so it cannot repeat one:
        the executable entry beside it stands, and the copy is refused for
        what it lacks."""
        loaded = load(tmp_path, [occurrence_entry(), occurrence_entry(occurrence_id=None)], version=6)
        assert [e.occurrence_id for e in loaded.entries] == ["oc-000000000001"]
        assert [why for _, why in loaded.malformed] == ["entry carries no occurrence_id"]


# ---------------------------------------------------------------------------
# End to end: every location is applied
# ---------------------------------------------------------------------------


def build_spec(path, paragraphs):
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    document.save(path)
    return path


_SPEC = [
    "SECTION 21 05 00",
    "Provide a gate valve at each branch line.",
    "Provide ball valves at each riser.",
    "Install a gate valve at the connection.",
]


class TestEveryLocationIsApplied:
    def test_the_same_fix_at_two_places_in_one_file_is_two_tracked_changes(self, tmp_path):
        build_spec(tmp_path / "spec.docx", _SPEC)
        sidecar = load(
            tmp_path,
            [
                occurrence_entry("oc-000000000001", element="p1"),
                occurrence_entry("oc-000000000003", element="p3"),
            ],
            version=6,
        )
        results = apply_sidecar(sidecar, [tmp_path / "spec.docx"], RunSettings())
        statuses = {o.entry.occurrence_id: o.status for r in results for o in r.outcomes}
        assert statuses == {
            "oc-000000000001": OutcomeStatus.APPLIED,
            "oc-000000000003": OutcomeStatus.APPLIED,
        }
        edited = Document(tmp_path / "spec.applied.docx")
        texts = [_accept_all_paragraph_text(p._p) for p in edited.paragraphs]
        assert texts[1] == "Provide a ball valve at each branch line."
        assert texts[3] == "Install a ball valve at the connection."
        assert texts[2] == _SPEC[2]
        assert extract_text_from_docx(tmp_path / "spec.applied.docx").tracked_changes_detected

    def test_one_addition_in_two_cells_of_a_row_is_two_new_paragraphs(self, tmp_path):
        """A row is one element with a paragraph per cell; the anchor chooses
        the cell. The occurrence model keeps these two apart (its test is
        ``test_additions_anchored_in_two_cells_of_one_row_are_two_occurrences``),
        and here they are two places for the applier too."""
        document = Document()
        document.add_paragraph("SECTION 21 05 00")
        row = document.add_table(rows=1, cols=2).rows[0]
        row.cells[0].text = "Gate valve, bronze"
        row.cells[1].text = "Check valve, swing"
        document.save(tmp_path / "spec.docx")
        element = next(
            m.element_id
            for m in extract_text_from_docx(tmp_path / "spec.docx").paragraph_map
            if "Gate valve" in m.text
        )
        text = "Provide a tamper switch."

        def cell_addition(occurrence_id, anchor):
            return occurrence_entry(
                occurrence_id,
                element=element,
                edit_proposal=proposal(
                    action_type="ADD",
                    existing_text=None,
                    replacement_text=text,
                    anchor_text=anchor,
                    insert_position="after",
                    target_element_id=element,
                ),
            )

        sidecar = load(
            tmp_path,
            [cell_addition("oc-000000000001", "Gate valve"), cell_addition("oc-000000000002", "Check valve")],
            version=6,
        )
        (result,) = apply_sidecar(sidecar, [tmp_path / "spec.docx"], RunSettings())
        assert [o.status for o in result.outcomes] == [OutcomeStatus.APPLIED] * 2
        cells = Document(tmp_path / "spec.applied.docx").tables[0].rows[0].cells
        assert [[_accept_all_paragraph_text(p._p) for p in cell.paragraphs] for cell in cells] == [
            ["Gate valve, bronze", text],
            ["Check valve, swing", text],
        ]

    def test_a_program_sidecar_applies_each_modules_places(self, tmp_path):
        build_spec(tmp_path / "spec.docx", _SPEC)
        sidecar = load(
            tmp_path,
            [
                occurrence_entry("oc-000000000001", element="p1", module_id="datacenter_fire"),
                occurrence_entry("oc-000000000003", element="p3", module_id="datacenter_electrical"),
            ],
            version=7,
            program_id="hyperscale_datacenter",
        )
        (result,) = apply_sidecar(sidecar, [tmp_path / "spec.docx"], RunSettings())
        assert result.applied == 2
        assert {o.entry.module_id for o in result.outcomes} == {"datacenter_fire", "datacenter_electrical"}

    def test_an_unresolved_place_is_found_by_its_text(self, tmp_path):
        build_spec(tmp_path / "spec.docx", ["Intro.", "Provide a gate valve here."])
        sidecar = load(tmp_path, [occurrence_entry(element=None, basis="unresolved")], version=6)
        (result,) = apply_sidecar(sidecar, [tmp_path / "spec.docx"], RunSettings())
        (outcome,) = result.outcomes
        assert outcome.status is OutcomeStatus.APPLIED
        assert outcome.location.status.value == "RESOLVED_BY_UNIQUE_TEXT"

    def test_an_unresolved_place_whose_text_repeats_is_refused(self, tmp_path):
        build_spec(tmp_path / "spec.docx", _SPEC)
        sidecar = load(tmp_path, [occurrence_entry(element=None, basis="unresolved")], version=6)
        (result,) = apply_sidecar(sidecar, [tmp_path / "spec.docx"], RunSettings())
        (outcome,) = result.outcomes
        assert outcome.status is OutcomeStatus.UNLOCATED
        assert outcome.location.status.value == "AMBIGUOUS"

    def test_a_missing_original_edit_is_located_by_text_alone(self, tmp_path):
        build_spec(tmp_path / "spec.docx", ["Intro.", "Provide a gate valve here."])
        sidecar = load(tmp_path, [occurrence_entry(element=None, basis="missing_original")], version=6)
        (result,) = apply_sidecar(sidecar, [tmp_path / "spec.docx"], RunSettings())
        assert result.outcomes[0].status is OutcomeStatus.APPLIED

    def test_the_receipt_carries_the_occurrence(self, tmp_path):
        build_spec(tmp_path / "spec.docx", _SPEC)
        sidecar = load(tmp_path, [occurrence_entry("oc-000000000001", element="p1")], version=6)
        results = apply_sidecar(sidecar, [tmp_path / "spec.docx"], RunSettings())
        receipt = build_receipt(
            sidecar=sidecar, file_results=results, settings=RunSettings().to_dict(), entries_in=1
        )
        (outcome,) = receipt["files"][0]["outcomes"]
        assert outcome["occurrence_id"] == "oc-000000000001"
        assert outcome["location_basis"] == "validated"
        assert outcome["related_entries"] == []
        assert receipt["sidecar"]["schema_version"] == 6
        assert receipt["accounting"]["balanced"] is True
        assert "[HIGH] rf-abc123 oc-000000000001 EDIT @p1" in render_summary(receipt, results)

    def test_the_cli_runs_a_schema_7_sidecar(self, tmp_path, capsys):
        build_spec(tmp_path / "spec.docx", _SPEC)
        path = write(
            tmp_path,
            [occurrence_entry(module_id="datacenter_fire")],
            version=7,
            program_id="hyperscale_datacenter",
        )
        assert main([str(path), "--specs", str(tmp_path), "--quiet"]) == EXIT_OK
        assert "(schema v7)" in capsys.readouterr().out


class TestTheKeyDecidesNothing:
    """S11's audit of ``EditEntry.key`` found no production consumer, and the
    documentation says conflicts are found from where entries resolve, never
    from a key. A key-based decision would treat two findings' edits at one
    place as unrelated, and one finding's edits at two places as one."""

    def test_no_applier_module_reads_the_key(self):
        import ast
        from pathlib import Path

        import applier

        readers = []
        for path in sorted(Path(applier.__file__).parent.glob("*.py")):
            tree = ast.parse(path.read_text(encoding="utf-8"))
            for node in ast.walk(tree):
                if isinstance(node, ast.Attribute) and node.attr == "key" and isinstance(node.ctx, ast.Load):
                    readers.append(f"{path.name}:{node.lineno}")
        # The reader's repeated-key rule is the one place a key is checked.
        assert readers and all(reader.startswith("sidecar.py:") for reader in readers), readers


class TestOnlySelectsAnOccurrence:
    def _run(self, tmp_path, only):
        build_spec(tmp_path / "spec.docx", _SPEC)
        path = write(
            tmp_path,
            [
                occurrence_entry("oc-000000000001", element="p1"),
                occurrence_entry("oc-000000000003", element="p3"),
            ],
            version=6,
        )
        main([str(path), "--specs", str(tmp_path), "--quiet", *sum((["--only", i] for i in only), [])])
        receipt = json.loads((tmp_path / "report.edits.applied.json").read_text(encoding="utf-8"))
        return {o["occurrence_id"]: o["outcome"] for o in receipt["files"][0]["outcomes"]}

    def test_a_finding_id_selects_every_place(self, tmp_path):
        assert set(self._run(tmp_path, ["rf-abc123"]).values()) == {"APPLIED"}

    def test_an_occurrence_id_selects_one_place(self, tmp_path):
        assert self._run(tmp_path, ["oc-000000000003"]) == {
            "oc-000000000001": "HELD_BY_POLICY",
            "oc-000000000003": "APPLIED",
        }
