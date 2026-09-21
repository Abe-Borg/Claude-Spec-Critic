"""The writer: tracked changes that Word can accept or reject, or nothing.

Two properties matter more than the rest and are asserted repeatedly:

* **Accept-all must equal the intended text.** Asserted with the extractor's
  own accept-all reader, which is the code that produced the text the review
  model saw — so a green test means Spec Critic would read the edited document
  as saying what the finding asked for.
* **Reject must restore the original.** A revision whose rejection leaves
  debris behind is worse than no revision at all, because the reviewer
  believes they undid it.
"""
from __future__ import annotations

import re

import pytest
from docx import Document
from docx.oxml.ns import qn

from applier.docx_edit import (
    DIRECT,
    TRACKED,
    DocumentEditor,
    EditError,
    REVISION_AUTHOR,
)
from applier.models import EditEntry, ElementKind, Location, LocationStatus
from src.input.extractor import _accept_all_paragraph_text, extract_text_from_docx


def entry(**overrides):
    base = dict(
        finding_id="rf-1",
        file_name="215000.docx",
        action_type="EDIT",
        existing_text=None,
        replacement_text=None,
        anchor_text=None,
        insert_position=None,
        target_element_id=None,
        evidence_element_id=None,
        edit_confidence=0.9,
        report_status="VERIFIED_SUPPORTED",
        verification_verdict="CONFIRMED",
        severity="HIGH",
        section="21 13 13",
        issue="…",
        code_reference=None,
        has_per_file_original=True,
    )
    base.update(overrides)
    return EditEntry(**base)


def at(element_id, kind=ElementKind.BODY_PARAGRAPH):
    return Location(
        status=LocationStatus.RESOLVED_BY_ID, element_id=element_id, kind=kind
    )


@pytest.fixture
def document():
    """A small spec whose key clause is split across three runs, as Word does."""
    doc = Document()
    doc.add_paragraph("SECTION 21 13 13")
    paragraph = doc.add_paragraph("Sprinklers shall be listed per ")
    paragraph.add_run("NFPA 13").bold = True
    paragraph.add_run(", 2019 edition.")
    doc.add_paragraph("Provide hangers per manufacturer instructions.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Hazard"
    table.cell(0, 1).text = "Density"
    table.cell(1, 0).text = "Ordinary Group 1"
    table.cell(1, 1).text = "0.15 gpm/sf"
    return doc


def accept_all(document, index):
    return _accept_all_paragraph_text(document.paragraphs[index]._p)


def reject_view(document, index):
    """What survives a Reject All: ``w:t`` outside ``w:ins``, plus ``w:delText``."""
    parts = []
    for node in document.paragraphs[index]._p.iter():
        if node.tag == qn("w:delText"):
            parts.append(node.text or "")
        elif node.tag == qn("w:t"):
            inside_insert = any(
                ancestor.tag == qn("w:ins") for ancestor in node.iterancestors()
            )
            if not inside_insert:
                parts.append(node.text or "")
    return "".join(parts)


class TestTrackedEdit:
    def test_accept_all_reads_as_the_replacement(self, document):
        DocumentEditor(document, mode=TRACKED).apply(
            entry(
                existing_text="NFPA 13, 2019 edition",
                replacement_text="NFPA 13, 2025 edition as amended by California",
            ),
            at("p1"),
        )
        assert accept_all(document, 1) == (
            "Sprinklers shall be listed per NFPA 13, 2025 edition as amended "
            "by California."
        )

    def test_reject_restores_the_original_exactly(self, document):
        original = accept_all(document, 1)
        DocumentEditor(document, mode=TRACKED).apply(
            entry(
                existing_text="NFPA 13, 2019 edition",
                replacement_text="NFPA 13, 2025 edition",
            ),
            at("p1"),
        )
        assert reject_view(document, 1) == original

    def test_only_the_matched_characters_are_touched(self, document):
        DocumentEditor(document, mode=TRACKED).apply(
            entry(existing_text="2019", replacement_text="2025"), at("p1")
        )
        deleted = [
            node.text
            for node in document.paragraphs[1]._p.iter()
            if node.tag == qn("w:delText")
        ]
        assert deleted == ["2019"]

    def test_the_revision_is_attributed(self, document):
        DocumentEditor(document, mode=TRACKED).apply(
            entry(existing_text="2019", replacement_text="2025"), at("p1")
        )
        xml = document.paragraphs[1]._p.xml
        assert REVISION_AUTHOR in xml
        assert "<w:del " in xml and "<w:ins " in xml

    def test_revision_ids_are_unique_within_a_run(self, document):
        editor = DocumentEditor(document, mode=TRACKED)
        editor.apply(entry(existing_text="2019", replacement_text="2025"), at("p1"))
        editor.apply(
            entry(existing_text="hangers", replacement_text="seismic hangers"), at("p2")
        )
        ids = re.findall(r'w:id="(\d+)"', document.element.body.xml)
        assert len(ids) == len(set(ids))

    def test_the_replacement_inherits_the_formatting_it_replaces(self, document):
        DocumentEditor(document, mode=TRACKED).apply(
            entry(existing_text="NFPA 13", replacement_text="NFPA 13 (2025)"), at("p1")
        )
        insertion = re.search(
            r"<w:ins\b.*?</w:ins>", document.paragraphs[1]._p.xml, re.S
        ).group(0)
        assert "<w:b" in insertion, "the bold run's formatting was dropped"

    def test_whitespace_tolerant_matching_reaches_a_wrapped_clause(self):
        doc = Document()
        paragraph = doc.add_paragraph("Comply with the ")
        paragraph.add_run("2019")
        paragraph.add_run(" edition of NFPA 13.")
        DocumentEditor(doc, mode=TRACKED).apply(
            entry(
                existing_text="2019   edition of NFPA 13",
                replacement_text="2025 edition of NFPA 13",
            ),
            at("p0"),
        )
        assert accept_all(doc, 0) == "Comply with the 2025 edition of NFPA 13."


class TestTrackedDelete:
    def test_accept_all_drops_the_text(self, document):
        DocumentEditor(document, mode=TRACKED).apply(
            entry(action_type="DELETE", existing_text=", 2019 edition"), at("p1")
        )
        assert accept_all(document, 1) == "Sprinklers shall be listed per NFPA 13."

    def test_reject_restores_it(self, document):
        original = accept_all(document, 1)
        DocumentEditor(document, mode=TRACKED).apply(
            entry(action_type="DELETE", existing_text=", 2019 edition"), at("p1")
        )
        assert reject_view(document, 1) == original

    def test_no_insertion_is_written_for_a_delete(self, document):
        DocumentEditor(document, mode=TRACKED).apply(
            entry(action_type="DELETE", existing_text=", 2019 edition"), at("p1")
        )
        assert "<w:ins " not in document.paragraphs[1]._p.xml


class TestTrackedAdd:
    def test_the_new_paragraph_lands_after_the_anchor(self, document):
        DocumentEditor(document, mode=TRACKED).apply(
            entry(
                action_type="ADD",
                replacement_text="Sprinklers shall be quick-response type.",
                anchor_text="Provide hangers",
                insert_position="after",
            ),
            at("p2"),
        )
        texts = [_accept_all_paragraph_text(p._p) for p in document.paragraphs]
        assert texts[2] == "Provide hangers per manufacturer instructions."
        assert texts[3] == "Sprinklers shall be quick-response type."

    def test_before_is_honoured(self, document):
        DocumentEditor(document, mode=TRACKED).apply(
            entry(
                action_type="ADD",
                replacement_text="New requirement.",
                anchor_text="Provide hangers",
                insert_position="before",
            ),
            at("p2"),
        )
        texts = [_accept_all_paragraph_text(p._p) for p in document.paragraphs]
        assert texts[2] == "New requirement."
        assert texts[3] == "Provide hangers per manufacturer instructions."

    def test_the_paragraph_mark_is_itself_marked_inserted(self, document):
        """Otherwise rejecting the change leaves an empty numbered paragraph."""
        DocumentEditor(document, mode=TRACKED).apply(
            entry(
                action_type="ADD",
                replacement_text="New requirement.",
                anchor_text="Provide hangers",
                insert_position="after",
            ),
            at("p2"),
        )
        xml = document.paragraphs[3]._p.xml
        assert re.search(r"<w:pPr>.*?<w:rPr>.*?<w:ins\b", xml, re.S)

    def test_the_new_paragraph_inherits_the_anchors_style(self):
        doc = Document()
        doc.add_paragraph("Body text")
        doc.add_paragraph("Anchor clause here.", style="List Bullet")
        DocumentEditor(doc, mode=TRACKED).apply(
            entry(
                action_type="ADD",
                replacement_text="Added clause.",
                anchor_text="Anchor clause",
                insert_position="after",
            ),
            at("p1"),
        )
        assert doc.paragraphs[2].style.name == "List Bullet"


class TestDirectMode:
    def test_text_is_replaced_with_no_revision_marks(self, document):
        DocumentEditor(document, mode=DIRECT).apply(
            entry(existing_text="2019 edition", replacement_text="2025 edition"),
            at("p1"),
        )
        xml = document.paragraphs[1]._p.xml
        assert "<w:ins " not in xml and "<w:del " not in xml
        assert document.paragraphs[1].text == (
            "Sprinklers shall be listed per NFPA 13, 2025 edition."
        )

    def test_delete_removes_the_text(self, document):
        DocumentEditor(document, mode=DIRECT).apply(
            entry(action_type="DELETE", existing_text=", 2019 edition"), at("p1")
        )
        assert document.paragraphs[1].text == "Sprinklers shall be listed per NFPA 13."


class TestTables:
    def test_a_cell_inside_a_row_is_edited(self, document):
        DocumentEditor(document, mode=TRACKED).apply(
            entry(existing_text="0.15 gpm/sf", replacement_text="0.20 gpm/sf"),
            at("t0r1", ElementKind.TABLE_ROW),
        )
        cell = document.tables[0].cell(1, 1)
        assert _accept_all_paragraph_text(cell.paragraphs[0]._p) == "0.20 gpm/sf"

    def test_the_other_cells_are_untouched(self, document):
        before = document.tables[0].cell(1, 0).text
        DocumentEditor(document, mode=TRACKED).apply(
            entry(existing_text="0.15 gpm/sf", replacement_text="0.20 gpm/sf"),
            at("t0r1", ElementKind.TABLE_ROW),
        )
        assert document.tables[0].cell(1, 0).text == before

    def test_a_missing_row_is_reported_as_change_not_crash(self, document):
        with pytest.raises(EditError) as excinfo:
            DocumentEditor(document).apply(
                entry(existing_text="x", replacement_text="y"),
                at("t0r9", ElementKind.TABLE_ROW),
            )
        assert "changed since the review" in str(excinfo.value)


class TestRefusals:
    def test_a_split_boundary_inside_a_tab_run_is_refused(self):
        doc = Document()
        doc.add_paragraph().add_run("Zone\tA-1 remote area")
        with pytest.raises(EditError) as excinfo:
            DocumentEditor(doc).apply(
                entry(existing_text="A-1", replacement_text="A-2"), at("p0")
            )
        assert "by hand" in str(excinfo.value)

    def test_an_out_of_range_paragraph_is_refused(self, document):
        with pytest.raises(EditError) as excinfo:
            DocumentEditor(document).apply(
                entry(existing_text="x", replacement_text="y"), at("p999")
            )
        assert "changed since the review" in str(excinfo.value)

    def test_an_id_that_is_now_a_table_is_refused(self, document):
        # p3 is the table in this fixture, not a paragraph.
        with pytest.raises(EditError) as excinfo:
            DocumentEditor(document).apply(
                entry(existing_text="Hazard", replacement_text="Occupancy"), at("p3")
            )
        assert "no longer a paragraph" in str(excinfo.value)

    def test_text_absent_from_the_located_element_is_refused(self, document):
        with pytest.raises(EditError) as excinfo:
            DocumentEditor(document).apply(
                entry(existing_text="NFPA 25", replacement_text="NFPA 13"), at("p1")
            )
        assert "not found in the located element" in str(excinfo.value)

    def test_a_malformed_id_is_refused(self, document):
        with pytest.raises(EditError):
            DocumentEditor(document).apply(
                entry(existing_text="x", replacement_text="y"),
                at("nonsense", ElementKind.BODY_PARAGRAPH),
            )

    def test_an_unsupported_kind_is_refused(self, document):
        with pytest.raises(EditError) as excinfo:
            DocumentEditor(document).apply(
                entry(existing_text="x", replacement_text="y"),
                at("tb0p1", ElementKind.UNSUPPORTED),
            )
        assert "does not write to" in str(excinfo.value)


class TestRepeatedTargetsAreRefused:
    """An element id names a paragraph or a row — never which occurrence
    inside it. Taking the first match would silently edit the wrong clause,
    irreversibly under --mode direct. This is the same ambiguity the locator
    refuses one level up."""

    def test_two_occurrences_in_one_paragraph(self):
        doc = Document()
        doc.add_paragraph("Install per NFPA 13 and test per NFPA 13 yearly.")
        with pytest.raises(EditError) as excinfo:
            DocumentEditor(doc).apply(
                entry(existing_text="per NFPA 13", replacement_text="per NFPA 13 (2025)"),
                at("p0"),
            )
        assert "occurs 2 times" in str(excinfo.value)
        assert "by hand" in str(excinfo.value)

    def test_two_occurrences_across_one_table_rows_cells(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Density 0.15 gpm/sf"
        table.cell(0, 1).text = "Remote area 0.15 gpm/sf"
        with pytest.raises(EditError) as excinfo:
            DocumentEditor(doc).apply(
                entry(existing_text="0.15 gpm/sf", replacement_text="0.20 gpm/sf"),
                at("t0r0", ElementKind.TABLE_ROW),
            )
        assert "occurs 2 times" in str(excinfo.value)

    def test_whitespace_variants_count_as_the_same_occurrence(self):
        """Counting normalized is the conservative choice — it can only ever
        find more candidates, and finding more is what makes this refuse."""
        doc = Document()
        doc.add_paragraph("Install per NFPA 13 and test per  NFPA   13 yearly.")
        with pytest.raises(EditError):
            DocumentEditor(doc).apply(
                entry(existing_text="per NFPA 13", replacement_text="x"), at("p0")
            )

    def test_a_single_occurrence_still_applies(self):
        doc = Document()
        doc.add_paragraph("Install per NFPA 13 and test annually.")
        DocumentEditor(doc, mode=TRACKED).apply(
            entry(existing_text="per NFPA 13", replacement_text="per NFPA 13 (2025)"),
            at("p0"),
        )
        assert accept_all(doc, 0) == "Install per NFPA 13 (2025) and test annually."

    def test_an_ambiguous_add_anchor_is_refused_too(self):
        doc = Document()
        doc.add_paragraph("Provide hangers. Provide hangers again.")
        with pytest.raises(EditError) as excinfo:
            DocumentEditor(doc).apply(
                entry(
                    action_type="ADD",
                    replacement_text="New clause.",
                    anchor_text="Provide hangers",
                    insert_position="after",
                ),
                at("p0"),
            )
        assert "occurs 2 times" in str(excinfo.value)


class TestNestedTableIndexingMatchesTheExtractor:
    """The ``cN`` step of a nested-table id counts cells AFTER the extractor's
    merge dedup, and that dedup is stateful across the whole table. Indexing
    raw ``row.cells`` drifts from the id on any merged table."""

    @pytest.fixture
    def merged(self, tmp_path):
        doc = Document()
        doc.add_paragraph("SECTION 21 13 13")
        table = doc.add_table(rows=2, cols=3)
        first = table.rows[0]
        first.cells[0].merge(first.cells[1])
        first.cells[0].text = "Merged heading"
        nested = first.cells[2].add_table(rows=1, cols=1)
        nested.cell(0, 0).text = "Density 0.15 gpm/sf"
        table.cell(1, 0).text = "A"
        table.cell(1, 1).text = "B"
        table.cell(1, 2).text = "C"
        path = tmp_path / "merged.docx"
        doc.save(path)
        return path

    def test_the_id_the_extractor_mints_resolves_to_the_nested_cell(self, merged):
        extracted = extract_text_from_docx(merged)
        nested_id = next(
            mapping.element_id
            for mapping in extracted.paragraph_map
            if "Density" in mapping.text
        )
        # Deduped index 1, though raw row.cells puts that cell at index 2.
        assert nested_id == "t0r0c1t0r0"

        document = Document(merged)
        DocumentEditor(document, mode=TRACKED).apply(
            entry(existing_text="0.15 gpm/sf", replacement_text="0.20 gpm/sf"),
            at(nested_id, ElementKind.TABLE_ROW),
        )
        nested_cell = document.tables[0].rows[0].cells[2].tables[0].cell(0, 0)
        assert _accept_all_paragraph_text(nested_cell.paragraphs[0]._p) == (
            "Density 0.20 gpm/sf"
        )

    def test_the_raw_index_would_have_pointed_at_the_merged_duplicate(self, merged):
        """Pins why the dedup is needed: raw indexing selects a different cell."""
        document = Document(merged)
        raw_cells = document.tables[0].rows[0].cells
        assert len(raw_cells) == 3
        assert raw_cells[1]._tc is raw_cells[0]._tc  # the merged duplicate
        assert raw_cells[1].tables == []  # ... and it holds no nested table

    def test_a_merged_row_collects_each_cell_once(self, merged):
        document = Document(merged)
        paragraphs = DocumentEditor(document).resolve_paragraphs(
            "t0r0", ElementKind.TABLE_ROW
        )
        texts = [_accept_all_paragraph_text(p) for p in paragraphs]
        assert texts.count("Merged heading") == 1


class TestBatchOrdering:
    """Element ids are positional, so an insertion renumbers everything after
    it. Resolving every edit before applying any is what keeps a batch
    landing where the locator said it would."""

    def test_an_insertion_does_not_displace_a_later_edit(self, document):
        editor = DocumentEditor(document, mode=TRACKED)
        add = entry(
            action_type="ADD",
            replacement_text="Inserted clause.",
            anchor_text="Sprinklers shall be listed",
            insert_position="after",
        )
        later = entry(
            existing_text="Provide hangers per manufacturer instructions.",
            replacement_text="Provide seismic hangers per NFPA 13.",
        )
        resolved_add = editor.resolve(at("p1"))
        resolved_later = editor.resolve(at("p2"))

        editor.apply_resolved(add, resolved_add)
        editor.apply_resolved(later, resolved_later)

        texts = [_accept_all_paragraph_text(p._p) for p in document.paragraphs]
        assert texts[2] == "Inserted clause."
        assert texts[3] == "Provide seismic hangers per NFPA 13."

    def test_resolving_lazily_after_an_insertion_would_hit_the_wrong_paragraph(
        self, document
    ):
        """The failure the two-pass design exists to prevent, pinned so the
        design cannot be 'simplified' back into the bug."""
        editor = DocumentEditor(document, mode=TRACKED)
        editor.apply(
            entry(
                action_type="ADD",
                replacement_text="Inserted clause.",
                anchor_text="Sprinklers shall be listed",
                insert_position="after",
            ),
            at("p1"),
        )
        # p2 was "Provide hangers…" before the insertion; it is now the new
        # paragraph, so a lazy resolve targets the wrong clause.
        assert _accept_all_paragraph_text(
            editor.resolve(at("p2"))[0]
        ) == "Inserted clause."


class TestSourceIsNeverTouched:
    def test_the_editor_only_mutates_the_in_memory_document(self, tmp_path):
        doc = Document()
        doc.add_paragraph("Comply with NFPA 13, 2019 edition.")
        path = tmp_path / "215000.docx"
        doc.save(path)
        before = path.read_bytes()

        opened = Document(path)
        DocumentEditor(opened, mode=TRACKED).apply(
            entry(existing_text="2019", replacement_text="2025"), at("p0")
        )
        assert path.read_bytes() == before

        output = tmp_path / "215000.applied.docx"
        opened.save(output)
        assert path.read_bytes() == before
        assert extract_text_from_docx(output).paragraph_map[0].text == (
            "Comply with NFPA 13, 2025 edition."
        )
        assert extract_text_from_docx(output).tracked_changes_detected is True
