"""Self-tests for the shared DOCX fixture builders (``tests/fixtures/spec_docx.py``).

Implementation plan WP-01 (chunk S01). Later chunks judge detectors,
extraction, routing, and the applier against these documents, so the
documents themselves are pinned here first: a builder that silently stopped
emitting a content control, or a "clean" fixture that quietly acquired a
second defect, would make every downstream test prove less than it claims.

What is checked is the *document*, never the code under test: the XML shape
after a save-and-reopen round trip, determinism, and the ground truth each
fixture declares (which paragraphs are headings, which single defect a
mutation introduces). The ground truth is computed from the fixture's block
model — its declared roles — not by parsing text, so it cannot share a bug
with the heading detector it will be used to judge.
"""
from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from tests.fixtures import spec_docx as fx

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
R_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"

_WRAPPER_BUILDERS = {
    "block_control": fx.build_block_control_spec,
    "inline_controls": fx.build_inline_controls_spec,
    "fields": fx.build_fields_spec,
    "smart_tag": fx.build_smart_tag_spec,
    "hyperlinks": fx.build_hyperlink_spec,
    "tracked_changes": fx.build_tracked_changes_spec,
    "merged_and_nested_tables": fx.build_merged_and_nested_tables_spec,
}

_ALL_BUILDERS = {
    **{
        f"three_part:{variant.name}": (lambda blocks=variant.blocks: fx.build_blocks(blocks))
        for variant in fx.three_part_variants()
    },
    **_WRAPPER_BUILDERS,
}


def _part_xml(path: Path, member: str) -> bytes:
    with zipfile.ZipFile(path) as archive:
        return archive.read(member)


def _saved_body(builder, tmp_path: Path, name: str = "fixture.docx"):
    """Save, reopen with python-docx, and return the reopened body element.

    Every structural assertion runs on the *reopened* document, so a
    fragment that only looked right in memory (for example one python-docx
    would drop or reorder on save) fails here.
    """
    path = fx.save_docx(builder, tmp_path, name)
    return Document(path).element.body, path


# ---------------------------------------------------------------------------
# Determinism
# ---------------------------------------------------------------------------


class TestDeterminism:
    @pytest.mark.parametrize("name", sorted(_ALL_BUILDERS))
    def test_building_twice_gives_identical_parts(self, name, tmp_path):
        build = _ALL_BUILDERS[name]
        first = fx.save_docx(build(), tmp_path / "a", "fixture.docx")
        second = fx.save_docx(build(), tmp_path / "b", "fixture.docx")
        for member in (
            "word/document.xml",
            "word/numbering.xml",
            "word/_rels/document.xml.rels",
        ):
            assert _part_xml(first, member) == _part_xml(second, member), member

    def test_ids_are_per_document_not_global(self):
        # Revision and control ids come from the builder, so building an
        # unrelated document first cannot shift them.
        fx.build_tracked_changes_spec()
        fx.build_inline_controls_spec()
        after_others = etree.tostring(fx.build_tracked_changes_spec().document.element)
        alone = etree.tostring(fx.build_tracked_changes_spec().document.element)
        assert after_others == alone


# ---------------------------------------------------------------------------
# Clean fixtures versus single-defect mutations
# ---------------------------------------------------------------------------


def _structural_truth(blocks) -> list[tuple[str, str]]:
    """The empty and duplicated headings in ``blocks``, from declared roles.

    A heading's subtree runs to the next heading of the same or a higher
    level (``part`` > ``article``); it is empty when the subtree holds no
    body paragraph and no table. An empty heading is reported only when its
    parent heading is not empty (the ancestor-versus-leaf policy S03 defined
    for plan WP-04A): a PART whose articles are all empty is one defect, the
    PART's. A heading repeated verbatim is a duplicate for every occurrence
    after the first.
    """
    rank = {"part": 0, "article": 1}
    headings = [(index, block) for index, block in enumerate(blocks) if block.is_heading]
    empty: dict[int, bool] = {}
    for index, block in headings:
        has_content = False
        for later in blocks[index + 1 :]:
            if later.is_heading and rank[later.role] <= rank[block.role]:
                break
            if not later.is_heading:
                has_content = True
                break
        empty[index] = not has_content
    defects: list[tuple[str, str]] = []
    seen: set[str] = set()
    for index, block in headings:
        parent = next(
            (
                earlier
                for earlier, candidate in reversed(headings)
                if earlier < index and rank[candidate.role] < rank[block.role]
            ),
            None,
        )
        if empty[index] and (parent is None or not empty[parent]):
            defects.append(("empty_section", block.displayed_text))
        if block.displayed_text in seen:
            defects.append(("duplicate_heading", block.displayed_text))
        seen.add(block.displayed_text)
    return defects


class TestCleanAndMutationSeparation:
    @pytest.mark.parametrize(
        "variant",
        [v for v in fx.three_part_variants() if v.clean],
        ids=lambda v: v.name,
    )
    def test_clean_variants_have_no_structural_defect(self, variant):
        assert _structural_truth(variant.blocks) == []

    @pytest.mark.parametrize(
        "variant",
        [v for v in fx.three_part_variants() if not v.clean],
        ids=lambda v: v.name,
    )
    def test_each_mutation_carries_exactly_its_named_defect(self, variant):
        assert _structural_truth(variant.blocks) == [
            (variant.expected_rule, variant.expected_match)
        ]

    @pytest.mark.parametrize(
        "variant",
        [v for v in fx.three_part_variants() if not v.clean],
        ids=lambda v: v.name,
    )
    def test_each_mutation_differs_from_the_clean_fixture_by_one_block(self, variant):
        changes = fx.changed_blocks(fx.clean_three_part_blocks(), variant.blocks)
        assert len(changes) == 1, changes

    def test_the_empty_article_mutation_removes_only_that_body_paragraph(self):
        ((kind, old, new),) = fx.changed_blocks(
            fx.clean_three_part_blocks(), fx.empty_article_blocks()
        )
        assert kind == "removed"
        assert old == fx.Para("A. Submit product data before fabrication.", "body")

    def test_the_duplicate_mutation_repeats_a_heading_verbatim(self):
        ((kind, old, new),) = fx.changed_blocks(
            fx.clean_three_part_blocks(), fx.duplicate_heading_blocks()
        )
        assert kind == "replaced"
        assert old.text == "1.02 SUBMITTALS"
        assert new == fx.clean_three_part_blocks()[1]  # 1.01 SUMMARY

    def test_the_table_only_variant_swaps_one_body_paragraph_for_a_table(self):
        ((kind, old, new),) = fx.changed_blocks(
            fx.clean_three_part_blocks(), fx.table_only_article_blocks()
        )
        assert kind == "replaced"
        assert old.role == "body"
        assert isinstance(new, fx.TableBlock)
        # No cell starts with a digit, so the variant cannot also exercise the
        # separate rule that a quantity line is body text, not a heading.
        assert not any(cell[:1].isdigit() for row in new.rows for cell in row)

    def test_the_clean_fixture_is_plan_appendix_a_verbatim(self):
        assert [block.text for block in fx.clean_three_part_blocks()] == [
            "PART 1 GENERAL",
            "1.01 SUMMARY",
            "A. Provide the specified piping system.",
            "1.02 SUBMITTALS",
            "A. Submit product data before fabrication.",
            "PART 2 PRODUCTS",
            "2.01 MATERIALS",
            "A. Provide materials meeting the scheduled requirements.",
            "PART 3 EXECUTION",
            "3.01 INSTALLATION",
            "A. Install in accordance with the approved product instructions.",
        ]

    def test_variant_names_are_unique(self):
        names = [variant.name for variant in fx.three_part_variants()]
        assert len(names) == len(set(names))

    def test_with_section_heading_prefixes_only_the_heading(self):
        blocks = fx.with_section_heading(
            fx.clean_three_part_blocks(), "21 05 00", "COMMON WORK RESULTS"
        )
        assert [b.text for b in blocks[:2]] == ["SECTION 21 05 00", "COMMON WORK RESULTS"]
        assert blocks[2:] == fx.clean_three_part_blocks()


# ---------------------------------------------------------------------------
# Automatic numbering
# ---------------------------------------------------------------------------


class TestAutomaticNumbering:
    def test_displayed_text_is_the_clean_fixture(self):
        assert fx.displayed_text(fx.auto_numbered_blocks()) == fx.blocks_text(
            fx.clean_three_part_blocks()
        )

    def test_no_label_is_typed_into_the_literal_text(self):
        for block in fx.auto_numbered_blocks():
            assert block.auto_label is not None
            assert block.auto_label not in block.text

    def test_every_paragraph_is_numbered_at_its_role_level(self, tmp_path):
        body, _ = _saved_body(fx.build_auto_numbered_three_part(), tmp_path)
        paragraphs = body.findall(qn("w:p"))
        blocks = fx.auto_numbered_blocks()
        assert len(paragraphs) == len(blocks)
        num_ids = set()
        for p_el, block in zip(paragraphs, blocks):
            num_pr = p_el.find(f"{qn('w:pPr')}/{qn('w:numPr')}")
            assert num_pr is not None, block.text
            assert int(num_pr.find(qn("w:ilvl")).get(qn("w:val"))) == block.level
            num_ids.add(num_pr.find(qn("w:numId")).get(qn("w:val")))
        assert len(num_ids) == 1  # one list instance for the whole document

    def test_the_list_definition_renders_csi_labels(self, tmp_path):
        """``PART %1`` / ``%1.%2`` zero-padded / ``%3.`` upper letter, each
        followed by a space — the definition that displays ``1.01 SUMMARY``."""
        path = fx.save_docx(fx.build_auto_numbered_three_part(), tmp_path, "auto.docx")
        numbering = etree.fromstring(_part_xml(path, "word/numbering.xml"))
        body = Document(path).element.body
        num_id = body.find(f".//{qn('w:numId')}").get(qn("w:val"))
        (num,) = [n for n in numbering.findall(qn("w:num")) if n.get(qn("w:numId")) == num_id]
        abstract_id = num.find(qn("w:abstractNumId")).get(qn("w:val"))
        (abstract,) = [
            a
            for a in numbering.findall(qn("w:abstractNum"))
            if a.get(qn("w:abstractNumId")) == abstract_id
        ]
        levels = {
            int(lvl.get(qn("w:ilvl"))): (
                lvl.find(qn("w:numFmt")).get(qn("w:val")),
                lvl.find(qn("w:lvlText")).get(qn("w:val")),
                lvl.find(qn("w:suff")).get(qn("w:val")),
                lvl.find(qn("w:start")).get(qn("w:val")),
            )
            for lvl in abstract.findall(qn("w:lvl"))
        }
        assert levels == {
            0: ("decimal", "PART %1", "space", "1"),
            1: ("decimalZero", "%1.%2", "space", "1"),
            2: ("upperLetter", "%3.", "space", "1"),
        }

    def test_numbering_part_keeps_schema_order(self, tmp_path):
        # Every w:abstractNum must precede every w:num, or Word rejects it.
        path = fx.save_docx(fx.build_auto_numbered_three_part(), tmp_path, "auto.docx")
        numbering = etree.fromstring(_part_xml(path, "word/numbering.xml"))
        tags = [child.tag for child in numbering if child.tag in (qn("w:abstractNum"), qn("w:num"))]
        first_num = tags.index(qn("w:num"))
        assert qn("w:abstractNum") not in tags[first_num:]

    def test_the_new_list_does_not_reuse_a_template_id(self):
        template = Document().part.numbering_part.element
        template_abstract = {a.get(qn("w:abstractNumId")) for a in template.findall(qn("w:abstractNum"))}
        template_num = {n.get(qn("w:numId")) for n in template.findall(qn("w:num"))}
        numbering = fx.build_auto_numbered_three_part().document.part.numbering_part.element
        added_nums = {n.get(qn("w:numId")) for n in numbering.findall(qn("w:num"))} - template_num
        added_abstracts = {
            a.get(qn("w:abstractNumId")) for a in numbering.findall(qn("w:abstractNum"))
        } - template_abstract
        assert len(added_nums) == 1 and len(added_abstracts) == 1


# ---------------------------------------------------------------------------
# Wrapped Word content: each sentinel sits in exactly the intended container
# ---------------------------------------------------------------------------


def _all_text_nodes(root) -> list[str]:
    return [node.text or "" for node in root.iter(qn("w:t"))]


def _count_in_text_nodes(root, needle: str) -> int:
    return sum(text.count(needle) for text in _all_text_nodes(root))


class TestWrappedContentShape:
    def test_block_control_wraps_a_paragraph_and_a_table(self, tmp_path):
        body, _ = _saved_body(fx.build_block_control_spec(), tmp_path)
        children = [child.tag for child in body]
        assert children == [
            qn("w:p"), qn("w:sdt"), qn("w:p"), qn("w:tbl"), qn("w:sectPr")
        ]
        content = body[1].find(qn("w:sdtContent"))
        assert [child.tag for child in content] == [qn("w:p"), qn("w:tbl")]
        assert fx.BLOCK_CONTROL_PARAGRAPH in "".join(_all_text_nodes(content[0]))
        assert _count_in_text_nodes(body, fx.BLOCK_CONTROL_PARAGRAPH) == 1
        assert body[1].find(f"{qn('w:sdtPr')}/{qn('w:tag')}").get(qn("w:val")) == "backflow_block"

    def test_block_control_can_omit_its_table(self, tmp_path):
        body, _ = _saved_body(fx.build_block_control_spec(wrapped_table=False), tmp_path)
        content = body[1].find(qn("w:sdtContent"))
        assert [child.tag for child in content] == [qn("w:p")]

    def test_inline_control_sits_between_ordinary_runs(self, tmp_path):
        body, _ = _saved_body(fx.build_inline_controls_spec(), tmp_path)
        p0 = body.findall(qn("w:p"))[0]
        assert [child.tag for child in p0] == [qn("w:r"), qn("w:sdt"), qn("w:r")]
        sdt_text = "".join(_all_text_nodes(p0[1]))
        assert sdt_text == fx.INLINE_CONTROL_TEXT
        assert _count_in_text_nodes(body, fx.INLINE_CONTROL_TEXT) == 1

    def test_dropdowns_list_their_items_and_store_the_shown_text(self, tmp_path):
        body, _ = _saved_body(fx.build_inline_controls_spec(), tmp_path)
        chosen, unresolved = body.findall(qn("w:p"))[1:3]
        chosen_sdt = chosen.find(qn("w:sdt"))
        items = chosen_sdt.findall(f".//{qn('w:listItem')}")
        assert [i.get(qn("w:displayText")) for i in items] == [t for t, _ in fx.DROPDOWN_ITEMS]
        assert "".join(_all_text_nodes(chosen_sdt.find(qn("w:sdtContent")))) == fx.DROPDOWN_CHOSEN
        assert chosen_sdt.find(f"{qn('w:sdtPr')}/{qn('w:showingPlcHdr')}") is None
        unresolved_sdt = unresolved.find(qn("w:sdt"))
        assert unresolved_sdt.find(f"{qn('w:sdtPr')}/{qn('w:showingPlcHdr')}") is not None
        assert (
            "".join(_all_text_nodes(unresolved_sdt.find(qn("w:sdtContent"))))
            == fx.UNRESOLVED_DROPDOWN_PLACEHOLDER
        )

    def test_control_ids_are_unique_within_the_document(self, tmp_path):
        body, _ = _saved_body(fx.build_inline_controls_spec(), tmp_path)
        ids = [el.get(qn("w:val")) for el in body.iter(qn("w:id")) if el.getparent().tag == qn("w:sdtPr")]
        assert len(ids) == 3 and len(set(ids)) == 3

    def test_simple_field_stores_its_result_and_instruction(self, tmp_path):
        body, _ = _saved_body(fx.build_fields_spec(), tmp_path)
        field = body.find(f".//{qn('w:fldSimple')}")
        assert field.get(qn("w:instr")) == fx.REF_FIELD_INSTRUCTION
        assert "".join(_all_text_nodes(field)) == fx.REF_FIELD_RESULT

    def test_complex_field_has_begin_separate_end_and_code(self, tmp_path):
        body, _ = _saved_body(fx.build_fields_spec(), tmp_path)
        p1 = body.findall(qn("w:p"))[1]
        kinds = [el.get(qn("w:fldCharType")) for el in p1.iter(qn("w:fldChar"))]
        assert kinds == ["begin", "separate", "end"]
        (instr,) = list(p1.iter(qn("w:instrText")))
        assert instr.text == fx.COMPLEX_FIELD_INSTRUCTION
        # The instruction is field code, never a w:t text node.
        assert _count_in_text_nodes(body, "REF sec_211313") == 0

    def test_smart_tag_wraps_its_text(self, tmp_path):
        body, _ = _saved_body(fx.build_smart_tag_spec(), tmp_path)
        tag = body.find(f".//{qn('w:smartTag')}")
        assert "".join(_all_text_nodes(tag)) == fx.SMART_TAG_TEXT
        assert _count_in_text_nodes(body, fx.SMART_TAG_TEXT) == 1

    def test_hyperlinks_point_at_a_real_external_relationship(self, tmp_path):
        body, path = _saved_body(fx.build_hyperlink_spec(), tmp_path)
        links = list(body.iter(qn("w:hyperlink")))
        assert len(links) == 2
        rels = Document(path).part.rels
        for link in links:
            rel = rels[link.get(R_ID)]
            assert rel.is_external and rel.target_ref == fx.HYPERLINK_URL
        # The second link carries a tracked insertion *inside* the hyperlink.
        inserted = links[1].find(qn("w:ins"))
        assert inserted is not None
        assert "".join(_all_text_nodes(inserted)) == fx.HYPERLINK_INSERTED_TEXT

    def test_tracked_changes_cover_every_revision_kind(self, tmp_path):
        body, _ = _saved_body(fx.build_tracked_changes_spec(), tmp_path)
        for tag in ("w:ins", "w:del", "w:moveFrom", "w:moveTo"):
            found = list(body.iter(qn(tag)))
            assert found, tag
            for revision in found:
                assert revision.get(qn("w:author")) == fx.REVISION_AUTHOR
                assert revision.get(qn("w:date")) == fx.REVISION_DATE
        revision_ids = [
            el.get(qn("w:id"))
            for tag in ("w:ins", "w:del", "w:moveFrom", "w:moveTo")
            for el in body.iter(qn(tag))
        ]
        assert len(revision_ids) == len(set(revision_ids))
        # A deletion's text is w:delText, never a live w:t.
        for deletion in body.iter(qn("w:del")):
            assert not list(deletion.iter(qn("w:t")))

    def test_tables_carry_real_merges_and_a_nested_table(self, tmp_path):
        body, _ = _saved_body(fx.build_merged_and_nested_tables_spec(), tmp_path)
        schedule, layout = body.findall(qn("w:tbl"))
        assert schedule.find(f".//{qn('w:gridSpan')}").get(qn("w:val")) == "3"
        v_merges = [el.get(qn("w:val")) for el in schedule.iter(qn("w:vMerge"))]
        assert v_merges[0] == "restart" and len(v_merges) == 2
        nested = layout.findall(f".//{qn('w:tc')}/{qn('w:tbl')}")
        assert len(nested) == 1

    @pytest.mark.parametrize("name", sorted(_WRAPPER_BUILDERS))
    def test_saved_documents_reopen(self, name, tmp_path):
        path = fx.save_docx(_WRAPPER_BUILDERS[name](), tmp_path, f"{name}.docx")
        Document(path)  # a malformed package raises here


# ---------------------------------------------------------------------------
# Representative file names
# ---------------------------------------------------------------------------


class TestFilenameExamples:
    def test_every_naming_style_is_represented(self):
        styles = {example.style for example in fx.FILENAME_EXAMPLES}
        assert styles == {"separated", "dashed", "compact", "section_prefixed", "unrecognized"}

    def test_names_are_unique_even_ignoring_case(self):
        names = [example.name.casefold() for example in fx.FILENAME_EXAMPLES]
        assert len(names) == len(set(names))

    def test_every_name_is_a_docx_in_some_case(self):
        assert all(e.name.lower().endswith(".docx") for e in fx.FILENAME_EXAMPLES)
        assert any(e.name.endswith(".DOCX") for e in fx.FILENAME_EXAMPLES)

    def test_section_numbers_are_normalized_or_absent(self):
        for example in fx.FILENAME_EXAMPLES:
            if example.style == "unrecognized":
                assert example.section is None, example.name
            else:
                digits = example.section.split(" ")
                assert [len(part) for part in digits] == [2, 2, 2], example.name
                assert example.section.replace(" ", "") in example.name.replace(" ", "").replace("-", "")

    def test_filtering_by_an_unknown_style_is_an_error(self):
        with pytest.raises(ValueError):
            fx.filename_examples("roman_numerals")
