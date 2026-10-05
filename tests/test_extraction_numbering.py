"""Word's automatic numbering reaches review as display text (plan WP-03, chunk S14).

Before S14 a paragraph Word numbers automatically lost its number: the review
saw "SUMMARY" and "Provide the specified piping system." with no "1.01" or
"A.", section attribution never saw a heading, and the structure checks never
saw a numbered heading at all. These tests pin the contract that replaced it:

* **Resolution** — list, level, and definition from direct properties and
  paragraph styles; counters per list per document; starts, overrides,
  restarts, formats, and suffixes as Word defines them.
* **No guessing** — where readers part (a restarted list that interleaves, a
  level first shown inside a deeper label, a list also used outside the main
  text, a style linked to a level it does not name) the paragraph gets no
  number and the spec a warning; so do undefined and unsupported definitions.
* **Displayed text, recorded spans** — ``content`` is what Word shows; every
  synthetic label is recorded on its element (``label_spans``) and in content
  coordinates (``ExtractedSpec.label_spans``), so the document's own text is
  always recoverable (``source_text``) and the element ids never change.
* **Consumers** — the review prompt, section attribution, and the structure
  checks see the numbers; the text checks read only authored text.

Where a behavior was checked against a real word processor it says so; the
resolver's module docstring lists what LibreOffice Writer 24.2 agrees with.
Hermetic: documents are built in memory and saved only under ``tmp_path``.
"""
from __future__ import annotations

import re

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from src.core.code_cycles import CALIFORNIA_2025
from src.input import extractor as extractor_module
from src.input.extractor import (
    ParagraphMapping,
    content_label_spans,
    extract_context_text,
    extract_multiple_specs,
    extract_text_from_docx,
)
from src.input.numbering import (
    NumberingLabel,
    REASON_AMBIGUOUS,
    REASON_UNDEFINED,
    REASON_UNSUPPORTED_FORMAT,
    format_number,
    labeled_text,
    resolve_numbering,
)
from src.input.preprocessor import (
    HEADING_PROVENANCE_AUTOMATIC,
    HEADING_PROVENANCE_TYPED,
    heading_candidates,
    preprocess_spec,
)
from src.review.review_request_builder import ReviewRequestSpec, build_user_message
from tests.fixtures import spec_docx as fx

lvl = fx.numbering_level


def _extract(builder, tmp_path, name: str = "230500.docx"):
    return extract_text_from_docx(fx.save_docx(builder, tmp_path, name))


def _texts(spec) -> list[str]:
    return [m.text for m in spec.paragraph_map]


def _labels(builder) -> list[tuple[str, str | None, str | None]]:
    """``(literal text, label, unresolved reason)`` for every body paragraph."""
    document = builder.document
    numbering = resolve_numbering(document)
    out = []
    for p_el in document.element.body.iter(qn("w:p")):
        literal = "".join(t.text or "" for t in p_el.iter(qn("w:t")))
        label = numbering.label(p_el)
        out.append((literal, label.text if label else None, numbering.unresolved_reason(p_el)))
    return out


def _shown(builder) -> list[str]:
    """Each body paragraph as the resolver says Word shows it."""
    return [f"{label} {text}" if label else text for text, label, _ in _labels(builder)]


def _two_level(*, level1_text: str = "%1.%2", restart: int | None = None) -> tuple[str, ...]:
    return (lvl(0), lvl(1, "decimal", level1_text, restart=restart))


# ===========================================================================
# 1. Resolution
# ===========================================================================


class TestTheCsiListFixture:
    """The WP-01 fixture: PART / article / paragraph, numbered by Word."""

    def test_every_label_is_the_one_word_shows(self):
        blocks = fx.auto_numbered_blocks()
        assert _shown(fx.build_auto_numbered_three_part()) == [b.displayed_text for b in blocks]

    def test_the_numbered_spec_extracts_exactly_like_the_typed_one(self, tmp_path):
        typed = _extract(fx.build_clean_three_part(), tmp_path, "typed.docx")
        numbered = _extract(fx.build_auto_numbered_three_part(), tmp_path, "numbered.docx")
        assert numbered.content == typed.content
        assert _texts(numbered) == _texts(typed)
        assert [m.section_id for m in numbered.paragraph_map] == [
            m.section_id for m in typed.paragraph_map
        ]

    def test_element_ids_do_not_depend_on_labels(self, tmp_path):
        typed = _extract(fx.build_clean_three_part(), tmp_path, "typed.docx")
        numbered = _extract(fx.build_auto_numbered_three_part(), tmp_path, "numbered.docx")
        assert [m.element_id for m in numbered.paragraph_map] == [
            m.element_id for m in typed.paragraph_map
        ]


class TestWhereNumberingComesFrom:
    def _styled_list(self) -> fx.SpecDocBuilder:
        builder = fx.SpecDocBuilder()
        abstract = builder.define_list(
            lvl(0, "decimal", "PART %1", style="PRT"),
            lvl(1, "decimalZero", "%1.%2", style="ART"),
            lvl(2, "upperLetter", "%3.", style="PR1"),
        )
        num = builder.define_num(abstract)
        builder.add_style("PRT", num_id=num)
        builder.add_style("ART", based_on="PRT", ilvl=1)  # list from PRT, level its own
        builder.add_style("PR1", based_on="ART", ilvl=2)
        builder.add_style("FLAT", num_id=num)  # a list but no level: level 0
        self.num = num
        return builder

    def test_style_inherited_numbering_follows_based_on(self):
        builder = self._styled_list()
        for text, style in (
            ("GENERAL", "PRT"), ("SUMMARY", "ART"), ("Provide x.", "PR1"),
            ("Provide y.", "PR1"), ("SUBMITTALS", "ART"), ("Submit z.", "PR1"),
            ("PRODUCTS", "PRT"), ("MATERIALS", "ART"),
        ):
            builder.add_numbered(text, style=style)
        # Checked against LibreOffice Writer 24.2.
        assert _shown(builder) == [
            "PART 1 GENERAL", "1.01 SUMMARY", "A. Provide x.", "B. Provide y.",
            "1.02 SUBMITTALS", "A. Submit z.", "PART 2 PRODUCTS", "2.01 MATERIALS",
        ]

    def test_list_and_level_are_inherited_independently(self):
        builder = self._styled_list()
        builder.add_numbered("GENERAL", style="PRT")
        builder.add_numbered("level only, list from the style", None, 2, style="PRT")
        builder.add_numbered("list only, no level anywhere", self.num)
        assert _shown(builder) == [
            "PART 1 GENERAL",
            "A. level only, list from the style",
            "PART 2 list only, no level anywhere",
        ]

    def test_a_style_numbering_with_no_level_is_level_0(self):
        builder = self._styled_list()
        builder.add_numbered("flat", style="FLAT")
        assert _shown(builder) == ["PART 1 flat"]

    def test_num_id_zero_turns_numbering_off(self, tmp_path):
        builder = self._styled_list()
        builder.add_numbered("SUMMARY", style="ART")
        builder.add_numbered("numbering removed", 0, style="PR1")
        builder.add_numbered("Provide x.", style="PR1")
        assert _shown(builder) == ["1.01 SUMMARY", "numbering removed", "A. Provide x."]
        # Off, not undefined: no reason, and so no warning.
        assert [reason for _, _, reason in _labels(builder)] == [None, None, None]
        assert _extract(builder, tmp_path).extraction_warnings == []

    def test_the_default_paragraph_style_can_carry_numbering(self):
        builder = fx.SpecDocBuilder()
        num = builder.define_num(builder.define_list(lvl(0)))
        styles = builder.document.part.styles.element
        for style in styles.findall(qn("w:style")):
            if style.get(qn("w:default")) == "1" and style.get(qn("w:type")) == "paragraph":
                style.append(parse_xml(
                    f'<w:pPr {nsdecls("w")}><w:numPr><w:numId w:val="{num}"/></w:numPr></w:pPr>'
                ))
        builder.add_text("first")
        builder.add_text("second")
        assert _shown(builder) == ["1. first", "2. second"]

    def test_a_numbering_style_link_is_followed(self):
        builder = fx.SpecDocBuilder()
        target = builder.define_list(
            lvl(0, "upperRoman", "%1."), lvl(1, "decimal", "%1.%2"), style_link="MyList"
        )
        direct = builder.define_num(target)
        builder.add_style("MyList", style_type="numbering", num_id=direct)
        linked = builder.define_num(builder.define_list(num_style_link="MyList"))
        builder.add_numbered("via link a", linked, 0)
        builder.add_numbered("via link b", linked, 1)
        builder.add_numbered("direct", direct, 0)
        # One list: the link and the definition share counters (LibreOffice agrees).
        assert _shown(builder) == ["I. via link a", "I.1 via link b", "II. direct"]


class TestCounters:
    def test_every_num_of_one_list_continues_it(self):
        builder = fx.SpecDocBuilder()
        abstract = builder.define_list(lvl(0))
        first, second = builder.define_num(abstract), builder.define_num(abstract)
        builder.add_numbered("one-a", first, 0)
        builder.add_numbered("one-b", first, 0)
        builder.add_text("plain")
        builder.add_numbered("two-a", second, 0)
        builder.add_numbered("one-c", first, 0)
        # Word's "Continue numbering"; LibreOffice agrees.
        assert _shown(builder) == ["1. one-a", "2. one-b", "plain", "3. two-a", "4. one-c"]

    def test_different_lists_count_independently(self):
        builder = fx.SpecDocBuilder()
        first = builder.define_num(builder.define_list(lvl(0)))
        second = builder.define_num(builder.define_list(lvl(0, "upperLetter")))
        for text, num in (("a", first), ("b", second), ("c", first), ("d", second)):
            builder.add_numbered(text, num, 0)
        assert _shown(builder) == ["1. a", "A. b", "2. c", "B. d"]

    def test_a_deeper_level_restarts_when_a_shallower_one_is_used(self):
        builder = fx.SpecDocBuilder()
        num = builder.define_num(builder.define_list(*_two_level()))
        for text, level in (("p1", 0), ("a", 1), ("b", 1), ("p2", 0), ("c", 1)):
            builder.add_numbered(text, num, level)
        assert _shown(builder) == ["1. p1", "1.1 a", "1.2 b", "2. p2", "2.1 c"]

    def test_lvl_restart_zero_keeps_a_level_counting(self):
        """``w:lvlRestart w:val="0"``: never restart (the standard, and Word's
        "Restart list after" left unticked). LibreOffice restarts it anyway,
        the one point where the resolver follows the standard over it."""
        builder = fx.SpecDocBuilder()
        num = builder.define_num(builder.define_list(*_two_level(restart=0)))
        for text, level in (("p1", 0), ("a", 1), ("b", 1), ("p2", 0), ("c", 1)):
            builder.add_numbered(text, num, level)
        assert _shown(builder) == ["1. p1", "1.1 a", "1.2 b", "2. p2", "2.3 c"]

    def test_lvl_restart_n_restarts_only_at_level_n(self):
        builder = fx.SpecDocBuilder()
        num = builder.define_num(builder.define_list(
            lvl(0), lvl(1, "decimal", "%1.%2"), lvl(2, "lowerLetter", "%3)", restart=2)
        ))
        for text, level in (("p1", 0), ("a1", 1), ("x", 2), ("y", 2), ("a2", 1), ("z", 2)):
            builder.add_numbered(text, num, level)
        assert _shown(builder) == ["1. p1", "1.1 a1", "a) x", "b) y", "1.2 a2", "a) z"]

    def test_lvl_restart_naming_a_later_level_is_ignored(self):
        builder = fx.SpecDocBuilder()
        num = builder.define_num(builder.define_list(lvl(0), lvl(1, "decimal", "%1.%2", restart=5)))
        for text, level in (("p1", 0), ("a", 1), ("p2", 0), ("b", 1)):
            builder.add_numbered(text, num, level)
        assert _shown(builder) == ["1. p1", "1.1 a", "2. p2", "2.1 b"]

    def test_a_missing_start_is_zero(self):
        builder = fx.SpecDocBuilder()
        num = builder.define_num(builder.define_list(lvl(0, start=None)))
        builder.add_numbered("a", num, 0)
        builder.add_numbered("b", num, 0)
        assert _shown(builder) == ["0. a", "1. b"]  # LibreOffice agrees

    def test_a_start_override_on_the_first_level_is_a_restart(self):
        builder = fx.SpecDocBuilder()
        abstract = builder.define_list(*_two_level())
        plain = builder.define_num(abstract)
        first_restart = builder.define_num(abstract, start_overrides={0: 1})
        second_restart = builder.define_num(abstract, start_overrides={0: 5})
        for text, num, level in (
            ("a", plain, 0), ("a1", plain, 1), ("b", plain, 0),
            ("r1", first_restart, 0), ("r1a", first_restart, 1), ("r2", first_restart, 0),
            ("s1", second_restart, 0), ("s1a", second_restart, 1),
        ):
            builder.add_numbered(text, num, level)
        assert _shown(builder) == [
            "1. a", "1.1 a1", "2. b", "1. r1", "1.1 r1a", "2. r2", "5. s1", "5.1 s1a",
        ]

    def test_a_level_override_redefines_the_level_for_its_num(self):
        builder = fx.SpecDocBuilder()
        abstract = builder.define_list(lvl(0))
        num = builder.define_num(abstract, level_overrides={0: lvl(0, "upperRoman", "%1)")})
        builder.add_numbered("a", num, 0)
        builder.add_numbered("b", num, 0)
        assert _shown(builder) == ["I) a", "II) b"]

    def test_paragraphs_in_tables_and_controls_count_in_document_order(self):
        builder = fx.SpecDocBuilder()
        num = builder.define_num(builder.define_list(lvl(0)))
        props = fx.numbered_properties(num, 0)
        builder.add_numbered("before", num, 0)
        builder.add_xml(
            "<w:tbl><w:tblPr/><w:tblGrid><w:gridCol/></w:tblGrid><w:tr><w:tc>"
            + fx.paragraph("in cell", properties=props)
            + "</w:tc></w:tr></w:tbl>"
        )
        builder.add_xml(fx.block_control(fx.paragraph("in control", properties=props), tag="c", control_id=7))
        builder.add_numbered("after", num, 0)
        assert _shown(builder) == ["1. before", "2. in cell", "3. in control", "4. after"]

    def test_counters_belong_to_one_document(self, tmp_path):
        paths = []
        for name in ("a.docx", "b.docx", "c.docx"):
            paths.append(fx.save_docx(fx.build_auto_numbered_three_part(), tmp_path, name))
        specs = extract_multiple_specs(paths, max_workers=3)
        typed = _extract(fx.build_clean_three_part(), tmp_path, "typed.docx")
        assert [spec.content for spec in specs] == [typed.content] * 3


class TestFormatsAndText:
    @pytest.mark.parametrize(
        "value, fmt, expected",
        [
            (1, "decimal", "1"), (12, "decimal", "12"),
            (1, "decimalZero", "01"), (9, "decimalZero", "09"), (10, "decimalZero", "10"),
            (1, "upperLetter", "A"), (26, "upperLetter", "Z"), (27, "upperLetter", "AA"),
            (28, "upperLetter", "BB"), (53, "upperLetter", "AAA"), (2, "lowerLetter", "b"),
            (4, "upperRoman", "IV"), (14, "lowerRoman", "xiv"), (1999, "upperRoman", "MCMXCIX"),
            (7, "none", ""),
            (0, "upperLetter", None), (0, "upperRoman", None), (4000, "upperRoman", None),
            (3, "ordinal", None), (3, "cardinalText", None), (-1, "decimal", None),
        ],
    )
    def test_format_number(self, value, fmt, expected):
        assert format_number(value, fmt) == expected

    def test_letters_past_z_and_zero_padded_articles(self):
        """Checked against LibreOffice Writer 24.2."""
        builder = fx.SpecDocBuilder()
        num = builder.define_num(builder.define_list(lvl(0, "decimalZero"), lvl(1, "upperLetter", "%2.")))
        for index in range(11):
            builder.add_numbered(f"d{index}", num, 0)
        for index in range(28):
            builder.add_numbered(f"l{index}", num, 1)
        shown = _shown(builder)
        assert shown[8:11] == ["09. d8", "10. d9", "11. d10"]
        assert shown[-3:] == ["Z. l25", "AA. l26", "BB. l27"]

    def test_legal_numbering_writes_every_level_in_arabic(self):
        builder = fx.SpecDocBuilder()
        num = builder.define_num(builder.define_list(
            lvl(0, "upperRoman", "%1."), lvl(1, "decimal", "%1.%2", legal=True)
        ))
        for text, level in (("p1", 0), ("p2", 0), ("a", 1)):
            builder.add_numbered(text, num, level)
        assert _shown(builder) == ["I. p1", "II. p2", "2.1 a"]  # LibreOffice agrees

    def test_bullets_carry_no_label_and_no_warning(self, tmp_path):
        builder = fx.SpecDocBuilder()
        num = builder.define_num(builder.define_list(lvl(0, "bullet", "")))
        builder.add_numbered("Provide hangers.", num, 0)
        spec = _extract(builder, tmp_path)
        assert _texts(spec) == ["Provide hangers."]
        assert spec.extraction_warnings == []

    def test_a_none_format_level_shows_only_its_literal_text(self):
        builder = fx.SpecDocBuilder()
        num = builder.define_num(builder.define_list(lvl(0, "none", ""), lvl(1, "none", "Note:")))
        builder.add_numbered("hidden", num, 0)
        builder.add_numbered("literal", num, 1)
        assert _shown(builder) == ["hidden", "Note: literal"]

    @pytest.mark.parametrize(
        "suffix, expected",
        [("space", "1. text"), ("tab", "1. text"), (None, "1. text"), ("nothing", "1.text")],
    )
    def test_the_suffix(self, tmp_path, suffix, expected):
        builder = fx.SpecDocBuilder()
        num = builder.define_num(builder.define_list(lvl(0, suffix=suffix)))
        builder.add_numbered("text", num, 0)
        assert _texts(_extract(builder, tmp_path)) == [expected]


# ===========================================================================
# 2. No guessing: ambiguous, undefined, unsupported
# ===========================================================================


class TestAmbiguousCountersAreNotGuessed:
    """Each case: the labels before the point where readings part are shown,
    the rest of that list gets none, and the reason is ambiguous."""

    def _reasons(self, builder) -> list[str | None]:
        return [reason for _, _, reason in _labels(builder)]

    def test_returning_to_a_list_after_a_restart(self):
        builder = fx.SpecDocBuilder()
        abstract = builder.define_list(lvl(0), lvl(1, "lowerLetter", "%2)"))
        first = builder.define_num(abstract)
        restart = builder.define_num(abstract, start_overrides={0: 1})
        for text, num in (("one-a", first), ("one-b", first), ("two-a", restart),
                          ("two-b", restart), ("one-c", first), ("two-c", restart)):
            builder.add_numbered(text, num, 0)
        # LibreOffice shows 3. and 4. for the last two; a reading in which the
        # restarted list is its own shows 3. and 3.
        assert _shown(builder) == ["1. one-a", "2. one-b", "1. two-a", "2. two-b", "one-c", "two-c"]
        assert self._reasons(builder)[-2:] == [REASON_AMBIGUOUS, REASON_AMBIGUOUS]

    def test_a_plain_num_after_a_restart(self):
        builder = fx.SpecDocBuilder()
        abstract = builder.define_list(*_two_level())
        plain = builder.define_num(abstract)
        restart = builder.define_num(abstract, start_overrides={0: 5})
        later = builder.define_num(abstract)
        for text, num in (("a", plain), ("b", plain), ("o1", restart), ("o2", restart), ("plain3", later)):
            builder.add_numbered(text, num, 0)
        assert _shown(builder) == ["1. a", "2. b", "5. o1", "6. o2", "plain3"]

    def test_a_restart_first_used_below_its_first_level(self):
        builder = fx.SpecDocBuilder()
        abstract = builder.define_list(*_two_level())
        plain = builder.define_num(abstract)
        restart = builder.define_num(abstract, start_overrides={0: 1})
        for text, num, level in (("a", plain, 0), ("b1", plain, 1), ("two first at lvl1", restart, 1),
                                 ("two lvl0", restart, 0)):
            builder.add_numbered(text, num, level)
        assert _shown(builder) == ["1. a", "1.1 b1", "two first at lvl1", "two lvl0"]

    def test_a_restart_that_keeps_a_deeper_level_counting(self):
        builder = fx.SpecDocBuilder()
        abstract = builder.define_list(*_two_level(restart=0))
        plain = builder.define_num(abstract)
        restart = builder.define_num(abstract, start_overrides={0: 1})
        builder.add_numbered("a", plain, 0)
        builder.add_numbered("r", restart, 0)
        assert _shown(builder) == ["1. a", "r"]

    def test_a_level_first_shown_inside_a_deeper_label(self):
        builder = fx.SpecDocBuilder()
        num = builder.define_num(builder.define_list(
            lvl(0, "decimal", "%1.", start=3), lvl(1, "decimal", "%1.%2")
        ))
        for text, level in (("deep-first", 1), ("top", 0), ("deep", 1)):
            builder.add_numbered(text, num, level)
        # LibreOffice: "3.1", "4.", "4.1" — it counts the first showing as a use.
        assert _shown(builder) == ["3.1 deep-first", "top", "deep"]
        assert self._reasons(builder) == [None, REASON_AMBIGUOUS, REASON_AMBIGUOUS]

    def test_an_earlier_level_than_lvl_restart_names(self):
        builder = fx.SpecDocBuilder()
        num = builder.define_num(builder.define_list(
            lvl(0), lvl(1, "decimal", "%1.%2"), lvl(2, "lowerLetter", "%3)", restart=2)
        ))
        for text, level in (("p1", 0), ("a1", 1), ("x", 2), ("p2", 0), ("y", 2)):
            builder.add_numbered(text, num, level)
        # Using level 1 restarts it; using level 0 — earlier than named — the
        # standard does not settle.
        assert _shown(builder) == ["1. p1", "1.1 a1", "a) x", "2. p2", "y"]

    def test_a_style_linked_to_a_level_it_does_not_name(self):
        builder = fx.SpecDocBuilder()
        num = builder.define_num(builder.define_list(
            lvl(0, style="PRT"), lvl(1, "decimal", "%1.%2", style="ART")
        ))
        builder.add_style("PRT", num_id=num)
        builder.add_style("ART", num_id=num)  # linked to level 1, names none
        builder.add_numbered("GENERAL", style="PRT")
        builder.add_numbered("SUMMARY", style="ART")
        builder.add_numbered("PRODUCTS", style="PRT")
        # LibreOffice takes level 0 ("2. SUMMARY", then "3. PRODUCTS"); Word's
        # level link says level 1 ("1.1", then "2."). Either way the unplaced
        # paragraph was counted, so the rest of the list is not numbered.
        assert _shown(builder) == ["1. GENERAL", "SUMMARY", "PRODUCTS"]
        assert self._reasons(builder)[1:] == [REASON_AMBIGUOUS, REASON_AMBIGUOUS]

    def test_a_level_out_of_range_stops_its_list(self):
        builder = fx.SpecDocBuilder()
        num = builder.define_num(builder.define_list(lvl(0)))
        builder.add_numbered("a", num, 0)
        builder.add_numbered("x", num, 9)
        builder.add_numbered("b", num, 0)
        assert _shown(builder) == ["1. a", "x", "b"]
        assert self._reasons(builder) == [None, REASON_UNDEFINED, REASON_AMBIGUOUS]

    def test_an_unplaced_paragraph_outside_the_main_text_still_shares_its_list(self, tmp_path):
        builder = fx.SpecDocBuilder()
        num = builder.define_num(builder.define_list(
            lvl(0, style="PRT"), lvl(1, "decimal", "%1.%2", style="ART")
        ))
        builder.add_style("ART", num_id=num)  # linked to level 1, names none
        builder.add_numbered("body", num, 0)
        header = builder.document.sections[0].header._element
        header.append(parse_xml(
            f'<w:p {nsdecls("w")}>{fx.numbered_properties(style="ART")}<w:r><w:t>head</w:t></w:r></w:p>'
        ))
        assert _texts(_extract(builder, tmp_path))[0] == "body"

    @pytest.mark.parametrize("where", ["header", "footnote", "text box"])
    def test_a_list_also_used_outside_the_main_text(self, tmp_path, where):
        builder = fx.SpecDocBuilder()
        abstract = builder.define_list(lvl(0))
        shared = builder.define_num(abstract)
        props = fx.numbered_properties(shared, 0)
        builder.add_numbered("body", shared, 0)
        if where == "header":
            header = builder.document.sections[0].header._element
            header.append(parse_xml(f'<w:p {nsdecls("w")}>{props}<w:r><w:t>head</w:t></w:r></w:p>'))
        elif where == "footnote":
            _add_footnote(builder, fx.paragraph("note", properties=props))
        else:
            builder.add_paragraph(_text_box(fx.paragraph("boxed", properties=props)))
        spec = _extract(builder, tmp_path)
        assert _texts(spec)[0] == "body"
        warnings = [w for w in spec.extraction_warnings if "numbered" in w]
        assert len(warnings) == 2
        assert "does not settle" in warnings[0] and warnings[0].startswith("Spec contains 1 ")
        assert "headers, footers, text boxes, or notes" in warnings[1]

    def test_numbering_outside_the_main_text_is_not_resolved(self):
        """A text box paragraph sits inside a body paragraph's run, so it is in
        the body's XML tree — but it is not main text: the resolver neither
        numbers it nor gives it a reason (the extractor warns about it)."""
        builder = fx.SpecDocBuilder()
        boxed = builder.define_num(builder.define_list(lvl(0, "upperLetter")))
        builder.add_paragraph(_text_box(fx.paragraph("boxed", properties=fx.numbered_properties(boxed, 0))))
        numbering = resolve_numbering(builder.document)
        (box,) = builder.document.element.body.iter(qn("w:txbxContent"))
        (box_paragraph,) = box.findall(qn("w:p"))
        assert numbering.label(box_paragraph) is None
        assert numbering.unresolved_reason(box_paragraph) is None
        assert numbering.numbered_outside_main_text(box_paragraph)

    def test_a_list_only_used_outside_the_main_text_does_not_touch_body_lists(self, tmp_path):
        builder = fx.SpecDocBuilder()
        body_list = builder.define_num(builder.define_list(lvl(0)))
        header_list = builder.define_num(builder.define_list(lvl(0, "upperLetter")))
        builder.add_numbered("body", body_list, 0)
        header = builder.document.sections[0].header._element
        header.append(parse_xml(
            f'<w:p {nsdecls("w")}>{fx.numbered_properties(header_list, 0)}<w:r><w:t>head</w:t></w:r></w:p>'
        ))
        spec = _extract(builder, tmp_path)
        assert _texts(spec)[0] == "1. body"
        assert "[Header] head" in _texts(spec)  # its own number is not shown
        assert spec.extraction_warnings == [
            "Spec contains 1 automatically numbered paragraph in headers, footers, text "
            "boxes, or notes; its number was not extracted for review. Verify visually."
        ]


def _text_box(*paragraphs: str) -> str:
    return (
        '<w:r><w:pict><v:shape xmlns:v="urn:schemas-microsoft-com:vml"><v:textbox>'
        f"<w:txbxContent>{''.join(paragraphs)}</w:txbxContent>"
        "</v:textbox></v:shape></w:pict></w:r>"
    )


def _add_footnote(builder, paragraph_xml: str) -> None:
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part

    blob = (
        f'<w:footnotes {nsdecls("w")}><w:footnote w:id="1">{paragraph_xml}</w:footnote></w:footnotes>'
    ).encode("utf-8")
    part = Part(
        PackURI("/word/footnotes.xml"),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
        blob,
        builder.document.part.package,
    )
    builder.document.part.relate_to(part, RT.FOOTNOTES)


class TestUndefinedAndUnsupported:
    def _reason_of_only(self, builder) -> str | None:
        ((_, label, reason),) = _labels(builder)
        assert label is None
        return reason

    def test_an_unknown_num_id(self):
        builder = fx.SpecDocBuilder()
        builder.add_numbered("x", 99, 0)
        assert self._reason_of_only(builder) == REASON_UNDEFINED

    def test_a_missing_abstract_definition(self):
        builder = fx.SpecDocBuilder()
        builder.add_numbered("x", builder.define_num("777"), 0)
        assert self._reason_of_only(builder) == REASON_UNDEFINED

    def test_an_undefined_level(self):
        builder = fx.SpecDocBuilder()
        num = builder.define_num(builder.define_list(lvl(0)))
        builder.add_numbered("x", num, 3)
        assert self._reason_of_only(builder) == REASON_UNDEFINED

    def test_a_paragraph_at_an_undefined_level_stops_its_list(self):
        builder = fx.SpecDocBuilder()
        num = builder.define_num(builder.define_list(lvl(0)))
        builder.add_numbered("a", num, 0)
        builder.add_numbered("x", num, 3)
        builder.add_numbered("b", num, 0)
        assert [reason for _, _, reason in _labels(builder)] == [None, REASON_UNDEFINED, REASON_AMBIGUOUS]

    def test_a_label_naming_a_deeper_level(self):
        builder = fx.SpecDocBuilder()
        num = builder.define_num(builder.define_list(lvl(0, text="%1.%2"), lvl(1)))
        builder.add_numbered("top", num, 0)
        assert self._reason_of_only(builder) == REASON_UNDEFINED  # LibreOffice prints "1.%2%"

    def test_a_malformed_value(self):
        builder = fx.SpecDocBuilder()
        num = builder.define_num(builder.define_list(lvl(0).replace('w:val="1"', 'w:val="one"', 1)))
        builder.add_numbered("x", num, 0)
        assert self._reason_of_only(builder) == REASON_UNDEFINED

    def test_a_missing_level_text(self):
        builder = fx.SpecDocBuilder()
        level = re.sub(r"<w:lvlText [^>]*/>", "", lvl(0))
        builder.add_numbered("x", builder.define_num(builder.define_list(level)), 0)
        assert self._reason_of_only(builder) == REASON_UNDEFINED

    @pytest.mark.parametrize("fmt", ["ordinal", "cardinalText", "decimalEnclosedCircle"])
    def test_an_unsupported_format(self, fmt):
        builder = fx.SpecDocBuilder()
        builder.add_numbered("x", builder.define_num(builder.define_list(lvl(0, fmt))), 0)
        assert self._reason_of_only(builder) == REASON_UNSUPPORTED_FORMAT

    @pytest.mark.parametrize("name", ["custom", "decimalZero"])
    def test_a_custom_format_string(self, name):
        """Word writes ``w:val="custom"`` beside ``w:format``; a format string
        beside a named format is refused too, rather than rendered as the
        named format it may override."""
        builder = fx.SpecDocBuilder()
        level = lvl(0).replace(
            '<w:numFmt w:val="decimal"/>', f'<w:numFmt w:val="{name}" w:format="001, 002, 003"/>'
        )
        builder.add_numbered("x", builder.define_num(builder.define_list(level)), 0)
        assert self._reason_of_only(builder) == REASON_UNSUPPORTED_FORMAT

    def test_an_unsupported_format_still_counts(self):
        builder = fx.SpecDocBuilder()
        num = builder.define_num(builder.define_list(lvl(0, "ordinal"), lvl(1, "decimal", "%2.")))
        builder.add_numbered("x", num, 0)
        builder.add_numbered("a", num, 1)
        assert _shown(builder) == ["x", "1. a"]


class TestRevisions:
    def test_a_deleted_paragraph_is_not_numbered_and_does_not_count(self, tmp_path):
        builder = fx.SpecDocBuilder()
        num = builder.define_num(builder.define_list(lvl(0, "upperLetter")))
        deleted_mark = (
            f'<w:rPr><w:del w:id="5" w:author="{fx.REVISION_AUTHOR}" w:date="{fx.REVISION_DATE}"/></w:rPr>'
        )
        builder.add_numbered("Keep one.", num, 0)
        builder.add_paragraph(
            fx.deleted(fx.deleted_run("Removed requirement."), revision_id=6),
            properties=fx.numbered_properties(num, 0, extra=deleted_mark),
        )
        builder.add_numbered("Keep two.", num, 0)
        spec = _extract(builder, tmp_path)
        assert _texts(spec) == ["A. Keep one.", "B. Keep two."]
        assert "Removed requirement." not in spec.content

    def test_an_inserted_numbered_paragraph_counts(self, tmp_path):
        builder = fx.SpecDocBuilder()
        num = builder.define_num(builder.define_list(lvl(0, "upperLetter")))
        inserted_mark = (
            f'<w:rPr><w:ins w:id="5" w:author="{fx.REVISION_AUTHOR}" w:date="{fx.REVISION_DATE}"/></w:rPr>'
        )
        builder.add_numbered("Keep one.", num, 0)
        builder.add_paragraph(
            fx.inserted(fx.run("New requirement."), revision_id=6),
            properties=fx.numbered_properties(num, 0, extra=inserted_mark),
        )
        builder.add_numbered("Keep two.", num, 0)
        assert _texts(_extract(builder, tmp_path)) == [
            "A. Keep one.", "B. New requirement.", "C. Keep two.",
        ]


# ===========================================================================
# 3. The extraction contract
# ===========================================================================


class TestDisplayedTextAndSpans:
    def test_labels_are_recorded_where_they_stand(self, tmp_path):
        spec = _extract(fx.build_auto_numbered_three_part(), tmp_path)
        for mapping, block in zip(spec.paragraph_map, fx.auto_numbered_blocks()):
            ((start, end),) = mapping.label_spans
            assert (start, mapping.text[start:end]) == (0, f"{block.auto_label} ")
            assert mapping.source_text == block.text

    def test_content_spans_point_at_the_labels(self, tmp_path):
        spec = _extract(fx.build_auto_numbered_three_part(), tmp_path)
        labels = [spec.content[start:end] for start, end in spec.label_spans]
        assert labels == [f"{b.auto_label} " for b in fx.auto_numbered_blocks()]
        assert content_label_spans(spec.paragraph_map) == spec.label_spans

    def test_the_source_text_is_the_literal_text(self, tmp_path):
        spec = _extract(fx.build_auto_numbered_three_part(), tmp_path)
        literal = [
            extractor_module._accept_all_paragraph_text(p).strip()
            for p in Document(spec.source_path).element.body.iter(qn("w:p"))
        ]
        assert [m.source_text for m in spec.paragraph_map] == literal

    def test_typed_numbering_has_no_spans(self, tmp_path):
        spec = _extract(fx.build_clean_three_part(), tmp_path)
        assert spec.label_spans == ()
        assert all(m.source_text == m.text for m in spec.paragraph_map)

    def test_a_table_row_keeps_one_span_per_numbered_cell_paragraph(self, tmp_path):
        builder = fx.SpecDocBuilder()
        num = builder.define_num(builder.define_list(lvl(0, "upperLetter")))
        props = fx.numbered_properties(num, 0)
        cell = lambda *xml: f"<w:tc>{''.join(xml)}</w:tc>"
        builder.add_xml(
            "<w:tbl><w:tblPr/><w:tblGrid><w:gridCol/><w:gridCol/></w:tblGrid><w:tr>"
            + cell(fx.paragraph("  Pipe", properties=props), fx.paragraph("Valve", properties=props))
            + cell(fx.paragraph("Copper"))
            + "</w:tr></w:tbl>"
        )
        spec = _extract(builder, tmp_path)
        (row,) = spec.paragraph_map
        assert row.element_id == "t0r0"
        assert row.text == "A.   Pipe\nB. Valve | Copper"
        assert [row.text[a:b] for a, b in row.label_spans] == ["A. ", "B. "]
        assert row.source_text == "  Pipe\nValve | Copper"

    def test_a_cell_that_starts_with_an_empty_paragraph(self, tmp_path):
        """Stripping the cell's leading whitespace moves its labels with it."""
        builder = fx.SpecDocBuilder()
        num = builder.define_num(builder.define_list(lvl(0, "upperLetter")))
        builder.add_xml(
            "<w:tbl><w:tblPr/><w:tblGrid><w:gridCol/></w:tblGrid><w:tr><w:tc>"
            + fx.paragraph("")
            + fx.paragraph("  ")
            + fx.paragraph("Valve", properties=fx.numbered_properties(num, 0))
            + "</w:tc></w:tr></w:tbl>"
        )
        (row,) = _extract(builder, tmp_path).paragraph_map
        assert row.text == "A. Valve"
        assert row.label_spans == ((0, 3),)

    def test_the_reconstruction_check_rejects_a_bad_span(self, tmp_path, monkeypatch):
        builder = fx.build_auto_numbered_three_part()
        monkeypatch.setattr(
            extractor_module._Numbering, "label", lambda self, p_el, text: (text, ((0, 999),))
        )
        with pytest.raises(ValueError, match="label spans"):
            _extract(builder, tmp_path)

    def test_unreadable_numbering_never_sinks_extraction(self, tmp_path, monkeypatch):
        def broken(document):
            raise RuntimeError("corrupt numbering part")

        monkeypatch.setattr(extractor_module, "resolve_numbering", broken)
        spec = _extract(fx.build_auto_numbered_three_part(), tmp_path)
        assert _texts(spec)[1] == "SUMMARY"
        assert spec.extraction_warnings == [
            "Spec's automatic numbering could not be read; numbers Word shows in front "
            "of paragraphs were not extracted for review. Verify visually."
        ]

    def test_a_paragraph_mapping_built_without_spans_reads_as_literal(self):
        mapping = ParagraphMapping(0, "paragraph", "1.01 SUMMARY", None, None, None)
        assert mapping.label_spans == () and mapping.source_text == "1.01 SUMMARY"

    def test_labeled_text_is_the_one_rule(self):
        label = NumberingLabel("1.01", " ", "0", 1)
        assert labeled_text(label, "SUMMARY") == ("1.01 SUMMARY", ((0, 5),))
        assert labeled_text(label, "1.01 SUMMARY") == ("1.01 SUMMARY", ())  # typed repeat
        assert labeled_text(label, "   ") == ("   ", ())  # nothing shown, nothing labeled
        assert labeled_text(None, "SUMMARY") == ("SUMMARY", ())

    def test_labels_are_looked_up_in_the_tree_that_was_resolved(self):
        numbering = resolve_numbering(fx.build_auto_numbered_three_part().document)
        other = fx.build_auto_numbered_three_part().document
        assert numbering.label(next(other.element.body.iter(qn("w:p")))) is None


class TestTypedNumbersAreNotRepeated:
    def test_a_typed_copy_of_the_label_is_shown_once_and_warned(self, tmp_path):
        builder = fx.SpecDocBuilder()
        num = builder.define_num(builder.define_list(lvl(0, "decimal", "PART %1"), lvl(1, "decimalZero", "%1.%2")))
        builder.add_numbered("PART 1 GENERAL", num, 0)
        builder.add_numbered("1.01 SUMMARY", num, 1)
        builder.add_numbered("1.03 WRONG", num, 1)  # a different typed number is not a repeat
        spec = _extract(builder, tmp_path)
        assert _texts(spec) == ["PART 1 GENERAL", "1.01 SUMMARY", "1.02 1.03 WRONG"]
        assert spec.paragraph_map[1].label_spans == ()
        assert spec.extraction_warnings == [
            "Spec contains 2 paragraphs whose typed number repeats the automatic number "
            'Word shows in front of it (Word displays both, as in "1.01 1.01 SUMMARY"); '
            "the number was extracted once. Verify visually."
        ]

    def test_a_typed_number_glued_to_text_is_not_a_repeat(self, tmp_path):
        builder = fx.SpecDocBuilder()
        num = builder.define_num(builder.define_list(lvl(0, "upperLetter", "%1.")))
        builder.add_numbered("A.B. Smith Co. shall install.", num, 0)
        assert _texts(_extract(builder, tmp_path)) == ["A. A.B. Smith Co. shall install."]


class TestWarnings:
    def test_one_warning_per_kind_counting_emitted_paragraphs(self, tmp_path):
        builder = fx.SpecDocBuilder()
        ordinal = builder.define_num(builder.define_list(lvl(0, "ordinal")))
        builder.add_numbered("x", 99, 0)
        builder.add_numbered("y", 99, 0)
        builder.add_numbered("", 99, 0)  # empty: not emitted, not counted
        builder.add_numbered("z", ordinal, 0)
        spec = _extract(builder, tmp_path)
        assert spec.extraction_warnings == [
            "Spec contains 2 automatically numbered paragraphs whose list definition is "
            "missing or malformed; their numbers were not extracted for review. Verify visually.",
            "Spec contains 1 automatically numbered paragraph in a number format the "
            "extractor does not render (it renders decimal numbers, letters, and roman "
            "numerals); its number was not extracted for review. Verify visually.",
        ]

    def test_a_clean_numbered_spec_has_no_warning(self, tmp_path):
        assert _extract(fx.build_auto_numbered_three_part(), tmp_path).extraction_warnings == []


class TestContextAttachments:
    def test_a_context_docx_shows_its_numbers(self, tmp_path):
        path = fx.save_docx(fx.build_auto_numbered_three_part(), tmp_path, "narrative.docx")
        assert extract_context_text(path) == _extract(fx.build_clean_three_part(), tmp_path).content


# ===========================================================================
# 4. Consumers: review prompt, section attribution, detectors
# ===========================================================================


class TestReviewPrompt:
    def test_the_prompt_shows_numbered_headings_and_body(self, tmp_path):
        spec = _extract(fx.build_auto_numbered_three_part(), tmp_path)
        message = build_user_message(
            ReviewRequestSpec(
                spec_content=spec.content, filename=spec.filename,
                model="claude-opus-5", paragraph_map=spec.paragraph_map,
            )
        )
        assert '<heading id="p1">1.01 SUMMARY</heading>' in message
        assert '<para id="p2">A. Provide the specified piping system.</para>' in message


class TestSectionAttribution:
    def _sections(self, tmp_path, *texts: str) -> list[str]:
        builder = fx.SpecDocBuilder()
        for text in texts:
            builder.add_text(text)
        return [m.section_id for m in _extract(builder, tmp_path).paragraph_map]

    def test_list_items_and_quantities_are_not_headings(self, tmp_path):
        """The S03 finding: the old heuristic took "1.5 inches minimum cover"
        (and, once labels show, every "1." list item) as a heading."""
        assert self._sections(
            tmp_path,
            "1.01 SUMMARY",
            "1. Provide pipe hangers.",
            "1.5 inches minimum cover.",
            "2.3.A Provide valves.",
            "Body text.",
        ) == ["1.01 SUMMARY"] * 5

    def test_part_article_and_section_lines_start_sections(self, tmp_path):
        assert self._sections(
            tmp_path, "SECTION 23 05 00 - COMMON WORK RESULTS", "PART 1 - GENERAL",
            "1.01 SUMMARY", "A. Text.", "See Section 21 13 13 for sprinklers.",
        ) == [
            "SECTION 23 05 00 - COMMON WORK RESULTS", "PART 1 - GENERAL",
            "1.01 SUMMARY", "1.01 SUMMARY", "1.01 SUMMARY",
        ]

    def test_a_table_takes_the_numbered_article_it_follows(self, tmp_path):
        spec = _extract(fx.build_blocks(fx.auto_numbered(fx.table_only_article_blocks())), tmp_path)
        rows = [m for m in spec.paragraph_map if m.element_type == "table_cell"]
        assert {m.section_id for m in rows} == {"2.01 MATERIALS"}


class TestDetectors:
    def test_numbered_headings_are_headings_with_their_provenance(self, tmp_path):
        spec = _extract(fx.build_auto_numbered_three_part(), tmp_path)
        candidates = heading_candidates(spec.content, label_spans=spec.label_spans)
        assert [c.label for c in candidates] == [
            "PART 1 GENERAL", "1.01 SUMMARY", "1.02 SUBMITTALS", "PART 2 PRODUCTS",
            "2.01 MATERIALS", "PART 3 EXECUTION", "3.01 INSTALLATION",
        ]
        assert {c.provenance for c in candidates} == {HEADING_PROVENANCE_AUTOMATIC}
        typed = _extract(fx.build_clean_three_part(), tmp_path, "typed.docx")
        assert {c.provenance for c in heading_candidates(typed.content)} == {HEADING_PROVENANCE_TYPED}

    def _structural(self, spec) -> list[tuple[str, str]]:
        result = preprocess_spec(spec.content, spec.filename, cycle=CALIFORNIA_2025, label_spans=spec.label_spans)
        return [(a["deterministic_rule"], a["match"]) for a in result.structural_alerts]

    def test_an_empty_numbered_article_is_reported(self, tmp_path):
        spec = _extract(fx.build_blocks(fx.auto_numbered(fx.empty_article_blocks())), tmp_path)
        assert self._structural(spec) == [("empty_section", "1.02 SUBMITTALS")]

    def test_an_empty_numbered_part_is_reported_once(self, tmp_path):
        spec = _extract(fx.build_blocks(fx.auto_numbered(fx.empty_part_blocks())), tmp_path)
        assert self._structural(spec) == [("empty_section", "PART 2 PRODUCTS")]

    def test_word_numbers_a_repeated_heading_on_so_it_is_no_duplicate(self, tmp_path):
        """Typed, a repeated "1.01 SUMMARY" is a duplicate heading. Numbered
        by Word, the copy is 1.02 — the detector reads what Word shows."""
        spec = _extract(fx.build_blocks(fx.auto_numbered(fx.duplicate_heading_blocks())), tmp_path)
        assert "1.02 SUMMARY" in spec.content
        assert self._structural(spec) == []

    def test_text_checks_ignore_a_label_and_read_the_same_text_typed(self, tmp_path):
        builder = fx.SpecDocBuilder()
        num = builder.define_num(builder.define_list(lvl(0, "upperRoman", "%1.", start=30)))
        builder.add_numbered("Provide valves.", num, 0)  # shown "XXX. Provide valves."
        builder.add_text("XXX. Typed marker.")
        spec = _extract(builder, tmp_path)
        assert _texts(spec)[0] == "XXX. Provide valves."
        markers = preprocess_spec(
            spec.content, spec.filename, cycle=CALIFORNIA_2025, label_spans=spec.label_spans
        ).template_marker_alerts
        assert [a["position"] for a in markers] == [spec.content.index("XXX. Typed")]
        unlabeled = preprocess_spec(spec.content, spec.filename, cycle=CALIFORNIA_2025)
        assert len(unlabeled.template_marker_alerts) == 2  # without spans: as before

    def test_duplicate_paragraphs_compare_authored_text(self, tmp_path):
        sentence = "Provide seismic bracing for every riser and branch line as scheduled on the drawings."
        builder = fx.SpecDocBuilder()
        num = builder.define_num(builder.define_list(lvl(0, "upperLetter")))
        builder.add_numbered(sentence, num, 0)
        builder.add_numbered(sentence, num, 0)
        spec = _extract(builder, tmp_path)
        alerts = preprocess_spec(
            spec.content, spec.filename, cycle=CALIFORNIA_2025, label_spans=spec.label_spans
        ).duplicate_paragraph_alerts
        assert [a["match"] for a in alerts] == [f"B. {sentence}"]

    def test_typed_numbers_still_make_paragraphs_different(self, tmp_path):
        sentence = "Provide seismic bracing for every riser and branch line as scheduled on the drawings."
        builder = fx.SpecDocBuilder()
        builder.add_text(f"A. {sentence}")
        builder.add_text(f"B. {sentence}")
        spec = _extract(builder, tmp_path)
        result = preprocess_spec(spec.content, spec.filename, cycle=CALIFORNIA_2025, label_spans=spec.label_spans)
        assert result.duplicate_paragraph_alerts == []


class TestThePipelinePassesTheSpans:
    """Both places the pipeline runs the detectors — preparing a run and
    rebuilding a repair's ``<pre_detected>`` block — hand them the spec's
    label spans, so a label never becomes a text alert in a real run."""

    def _numbered_spec(self, tmp_path):
        builder = fx.SpecDocBuilder()
        num = builder.define_num(builder.define_list(lvl(0, "upperRoman", "%1.", start=30)))
        sentence = "Provide seismic bracing for every riser and branch line as scheduled on the drawings."
        builder.add_numbered(sentence, num, 0)  # "XXX. Provide ..."
        builder.add_numbered(sentence, num, 0)  # "XXXI. Provide ..."
        return fx.save_docx(builder, tmp_path, "230500.docx")

    def _rules(self, alerts) -> list[str]:
        return [alert["deterministic_rule"] for alert in alerts]

    def test_preparing_a_run(self, tmp_path):
        from src.orchestration import pipeline

        path = self._numbered_spec(tmp_path)
        prepared = pipeline._prepare_specs(input_dir=tmp_path, files=[path], preflight=False)
        assert prepared.template_marker_alerts == []
        assert self._rules(prepared.duplicate_paragraph_alerts) == ["duplicate_paragraph"]

    def test_rebuilding_a_repair(self, tmp_path):
        from types import SimpleNamespace

        from src.modules import get_module
        from src.orchestration import pipeline

        spec = extract_text_from_docx(self._numbered_spec(tmp_path))
        submission = SimpleNamespace(prepared_specs=[spec], project_profile=None)
        alerts = pipeline._repair_pre_detected_alerts(
            submission, [spec], module=get_module("california_k12_mep")
        )[spec.filename]
        assert self._rules(alerts) == ["duplicate_paragraph"]
