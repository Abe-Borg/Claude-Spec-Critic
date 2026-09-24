"""Word content controls, fields, and smart tags reach review (plan WP-02, chunk S10).

Before S10 the extractor read a paragraph's direct runs and hyperlinks, the
body's direct paragraphs and tables, and nothing else. Text Word shows from
inside a content control (a template's fill-in value, a drop-down's chosen
entry, an unfilled control's placeholder), a simple field's stored result
(a cross-reference such as "23 05 00"), or a smart tag never reached the
review, and neither did an insertion tracked inside a hyperlink.

What these tests hold the extractor to:

* **Read in place, in order, once.** Inline wrappers are read where they sit
  in their paragraph; block controls where they sit in the body, a cell, a
  header, a text box, or a note. Accept-All applies at every depth.
* **Results, never instructions.** A field's stored result is read as stored;
  its instruction never is — nor a nested field's result that is only an
  argument inside an outer instruction.
* **Legacy ids keep their meaning.** ``pN`` is still physical body child
  ``N``; ``tN`` still the ``N``-th direct body table; ``r<n>`` still the
  table's ``n``-th own row. Text read from inside a block control gets ids
  of its own, each with a ``cc<n>`` step.
* **Honest about what is left.** Structures the walk meets but cannot read
  (equations, embedded documents, legacy drop-down form fields, tables inside
  table controls) produce a warning naming what and how many — never an
  invented measure of lost text. A Word table of contents is skipped on
  purpose: it repeats the headings, which are read where they stand.

The applier's side of the boundary (readable is not writable) is in
``test_applier_wrapped_content.py``.
"""
from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.opc.packuri import PackURI
from docx.opc.part import Part

from src.core.code_cycles import CALIFORNIA_2025
from src.input.extractor import (
    _ENDNOTES_CONTENT_TYPE,
    _FOOTNOTES_CONTENT_TYPE,
    _accept_all_paragraph_text,
    extract_text_from_docx,
)
from src.input.preprocessor import detect_placeholders, preprocess_spec
from src.review.prompt_serialization import render_spec_with_ids
from tests.fixtures import spec_docx as fx

_MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _extract(builder, tmp_path: Path, name: str = "spec.docx"):
    spec = extract_text_from_docx(fx.save_docx(builder, tmp_path, name))
    assert "\n\n".join(m.text for m in spec.paragraph_map) == spec.content
    return spec


def _ids(spec) -> dict[str, str]:
    return {m.element_id: m.text for m in spec.paragraph_map}


def _paragraph(*inline: str) -> fx.SpecDocBuilder:
    builder = fx.SpecDocBuilder()
    builder.add_paragraph(*inline)
    return builder


def _text_of(*inline: str) -> str:
    """The Accept-All text the extractor reads for one paragraph."""
    (element,) = fx.parse_fragments(fx.paragraph(*inline))
    return _accept_all_paragraph_text(element)


def _begin(ff_data: str = "") -> str:
    return f'<w:r><w:fldChar w:fldCharType="begin">{ff_data}</w:fldChar></w:r>'


def _instr(code: str) -> str:
    return f'<w:r><w:instrText xml:space="preserve">{code}</w:instrText></w:r>'


_SEPARATE = '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
_END = '<w:r><w:fldChar w:fldCharType="end"/></w:r>'


# ---------------------------------------------------------------------------
# Inline wrappers
# ---------------------------------------------------------------------------


class TestInlineWrappers:
    def test_control_text_is_read_in_place(self, tmp_path):
        spec = _extract(fx.build_inline_controls_spec(), tmp_path)
        assert _ids(spec)["p0"] == (
            fx.INLINE_CONTROL_BEFORE + fx.INLINE_CONTROL_TEXT + fx.INLINE_CONTROL_AFTER
        )

    def test_a_drop_down_reads_its_chosen_entry_only(self, tmp_path):
        spec = _extract(fx.build_inline_controls_spec(), tmp_path)
        assert _ids(spec)["p1"] == fx.DROPDOWN_LEAD + fx.DROPDOWN_CHOSEN
        # The other list items are control properties, not visible text.
        assert "Steel Schedule 40" not in spec.content
        assert "Galvanized" not in spec.content and "Painted" not in spec.content

    def test_an_unfilled_control_reads_its_placeholder_for_the_placeholder_check(self, tmp_path):
        spec = _extract(fx.build_inline_controls_spec(), tmp_path)
        assert _ids(spec)["p2"] == fx.UNRESOLVED_DROPDOWN_LEAD + fx.UNRESOLVED_DROPDOWN_PLACEHOLDER
        matches = [a["match"] for a in detect_placeholders(spec.content, spec.filename)]
        assert matches == [fx.UNRESOLVED_DROPDOWN_PLACEHOLDER]

    def test_control_properties_are_never_read(self):
        text = _text_of(
            "Pipe: ",
            fx.inline_control(fx.run("copper"), tag="SECRET_TAG", control_id=7, alias="SECRET_ALIAS"),
        )
        assert text == "Pipe: copper"

    def test_smart_tag_text_is_read(self, tmp_path):
        spec = _extract(fx.build_smart_tag_spec(), tmp_path)
        assert _ids(spec)["p0"] == "Obtain approval from the City of Oakland fire marshal."

    def test_custom_xml_text_is_read_and_its_properties_are_not(self):
        custom = (
            '<w:customXml w:uri="urn:fixture" w:element="material">'
            '<w:customXmlPr><w:attr w:name="HIDDEN_ATTR" w:val="HIDDEN_VALUE"/></w:customXmlPr>'
            f"{fx.run('ductile iron')}</w:customXml>"
        )
        assert _text_of("Pipe: ", custom, " fittings.") == "Pipe: ductile iron fittings."

    def test_bidirectional_containers_are_read(self):
        text = _text_of(
            "Tag ",
            f'<w:dir w:val="rtl">{fx.run("A-1")}</w:dir>',
            " and ",
            f'<w:bdo w:val="ltr">{fx.run("B-2")}</w:bdo>',
            ".",
        )
        assert text == "Tag A-1 and B-2."

    def test_an_insertion_inside_a_hyperlink_is_accepted(self, tmp_path):
        spec = _extract(fx.build_hyperlink_spec(), tmp_path)
        assert _ids(spec) == {
            "p0": "Verify each listing in the listing directory.",
            "p1": "Submit manufacturer data sheets.",
        }

    def test_a_deletion_inside_a_hyperlink_is_dropped(self):
        builder = fx.SpecDocBuilder()
        builder.add_paragraph(
            "See ",
            builder.hyperlink(
                fx.HYPERLINK_URL,
                fx.deleted(fx.deleted_run("the old "), revision_id=builder.next_id()),
                fx.run("the guide"),
            ),
            ".",
        )
        assert _accept_all_paragraph_text(builder.body[0]) == "See the guide."

    def test_nested_controls_honor_revisions_at_the_inner_depth(self, tmp_path):
        spec = _extract(fx.build_nested_controls_spec(), tmp_path)
        assert _ids(spec)["p0"] == fx.NESTED_INLINE_TEXT
        assert spec.tracked_changes_detected is True

    def test_a_control_inside_a_deletion_disappears_and_inside_an_insertion_stays(self):
        gone = fx.deleted(
            fx.inline_control(fx.deleted_run("zinc"), tag="old", control_id=1), revision_id=2
        )
        kept = fx.inserted(fx.inline_control(fx.run("galvanized"), tag="new", control_id=3), revision_id=4)
        assert _text_of("Finish: ", gone, kept) == "Finish: galvanized"

    def test_wrappers_nest_in_any_order(self):
        builder = fx.SpecDocBuilder()
        inner = fx.smart_tag(fx.simple_field(" REF bm ", "Oakland"))
        control = fx.inline_control(inner, tag="city", control_id=builder.next_id())
        builder.add_paragraph("City of ", builder.hyperlink(fx.HYPERLINK_URL, control), ".")
        assert _accept_all_paragraph_text(builder.body[0]) == "City of Oakland."

    def test_a_paragraph_without_wrappers_reads_exactly_as_python_docx_does(self, tmp_path):
        """Byte-identity for ordinary documents: the structured walk adds
        nothing where there is nothing to add (hyperlinks, tabs, and breaks
        included)."""
        builder = fx.SpecDocBuilder()
        builder.add_paragraph("Plain text.")
        builder.add_paragraph("Link to ", builder.hyperlink(fx.HYPERLINK_URL, "the guide"), ".")
        builder.add_paragraph("Tab", fx.tab_run(), "stop", '<w:r><w:br/></w:r>', "break")
        builder.add_paragraph('<w:r><w:t>a</w:t><w:noBreakHyphen/><w:t>b</w:t></w:r>')
        path = fx.save_docx(builder, tmp_path, "plain.docx")
        for paragraph in Document(path).paragraphs:
            assert _accept_all_paragraph_text(paragraph._p) == paragraph.text


# ---------------------------------------------------------------------------
# Fields: stored results only, never instructions
# ---------------------------------------------------------------------------


class TestFields:
    def test_a_simple_fields_stored_result_is_read_and_its_instruction_is_not(self, tmp_path):
        spec = _extract(fx.build_fields_spec(), tmp_path)
        assert _ids(spec)["p0"] == "Refer to Section 23 05 00 for common work results."
        assert "REF" not in spec.content and "sec_230500" not in spec.content

    def test_a_complex_fields_result_is_read_and_its_instruction_is_not(self, tmp_path):
        spec = _extract(fx.build_fields_spec(), tmp_path)
        assert _ids(spec)["p1"] == "Coordinate with Section 21 13 13."
        assert "sec_211313" not in spec.content

    def test_a_nested_fields_result_inside_an_instruction_is_not_read(self):
        """``{ IF { REF bm } = "x" "copper" "steel" }``: the REF result is an
        argument of the IF, not display text; only the IF's result shows."""
        text = _text_of(
            "Use ",
            _begin() + _instr(" IF ")
            + _begin() + _instr(" REF bm ") + _SEPARATE + fx.run("ARGUMENT") + _END
            + _instr(' = "x" "copper" "steel" ') + _SEPARATE + fx.run("copper") + _END,
            " pipe.",
        )
        assert text == "Use copper pipe."

    def test_a_simple_field_inside_an_instruction_is_not_read(self):
        text = _text_of(
            "Use ",
            _begin() + _instr(" IF ") + fx.simple_field(" MERGEFIELD m ", "ARGUMENT")
            + _instr(' = "x" "steel" "iron" ') + _SEPARATE + fx.run("steel") + _END,
            " pipe.",
        )
        assert text == "Use steel pipe."

    def test_a_nested_field_in_a_result_is_read(self):
        """A table of contents entry: a PAGEREF field inside the TOC field's
        result. Both results are display text."""
        text = _text_of(
            _begin() + _instr(" TOC ") + _SEPARATE + fx.run("PART 1 GENERAL ")
            + _begin() + _instr(" PAGEREF _Toc1 ") + _SEPARATE + fx.run("3") + _END + _END
        )
        assert text == "PART 1 GENERAL 3"

    def test_a_field_with_no_stored_result_shows_nothing(self):
        assert _text_of("Page ", _begin() + _instr(" PAGE ") + _END, ".") == "Page ."

    def test_a_field_carried_over_from_an_earlier_paragraph_reads_normally(self):
        """A multi-paragraph field (a table of contents): a later paragraph
        holds only result text and the closing character."""
        assert _text_of(fx.run("2.01 MATERIALS 4"), _END) == "2.01 MATERIALS 4"

    def test_a_deleted_field_disappears(self):
        field = (
            '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
            '<w:r><w:delInstrText xml:space="preserve"> REF x </w:delInstrText></w:r>'
            '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
            + fx.deleted_run("GONE")
            + '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
        )
        text = _text_of("See ", fx.deleted(field, revision_id=1), "Section 2.")
        assert text == "See Section 2."

    def test_a_legacy_text_form_fields_result_is_read(self):
        form = _begin("<w:ffData><w:name w:val=\"Text1\"/><w:textInput/></w:ffData>")
        text = _text_of("Owner: ", form + _instr(" FORMTEXT ") + _SEPARATE + fx.run("Oakland USD") + _END)
        assert text == "Owner: Oakland USD"


# ---------------------------------------------------------------------------
# Block controls in the body
# ---------------------------------------------------------------------------


class TestBlockControls:
    def test_a_block_control_yields_its_paragraph_and_table_in_order(self, tmp_path):
        spec = _extract(fx.build_block_control_spec(), tmp_path)
        assert [(m.element_id, m.text) for m in spec.paragraph_map] == [
            ("p0", fx.BEFORE_CONTROL),
            ("cc1p0", fx.BLOCK_CONTROL_PARAGRAPH),
            ("cc1t1r0", "Assembly | Listing"),
            ("cc1t1r1", "Backflow preventer | Listed double check"),
            ("p2", fx.AFTER_CONTROL),
            ("t0r0", "Service | Pipe size"),
            ("t0r1", "Riser | Four inch"),
        ]

    def test_text_read_from_a_control_is_marked_as_such(self, tmp_path):
        spec = _extract(fx.build_block_control_spec(), tmp_path)
        wrapped = {m.element_id for m in spec.paragraph_map if m.container_type == "content_control"}
        assert wrapped == {"cc1p0", "cc1t1r0", "cc1t1r1"}
        body_index = {m.element_id: m.body_index for m in spec.paragraph_map}
        assert body_index["cc1p0"] == body_index["cc1t1r0"] == 1

    def test_nested_block_controls_get_nested_ids(self, tmp_path):
        spec = _extract(fx.build_nested_controls_spec(), tmp_path)
        assert [(m.element_id, m.text) for m in spec.paragraph_map] == [
            ("p0", fx.NESTED_INLINE_TEXT),
            ("cc1p0", fx.NESTED_BLOCK_PARAGRAPH),
            ("cc1cc1p0", fx.DOUBLY_NESTED_PARAGRAPH),
            ("p2", fx.AFTER_CONTROL),
        ]

    def test_a_custom_xml_block_is_read_like_a_control(self, tmp_path):
        builder = fx.SpecDocBuilder()
        builder.add_paragraph("Before.")
        builder.add_xml(
            '<w:customXml w:uri="urn:fixture" w:element="clause">'
            "<w:customXmlPr/>" + fx.paragraph("Custom XML paragraph.") + "</w:customXml>"
        )
        spec = _extract(builder, tmp_path)
        assert _ids(spec) == {"p0": "Before.", "cc1p0": "Custom XML paragraph."}

    def test_a_heading_inside_a_control_sets_the_section_that_follows(self, tmp_path):
        """Section attribution continues through a control (routing reads the
        SECTION heading, which templates sometimes put in one — plan S13)."""
        builder = fx.SpecDocBuilder()
        builder.add_xml(
            fx.block_control(fx.paragraph("PART 2 PRODUCTS"), tag="part", control_id=builder.next_id())
        )
        builder.add_text("A. Provide listed valves.")
        spec = _extract(builder, tmp_path)
        by_id = {m.element_id: m for m in spec.paragraph_map}
        assert by_id["p1"].section_id == "PART 2 PRODUCTS"
        rendered = render_spec_with_ids(spec.content, spec.paragraph_map, filename=spec.filename)
        assert '<heading id="cc0p0">PART 2 PRODUCTS</heading>' in rendered
        assert '<para id="p1" section="PART 2 PRODUCTS">A. Provide listed valves.</para>' in rendered

    def test_an_empty_control_contributes_nothing(self, tmp_path):
        builder = fx.SpecDocBuilder()
        builder.add_paragraph("Only text.")
        builder.add_xml("<w:sdt><w:sdtPr><w:id w:val=\"9\"/></w:sdtPr><w:sdtContent/></w:sdt>")
        builder.add_xml("<w:sdt><w:sdtPr><w:id w:val=\"8\"/></w:sdtPr></w:sdt>")
        spec = _extract(builder, tmp_path)
        assert spec.content == "Only text." and spec.extraction_warnings == []

    def test_word_count_includes_control_text(self, tmp_path):
        spec = _extract(fx.build_block_control_spec(), tmp_path)
        assert spec.word_count == len(spec.content.split())
        assert spec.content.count(fx.BLOCK_CONTROL_PARAGRAPH) == 1

    def test_the_table_of_contents_is_skipped_so_headings_are_not_duplicated(self, tmp_path):
        spec = _extract(fx.build_table_of_contents_spec(), tmp_path)
        assert spec.content == fx.blocks_text(fx.CLEAN_THREE_PART)
        assert "Contents" not in spec.content
        result = preprocess_spec(spec.content, spec.filename, cycle=CALIFORNIA_2025)
        assert result.structural_alerts == []
        assert spec.extraction_warnings == []

    def test_a_control_of_another_gallery_is_read(self, tmp_path):
        """Only the table of contents is skipped: a cover page or a quick
        part is authored text."""
        builder = fx.SpecDocBuilder()
        builder.add_xml(
            '<w:sdt><w:sdtPr><w:id w:val="5"/><w:docPartObj>'
            '<w:docPartGallery w:val="Cover Pages"/></w:docPartObj></w:sdtPr>'
            f"<w:sdtContent>{fx.paragraph('Project Manual, Volume 2')}</w:sdtContent></w:sdt>"
        )
        spec = _extract(builder, tmp_path)
        assert _ids(spec) == {"cc0p0": "Project Manual, Volume 2"}


# ---------------------------------------------------------------------------
# Controls inside tables
# ---------------------------------------------------------------------------


class TestTableControls:
    def test_controls_in_cells_honor_revisions_and_join_their_row(self, tmp_path):
        spec = _extract(fx.build_table_controls_spec(), tmp_path)
        assert _ids(spec)["t0r1"] == f"Riser | {fx.CELL_CONTROL_TEXT}\n{fx.CELL_CONTROL_REVISED}"
        assert "zinc" not in spec.content
        assert spec.tracked_changes_detected is True

    def test_a_wrapped_cell_joins_its_row_in_document_order(self, tmp_path):
        spec = _extract(fx.build_table_controls_spec(), tmp_path)
        assert _ids(spec)["t0r2"] == f"Branch | {fx.WRAPPED_CELL_TEXT}"

    def test_wrapped_rows_get_their_own_ids_and_leave_the_others_numbered(self, tmp_path):
        spec = _extract(fx.build_table_controls_spec(), tmp_path)
        assert list(_ids(spec)) == ["t0r0", "t0r1", "t0r2", "t0cc5r0"]
        by_id = {m.element_id: m for m in spec.paragraph_map}
        assert by_id["t0cc5r0"].container_type == "content_control"
        assert by_id["t0r2"].container_type is None
        rendered = render_spec_with_ids(spec.content, spec.paragraph_map, filename=spec.filename)
        assert '<row id="t0cc5r0">Main | Six inch welded</row>' in rendered

    def test_a_control_between_rows_leaves_the_rows_after_it_numbered(self, tmp_path):
        """``r<n>`` counts the table's own ``<w:tr>`` children, so a row-level
        control in the middle of a table renumbers nothing after it."""
        builder = fx.SpecDocBuilder()
        wrapped = fx._row_control(f"<w:tr>{fx._cell_xml('B')}</w:tr>", tag="r", control_id=1)
        builder.add_xml(
            '<w:tbl><w:tblPr><w:tblW w:w="0" w:type="auto"/></w:tblPr>'
            '<w:tblGrid><w:gridCol w:w="2000"/></w:tblGrid>'
            f"<w:tr>{fx._cell_xml('A')}</w:tr>{wrapped}<w:tr>{fx._cell_xml('C')}</w:tr></w:tbl>"
        )
        spec = _extract(builder, tmp_path)
        assert list(_ids(spec).items()) == [("t0r0", "A"), ("t0cc3r0", "B"), ("t0r1", "C")]

    def test_a_cell_control_inside_a_wrapped_row_is_read(self, tmp_path):
        builder = fx.SpecDocBuilder()
        wrapped_cell = fx.block_control(fx._cell_xml("Six inch welded"), tag="c", control_id=2)
        wrapped = fx._row_control(
            f"<w:tr>{fx._cell_xml('Main')}{wrapped_cell}</w:tr>", tag="r", control_id=1
        )
        builder.add_xml(
            '<w:tbl><w:tblPr><w:tblW w:w="0" w:type="auto"/></w:tblPr>'
            '<w:tblGrid><w:gridCol w:w="2000"/><w:gridCol w:w="2000"/></w:tblGrid>'
            f"{wrapped}</w:tbl>"
        )
        assert _ids(_extract(builder, tmp_path)) == {"t0cc2r0": "Main | Six inch welded"}

    def test_an_inline_control_in_a_plain_cell_is_read(self, tmp_path):
        builder = fx.SpecDocBuilder()
        cell = fx._cell_of(
            fx.paragraph("Size: ", fx.inline_control(fx.run("4 inch"), tag="s", control_id=builder.next_id()))
        )
        builder.add_xml(
            '<w:tbl><w:tblPr><w:tblW w:w="0" w:type="auto"/></w:tblPr>'
            f'<w:tblGrid><w:gridCol w:w="2000"/></w:tblGrid><w:tr>{cell}</w:tr></w:tbl>'
        )
        assert _ids(_extract(builder, tmp_path)) == {"t0r0": "Size: 4 inch"}

    def test_a_wrapped_cell_continuing_a_vertical_merge_adds_nothing(self, tmp_path):
        """Word shows a merged region with its origin cell's content, and the
        walk never reads a continuation cell's own content — for an ordinary
        cell (python-docx maps it to its origin) or a wrapped one."""
        builder = fx.SpecDocBuilder()
        continuation = (
            '<w:sdt><w:sdtPr><w:id w:val="4"/></w:sdtPr><w:sdtContent>'
            '<w:tc><w:tcPr><w:vMerge/></w:tcPr><w:p><w:r><w:t>STALE</w:t></w:r></w:p></w:tc>'
            "</w:sdtContent></w:sdt>"
        )
        builder.add_xml(
            '<w:tbl><w:tblPr><w:tblW w:w="0" w:type="auto"/></w:tblPr>'
            '<w:tblGrid><w:gridCol w:w="2000"/><w:gridCol w:w="2000"/></w:tblGrid>'
            f"<w:tr>{fx._cell_xml('A')}{fx._cell_xml('B')}</w:tr>"
            f"<w:tr>{fx._cell_xml('C')}{continuation}</w:tr></w:tbl>"
        )
        assert _ids(_extract(builder, tmp_path)) == {"t0r0": "A | B", "t0r1": "C"}

    def test_a_nested_table_in_an_ordinary_cell_is_still_read(self, tmp_path):
        spec = _extract(fx.build_merged_and_nested_tables_spec(), tmp_path)
        assert "t1r0c1t0r1" in _ids(spec) and spec.extraction_warnings == []

    @pytest.mark.parametrize("where", ["cell_control", "wrapped_cell", "wrapped_row"])
    def test_a_table_inside_a_table_control_is_reported_not_read(self, where, tmp_path):
        builder = fx.SpecDocBuilder()
        inner = fx.table([("Hidden schedule", "Value")])
        if where == "cell_control":
            cell = fx._cell_of(fx.block_control(inner, tag="c", control_id=1), fx.paragraph("x"))
            rows = f"<w:tr>{cell}</w:tr>"
        elif where == "wrapped_cell":
            cell = f"<w:sdt><w:sdtPr><w:id w:val=\"1\"/></w:sdtPr><w:sdtContent>{fx._cell_of(inner, fx.paragraph('x'))}</w:sdtContent></w:sdt>"
            rows = f"<w:tr>{cell}</w:tr>"
        else:
            rows = fx._row_control(f"<w:tr>{fx._cell_of(inner, fx.paragraph('x'))}</w:tr>", tag="r", control_id=1)
        builder.add_xml(
            '<w:tbl><w:tblPr><w:tblW w:w="0" w:type="auto"/></w:tblPr>'
            f'<w:tblGrid><w:gridCol w:w="2000"/></w:tblGrid>{rows}</w:tbl>'
        )
        spec = _extract(builder, tmp_path)
        assert "Hidden schedule" not in spec.content
        assert spec.extraction_warnings == [
            "Spec contains 1 table inside content controls within table rows or cells; "
            "their text was not extracted for review. Verify visually."
        ]


# ---------------------------------------------------------------------------
# Headers, footers, text boxes, and notes
# ---------------------------------------------------------------------------


def _attach_notes(document, *, footnotes: str = "", endnotes: str = "") -> None:
    """Attach footnote / endnote parts holding the given ``<w:footnote>`` /
    ``<w:endnote>`` fragments (plus Word's structural separator)."""
    for kind, fragments, content_type in (
        ("footnote", footnotes, _FOOTNOTES_CONTENT_TYPE),
        ("endnote", endnotes, _ENDNOTES_CONTENT_TYPE),
    ):
        if not fragments:
            continue
        blob = (
            f'<?xml version="1.0"?><w:{kind}s xmlns:w="{_W_NS}">'
            f'<w:{kind} w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:{kind}>'
            f"{fragments}</w:{kind}s>"
        ).encode()
        part = Part(PackURI(f"/word/{kind}s.xml"), content_type, blob, document.part.package)
        document.part.relate_to(
            part,
            f"http://schemas.openxmlformats.org/officeDocument/2006/relationships/{kind}s",
        )


class TestSupplementalStories:
    def test_a_header_control_is_read_beside_the_headers_own_paragraphs(self, tmp_path):
        builder = fx.SpecDocBuilder()
        builder.add_text("Body.")
        header = builder.document.sections[0].header
        header.paragraphs[0].text = "Project 2401"
        for element in fx.parse_fragments(
            fx.block_control(fx.paragraph("Page 3 of 9"), tag="page", control_id=builder.next_id()),
            fx.paragraph("Issued for bid"),
        ):
            header._element.append(element)
        spec = _extract(builder, tmp_path)
        rows = [(m.element_id, m.text, m.container_type) for m in spec.paragraph_map if m.element_type == "header"]
        assert rows == [
            ("s0h0", "[Header] Project 2401", "header"),
            ("s0hcc1p0", "[Header] Page 3 of 9", "content_control"),
            ("s0h1", "[Header] Issued for bid", "header"),
        ]

    def test_a_revision_inside_a_header_control_is_detected(self, tmp_path):
        builder = fx.SpecDocBuilder()
        builder.add_text("Body.")
        control = fx.block_control(
            fx.paragraph(fx.inserted(fx.run("Revised title"), revision_id=1)), tag="t", control_id=2
        )
        for element in fx.parse_fragments(control):
            builder.document.sections[0].footer._element.append(element)
        spec = _extract(builder, tmp_path)
        assert "[Footer] Revised title" in spec.content
        assert spec.tracked_changes_detected is True

    def test_a_control_inside_a_text_box_is_read(self, tmp_path):
        box = (
            '<w:p xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"'
            ' xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
            ' xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">'
            '<w:r><w:drawing><wp:inline><a:graphic><a:graphicData uri="x"><wps:wsp><wps:txbx><w:txbxContent>'
            + fx.paragraph("Box text.")
            + fx.block_control(fx.paragraph("Boxed control text."), tag="b", control_id=1)
            + "</w:txbxContent></wps:txbx></wps:wsp></a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p>"
        )
        builder = fx.SpecDocBuilder()
        builder.add_xml(box)
        spec = _extract(builder, tmp_path)
        boxes = [(m.element_id, m.text) for m in spec.paragraph_map if m.element_type == "textbox"]
        assert boxes == [
            ("tb0p0", "[Text Box] Box text."),
            ("tb0cc1p0", "[Text Box] Boxed control text."),
        ]

    def test_a_control_inside_a_footnote_is_read(self, tmp_path):
        builder = fx.SpecDocBuilder()
        builder.add_text("Body.")
        _attach_notes(
            builder.document,
            footnotes=(
                '<w:footnote w:id="1">'
                + fx.paragraph("Note text.")
                + fx.block_control(fx.paragraph("Noted control text."), tag="n", control_id=1)
                + "</w:footnote>"
            ),
        )
        spec = _extract(builder, tmp_path)
        notes = [(m.element_id, m.text) for m in spec.paragraph_map if m.element_type == "footnote"]
        assert notes == [
            ("fn1p0", "[Footnote 1] Note text."),
            ("fn1cc1p0", "[Footnote 1] Noted control text."),
        ]


class TestTextBoxesAreReadOnce:
    """Word saves each text box twice — a DrawingML shape and a VML copy in
    ``<mc:AlternateContent>`` — and both copies used to be read."""

    _NS = (
        ' xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'
        ' xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"'
        ' xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
        ' xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"'
        ' xmlns:v="urn:schemas-microsoft-com:vml"'
    )

    def _alternate(self, choice_text: str | None, fallback_text: str) -> str:
        choice = (
            '<mc:Choice Requires="wps"><w:drawing><wp:anchor><a:graphic><a:graphicData uri="x">'
            + (
                f"<wps:wsp><wps:txbx><w:txbxContent>{fx.paragraph(choice_text)}</w:txbxContent></wps:txbx></wps:wsp>"
                if choice_text is not None
                else "<wps:wsp/>"
            )
            + "</a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice>"
        )
        fallback = (
            "<mc:Fallback><w:pict><v:shape><v:textbox><w:txbxContent>"
            f"{fx.paragraph(fallback_text)}</w:txbxContent></v:textbox></v:shape></w:pict></mc:Fallback>"
        )
        return f"<w:p{self._NS}><w:r><mc:AlternateContent>{choice}{fallback}</mc:AlternateContent></w:r></w:p>"

    def test_words_two_copies_are_read_once(self, tmp_path):
        builder = fx.SpecDocBuilder()
        builder.add_xml(self._alternate("First box.", "First box."), self._alternate("Second box.", "Second box."))
        spec = _extract(builder, tmp_path)
        assert [(m.element_id, m.text) for m in spec.paragraph_map if m.element_type == "textbox"] == [
            ("tb0p0", "[Text Box] First box."),
            ("tb1p0", "[Text Box] Second box."),
        ]

    def test_the_first_branch_holding_a_text_box_is_the_one_read(self, tmp_path):
        """Word writes the same text in both branches; which one is read is
        pinned with different texts so the rule cannot drift unseen."""
        builder = fx.SpecDocBuilder()
        builder.add_xml(self._alternate("From the DrawingML choice.", "From the VML fallback."))
        spec = _extract(builder, tmp_path)
        assert "[Text Box] From the DrawingML choice." in spec.content
        assert "VML fallback" not in spec.content

    def test_a_fallback_is_read_when_the_choice_has_no_text_box(self, tmp_path):
        builder = fx.SpecDocBuilder()
        builder.add_xml(self._alternate(None, "Only in the fallback."))
        spec = _extract(builder, tmp_path)
        assert "[Text Box] Only in the fallback." in spec.content

    def test_text_boxes_inside_a_grouped_shape_are_read(self, tmp_path):
        """Older gap lists named "grouped-shape text" as unread. A text box in
        a group is read like any other (it was before S10 too); SmartArt, whose
        text lives in a separate diagram part, is the gap that remains."""
        ns = self._NS + ' xmlns:wpg="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup"'
        shapes = "".join(
            f"<wps:wsp><wps:txbx><w:txbxContent>{fx.paragraph(text)}</w:txbxContent></wps:txbx></wps:wsp>"
            for text in ("First grouped box.", "Second grouped box.")
        )
        builder = fx.SpecDocBuilder()
        builder.add_xml(
            f"<w:p{ns}><w:r><w:drawing><wp:inline><a:graphic><a:graphicData uri=\"x\">"
            f"<wpg:wgp>{shapes}</wpg:wgp></a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p>"
        )
        spec = _extract(builder, tmp_path)
        assert [(m.element_id, m.text) for m in spec.paragraph_map if m.element_type == "textbox"] == [
            ("tb0p0", "[Text Box] First grouped box."),
            ("tb1p0", "[Text Box] Second grouped box."),
        ]


# ---------------------------------------------------------------------------
# What is left unread says so
# ---------------------------------------------------------------------------


def _math(text: str) -> str:
    return f'<m:oMath xmlns:m="{_MATH_NS}"><m:r><m:t>{text}</m:t></m:r></m:oMath>'


_LEGACY_DROP_DOWN = (
    _begin(
        '<w:ffData><w:name w:val="Dropdown1"/><w:enabled/><w:ddList>'
        '<w:result w:val="1"/><w:listEntry w:val="Copper"/><w:listEntry w:val="Steel"/>'
        "</w:ddList></w:ffData>"
    )
    + _instr(" FORMDROPDOWN ")
    + _END
)


class TestUnreadStructuresWarn:
    def test_equations_are_counted_not_read(self, tmp_path):
        builder = fx.SpecDocBuilder()
        # Raw paragraphs: ``fx.paragraph`` would treat an ``m:`` fragment as text.
        builder.add_xml(f"<w:p>{fx.run('Area ')}{_math('A=πr2')}{fx.run(' applies.')}</w:p>")
        builder.add_xml(f"<w:p>{_math('Q=VA')}</w:p>")
        spec = _extract(builder, tmp_path)
        assert spec.content == "Area  applies."
        assert spec.extraction_warnings == [
            "Spec contains 2 equations (Office Math) whose text was not extracted "
            "for review. Verify visually."
        ]

    def test_an_embedded_document_is_counted(self, tmp_path):
        builder = fx.SpecDocBuilder()
        builder.add_paragraph("Before.")
        builder.add_xml('<w:altChunk r:id="rIdChunk"/>')
        spec = _extract(builder, tmp_path)
        assert spec.extraction_warnings == [
            "Spec contains 1 embedded document (altChunk) whose text was not "
            "extracted for review. Verify visually."
        ]

    def test_a_legacy_drop_down_form_field_is_counted(self, tmp_path):
        builder = fx.SpecDocBuilder()
        builder.add_paragraph("Material: ", _LEGACY_DROP_DOWN, ".")
        spec = _extract(builder, tmp_path)
        assert spec.content == "Material: ."
        assert spec.extraction_warnings == [
            "Spec contains 1 legacy drop-down form field whose chosen entry was "
            "not extracted for review. Verify visually."
        ]

    def test_several_legacy_drop_downs_are_counted_in_one_warning(self, tmp_path):
        builder = fx.SpecDocBuilder()
        builder.add_paragraph("Material: ", _LEGACY_DROP_DOWN, ".")
        builder.add_paragraph("Finish: ", _LEGACY_DROP_DOWN, ".")
        assert _extract(builder, tmp_path).extraction_warnings == [
            "Spec contains 2 legacy drop-down form fields whose chosen entries were "
            "not extracted for review. Verify visually."
        ]

    def test_warnings_state_counts_never_measures_of_lost_text(self, tmp_path):
        builder = fx.SpecDocBuilder()
        builder.add_xml(f"<w:p>{_math('x')}{_LEGACY_DROP_DOWN}</w:p>")
        builder.add_xml('<w:altChunk r:id="rIdChunk"/>')
        warnings = _extract(builder, tmp_path).extraction_warnings
        assert len(warnings) == 3
        for warning in warnings:
            assert "%" not in warning and "characters" not in warning

    @pytest.mark.parametrize(
        "build",
        [
            fx.build_clean_three_part,
            fx.build_block_control_spec,
            fx.build_inline_controls_spec,
            fx.build_fields_spec,
            fx.build_smart_tag_spec,
            fx.build_hyperlink_spec,
            fx.build_table_controls_spec,
            fx.build_nested_controls_spec,
            fx.build_table_of_contents_spec,
        ],
        ids=lambda build: build.__name__,
    )
    def test_documents_with_nothing_unread_carry_no_warning(self, build, tmp_path):
        assert _extract(build(), tmp_path).extraction_warnings == []
