"""Shared jurisdiction requests, applicability, EX-05 and pending resume."""
from __future__ import annotations

import json
import threading
from collections import Counter
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document

from src.batch.batch import BatchJob
from src.compliance.compliance_checker import (
    _render_profile_block, expected_coverage_ids, run_compliance_check,
)
from src.core.project_profile import ProjectProfile
from src.input.extractor import ExtractedSpec
from src.modules import DEFAULT_MODULE, require_module, validate_module_registry
from src.orchestration import pipeline as pl, program_pipeline as pp
from src.orchestration.batch_resume import (
    PendingProgramRun, load_pending_run, save_pending_program_run,
)
from src.programs import HYPERSCALE_DATACENTER_PROGRAM, RoutingState, SpecAssignment, SpecRoutingDecision
from src.research import RequirementsProfile
from src.research import requirements_research as rr
from src.research.shared_jurisdiction import (
    DATACENTER_MODULE_IDS, JURISDICTION_SCOPE_ID, SupplementSignals,
    compose_module_profile, jurisdiction_research_plan,
)
from src.review.structured_schemas import requirements_research_tool
from tests.fixtures.fake_anthropic import research_tool_use_response, sample_research_profile_payload
from tests.test_requirements_research import FakeResearchClient


PROFILE = ProjectProfile("Ashburn", "VA", "US", "ExampleCo")


def _assignments(tmp_path, module_ids=DATACENTER_MODULE_IDS):
    assignments = []
    for index, module_id in enumerate(module_ids):
        path = tmp_path / f"{index} Specification.docx"
        doc = Document()
        doc.add_paragraph("1.01 SUMMARY\nProvide systems per NFPA 13-2019.")
        doc.save(path)
        assignments.append(SpecAssignment(
            source_path=str(path),
            decision=SpecRoutingDecision(
                spec_id=path.name, program_id=HYPERSCALE_DATACENTER_PROGRAM.program_id,
                automatic_state=RoutingState.SUPPORTED, automatic_module_ids=(module_id,),
                confidence=1.0, evidence=(),
            ),
        ))
    return assignments


@pytest.fixture
def research_env(monkeypatch, tmp_path):
    monkeypatch.delenv("SPEC_CRITIC_RESEARCH_CACHE", raising=False)
    monkeypatch.setenv("SPEC_CRITIC_RESEARCH_CACHE_PATH", str(tmp_path / "research-cache.json"))
    requests = Counter()
    lock = threading.Lock()
    plans = [jurisdiction_research_plan()] + [require_module(mid) for mid in DATACENTER_MODULE_IDS]
    briefs = {
        rr.build_dimension_user_message(plan, PROFILE, dimension): dimension.dimension_id
        for plan in plans for dimension in plan.research_dimensions
    }

    def route(kwargs):
        brief = kwargs["messages"][0]["content"].split("\n\n<corpus_signals>")[0]
        dimension_id = briefs[brief]
        with lock:
            requests[dimension_id] += 1
        item = sample_research_profile_payload()["items"][0]
        item["requirement"] = f"Requirement for {dimension_id}."
        # Some shared items apply only to electrical; all profiles still carry them.
        if dimension_id.startswith("jurisdiction_"):
            item["applicable_module_ids"] = (
                ["datacenter_electrical"] if dimension_id == "jurisdiction_client"
                else list(DATACENTER_MODULE_IDS)
            )
        return research_tool_use_response(payload={"summary": "Research", "items": [item]})

    client = FakeResearchClient(route)
    monkeypatch.setattr(rr, "_get_client", lambda **_kw: client)
    monkeypatch.setattr(rr, "context_within_token_cap", lambda text: (len(text.split()), True))
    monkeypatch.setattr("src.research.corpus_signals.count_tokens", lambda text: len(text.split()))
    # Exercise the real extraction/preprocessing/preparation, without live counts.
    prepare_specs = pl._prepare_specs
    monkeypatch.setattr(pl, "_prepare_specs", lambda **kw: prepare_specs(**kw, preflight=False))
    return SimpleNamespace(requests=requests, client=client, cache=tmp_path / "research-cache.json")


def _program(tmp_path, assignments):
    return pp.prepare_program_review(
        program_id=HYPERSCALE_DATACENTER_PROGRAM.program_id, assignments=assignments,
        input_dir=tmp_path, model="test-model", project_profile=PROFILE,
    )


def test_four_module_program_researches_core_once_and_composes_every_profile(tmp_path, research_env):
    prepared = _program(tmp_path, _assignments(tmp_path))
    expected = {
        dimension.dimension_id for dimension in jurisdiction_research_plan().research_dimensions
    } | {
        dimension.dimension_id for mid in DATACENTER_MODULE_IDS
        for dimension in require_module(mid).research_dimensions
    }
    assert research_env.requests == Counter({dimension: 1 for dimension in expected})
    assert len(research_env.client.calls) == 8  # formerly 4 + 4 + 5 + 5 = 18
    shared_ids = None
    for mid, child in prepared.partitions.items():
        profile = RequirementsProfile.from_dict(child.requirements_profile)
        shared = [item for item in profile.items if item.dimension_id.startswith("jurisdiction_")]
        assert len(shared) == 4 and len(profile.items) == 5
        assert all(item.grounded for item in profile.items)
        assert len(profile.dimension_statuses) == 5
        assert profile.module_id == mid
        if shared_ids is None:
            shared_ids = [item.item_id for item in shared]
        assert [item.item_id for item in shared] == shared_ids
        electrical_only = next(item for item in shared if item.dimension_id == "jurisdiction_client")
        assert (electrical_only.item_id in expected_coverage_ids(profile)) == (mid == "datacenter_electrical")
        assert (electrical_only.requirement in _render_profile_block(profile)) == (mid == "datacenter_electrical")
        basis_ids = {item["item_id"] for item in child.governing_basis["items"]}
        assert (electrical_only.item_id in basis_ids) == (mid == "datacenter_electrical")
        assert "PROJECT REQUIREMENTS PROFILE" in child.effective_context
    fire = RequirementsProfile.from_dict(prepared.partitions["datacenter_fire"].requirements_profile)
    assert "[CONTEXT ONLY]" in fire.render_text()
    assert not research_env.cache.exists()  # EX-05 remains opt-in


@pytest.mark.parametrize("mid", DATACENTER_MODULE_IDS)
def test_standalone_module_includes_the_shared_core(tmp_path, research_env, mid):
    assignment = _assignments(tmp_path, (mid,))[0]
    prepared = pl.prepare_batch_review(
        input_dir=tmp_path, files=[Path(assignment.source_path)], module=require_module(mid),
        project_profile=PROFILE, model="test-model",
    )
    assert len(research_env.client.calls) == 5
    assert len(prepared.requirements_profile["items"]) == 5


def test_core_cache_has_independent_key_and_supplements_key_the_supplied_core(tmp_path, research_env, monkeypatch):
    monkeypatch.setenv("SPEC_CRITIC_RESEARCH_CACHE", "reuse")
    assignments = _assignments(tmp_path)
    first = _program(tmp_path, assignments)
    assert len(research_env.client.calls) == 8
    entries = json.loads(research_env.cache.read_text())["entries"]
    assert Counter(entry["subject"]["module_id"] for entry in entries.values()) == Counter({
        scope: 1 for scope in (*DATACENTER_MODULE_IDS, JURISDICTION_SCOPE_ID)
    })
    second = _program(tmp_path, assignments)
    assert len(research_env.client.calls) == 8
    for child in second.partitions.values():
        profile = child.requirements_profile
        assert set(profile["reuse"]["scopes"]) == {JURISDICTION_SCOPE_ID, child.module.module_id}
        assert len(profile["research_components"]) == 2
    shared = RequirementsProfile.from_dict(first.partitions["datacenter_fire"].requirements_profile)
    module = require_module("datacenter_fire")
    key = rr.research_reuse_key(module, PROFILE, corpus_signals=SupplementSignals(shared, module.module_id))
    shared.items[0].requirement += " Changed adoption."
    changed = rr.research_reuse_key(module, PROFILE, corpus_signals=SupplementSignals(shared, module.module_id))
    assert key.key != changed.key
    # Selecting only one discipline retains the same core key.
    standalone = pl.prepare_batch_review(
        input_dir=tmp_path, files=[Path(assignments[0].source_path)], module=module,
        project_profile=PROFILE, model="test-model",
    )
    assert len(research_env.client.calls) == 8
    assert standalone.requirements_profile["reuse"]["scopes"] == [JURISDICTION_SCOPE_ID, module.module_id]


@pytest.mark.parametrize("value", [None, "datacenter_fire", {}, [42], [], ["other_module"]])
def test_shared_applicability_fails_closed_on_unassigned_or_malformed_payload(value):
    plan = jurisdiction_research_plan()
    item = sample_research_profile_payload()["items"][0]
    item["applicable_module_ids"] = value
    shared_item = rr._items_from_payload(
        {"items": [item]}, "jurisdiction_governing_codes",
        applicable_module_ids=plan.research_applicability_module_ids,
    )[0]
    assert shared_item.applicable_module_ids == []
    shared_item.grounded = True
    profile = RequirementsProfile(items=[shared_item], module_id="datacenter_fire")
    restored = RequirementsProfile.from_dict(profile.to_dict())
    assert expected_coverage_ids(restored) == ()


def test_pending_program_resume_preserves_scope_and_never_researches(tmp_path, research_env, monkeypatch):
    assignments = _assignments(tmp_path)
    prepared = _program(tmp_path, assignments)
    children = {}
    for mid, child in prepared.partitions.items():
        rid = f"review__{mid}__0"
        children[mid] = pl.BatchSubmission(
            job=BatchJob(batch_id=f"msgbatch_{mid}", job_type="review", created_at=1.0,
                         request_map={rid: {"filename": child.prepared.specs[0].filename, "index": 0, "type": "review"}}),
            files_reviewed=[child.prepared.specs[0].filename], review_request_ids=[rid],
            model="test-model", prepared_specs=child.prepared.specs,
            project_context=child.effective_context, module_id=mid, cycle_label=child.module.cycle.label,
            project_profile=child.project_profile, requirements_profile=child.requirements_profile,
            governing_basis=child.governing_basis,
        )
    submission = pp.ProgramSubmission(
        program_id=prepared.program_id, assignments=assignments, partitions=children,
        project_profile=PROFILE.to_dict(),
    )
    pending = PendingProgramRun.from_submission(submission, input_dir=tmp_path)
    path = tmp_path / "pending.json"
    assert save_pending_program_run(pending, path=path)
    loaded = load_pending_run(path=path)
    monkeypatch.setattr(rr, "_get_client", lambda **_kw: pytest.fail("resume reran research"))
    resumed = loaded.to_submission()
    assert len(research_env.client.calls) == 8
    for mid, child in resumed.partitions.items():
        assert child.requirements_profile == children[mid].requirements_profile
        assert child.governing_basis == children[mid].governing_basis


def test_bad_partition_is_rejected_before_shared_network_work(tmp_path, research_env):
    assignments = _assignments(tmp_path)
    Path(assignments[-1].source_path).write_bytes(b"corrupt docx")
    with pytest.raises(Exception):
        _program(tmp_path, assignments)
    assert not research_env.client.calls


def test_module_slots_and_california_are_preserved(tmp_path, research_env):
    validate_module_registry([require_module(mid) for mid in DATACENTER_MODULE_IDS])
    assert not DEFAULT_MODULE.research_dimensions
    assignment = _assignments(tmp_path, ("datacenter_fire",))[0]
    prepared = pl.prepare_batch_review(
        input_dir=tmp_path, files=[Path(assignment.source_path)], module=DEFAULT_MODULE,
        project_profile=PROFILE, model="test-model",
    )
    assert prepared.requirements_profile is None and prepared.governing_basis is None
    assert not research_env.client.calls


def test_inapplicable_core_creates_no_compliance_rows_or_additions(monkeypatch):
    from src.compliance import compliance_checker as cc
    item = rr.ResearchItem(
        item_id="r-electrical", dimension_id="jurisdiction_governing_codes",
        topic="Power equipment", category="governing_code", requirement="An electrical-only requirement.",
        grounded=True, accepted_sources=["https://codes.example.gov/adoption"],
        applicable_module_ids=["datacenter_electrical"],
    )
    profile = RequirementsProfile(items=[item], module_id="datacenter_fire")
    monkeypatch.setattr(cc, "_get_client", lambda **_kw: pytest.fail("no applicable items must not call compliance"))
    result = run_compliance_check(
        [ExtractedSpec(filename="Sprinklers.docx", content="Provide sprinklers.", word_count=2)],
        profile, [],
    )
    assert result.coverage == [] and result.findings == []
    assert result.coverage_completeness.state == "no_applicable_items"


def test_shared_schema_requires_explicit_scope_without_mutating_discipline_schema():
    ordinary = requirements_research_tool()
    shared = requirements_research_tool(applicable_module_ids=DATACENTER_MODULE_IDS)
    item = shared["input_schema"]["properties"]["items"]["items"]
    assert "applicable_module_ids" in item["required"]
    assert item["properties"]["applicable_module_ids"]["items"]["enum"] == list(DATACENTER_MODULE_IDS)
    assert "applicable_module_ids" not in ordinary["input_schema"]["properties"]["items"]["items"]["properties"]


def test_shared_profile_copies_and_component_reuse_remain_honest():
    shared = RequirementsProfile(
        items=[rr.ResearchItem("r-shared", "jurisdiction_site", "Hazard", "site_environment", "Original",
                               applicable_module_ids=list(DATACENTER_MODULE_IDS))],
        project=PROFILE.to_dict(), research_date="2026-09-20",
        reuse={"source": "research_cache", "age_days": 15, "research_date": "2026-09-20", "key": "shared-key"},
    )
    supplement = RequirementsProfile(project=PROFILE.to_dict(), research_date="2026-10-05")
    fire = compose_module_profile(shared, supplement, "datacenter_fire")
    electrical = compose_module_profile(shared, supplement, "datacenter_electrical")
    fire.items[0].requirement = "Changed"
    fire.items[0].applicable_module_ids.clear()
    assert electrical.items[0].requirement == shared.items[0].requirement == "Original"
    assert electrical.items[0].applicable_module_ids == list(DATACENTER_MODULE_IDS)
    assert fire.research_date == "2026-09-20"
    assert fire.research_components["datacenter_fire"]["research_date"] == "2026-10-05"
    from src.research.research_cache import reuse_notice
    assert fire.reuse["scopes"] == [JURISDICTION_SCOPE_ID]
    assert "Other components may be fresh" in reuse_notice(fire)


def test_word_and_html_requirement_tables_exclude_other_disciplines():
    from src.output.report_exporter import _write_requirements_section
    from src.output.html_report_exporter import _render_requirements_section
    from tests.test_golden_datacenter_surfaces import _golden_composed_profile
    profile = _golden_composed_profile()
    doc = Document()
    _write_requirements_section(doc, profile, None, require_module("datacenter_fire"))
    word = "\n".join(p.text for p in doc.paragraphs)
    html, lines = _render_requirements_section(profile, None)
    assert "roof warranty" not in word and "roof warranty" not in html
    assert "Module requirement tables include only applicable items" in word
    assert "paid calls are counted once in diagnostics" in html


@pytest.mark.parametrize("module_id", DATACENTER_MODULE_IDS)
@pytest.mark.parametrize("grounded", [True, False])
def test_shared_process_advisories_remain_visible_without_specification_applicability(
    module_id, grounded, monkeypatch,
):
    from src.compliance import compliance_checker as cc
    from src.output.report_exporter import _write_requirements_section
    from src.output.html_report_exporter import _render_requirements_section

    advisory = rr.ResearchItem(
        item_id="r-shared-permit-window", dimension_id="jurisdiction_ahj",
        topic="Permit hearing window", category="ahj_requirement",
        requirement="Permit hearings occur on the first Tuesday of each month.",
        actionability="process_advisory", applicable_module_ids=[],
        grounded=grounded,
        accepted_sources=["https://city.example.gov/permits"] if grounded else [],
    )
    shared = RequirementsProfile(
        items=[advisory], project=PROFILE.to_dict(), research_date="2026-10-05",
        dimension_statuses=[rr.DimensionStatus("jurisdiction_ahj", "completed")],
    )
    supplement = RequirementsProfile(project=PROFILE.to_dict(), research_date="2026-10-05")
    profile = RequirementsProfile.from_dict(compose_module_profile(shared, supplement, module_id).to_dict())
    assert not profile.item_applies(profile.items[0])
    assert expected_coverage_ids(profile) == ()
    assert advisory.requirement not in _render_profile_block(profile)
    monkeypatch.setattr(cc, "_get_client", lambda **_kw: pytest.fail("advisories must not call compliance"))
    compliance = run_compliance_check(
        [ExtractedSpec(filename="Spec.docx", content="Provide systems.", word_count=2)],
        profile, [],
    )
    assert compliance.coverage == [] and compliance.findings == []

    doc = Document()
    _write_requirements_section(doc, profile, compliance, require_module(module_id))
    word = "\n".join(p.text for p in doc.paragraphs)
    html, lines = _render_requirements_section(profile, compliance)
    assert "Process & Schedule Advisories" in word
    assert "Process &amp; Schedule Advisories" in html
    assert advisory.requirement in word and advisory.requirement in html
    assert advisory.requirement in "\n".join(lines)
