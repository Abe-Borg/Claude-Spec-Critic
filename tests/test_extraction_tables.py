"""Tests for merged-cell and nested-table extraction (B-28 a/b).

Two table-shaped gaps in ``extract_text_from_docx``:

* **Merged cells emitted more than once.** python-docx's ``_Row.cells``
  approximates a uniform grid, so a horizontally merged ``<w:tc>`` is
  returned once per grid column it spans and a vertically merged cell's
  continuation rows resolve to the origin cell above. The walk used to
  emit a 3-column merged heading three times on its row and a 3-row merged
  label once per spanned row. The XML stores the text once, so each
  ``<w:tc>`` now contributes its text exactly once, in its origin row.
* **Nested tables silently dropped.** ``_accept_all_cell_text`` reads only a
  cell's own paragraphs (by design, so nothing is double counted), and the
  walk never descended into ``cell.tables``. A schedule table nested inside
  a layout table — a common authoring pattern — was invisible to review.
  Nested rows are now emitted after the row that contains them under path
  ids (``t<table>r<row>c<cell>t<nested>r<row>``), depth-bounded, with the
  reconstruction invariant intact.

The content-loss warning is a separate, body-level heuristic and must be
unaffected by either change.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement

from src.input.extractor import (
    _NESTED_TABLE_DEPTH_WARNING,
    _NESTED_TABLE_MAX_DEPTH,
    extract_text_from_docx,
)
from src.review.prompt_serialization import render_spec_with_ids


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _fill(table, prefix: str = "") -> None:
    """Stamp every cell with a unique ``{prefix}r<row>c<col>`` label."""
    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            cell.text = f"{prefix}r{r}c{c}"


def _extract(doc, tmp_path: Path, name: str = "spec.docx"):
    out = tmp_path / name
    doc.save(out)
    return extract_text_from_docx(out)


def _table_rows(spec) -> list:
    return [m for m in spec.paragraph_map if m.element_type == "table_cell"]


def _assert_reconstructs(spec) -> None:
    """The invariant the extractor enforces with a ValueError — asserted
    explicitly here so a regression reads as a test failure, not a crash."""
    assert "\n\n".join(m.text for m in spec.paragraph_map) == spec.content


def _drawing_paragraph(doc) -> None:
    """Append a body paragraph whose only content is a ``<w:drawing>`` run."""
    para = doc.add_paragraph()
    run = OxmlElement("w:r")
    run.append(OxmlElement("w:drawing"))
    para._element.append(run)


# ---------------------------------------------------------------------------
# (a) Merged cells
# ---------------------------------------------------------------------------


class TestMergedCells:
    def test_horizontal_merge_emits_text_once(self, tmp_path: Path):
        doc = Document()
        table = doc.add_table(rows=2, cols=3)
        _fill(table)
        merged = table.cell(0, 0).merge(table.cell(0, 2))
        merged.text = "SCHEDULE HEADING"

        spec = _extract(doc, tmp_path)
        rows = _table_rows(spec)
        assert [m.text for m in rows] == ["SCHEDULE HEADING", "r1c0 | r1c1 | r1c2"]
        assert spec.content.count("SCHEDULE HEADING") == 1
        _assert_reconstructs(spec)

    def test_partial_horizontal_merge_keeps_unmerged_neighbour(self, tmp_path: Path):
        doc = Document()
        table = doc.add_table(rows=1, cols=3)
        _fill(table)
        table.cell(0, 0).merge(table.cell(0, 1)).text = "SPAN"

        spec = _extract(doc, tmp_path)
        assert [m.text for m in _table_rows(spec)] == ["SPAN | r0c2"]

    def test_vertical_merge_emits_text_once_in_origin_row(self, tmp_path: Path):
        # python-docx resolves a vMerge continuation cell to the origin
        # cell above it, so without table-wide dedupe the label repeated on
        # every spanned row. The text is stored once in the XML and is now
        # emitted once, in the row where it originates.
        doc = Document()
        table = doc.add_table(rows=3, cols=2)
        _fill(table)
        table.cell(0, 0).merge(table.cell(2, 0)).text = "PIPING"

        spec = _extract(doc, tmp_path)
        assert [m.text for m in _table_rows(spec)] == [
            "PIPING | r0c1",
            "r1c1",
            "r2c1",
        ]
        assert spec.content.count("PIPING") == 1
        _assert_reconstructs(spec)

    def test_combined_horizontal_and_vertical_merges(self, tmp_path: Path):
        doc = Document()
        table = doc.add_table(rows=3, cols=3)
        _fill(table)
        table.cell(0, 0).merge(table.cell(0, 2)).text = "TITLE"
        table.cell(1, 0).merge(table.cell(2, 0)).text = "GROUP"

        spec = _extract(doc, tmp_path)
        assert [m.text for m in _table_rows(spec)] == [
            "TITLE",
            "GROUP | r1c1 | r1c2",
            "r2c1 | r2c2",
        ]
        assert spec.content.count("TITLE") == 1
        assert spec.content.count("GROUP") == 1
        _assert_reconstructs(spec)

    def test_unmerged_table_output_unchanged(self, tmp_path: Path):
        # Regression pin for the common case: plain grids keep the exact
        # ``" | "`` join and ``t<table>r<row>`` ids they always had.
        doc = Document()
        doc.add_paragraph("Intro.")
        table = doc.add_table(rows=2, cols=2)
        _fill(table)

        spec = _extract(doc, tmp_path)
        rows = _table_rows(spec)
        assert [m.text for m in rows] == ["r0c0 | r0c1", "r1c0 | r1c1"]
        assert [m.element_id for m in rows] == ["t0r0", "t0r1"]
        assert all(m.container_type is None for m in rows)
        assert all(m.table_index == 0 for m in rows)
        assert [m.row_index for m in rows] == [0, 1]
        assert spec.content == "Intro.\n\nr0c0 | r0c1\n\nr1c0 | r1c1"
        _assert_reconstructs(spec)

    def test_dedupe_is_per_table_not_per_document(self, tmp_path: Path):
        # The seen-set is scoped to one table: two separate tables with
        # identical layouts must each emit their own text.
        doc = Document()
        for label in ("A", "B"):
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).merge(table.cell(0, 1)).text = f"{label}-merged"

        spec = _extract(doc, tmp_path)
        rows = _table_rows(spec)
        assert [m.text for m in rows] == ["A-merged", "B-merged"]
        assert [m.element_id for m in rows] == ["t0r0", "t1r0"]


# ---------------------------------------------------------------------------
# (b) Nested tables
# ---------------------------------------------------------------------------


def _nest(cell, rows: int, cols: int, prefix: str):
    """Add a ``rows x cols`` table inside ``cell`` and label its cells."""
    inner = cell.add_table(rows=rows, cols=cols)
    _fill(inner, prefix)
    return inner


class TestNestedTables:
    def test_nested_table_text_extracted_with_path_ids(self, tmp_path: Path):
        doc = Document()
        outer = doc.add_table(rows=1, cols=1)
        outer.cell(0, 0).text = "OUTER"
        _nest(outer.cell(0, 0), 1, 2, "in-")

        spec = _extract(doc, tmp_path)
        rows = _table_rows(spec)
        assert [m.text for m in rows] == ["OUTER", "in-r0c0 | in-r0c1"]
        assert [m.element_id for m in rows] == ["t0r0", "t0r0c0t0r0"]
        nested = rows[1]
        assert nested.element_type == "table_cell"
        assert nested.container_type == "nested_table"
        assert nested.table_index == 0
        assert nested.row_index == 0
        assert nested.body_index == rows[0].body_index
        assert spec.content == "OUTER\n\nin-r0c0 | in-r0c1"
        # The invariant, stated explicitly.
        assert "\n\n".join(m.text for m in spec.paragraph_map) == spec.content

    def test_nested_rows_follow_parent_row_in_cell_order(self, tmp_path: Path):
        doc = Document()
        outer = doc.add_table(rows=2, cols=2)
        _fill(outer, "o-")
        _nest(outer.cell(0, 0), 1, 1, "a-")
        _nest(outer.cell(0, 1), 2, 1, "b-")
        _nest(outer.cell(1, 1), 1, 1, "c-")

        spec = _extract(doc, tmp_path)
        rows = _table_rows(spec)
        assert [(m.element_id, m.text) for m in rows] == [
            ("t0r0", "o-r0c0 | o-r0c1"),
            ("t0r0c0t0r0", "a-r0c0"),
            ("t0r0c1t0r0", "b-r0c0"),
            ("t0r0c1t0r1", "b-r1c0"),
            ("t0r1", "o-r1c0 | o-r1c1"),
            ("t0r1c1t0r0", "c-r0c0"),
        ]
        _assert_reconstructs(spec)

    def test_nested_text_is_not_double_counted(self, tmp_path: Path):
        # The parent row's text is its cells' own paragraphs only; the
        # nested table's text appears exactly once, in its own row.
        doc = Document()
        outer = doc.add_table(rows=1, cols=1)
        outer.cell(0, 0).text = "OUTER"
        _nest(outer.cell(0, 0), 1, 1, "INNER-")

        spec = _extract(doc, tmp_path)
        assert spec.content.count("INNER-r0c0") == 1
        assert _table_rows(spec)[0].text == "OUTER"

    def test_cell_holding_only_a_nested_table_still_surfaces_it(self, tmp_path: Path):
        # A layout cell with no text of its own emits no row, but the
        # table it holds is still walked (and its ids still hang off the
        # parent row path).
        doc = Document()
        outer = doc.add_table(rows=1, cols=1)
        _nest(outer.cell(0, 0), 1, 1, "in-")

        spec = _extract(doc, tmp_path)
        rows = _table_rows(spec)
        assert [(m.element_id, m.text) for m in rows] == [("t0r0c0t0r0", "in-r0c0")]
        assert spec.content == "in-r0c0"
        _assert_reconstructs(spec)

    def test_two_nested_tables_in_one_cell_get_distinct_indices(self, tmp_path: Path):
        doc = Document()
        outer = doc.add_table(rows=1, cols=1)
        _nest(outer.cell(0, 0), 1, 1, "first-")
        _nest(outer.cell(0, 0), 1, 1, "second-")

        spec = _extract(doc, tmp_path)
        assert [m.element_id for m in _table_rows(spec)] == [
            "t0r0c0t0r0",
            "t0r0c0t1r0",
        ]

    def test_merged_cell_inside_nested_table_deduped(self, tmp_path: Path):
        doc = Document()
        outer = doc.add_table(rows=1, cols=1)
        inner = _nest(outer.cell(0, 0), 2, 3, "in-")
        inner.cell(0, 0).merge(inner.cell(0, 2)).text = "INNER HEADING"

        spec = _extract(doc, tmp_path)
        assert [m.text for m in _table_rows(spec)] == [
            "INNER HEADING",
            "in-r1c0 | in-r1c1 | in-r1c2",
        ]
        assert spec.content.count("INNER HEADING") == 1

    def test_nested_under_vertically_merged_parent_uses_deduped_cell_index(
        self, tmp_path: Path
    ):
        # In a continuation row the merged origin cell is skipped, so the
        # remaining cells' indices in the id path are their positions among
        # the row's *distinct* cells — deterministic and collision-free.
        doc = Document()
        outer = doc.add_table(rows=2, cols=2)
        _fill(outer, "o-")
        outer.cell(0, 0).merge(outer.cell(1, 0)).text = "SPAN"
        _nest(outer.cell(1, 1), 1, 1, "in-")

        spec = _extract(doc, tmp_path)
        assert [(m.element_id, m.text) for m in _table_rows(spec)] == [
            ("t0r0", "SPAN | o-r0c1"),
            ("t0r1", "o-r1c1"),
            ("t0r1c0t0r0", "in-r0c0"),
        ]
        _assert_reconstructs(spec)

    def test_recursion_walks_to_the_depth_bound(self, tmp_path: Path):
        doc = Document()
        cell = doc.add_table(rows=1, cols=1).cell(0, 0)
        cell.text = "LEVEL1"
        for level in range(2, _NESTED_TABLE_MAX_DEPTH + 1):
            cell = cell.add_table(rows=1, cols=1).cell(0, 0)
            cell.text = f"LEVEL{level}"

        spec = _extract(doc, tmp_path)
        for level in range(1, _NESTED_TABLE_MAX_DEPTH + 1):
            assert f"LEVEL{level}" in spec.content
        assert spec.extraction_warnings == []
        # Deepest id is a 4-level path.
        assert _table_rows(spec)[-1].element_id == "t0r0" + "c0t0r0" * (
            _NESTED_TABLE_MAX_DEPTH - 1
        )
        _assert_reconstructs(spec)

    def test_beyond_depth_bound_stops_and_warns(self, tmp_path: Path):
        doc = Document()
        cell = doc.add_table(rows=1, cols=1).cell(0, 0)
        cell.text = "LEVEL1"
        for level in range(2, _NESTED_TABLE_MAX_DEPTH + 2):
            cell = cell.add_table(rows=1, cols=1).cell(0, 0)
            cell.text = f"LEVEL{level}"

        spec = _extract(doc, tmp_path)
        assert f"LEVEL{_NESTED_TABLE_MAX_DEPTH}" in spec.content
        assert f"LEVEL{_NESTED_TABLE_MAX_DEPTH + 1}" not in spec.content
        # Honest, single warning — not a silent drop.
        assert spec.extraction_warnings == [_NESTED_TABLE_DEPTH_WARNING]
        assert "Verify visually" in _NESTED_TABLE_DEPTH_WARNING
        _assert_reconstructs(spec)

    def test_all_element_ids_unique_with_merges_and_nesting(self, tmp_path: Path):
        doc = Document()
        doc.add_paragraph("1.01 SUMMARY")
        outer = doc.add_table(rows=3, cols=3)
        _fill(outer, "o-")
        outer.cell(0, 0).merge(outer.cell(0, 2)).text = "TITLE"
        outer.cell(1, 0).merge(outer.cell(2, 0)).text = "GROUP"
        _nest(outer.cell(1, 1), 2, 2, "n1-")
        _nest(outer.cell(2, 2), 1, 1, "n2-")
        doc.add_paragraph("Closing paragraph.")
        second = doc.add_table(rows=1, cols=1)
        _nest(second.cell(0, 0), 1, 1, "t1-")

        spec = _extract(doc, tmp_path)
        ids = [m.element_id for m in spec.paragraph_map]
        assert all(ids)
        assert len(ids) == len(set(ids))
        # Body paragraph ids still follow the body index, unaffected by
        # how many nested rows the tables emitted.
        assert ids[0] == "p0" and "p2" in ids
        assert "t1r0c0t0r0" in ids
        _assert_reconstructs(spec)

    def test_nested_rows_render_as_row_tags_with_section(self, tmp_path: Path):
        # Downstream: a nested row is still a table row to the prompt
        # renderer (``<row id=…>``), and carries the enclosing section.
        doc = Document()
        doc.add_paragraph("2.01 EQUIPMENT")
        outer = doc.add_table(rows=1, cols=1)
        _nest(outer.cell(0, 0), 1, 1, "in-")

        spec = _extract(doc, tmp_path)
        nested = _table_rows(spec)[0]
        assert nested.section_id == "2.01 EQUIPMENT"
        rendered = render_spec_with_ids(spec.content, spec.paragraph_map)
        assert '<row id="t0r0c0t0r0" section="2.01 EQUIPMENT">' in rendered

    def test_word_count_includes_nested_text(self, tmp_path: Path):
        doc = Document()
        outer = doc.add_table(rows=1, cols=1)
        outer.cell(0, 0).text = "one two"
        _nest(outer.cell(0, 0), 1, 1, "three-")

        spec = _extract(doc, tmp_path)
        assert spec.word_count == len(spec.content.split()) == 3


# ---------------------------------------------------------------------------
# Content-loss warning is unaffected
# ---------------------------------------------------------------------------


class TestContentLossWarningUnaffected:
    def test_table_heavy_spec_has_no_content_loss_warning(self, tmp_path: Path):
        doc = Document()
        outer = doc.add_table(rows=2, cols=3)
        _fill(outer)
        outer.cell(0, 0).merge(outer.cell(0, 2)).text = "TITLE"
        _nest(outer.cell(1, 1), 2, 2, "in-")

        spec = _extract(doc, tmp_path)
        assert spec.extraction_warnings == []

    def test_drawing_heavy_spec_still_warns_with_tables_present(self, tmp_path: Path):
        # 2 drawing paragraphs out of 4 body children (1 text paragraph,
        # 2 drawing paragraphs, 1 table) = 50% > the 20% threshold. The
        # table (with its nested table) counts as one clean body child, as
        # before — the heuristic is body-level and knows nothing about the
        # table walk.
        doc = Document()
        doc.add_paragraph("Body text.")
        _drawing_paragraph(doc)
        _drawing_paragraph(doc)
        outer = doc.add_table(rows=1, cols=1)
        _nest(outer.cell(0, 0), 1, 1, "in-")

        spec = _extract(doc, tmp_path)
        assert len(spec.extraction_warnings) == 1
        assert spec.extraction_warnings[0].startswith(
            "Spec contains 50% non-text elements (2 drawings, 0 pictures, 0 OLE objects)."
        )
