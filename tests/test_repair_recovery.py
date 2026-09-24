"""Paid review repair batches stay recoverable (plan WP-14).

A repair batch is a paid re-run of the specs whose first review failed in a
retryable way. While one is still running — or out of reach — the saved run
record is the only handle to it, and every stage downstream of the review
(verification, cross-check, compliance, drawing impact) would pay again once
it lands. Before this, the GUI deleted the record after any collection that
returned, the recovery CLI deleted it whenever one spec had succeeded, and
both ran every dependent stage on the primary results alone.

Locked in here (the GUI entry points are in ``test_repair_recovery_gui.py``):

* **The outcome contract** — ``CollectionOutcome`` separates "there is
  something to report" from "the remote work is finished", and
  ``RepairOutcome`` names six states.
* **One cleanup rule** — ``decide_saved_state_cleanup``: kept while a repair
  is pending or unreachable, while a module could not be collected, when no
  outcome was recorded, or when every spec failed; cleared otherwise.
* **Identity** — a record is deleted only when it is the run's own and names
  no repair the run did not settle; a stale completion cannot delete a newer
  run's record.
* **Atomic stamps** — concurrent module collections cannot overwrite each
  other's repair stamps on one program manifest.
* **Deferral** — a provisional collection runs no dependent paid stage.
* **Plan §28 scenario C**, through ``scripts/recover_batch.py``: a pending
  repair survives a restart, is collected (never resubmitted), and the
  downstream stages run exactly once — counted at the paid boundaries.
* **Program children** keep their unresolved state independently.
* **Reports** say the result is provisional and which stages are waiting.
* **Kept records name every repair** — a repair id whose first save failed
  is re-stamped by the shared cleanup step, for single-module runs and each
  program child, on every entry point (found in review).
"""
from __future__ import annotations

import importlib.util
import json
import threading
from pathlib import Path

import pytest

from src.batch.batch_runtime import PollOutcome
from src.orchestration import batch_resume as br
from src.orchestration import pipeline as pl
from src.orchestration import program_pipeline as pp
from src.orchestration.batch_resume import (
    CLEAR_ABSENT,
    CLEAR_CLEARED,
    CLEAR_FOREIGN,
    CLEAR_UNRECOGNIZED,
    PendingBatch,
    PendingProgramRun,
    adopt_outstanding_run,
    apply_saved_state_cleanup,
    clear_saved_state_for,
    discard_saved_state,
    load_pending_batch,
    load_pending_run,
    record_identity,
    restamp_repair_batches,
    save_pending_batch,
    save_pending_program_run,
)
from src.orchestration.collection_outcome import (
    REPAIR_CONSUMED,
    REPAIR_NOT_NEEDED,
    REPAIR_NOT_SUBMITTED,
    REPAIR_PENDING,
    REPAIR_STATES,
    REPAIR_UNREACHABLE,
    REPAIR_UNUSABLE,
    STAGE_COMPLIANCE,
    STAGE_CROSS_CHECK,
    STAGE_DRAWING_IMPACT,
    STAGE_VERIFICATION,
    CollectionOutcome,
    RepairOutcome,
    decide_saved_state_cleanup,
    provisional_notice,
    record_owned_by,
    saved_state_identity,
)
from src.orchestration.pipeline import PipelineResult
from src.output.edit_sidecar import build_edit_instructions
from src.output.html_report_exporter import render_html_report
from src.output.report_exporter import (
    _aggregate_run_diagnostics,
    _summarize_run_diagnostics,
    export_report,
)
from src.programs import (
    HYPERSCALE_DATACENTER_PROGRAM,
    RoutingState,
    SpecAssignment,
    SpecRoutingDecision,
)
from src.review.reviewer import ReviewResult
from tests.fixtures.batch_service import (
    PARITY_NAMES,
    PARITY_SCENARIOS,
    PRIMARY_ID,
    FakeBatchService,
    parity_primary,
    request_id,
    review_ok,
    review_refused,
    review_truncated,
    submission,
    write_docx_specs,
)

NAMES = ["A.docx", "B.docx"]


class _Log:
    def __init__(self):
        self.lines: list[tuple[str, str]] = []

    def __call__(self, msg: str, *, level: str = "info", **_kw):
        self.lines.append((str(msg), level))

    def text(self, level: str | None = None) -> str:
        return "\n".join(m for m, lvl in self.lines if level is None or lvl == level)


@pytest.fixture
def state_path(tmp_path, monkeypatch):
    path = tmp_path / "state" / "pending_batch.json"
    monkeypatch.setenv("SPEC_CRITIC_PENDING_BATCH_PATH", str(path))
    monkeypatch.setenv("SPEC_CRITIC_VERIFICATION_CACHE_PERSIST", "0")
    return path


def _outcome(
    state: str = REPAIR_NOT_NEEDED,
    *,
    batch_id: str = PRIMARY_ID,
    repair_id: str | None = None,
    submitted=NAMES,
    failed=(),
    transport: str = "batch",
    module_id: str = "california_k12_mep",
    replaced: str | None = None,
) -> CollectionOutcome:
    return CollectionOutcome(
        batch_id=batch_id,
        module_id=module_id,
        transport=transport,
        repair=RepairOutcome(
            state=state,
            batch_id=repair_id,
            specs=tuple(failed),
            replaced_batch_id=replaced,
        ),
        submitted_specs=tuple(submitted),
        failed_specs=tuple(failed),
    )


def _result(outcome: CollectionOutcome | None, **kwargs) -> PipelineResult:
    return PipelineResult(
        review_result=ReviewResult(findings=[]),
        files_reviewed=list(NAMES),
        collection_outcome=outcome,
        **kwargs,
    )


def _primary(*, failed_index: int | None = 1, refuse: bool = False) -> dict:
    results = {request_id(i): review_ok(name) for i, name in enumerate(NAMES)}
    if failed_index is not None:
        results[request_id(failed_index)] = (
            review_refused() if refuse else review_truncated()
        )
    return {PRIMARY_ID: results}


# ===========================================================================
# 1. The outcome contract
# ===========================================================================


class TestOutcomeContract:
    @pytest.mark.parametrize(
        "state,outstanding",
        [
            (REPAIR_NOT_NEEDED, False),
            (REPAIR_PENDING, True),
            (REPAIR_UNREACHABLE, True),
            (REPAIR_CONSUMED, False),
            (REPAIR_UNUSABLE, False),
            (REPAIR_NOT_SUBMITTED, False),
        ],
    )
    def test_only_pending_and_unreachable_are_outstanding(self, state, outstanding):
        assert RepairOutcome(state=state).outstanding is outstanding
        outcome = _outcome(state, failed=["B.docx"])
        assert outcome.provisional is outstanding
        assert outcome.remote_settled is not outstanding

    def test_the_six_states_are_the_closed_set(self):
        assert set(REPAIR_STATES) == {
            "not_needed", "pending", "unreachable", "consumed", "unusable",
            "not_submitted",
        }
        with pytest.raises(ValueError):
            RepairOutcome(state="finished")

    def test_reportable_is_separate_from_settled(self):
        # A usable primary result with the repair still running: reportable,
        # not settled — the distinction the plan asks for.
        outcome = _outcome(REPAIR_PENDING, repair_id="R", failed=["B.docx"])
        assert outcome.reportable and not outcome.remote_settled
        # Every spec failed and the repair is pending: nothing to report yet.
        pending_all = _outcome(REPAIR_PENDING, repair_id="R", failed=NAMES)
        assert not pending_all.reportable and pending_all.provisional

    def test_a_valid_review_with_zero_findings_is_reportable(self):
        outcome = _outcome(REPAIR_NOT_NEEDED)
        assert outcome.reportable and not outcome.all_failed

    def test_all_failed(self):
        assert _outcome(REPAIR_CONSUMED, repair_id="R", failed=NAMES).all_failed
        assert not _outcome(submitted=(), failed=()).all_failed

    def test_awaiting_repair_specs_only_while_outstanding(self):
        assert _outcome(REPAIR_PENDING, repair_id="R", failed=["B.docx"]).awaiting_repair_specs == ("B.docx",)
        assert _outcome(REPAIR_CONSUMED, repair_id="R", failed=["B.docx"]).awaiting_repair_specs == ()

    def test_settled_ids_include_a_replaced_repair(self):
        repair = RepairOutcome(state=REPAIR_CONSUMED, batch_id="FRESH", replaced_batch_id="OLD")
        assert repair.settled_batch_ids == {"FRESH", "OLD"}

    def test_to_dict_says_what_is_outstanding(self):
        data = _outcome(REPAIR_PENDING, repair_id="R", failed=["B.docx"]).to_dict()
        assert data["provisional"] is True
        assert data["repair"]["state"] == "pending"
        assert data["repair"]["batch_id"] == "R"
        assert data["awaiting_repair_specs"] == ["B.docx"]
        json.dumps(data)  # JSON-ready


# ===========================================================================
# 2. The one cleanup decision
# ===========================================================================


class TestCleanupDecision:
    @pytest.mark.parametrize("state", [REPAIR_PENDING, REPAIR_UNREACHABLE])
    def test_an_outstanding_repair_keeps_the_record(self, state):
        decision = decide_saved_state_cleanup(
            _result(_outcome(state, repair_id="msgbatch_R", failed=["B.docx"]))
        )
        assert not decision.clear and not decision.complete
        assert "msgbatch_R" in decision.reason

    @pytest.mark.parametrize(
        "state", [REPAIR_NOT_NEEDED, REPAIR_CONSUMED, REPAIR_UNUSABLE, REPAIR_NOT_SUBMITTED]
    )
    def test_a_settled_partial_run_may_clear(self, state):
        # One spec reviewed, one still failed after a settled repair: the
        # partial result is final, nothing is retrievable any more.
        failed = [] if state == REPAIR_NOT_NEEDED else ["B.docx"]
        decision = decide_saved_state_cleanup(
            _result(_outcome(state, repair_id="R" if failed else None, failed=failed))
        )
        assert decision.clear and decision.complete

    def test_a_valid_zero_findings_run_may_clear(self):
        decision = decide_saved_state_cleanup(_result(_outcome()))
        assert decision.clear and decision.complete

    def test_an_all_failed_run_stays_recoverable(self):
        decision = decide_saved_state_cleanup(
            _result(_outcome(REPAIR_CONSUMED, repair_id="R", failed=NAMES))
        )
        assert not decision.clear and not decision.complete
        assert "every submitted specification failed review" in decision.reason

    def test_a_missing_outcome_never_authorizes_cleanup(self):
        decision = decide_saved_state_cleanup(_result(None))
        assert not decision.clear and not decision.complete
        assert "no outcome" in decision.reason

    def test_keep_requested(self):
        decision = decide_saved_state_cleanup(_result(_outcome()), keep_requested=True)
        assert decision.complete and not decision.clear

    def test_a_real_time_run_never_touches_saved_state(self):
        decision = decide_saved_state_cleanup(_result(_outcome(transport="realtime")))
        assert not decision.clear and not decision.saved_state_applies
        assert decision.complete

    def test_program_module_errors_keep_the_record(self):
        result = _program_result(
            {"datacenter_fire": _outcome(module_id="datacenter_fire")},
            module_errors={"datacenter_architecture": "results endpoint timed out"},
        )
        decision = decide_saved_state_cleanup(result)
        assert not decision.clear and not decision.complete
        assert "could not be collected" in decision.reason

    def test_program_one_child_outstanding_keeps_the_whole_manifest(self):
        result = _program_result(
            {
                "datacenter_fire": _outcome(
                    REPAIR_PENDING, batch_id="F", repair_id="FR", failed=["B.docx"],
                    module_id="datacenter_fire",
                ),
                "datacenter_architecture": _outcome(
                    batch_id="A", module_id="datacenter_architecture"
                ),
            }
        )
        decision = decide_saved_state_cleanup(result)
        assert not decision.clear
        assert "datacenter_fire" in decision.reason and "FR" in decision.reason

    def test_a_provisional_program_is_partial_on_its_own(self):
        # Even with no failed-review spec recorded (a hand-built or legacy
        # result), an outstanding repair makes the program partial: its
        # dependent stages never ran.
        result = _program_result(
            {
                "datacenter_fire": _outcome(
                    REPAIR_PENDING, batch_id="F", repair_id="FR",
                    module_id="datacenter_fire",
                ),
                "datacenter_architecture": _outcome(
                    batch_id="A", module_id="datacenter_architecture"
                ),
            }
        )
        assert result.failed_review_specs == []
        assert result.provisional and result.status == "partial"

    def test_program_all_settled_clears(self):
        result = _program_result(
            {
                "datacenter_fire": _outcome(batch_id="F", module_id="datacenter_fire"),
                "datacenter_architecture": _outcome(
                    batch_id="A", module_id="datacenter_architecture"
                ),
            }
        )
        assert decide_saved_state_cleanup(result).clear

    def test_program_all_failed_across_modules_is_kept_but_one_module_failing_is_not(self):
        both_failed = _program_result(
            {
                "datacenter_fire": _outcome(
                    REPAIR_CONSUMED, batch_id="F", repair_id="FR", failed=NAMES,
                    module_id="datacenter_fire",
                ),
                "datacenter_architecture": _outcome(
                    REPAIR_NOT_SUBMITTED, batch_id="A", failed=NAMES,
                    module_id="datacenter_architecture",
                ),
            }
        )
        assert not decide_saved_state_cleanup(both_failed).clear
        one_failed = _program_result(
            {
                "datacenter_fire": _outcome(
                    REPAIR_CONSUMED, batch_id="F", repair_id="FR", failed=NAMES,
                    module_id="datacenter_fire",
                ),
                "datacenter_architecture": _outcome(
                    batch_id="A", module_id="datacenter_architecture"
                ),
            }
        )
        assert decide_saved_state_cleanup(one_failed).clear


def _program_assignments(
    names=("21 13 13 Wet.docx", "07 27 26 Air.docx"), *, spec_dir: Path | None = None
):
    modules = ("datacenter_fire", "datacenter_architecture")
    return tuple(
        SpecAssignment(
            source_path=str((spec_dir or Path("C:/specs")) / name),
            decision=SpecRoutingDecision(
                spec_id=name,
                program_id=HYPERSCALE_DATACENTER_PROGRAM.program_id,
                automatic_state=RoutingState.SUPPORTED,
                automatic_module_ids=(module_id,),
                confidence=0.95,
                evidence=(),
            ),
        )
        for name, module_id in zip(names, modules)
    )


def _program_result(outcomes: dict, *, module_errors=None, deferred_program_stages=()):
    results = {
        module_id: PipelineResult(
            review_result=ReviewResult(findings=[]),
            files_reviewed=[],
            module_id=module_id,
            cycle_label=pl.get_module(module_id).cycle.label,
            collection_outcome=outcome,
        )
        for module_id, outcome in outcomes.items()
    }
    return pp.ProgramPipelineResult(
        program_id=HYPERSCALE_DATACENTER_PROGRAM.program_id,
        assignments=_program_assignments(),
        module_results=results,
        module_errors=dict(module_errors or {}),
        deferred_program_stages=tuple(deferred_program_stages),
    )


# ===========================================================================
# 3. Identity: a record is cleared only when it is this run's own
# ===========================================================================


def _save_single(batch_id=PRIMARY_ID, repair_id=None, path=None):
    sub = submission(NAMES, batch_id=batch_id)
    pending = PendingBatch.from_submission(sub)
    pending.repair_batch_id = repair_id
    save_pending_batch(pending, path=path)
    return pending


class TestIdentityCheckedClear:
    def test_own_record_is_cleared(self, state_path):
        _save_single()
        status = clear_saved_state_for({PRIMARY_ID: frozenset()})
        assert status == CLEAR_CLEARED and not state_path.exists()

    def test_another_runs_record_is_kept(self, state_path):
        # The stale-completion case: this collection finished after another
        # run saved its record over the slot.
        _save_single(batch_id="msgbatch_NEWER_RUN")
        status = clear_saved_state_for({PRIMARY_ID: frozenset()})
        assert status == CLEAR_FOREIGN
        assert load_pending_batch().batch_id == "msgbatch_NEWER_RUN"

    def test_a_record_naming_a_repair_this_run_never_settled_is_kept(self, state_path):
        _save_single(repair_id="msgbatch_UNKNOWN_REPAIR")
        status = clear_saved_state_for({PRIMARY_ID: frozenset({"msgbatch_OTHER"})})
        assert status == CLEAR_FOREIGN and state_path.exists()

    def test_a_record_naming_the_replaced_repair_is_this_runs(self, state_path):
        # The saved repair ended unusable and was replaced, but the re-stamp
        # never reached the file: the record still names the old repair.
        _save_single(repair_id="msgbatch_OLD")
        identity = saved_state_identity(
            _result(
                _outcome(
                    REPAIR_CONSUMED, repair_id="msgbatch_FRESH", failed=["B.docx"],
                    replaced="msgbatch_OLD",
                )
            )
        )
        assert clear_saved_state_for(identity) == CLEAR_CLEARED

    def test_absent_and_unreadable_records(self, state_path):
        assert clear_saved_state_for({PRIMARY_ID: frozenset()}) == CLEAR_ABSENT
        state_path.parent.mkdir(parents=True, exist_ok=True)
        state_path.write_text("{not json", encoding="utf-8")
        assert clear_saved_state_for({PRIMARY_ID: frozenset()}) == CLEAR_UNRECOGNIZED
        assert state_path.exists()

    def test_a_newer_schema_record_is_never_deleted(self, state_path):
        state_path.parent.mkdir(parents=True, exist_ok=True)
        state_path.write_text(
            json.dumps({"schema_version": 99, "batch_id": PRIMARY_ID}), encoding="utf-8"
        )
        assert clear_saved_state_for({PRIMARY_ID: frozenset()}) == CLEAR_UNRECOGNIZED
        assert state_path.exists()

    def test_a_legacy_record_without_repair_fields_is_cleared_when_owned(self, state_path):
        state_path.parent.mkdir(parents=True, exist_ok=True)
        state_path.write_text(
            json.dumps({"schema_version": 1, "batch_id": PRIMARY_ID, "model": "m"}),
            encoding="utf-8",
        )
        assert clear_saved_state_for({PRIMARY_ID: frozenset()}) == CLEAR_CLEARED

    def test_program_manifest_needs_every_child(self, state_path):
        _save_manifest()
        # Recovering one child of the manifest never clears it.
        assert (
            clear_saved_state_for({"msgbatch_datacenter_fire": frozenset()})
            == CLEAR_FOREIGN
        )
        assert (
            clear_saved_state_for(
                {
                    "msgbatch_datacenter_fire": frozenset(),
                    "msgbatch_datacenter_architecture": frozenset(),
                }
            )
            == CLEAR_CLEARED
        )

    def test_a_manifest_child_naming_an_unsettled_repair_is_kept(self, state_path):
        # Another process stamped a repair on the fire child that this
        # collection never consumed: that billed work keeps the manifest.
        _save_manifest(fire_repair="msgbatch_UNSETTLED")
        status = clear_saved_state_for(
            {
                "msgbatch_datacenter_fire": frozenset(),
                "msgbatch_datacenter_architecture": frozenset(),
            }
        )
        assert status == CLEAR_FOREIGN and state_path.exists()
        assert (
            clear_saved_state_for(
                {
                    "msgbatch_datacenter_fire": frozenset({"msgbatch_UNSETTLED"}),
                    "msgbatch_datacenter_architecture": frozenset(),
                }
            )
            == CLEAR_CLEARED
        )

    def test_a_manifest_partition_without_a_batch_id_is_unidentifiable(self):
        manifest = {
            "record_type": "program",
            "partitions": {"datacenter_fire": {"module_id": "datacenter_fire"}},
        }
        assert record_identity(manifest) == {}
        assert not record_owned_by({}, {"x": frozenset()})

    def test_explicit_discard_deletes_only_the_record_it_was_shown(self, state_path):
        shown = _save_single()
        assert discard_saved_state(shown) == CLEAR_CLEARED
        # A record another run saved after the prompt loaded is not discarded.
        shown = _save_single()
        _save_single(batch_id="msgbatch_SAVED_LATER")
        assert discard_saved_state(shown) == CLEAR_FOREIGN
        assert load_pending_batch().batch_id == "msgbatch_SAVED_LATER"

    def test_apply_clears_only_complete_runs_and_logs_why(self, state_path):
        _save_single(repair_id="msgbatch_R")
        log = _Log()
        decision, status = apply_saved_state_cleanup(
            _result(_outcome(REPAIR_PENDING, repair_id="msgbatch_R", failed=["B.docx"])),
            log=log,
        )
        assert not decision.clear and status is None and state_path.exists()
        assert "Saved batch state kept" in log.text("warning")
        decision, status = apply_saved_state_cleanup(
            _result(_outcome(REPAIR_CONSUMED, repair_id="msgbatch_R", failed=())),
            log=log,
        )
        assert decision.clear and status == CLEAR_CLEARED and not state_path.exists()

    def test_apply_never_raises(self, state_path, monkeypatch):
        _save_single()

        def boom(*_a, **_k):
            raise RuntimeError("disk gone")

        monkeypatch.setattr(br, "clear_saved_state_for", boom)
        log = _Log()
        decision, status = apply_saved_state_cleanup(_result(_outcome()), log=log)
        assert status == CLEAR_UNRECOGNIZED
        assert "cleanup check failed" in log.text("warning")

    def test_real_time_results_leave_another_runs_record_alone(self, state_path):
        _save_single(batch_id="msgbatch_EARLIER_DETACHED")
        decision, status = apply_saved_state_cleanup(
            _result(_outcome(transport="realtime", batch_id="realtime"))
        )
        assert status is None and state_path.exists()


def _save_manifest(path=None, *, fire_repair: str | None = None):
    partitions = {
        "datacenter_fire": submission(
            ["21 13 13 Wet.docx"], batch_id="msgbatch_datacenter_fire",
            module_id="datacenter_fire",
        ),
        "datacenter_architecture": submission(
            ["07 27 26 Air.docx"], batch_id="msgbatch_datacenter_architecture",
            module_id="datacenter_architecture",
        ),
    }
    if fire_repair:
        partitions["datacenter_fire"].repair_batch_id = fire_repair
    program_submission = pp.ProgramSubmission(
        program_id=HYPERSCALE_DATACENTER_PROGRAM.program_id,
        assignments=_program_assignments(),
        partitions=partitions,
    )
    save_pending_program_run(
        PendingProgramRun.from_submission(program_submission), path=path
    )
    return program_submission


# ===========================================================================
# 4. Repair stamps on one manifest are atomic
# ===========================================================================


class TestConcurrentRepairStamps:
    def test_two_modules_stamping_at_once_both_survive(self, state_path, monkeypatch):
        """The first stamp pauses inside its save until the second has read
        the manifest (or half a second passes). Unguarded, the second reads
        the stale manifest and its save wins, losing the first stamp — the
        only handle to a billed repair."""
        _save_manifest()
        second_read = threading.Event()
        first_in_save = threading.Event()
        real_load = br.load_pending_run
        real_save = br.save_pending_program_run

        def load(**kw):
            if threading.current_thread().name == "second":
                second_read.set()
            return real_load(**kw)

        def save(pending, **kw):
            if threading.current_thread().name == "first":
                first_in_save.set()
                second_read.wait(timeout=0.5)
            return real_save(pending, **kw)

        monkeypatch.setattr(br, "load_pending_run", load)
        monkeypatch.setattr(br, "save_pending_program_run", save)

        def stamp(parent: str, repair: str):
            from src.batch.batch import BatchJob

            child = submission(["X.docx"], batch_id=parent)
            job = BatchJob(
                batch_id=repair,
                job_type="review",
                request_map={"r": {"filename": "X.docx", "index": 0}},
                created_at=0.0,
            )
            pl._persist_repair_batch(child, job)

        first = threading.Thread(
            target=stamp, name="first",
            args=("msgbatch_datacenter_fire", "msgbatch_FIRE_REPAIR"),
        )
        second = threading.Thread(
            target=stamp, name="second",
            args=("msgbatch_datacenter_architecture", "msgbatch_ARCH_REPAIR"),
        )
        first.start()
        assert first_in_save.wait(timeout=5)
        second.start()
        first.join(timeout=5)
        second.join(timeout=5)

        saved = json.loads(state_path.read_text(encoding="utf-8"))
        stamps = {
            module_id: child.get("repair_batch_id")
            for module_id, child in saved["partitions"].items()
        }
        assert stamps == {
            "datacenter_fire": "msgbatch_FIRE_REPAIR",
            "datacenter_architecture": "msgbatch_ARCH_REPAIR",
        }


# ===========================================================================
# 5. A saved repair is consumed without the source files
# ===========================================================================


class TestSavedRepairWithoutSourceFiles:
    def test_consumed_when_the_files_moved(self, monkeypatch):
        service = FakeBatchService(monkeypatch, primary=_primary())
        sub = submission(NAMES, prepared=False)
        sub.repair_batch_id = "msgbatch_SAVED"
        sub.repair_request_map = {
            "review__B__0": {"filename": "B.docx", "index": 0, "type": "review"}
        }
        results, repair = pl._recover_retryable_review_batch_results(
            sub, _primary()[PRIMARY_ID], log=_Log()
        )
        assert service.polls == ["msgbatch_SAVED"]
        assert repair.state == REPAIR_CONSUMED and repair.reattached
        assert results[request_id(1)].parse_status == "ok"
        assert service.repair_submits == []

    def test_still_running_without_files_is_pending_not_unrepairable(self, monkeypatch):
        service = FakeBatchService(monkeypatch, primary=_primary())
        service.status["msgbatch_SAVED"] = "processing"
        sub = submission(NAMES, prepared=False)
        sub.repair_batch_id = "msgbatch_SAVED"
        _results, repair = pl._recover_retryable_review_batch_results(
            sub, _primary()[PRIMARY_ID], log=_Log()
        )
        # Before the fix this read "fallback skipped: specs unavailable" and
        # the saved, billed repair was never polled at all.
        assert repair.state == REPAIR_PENDING
        assert service.polls == ["msgbatch_SAVED"]

    def test_a_legacy_id_only_record_rebuilds_the_map_from_the_request_map(self, monkeypatch):
        from src.batch.batch import _review_custom_id

        service = FakeBatchService(monkeypatch, primary=_primary())
        seen: dict = {}
        real_retrieve = service._retrieve

        def retrieve(job, *, model):
            seen[job.batch_id] = dict(job.request_map)
            return real_retrieve(job, model=model)

        monkeypatch.setattr(pl, "retrieve_review_results", retrieve)
        sub = submission(NAMES, prepared=False)
        sub.repair_batch_id = "msgbatch_SAVED"
        results, repair = pl._recover_retryable_review_batch_results(
            sub, _primary()[PRIMARY_ID], log=_Log()
        )
        assert seen["msgbatch_SAVED"] == {
            _review_custom_id("B.docx", 0): {"filename": "B.docx", "index": 0, "type": "review"}
        }
        assert repair.state == REPAIR_CONSUMED

    def test_no_saved_repair_and_no_files_is_not_submitted(self, monkeypatch):
        service = FakeBatchService(monkeypatch, primary=_primary())
        sub = submission(NAMES, prepared=False)
        _results, repair = pl._recover_retryable_review_batch_results(
            sub, _primary()[PRIMARY_ID], log=_Log()
        )
        assert repair.state == REPAIR_NOT_SUBMITTED and not repair.outstanding
        assert repair.specs == ("B.docx",)
        assert service.repair_submits == []

    def test_an_unusable_saved_repair_without_files_settles_as_unusable(self, monkeypatch):
        service = FakeBatchService(monkeypatch, primary=_primary())
        service.status["msgbatch_SAVED"] = "expired"
        sub = submission(NAMES, prepared=False)
        sub.repair_batch_id = "msgbatch_SAVED"
        _results, repair = pl._recover_retryable_review_batch_results(
            sub, _primary()[PRIMARY_ID], log=_Log()
        )
        assert repair.state == REPAIR_UNUSABLE
        assert repair.batch_id == "msgbatch_SAVED"
        assert service.repair_submits == []


# ===========================================================================
# 6. A provisional collection runs no dependent paid stage
# ===========================================================================


_DIGEST_CONTEXT = (
    "--- BEGIN ATTACHMENT: Construction Drawing Digest ---\n"
    "A101 shows the riser room. [set.pdf p.1]\n"
    "--- END ATTACHMENT: Construction Drawing Digest ---"
)


class TestHeadlessDeferral:
    @pytest.mark.parametrize("repair_status", ["processing", "unreachable"])
    def test_no_dependent_stage_runs_while_the_repair_is_outstanding(
        self, monkeypatch, state_path, repair_status
    ):
        service = FakeBatchService(monkeypatch, primary=_primary())
        service.default_repair_status = repair_status
        sub = submission(NAMES, project_context=_DIGEST_CONTEXT)
        result = pl.run_batch_collection_headless(sub, log=_Log())

        assert result.provisional
        assert service.repair_submits == [["B.docx"]]
        assert service.paid_downstream == 0
        outcome = result.collection_outcome
        assert outcome.repair.state == (
            REPAIR_PENDING if repair_status == "processing" else REPAIR_UNREACHABLE
        )
        assert outcome.deferred_stages == (
            STAGE_VERIFICATION, STAGE_CROSS_CHECK, STAGE_DRAWING_IMPACT,
        )
        # The deferred cross-check says so instead of reading as disabled.
        assert result.cross_check_result.cross_check_status == "skipped"
        assert "deferred until the review repair finishes" in result.cross_check_result.thinking
        assert result.drawing_impact_result is None
        # The primary findings are shown, unverified.
        assert [f.fileName for f in result.review_result.findings] == ["A.docx"]
        assert all(f.verification is None for f in result.review_result.findings)

    def test_a_settled_collection_runs_every_stage_once(self, monkeypatch, state_path):
        service = FakeBatchService(monkeypatch, primary=_primary())
        result = pl.run_batch_collection_headless(
            submission(NAMES, project_context=_DIGEST_CONTEXT), log=_Log()
        )
        assert not result.provisional
        assert result.collection_outcome.repair.state == REPAIR_CONSUMED
        assert service.verification_rounds == [2]
        assert service.cross_checks == 1 and service.drawing_impacts == 1
        assert result.collection_outcome.deferred_stages == ()

    def test_compliance_is_deferred_with_its_coverage_unassessed(self, monkeypatch, state_path):
        service = FakeBatchService(monkeypatch, primary=_primary())
        service.default_repair_status = "processing"
        sub = submission(NAMES, module_id="datacenter_fire")
        sub.requirements_profile = _requirements_profile()
        result = pl.run_batch_collection_headless(sub, log=_Log())
        assert STAGE_COMPLIANCE in result.collection_outcome.deferred_stages
        assert service.compliance_passes == 0
        compliance = result.compliance_result
        assert compliance.cross_check_status == "skipped"
        assert "deferred until the review repair finishes" in compliance.thinking
        completeness = compliance.coverage_completeness
        assert not completeness.complete
        assert completeness.omitted_ids == ("r-aaaaaaaaaaaa",)
        assert completeness.unassessed_specs == tuple(NAMES)

    def test_a_real_time_collection_is_never_provisional(self, monkeypatch):
        sub = submission(NAMES, transport="realtime")
        sub.realtime_results = {request_id(0): review_ok("A.docx"), request_id(1): review_truncated()}
        state = pl.collect_review_batch_results(sub)
        assert state.collection_outcome.transport == "realtime"
        assert state.collection_outcome.repair.state == REPAIR_NOT_NEEDED
        assert not state.collection_outcome.provisional

    def test_the_review_error_tells_the_operator_not_to_rerun_a_pending_spec(self, monkeypatch):
        service = FakeBatchService(monkeypatch, primary=_primary())
        service.default_repair_status = "processing"
        state = pl.collect_review_batch_results(submission(NAMES))
        assert "msgbatch_REPAIR_1" in state.review_result.error
        assert "rather than re-running the spec" in state.review_result.error

    def test_the_review_event_carries_the_outcome(self, monkeypatch):
        from src.orchestration.diagnostics import DiagnosticsReport

        service = FakeBatchService(monkeypatch, primary=_primary())
        service.default_repair_status = "processing"
        report = DiagnosticsReport()
        pl.collect_review_state_headless(submission(NAMES), diagnostics=report)
        events = [e for e in report.events if e.phase == "batch_collect"]
        assert events[0].data["collection"]["repair"]["state"] == REPAIR_PENDING


def _requirements_profile() -> dict:
    from src.research import DimensionStatus, RequirementsProfile, ResearchItem

    return RequirementsProfile(
        items=[
            ResearchItem(
                item_id="r-aaaaaaaaaaaa",
                dimension_id="governing_codes",
                topic="Topic",
                category="governing_code",
                requirement="Requirement A applies.",
                grounded=True,
                accepted_sources=["https://codes.example.gov/x"],
                confidence=0.8,
            )
        ],
        dimension_statuses=[
            DimensionStatus(dimension_id="governing_codes", status="completed")
        ],
        research_date="2026-09-24",
        project={"city": "Ashburn", "state_or_province": "VA", "country": "US",
                 "client_name": "ExampleCo"},
    ).to_dict()


# ===========================================================================
# 7. Plan §28 scenario C through scripts/recover_batch.py
# ===========================================================================


def _load_cli():
    script = Path(__file__).resolve().parent.parent / "scripts" / "recover_batch.py"
    spec = importlib.util.spec_from_file_location("recover_batch_s08_under_test", script)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)  # type: ignore[union-attr]
    return module


@pytest.fixture
def cli(monkeypatch):
    module = _load_cli()
    monkeypatch.setattr(module, "_ensure_api_key", lambda parser: None)
    monkeypatch.setattr(module, "export_report", lambda result, path: path)
    monkeypatch.setattr(module, "write_edit_instructions_sidecar", lambda result, path: path)
    monkeypatch.setattr(module, "write_requirements_profile_sidecar", lambda result, path: None)
    # The primary batches have ended; the repair's state is the fake
    # service's to decide.
    monkeypatch.setattr(
        module, "_poll_batches",
        lambda ids: {
            label: PollOutcome(terminal=True, terminal_status="ended") for label in ids
        },
    )
    return module


class TestScenarioCThroughTheRecoveryCli:
    """A truncated primary, a repair still running, a restart, then the
    repair lands. Counted at the paid boundaries: one repair submission ever,
    and each dependent stage exactly once, after the repair."""

    def _save_run(self, tmp_path, state_path):
        spec_dir = tmp_path / "specs"
        spec_dir.mkdir()
        files = write_docx_specs(spec_dir, NAMES)
        sub = submission(NAMES)
        save_pending_batch(
            PendingBatch.from_submission(sub, input_dir=spec_dir, files=files)
        )
        return files

    def test_pending_repair_survives_restart_and_is_collected_once(
        self, cli, tmp_path, state_path, monkeypatch
    ):
        self._save_run(tmp_path, state_path)
        service = FakeBatchService(monkeypatch, primary=_primary())
        service.default_repair_status = "processing"
        out = str(tmp_path / "report.docx")

        # Session 1: the repair is submitted and still running.
        assert cli.main(["-o", out]) == 2
        assert service.repair_submits == [["B.docx"]]
        assert service.paid_downstream == 0
        saved = load_pending_batch()
        assert saved is not None and saved.batch_id == PRIMARY_ID
        assert saved.repair_batch_id == "msgbatch_REPAIR_1"
        assert saved.repair_request_map  # enough to map results back

        # Session 2 (restart): still running — nothing new is paid for.
        assert cli.main(["-o", out]) == 2
        assert service.repair_submits == [["B.docx"]]
        assert service.paid_downstream == 0
        assert load_pending_batch().repair_batch_id == "msgbatch_REPAIR_1"

        # Session 3: the repair has ended. It is collected, not resubmitted,
        # and every dependent stage runs exactly once.
        service.status["msgbatch_REPAIR_1"] = "ended"
        assert cli.main(["-o", out]) == 0
        assert service.repair_submits == [["B.docx"]]
        assert service.polls.count("msgbatch_REPAIR_1") == 3
        assert service.verification_rounds == [2]
        assert service.cross_checks == 1
        # Only then is the record cleared.
        assert not state_path.exists()

    def test_a_repair_stamp_that_failed_to_save_is_retried_before_exit(
        self, cli, tmp_path, state_path, monkeypatch
    ):
        """A saved run whose repair stamp could not be written (after the
        writer's own retries) must not end with an unstamped record: the next
        run would rebuild the submission without the repair id and pay for a
        second repair. The CLI retries the stamp before it exits, as the GUI
        does."""
        self._save_run(tmp_path, state_path)
        service = FakeBatchService(monkeypatch, primary=_primary())
        service.default_repair_status = "processing"
        real_write = br._write_pending_state
        writes: list[str] = []

        def first_write_fails(payload, target, *, what):
            writes.append(what)
            if len(writes) == 1:
                return False  # the stamp's own save, after its retries
            return real_write(payload, target, what=what)

        monkeypatch.setattr(br, "_write_pending_state", first_write_fails)

        # Session 1: the repair is submitted, its first stamp is lost.
        assert cli.main(["-o", str(tmp_path / "r.docx")]) == 2
        assert service.repair_submits == [["B.docx"]]
        assert len(writes) == 2  # the failed stamp, then the retry
        saved = load_pending_batch()
        assert saved is not None and saved.repair_batch_id == "msgbatch_REPAIR_1"
        assert saved.repair_request_map

        # Session 2: the repair has ended. It is collected, not resubmitted.
        service.status["msgbatch_REPAIR_1"] = "ended"
        assert cli.main(["-o", str(tmp_path / "r2.docx")]) == 0
        assert service.repair_submits == [["B.docx"]]
        assert service.cross_checks == 1
        assert not state_path.exists()

    def test_a_record_that_vanished_mid_run_is_saved_again_with_its_inputs(
        self, cli, tmp_path, state_path, monkeypatch
    ):
        """If the saved record disappears while the run collects (deleted by
        hand, say), a provisional run saves it again from its own record's
        inputs, so a resume still re-reads the specs."""
        files = self._save_run(tmp_path, state_path)
        service = FakeBatchService(monkeypatch, primary=_primary())
        service.default_repair_status = "processing"
        poll = service._poll

        def delete_the_record_then_poll(batch_id, **kwargs):
            state_path.unlink(missing_ok=True)
            return poll(batch_id, **kwargs)

        monkeypatch.setattr(pl, "poll_batch_bounded", delete_the_record_then_poll)
        assert cli.main(["-o", str(tmp_path / "r.docx")]) == 2
        saved = load_pending_batch()
        assert saved is not None and saved.repair_batch_id == "msgbatch_REPAIR_1"
        assert saved.input_dir == str(tmp_path / "specs")
        assert saved.files == [str(f) for f in files]

    def test_a_temporary_retrieval_failure_keeps_state(
        self, cli, tmp_path, state_path, monkeypatch
    ):
        self._save_run(tmp_path, state_path)
        service = FakeBatchService(monkeypatch, primary=_primary())
        service.default_repair_status = "unreachable"
        assert cli.main(["-o", str(tmp_path / "r.docx")]) == 2
        assert state_path.exists()
        assert load_pending_batch().repair_batch_id == "msgbatch_REPAIR_1"
        assert service.paid_downstream == 0

    def test_keep_state_flag_keeps_a_complete_run(self, cli, tmp_path, state_path, monkeypatch):
        self._save_run(tmp_path, state_path)
        FakeBatchService(monkeypatch, primary=_primary())
        assert cli.main(["-o", str(tmp_path / "r.docx"), "--keep-state"]) == 0
        assert state_path.exists()

    def test_all_failed_run_is_kept_and_exits_2(self, cli, tmp_path, state_path, monkeypatch):
        self._save_run(tmp_path, state_path)
        both_refused = {
            PRIMARY_ID: {request_id(i): review_refused() for i in range(len(NAMES))}
        }
        FakeBatchService(monkeypatch, primary=both_refused)
        assert cli.main(["-o", str(tmp_path / "r.docx")]) == 2
        assert state_path.exists()

    def test_bare_id_recovery_saves_a_record_for_an_outstanding_repair(
        self, cli, tmp_path, state_path, monkeypatch
    ):
        """A batch recovered by id has no saved record. If its repair is still
        running, one is saved (the slot is free), so the next run resumes the
        repair instead of paying for a second."""
        spec_dir = tmp_path / "specs"
        spec_dir.mkdir()
        write_docx_specs(spec_dir, NAMES)
        service = FakeBatchService(monkeypatch, primary=_primary())
        service.default_repair_status = "processing"
        monkeypatch.setattr(cli, "ensure_batch_ended", lambda batch_id, **_k: None)
        monkeypatch.setattr(
            cli, "thin_submission_from_batch_results",
            lambda batch_id, **kw: pl.reconstruct_batch_submission(
                batch_id=batch_id,
                request_map={
                    request_id(i): {"filename": n, "index": i, "type": "review"}
                    for i, n in enumerate(NAMES)
                },
                review_request_ids=[request_id(i) for i in range(len(NAMES))],
                files_reviewed=list(NAMES),
                input_dir=kw["input_dir"],
                files=kw["files"],
                model="claude-opus-5",
                project_context="",
                module=kw["module"],
                cross_check_enabled=kw["cross_check_enabled"],
                created_at=0.0,
            ),
        )
        rc = cli.main([
            "--batch-id", PRIMARY_ID, "--module", "california_k12_mep",
            "--input-dir", str(spec_dir), "-o", str(tmp_path / "r.docx"),
        ])
        assert rc == 2
        saved = load_pending_batch()
        assert saved is not None and saved.batch_id == PRIMARY_ID
        assert saved.repair_batch_id == "msgbatch_REPAIR_1"
        assert saved.input_dir == str(spec_dir)
        assert len(saved.files) == 2

        # The next plain run resumes it: no second repair.
        service.status["msgbatch_REPAIR_1"] = "ended"
        assert cli.main(["-o", str(tmp_path / "r2.docx")]) == 0
        assert service.repair_submits == [["B.docx"]]
        assert service.cross_checks == 1
        assert not state_path.exists()

    def test_bare_id_recovery_never_overwrites_another_runs_record(
        self, cli, tmp_path, state_path, monkeypatch
    ):
        _save_single(batch_id="msgbatch_SOMEONE_ELSE")
        service = FakeBatchService(monkeypatch, primary=_primary())
        service.default_repair_status = "processing"
        monkeypatch.setattr(cli, "ensure_batch_ended", lambda batch_id, **_k: None)
        monkeypatch.setattr(
            cli, "thin_submission_from_batch_results",
            lambda batch_id, **kw: submission(NAMES, batch_id=batch_id),
        )
        printed: list[str] = []
        monkeypatch.setattr(cli, "_log", lambda msg, *, level="info": printed.append(msg))
        rc = cli.main([
            "--batch-id", PRIMARY_ID, "--module", "california_k12_mep",
            "-o", str(tmp_path / "r.docx"),
        ])
        assert rc == 2
        assert load_pending_batch().batch_id == "msgbatch_SOMEONE_ELSE"
        assert any("msgbatch_REPAIR_1" in line and "another run" in line for line in printed)


class TestEntryPointParity:
    """Every entry point applies the same rule to the same outcome.

    The GUI halves of this table are in ``test_repair_recovery_gui.py``; they
    read the same ``PARITY_SCENARIOS`` rows.
    """

    def _save(self, tmp_path):
        spec_dir = tmp_path / "specs"
        spec_dir.mkdir(exist_ok=True)
        files = write_docx_specs(spec_dir, PARITY_NAMES)
        save_pending_batch(
            PendingBatch.from_submission(
                submission(PARITY_NAMES), input_dir=spec_dir, files=files
            )
        )

    @pytest.mark.parametrize("name", list(PARITY_SCENARIOS))
    def test_recovery_cli(self, name, cli, tmp_path, state_path, monkeypatch):
        kind, repair_status, kept = PARITY_SCENARIOS[name]
        self._save(tmp_path)
        service = FakeBatchService(monkeypatch, primary=parity_primary(kind))
        service.default_repair_status = repair_status
        rc = cli.main(["-o", str(tmp_path / "r.docx")])
        assert state_path.exists() is kept
        assert rc == (2 if kept else 0)

    @pytest.mark.parametrize("name", list(PARITY_SCENARIOS))
    def test_headless_decision(self, name, tmp_path, state_path, monkeypatch):
        kind, repair_status, kept = PARITY_SCENARIOS[name]
        self._save(tmp_path)
        service = FakeBatchService(monkeypatch, primary=parity_primary(kind))
        service.default_repair_status = repair_status
        result = pl.run_batch_collection_headless(submission(PARITY_NAMES), log=_Log())
        apply_saved_state_cleanup(result)
        assert state_path.exists() is kept


# ===========================================================================
# 8. Program children keep their unresolved state independently
# ===========================================================================


def _program_primary(*, fire_fails: bool = True, arch_fails: bool = False) -> dict:
    fire = {request_id(0): review_truncated() if fire_fails else review_ok("21 13 13 Wet.docx")}
    arch = {request_id(0): review_truncated() if arch_fails else review_ok("07 27 26 Air.docx")}
    return {"msgbatch_datacenter_fire": fire, "msgbatch_datacenter_architecture": arch}


class TestProgramChildren:
    def _submission(self, *, project_context: str = "", spec_dir: Path | None = None):
        if spec_dir is not None:
            write_docx_specs(spec_dir, ["21 13 13 Wet.docx", "07 27 26 Air.docx"])
        partitions = {
            "datacenter_fire": submission(
                ["21 13 13 Wet.docx"], batch_id="msgbatch_datacenter_fire",
                module_id="datacenter_fire", project_context=project_context,
            ),
            "datacenter_architecture": submission(
                ["07 27 26 Air.docx"], batch_id="msgbatch_datacenter_architecture",
                module_id="datacenter_architecture", project_context=project_context,
            ),
        }
        return pp.ProgramSubmission(
            program_id=HYPERSCALE_DATACENTER_PROGRAM.program_id,
            assignments=_program_assignments(spec_dir=spec_dir),
            partitions=partitions,
        )

    def test_one_childs_pending_repair_holds_every_childs_paid_stages(
        self, monkeypatch, state_path
    ):
        service = FakeBatchService(monkeypatch, primary=_program_primary())
        service.default_repair_status = "processing"
        result = pp.collect_program_results(
            self._submission(project_context=_DIGEST_CONTEXT), log=_Log()
        )
        assert result.provisional and result.status == "partial"
        assert service.paid_downstream == 0
        # The settled sibling waited too, and says so.
        arch = result.module_results["datacenter_architecture"]
        assert not arch.collection_outcome.provisional
        assert STAGE_VERIFICATION in arch.collection_outcome.deferred_stages
        assert "Fire Suppression" in arch.cross_check_result.thinking
        assert result.deferred_program_stages == (STAGE_DRAWING_IMPACT,)
        assert result.drawing_impact_result is None
        assert not decide_saved_state_cleanup(result).clear

    def test_each_childs_repair_is_stamped_on_its_own_partition(self, monkeypatch, state_path):
        program_submission = self._submission()
        save_pending_program_run(PendingProgramRun.from_submission(program_submission))
        service = FakeBatchService(
            monkeypatch, primary=_program_primary(fire_fails=True, arch_fails=True)
        )
        service.default_repair_status = "processing"
        pp.collect_program_results(program_submission, log=_Log())
        saved = load_pending_run()
        stamps = {
            module_id: child.get("repair_batch_id")
            for module_id, child in saved.partitions.items()
        }
        assert set(stamps.values()) == {"msgbatch_REPAIR_1", "msgbatch_REPAIR_2"}

    def test_resume_rereads_a_consumed_repair_and_runs_downstream_once(
        self, monkeypatch, state_path, tmp_path
    ):
        """Fire's repair is pending while architecture's is consumed. On the
        next collection architecture's repair is re-read, not resubmitted;
        once fire's lands, both modules' stages run exactly once."""
        program_submission = self._submission(spec_dir=tmp_path)
        save_pending_program_run(PendingProgramRun.from_submission(program_submission))
        service = FakeBatchService(
            monkeypatch, primary=_program_primary(fire_fails=True, arch_fails=True)
        )
        service.default_repair_status = "ended"

        submit = service._submit

        def submit_fire_stays_running(specs, **kwargs):
            job = submit(specs, **kwargs)
            if specs[0].filename == "21 13 13 Wet.docx":
                service.status[job.batch_id] = "processing"
            return job

        monkeypatch.setattr(pl, "submit_review_batch", submit_fire_stays_running)
        first = pp.collect_program_results(program_submission, log=_Log())
        assert first.provisional
        assert len(service.repair_submits) == 2
        assert service.paid_downstream == 0

        # Restart from the manifest; the fire repair has now ended.
        for repair_id in list(service.status):
            service.status[repair_id] = "ended"
        resumed = load_pending_run().to_submission()
        second = pp.collect_program_results(resumed, log=_Log())
        assert not second.provisional
        assert len(service.repair_submits) == 2  # no duplicate for either module
        assert service.cross_checks == 2  # once per module
        assert len(service.verification_rounds) == 2
        assert decide_saved_state_cleanup(second).clear


# ===========================================================================
# 9. Reports say the result is provisional and which stages wait
# ===========================================================================


def _provisional_result(monkeypatch) -> PipelineResult:
    service = FakeBatchService(monkeypatch, primary=_primary())
    service.default_repair_status = "processing"
    return pl.run_batch_collection_headless(
        submission(NAMES, project_context=_DIGEST_CONTEXT), log=_Log()
    )


class TestProvisionalReports:
    def test_banner_row_and_notice_name_the_repair_and_the_waiting_stages(
        self, monkeypatch, tmp_path, state_path
    ):
        from docx import Document

        result = _provisional_result(monkeypatch)
        out = export_report(result, tmp_path / "r.docx")
        doc = Document(str(out))
        text = "\n".join(p.text for p in doc.paragraphs)
        rows = {
            row.cells[0].text: row.cells[1].text
            for table in doc.tables
            for row in table.rows
            if len(row.cells) == 2
        }
        assert rows["Provisional — review repair outstanding"] == (
            "1 spec awaiting a review repair; 3 stages deferred"
        )
        assert "PROVISIONAL REPORT" in text
        assert "msgbatch_REPAIR_1" in text
        assert (
            "finding verification, cross-spec coordination and drawing-impact "
            "analysis have not run"
        ) in text
        assert "None of the findings below has been verified" in text
        # The spec awaiting its repair is not told to be re-run separately.
        assert "Re-run this spec individually" not in text

    def test_a_spec_that_failed_for_good_keeps_its_rerun_hint(self, monkeypatch, tmp_path, state_path):
        from docx import Document

        names = ["A.docx", "B.docx", "C.docx"]
        primary = {
            PRIMARY_ID: {
                request_id(0): review_ok("A.docx"),
                request_id(1): review_truncated(),
                request_id(2): review_refused(),
            }
        }
        service = FakeBatchService(monkeypatch, primary=primary)
        service.default_repair_status = "processing"
        result = pl.run_batch_collection_headless(submission(names), log=_Log())
        text = "\n".join(
            p.text for p in Document(str(export_report(result, tmp_path / "r.docx"))).paragraphs
        )
        assert "1 spec failed review and was NOT reviewed: C.docx" in text

    def test_html_report_carries_the_same_notice(self, monkeypatch, state_path):
        html = render_html_report(_provisional_result(monkeypatch))
        assert "Provisional — review repair outstanding" in html
        assert "PROVISIONAL REPORT" in html

    def test_sidecar_is_marked_provisional(self, monkeypatch, state_path):
        payload = build_edit_instructions(_provisional_result(monkeypatch))
        assert payload["provisional"] is True
        assert payload["collection"]["repair"]["state"] == REPAIR_PENDING
        assert payload["collection"]["awaiting_repair_specs"] == ["B.docx"]

    def test_a_settled_run_is_unchanged(self, monkeypatch, state_path):
        FakeBatchService(monkeypatch, primary=_primary())
        result = pl.run_batch_collection_headless(submission(NAMES), log=_Log())
        summary = _summarize_run_diagnostics(
            findings=list(result.review_result.findings),
            status_counts={},
            edit_action_counts={},
            cross_check_result=result.cross_check_result,
            pipeline_result=result,
        )
        assert "collection" not in summary
        assert "PROVISIONAL" not in render_html_report(result)
        assert build_edit_instructions(result)["provisional"] is False

    def test_program_roll_up_names_the_module_and_the_program_stage(self):
        states = [
            (
                "Fire Suppression",
                {
                    "provisional": True,
                    "waiting": [
                        {"label": "", "repair_batch_id": "FR", "repair_state": "pending",
                         "specs": ["21 13 13 Wet.docx"]}
                    ],
                    "deferred_stages": [STAGE_VERIFICATION, STAGE_CROSS_CHECK],
                },
            ),
            (
                "Architecture",
                {"provisional": False, "waiting": [],
                 "deferred_stages": [STAGE_VERIFICATION]},
            ),
        ]
        aggregate = _aggregate_run_diagnostics(
            [(label, {"collection": state}) for label, state in states],
            deferred_program_stages=(STAGE_DRAWING_IMPACT,),
        )
        collection = aggregate["collection"]
        assert collection["waiting"][0]["label"] == "Fire Suppression"
        assert collection["deferred_stages"] == [
            STAGE_VERIFICATION, STAGE_CROSS_CHECK, STAGE_DRAWING_IMPACT,
        ]

    def test_program_report_opens_with_the_notice(self, monkeypatch, tmp_path, state_path):
        from docx import Document

        service = FakeBatchService(monkeypatch, primary=_program_primary())
        service.default_repair_status = "processing"
        result = pp.collect_program_results(
            TestProgramChildren()._submission(), log=_Log()
        )
        text = "\n".join(
            p.text for p in Document(str(export_report(result, tmp_path / "p.docx"))).paragraphs
        )
        fire = pl.get_module("datacenter_fire").display_name
        assert "PROVISIONAL REPORT" in text
        assert f"{fire}: review repair batch msgbatch_REPAIR_1" in text

    def test_provisional_notice_for_logs(self, monkeypatch, state_path):
        notice = provisional_notice(_provisional_result(monkeypatch))
        assert notice.startswith("Provisional result: review repair batch msgbatch_REPAIR_1")
        assert "not resubmitted" in notice


# ===========================================================================
# 10. Adoption of a provisional run with no record
# ===========================================================================


class TestAdoption:
    def _provisional_submission(self):
        sub = submission(NAMES)
        sub.repair_batch_id = "msgbatch_R"
        sub.repair_request_map = {"x": {"filename": "B.docx", "index": 0}}
        return sub

    def test_saves_when_the_slot_is_free(self, state_path):
        log = _Log()
        assert adopt_outstanding_run(self._provisional_submission(), log=log)
        saved = load_pending_batch()
        assert saved.batch_id == PRIMARY_ID and saved.repair_batch_id == "msgbatch_R"
        assert "resumed later without resubmitting" in log.text("info")

    def test_never_overwrites_another_run(self, state_path):
        _save_single(batch_id="msgbatch_OTHER")
        log = _Log()
        assert not adopt_outstanding_run(self._provisional_submission(), log=log)
        assert load_pending_batch().batch_id == "msgbatch_OTHER"
        assert "msgbatch_R" in log.text("warning")

    def test_restamps_its_own_record_whose_stamp_was_lost(self, state_path):
        _save_single(repair_id=None)
        assert adopt_outstanding_run(self._provisional_submission())
        assert load_pending_batch().repair_batch_id == "msgbatch_R"

    def test_real_time_is_a_no_op_and_failures_never_raise(self, state_path, monkeypatch):
        realtime = submission(NAMES, transport="realtime")
        assert not adopt_outstanding_run(realtime)
        assert not state_path.exists()

        def boom(*_a, **_k):
            raise RuntimeError("disk full")

        monkeypatch.setattr(br, "save_pending_batch", boom)
        log = _Log()
        assert not adopt_outstanding_run(self._provisional_submission(), log=log)
        assert "disk full" in log.text("warning")


# ===========================================================================
# 11. A kept record names every repair batch the run created
# ===========================================================================


def _with_repair(sub, repair_id="msgbatch_R"):
    """``sub`` as the collect step leaves it once a repair is submitted: the
    in-memory submission carries the id even when saving it failed."""
    sub.repair_batch_id = repair_id
    sub.repair_request_map = {"r0": {"filename": "B.docx", "index": 0, "type": "review"}}
    return sub


def _program_partitions():
    return {
        "datacenter_fire": submission(
            ["21 13 13 Wet.docx"], batch_id="msgbatch_datacenter_fire",
            module_id="datacenter_fire",
        ),
        "datacenter_architecture": submission(
            ["07 27 26 Air.docx"], batch_id="msgbatch_datacenter_architecture",
            module_id="datacenter_architecture",
        ),
    }


def _program_submission():
    return pp.ProgramSubmission(
        program_id=HYPERSCALE_DATACENTER_PROGRAM.program_id,
        assignments=_program_assignments(),
        partitions=_program_partitions(),
    )


class TestKeptRecordsNameEveryRepair:
    """Found in review (Codex, P1). Saving a repair's id can fail after the
    writer's own retries; the id then lives only on the in-memory submission.
    A record kept without it would let the next collection pay for a second
    repair, so the shared cleanup step re-stamps it — for every kept run, on
    every entry point, programs included — and never onto another run's
    record."""

    def test_a_kept_single_run_is_restamped(self, state_path):
        _save_single(repair_id=None)
        sub = _with_repair(submission(NAMES))
        # Kept because every spec failed, not because the repair is pending.
        result = _result(_outcome(REPAIR_CONSUMED, repair_id="msgbatch_R", failed=NAMES))
        log = _Log()
        decision, status = apply_saved_state_cleanup(result, submission=sub, log=log)
        assert not decision.clear and status is None
        saved = load_pending_batch()
        assert saved.repair_batch_id == "msgbatch_R"
        assert saved.repair_request_map == sub.repair_request_map
        assert "first save had failed" in log.text("info")

    def test_a_kept_program_manifest_is_restamped_per_child(self, state_path):
        program = _program_submission()
        save_pending_program_run(PendingProgramRun.from_submission(program))
        _with_repair(program.partitions["datacenter_fire"], "msgbatch_FIRE_R")
        result = _program_result(
            {"datacenter_fire": _outcome(
                REPAIR_CONSUMED, batch_id="msgbatch_datacenter_fire",
                repair_id="msgbatch_FIRE_R", module_id="datacenter_fire",
            )},
            module_errors={"datacenter_architecture": "results endpoint timed out"},
        )
        decision, _status = apply_saved_state_cleanup(result, submission=program)
        assert not decision.clear
        saved = load_pending_run()
        assert saved.partitions["datacenter_fire"]["repair_batch_id"] == "msgbatch_FIRE_R"
        assert saved.partitions["datacenter_architecture"].get("repair_batch_id") is None

    def test_a_record_that_already_names_the_repair_is_not_rewritten(
        self, state_path, monkeypatch
    ):
        _save_single(repair_id="msgbatch_R")
        writes: list[str] = []
        monkeypatch.setattr(
            br, "_write_pending_state", lambda *_a, what, **_k: writes.append(what) or True
        )
        assert restamp_repair_batches(_with_repair(submission(NAMES)))
        assert writes == []

    def test_another_runs_record_is_never_stamped_and_none_is_created(self, state_path):
        _save_single(batch_id="msgbatch_OTHER")
        log = _Log()
        assert not restamp_repair_batches(_with_repair(submission(NAMES)), log=log)
        other = load_pending_batch()
        assert other.batch_id == "msgbatch_OTHER" and other.repair_batch_id is None

        state_path.unlink()
        assert not restamp_repair_batches(_with_repair(submission(NAMES)), log=log)
        assert not state_path.exists()
        # Nothing of this run's was on disk, so there was nothing to fail at.
        assert log.lines == []

    def test_a_complete_run_is_cleared_not_restamped(self, state_path):
        _save_single(repair_id=None)
        sub = _with_repair(submission(NAMES))
        result = _result(
            _outcome(REPAIR_CONSUMED, repair_id="msgbatch_R", failed=["B.docx"])
        )
        decision, status = apply_saved_state_cleanup(result, submission=sub)
        assert decision.clear and status == CLEAR_CLEARED
        assert not state_path.exists()

    def test_a_failed_restamp_is_reported_and_never_raises(self, state_path, monkeypatch):
        _save_single(repair_id=None)
        monkeypatch.setattr(br, "_write_pending_state", lambda *_a, **_k: False)
        log = _Log()
        assert not restamp_repair_batches(_with_repair(submission(NAMES)), log=log)
        assert "msgbatch_R" in log.text("warning") and PRIMARY_ID in log.text("warning")

        def boom(*_a, **_k):
            raise RuntimeError("disk full")

        monkeypatch.setattr(br, "record_repair_batch", boom)
        assert not restamp_repair_batches(_with_repair(submission(NAMES)), log=log)

    def test_a_program_childs_lost_stamp_is_restamped_and_resumed_once(
        self, monkeypatch, state_path, tmp_path
    ):
        """End to end through the program collector: the fire child's repair
        stamp is lost, the cleanup step re-stamps it, and the resumed run
        re-attaches to that repair instead of paying for another."""
        program = TestProgramChildren()._submission(spec_dir=tmp_path)
        save_pending_program_run(PendingProgramRun.from_submission(program))
        service = FakeBatchService(monkeypatch, primary=_program_primary())
        service.default_repair_status = "processing"
        real_write = br._write_pending_state
        writes: list[str] = []

        def first_write_fails(payload, target, *, what):
            writes.append(what)
            if len(writes) == 1:
                return False
            return real_write(payload, target, what=what)

        monkeypatch.setattr(br, "_write_pending_state", first_write_fails)
        first = pp.collect_program_results(program, log=_Log())
        assert first.provisional
        assert load_pending_run().partitions["datacenter_fire"].get("repair_batch_id") is None
        apply_saved_state_cleanup(first, submission=program)
        saved = load_pending_run()
        assert saved.partitions["datacenter_fire"]["repair_batch_id"] == "msgbatch_REPAIR_1"

        service.status["msgbatch_REPAIR_1"] = "ended"
        second = pp.collect_program_results(saved.to_submission(), log=_Log())
        assert not second.provisional
        assert service.repair_submits == [["21 13 13 Wet.docx"]]
