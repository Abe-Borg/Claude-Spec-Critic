"""The applier's side of plan WP-02: readable is not writable (chunk S10).

S10 made text inside content controls, fields, smart tags, custom XML
elements, and hyperlinks readable, so it reaches the review and can be
quoted in an edit instruction. The writer still edits only plain runs that
are direct children of an ordinary paragraph. These tests hold the boundary:

* **Every wrapped target is refused with its own reason** — content control,
  field result, smart tag, custom XML, bidirectional container, hyperlink,
  another author's revision — and nothing is written.
* **The match is made against the text the review saw.** A target is never
  found by skipping over a control's text to runs that happen to read the
  same without it, and a match that would have to move a control, field,
  revision, or anchored object is refused.
* **A copy inside a wrapper counts.** Identical text inside and outside a
  control, in one paragraph or one row, is ambiguous; with no confirming
  element id, a copy in any unwritable element makes the whole document's
  match ambiguous (the locator), and the assist tier cannot choose it.
* **Ids with a ``cc`` step are explicitly unsupported** for writing.

Hermetic: documents are built with the shared fixture builders and saved
only under ``tmp_path``.
"""
from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from applier.assist import AssistConfig, assist_location
from applier.docx_edit import DIRECT, TRACKED, DocumentEditor, EditError
from applier.locator import build_candidates, classify_element_id, locate
from applier.models import (
    EditEntry,
    ElementKind,
    Location,
    LocationStatus,
    OutcomeStatus,
)
from applier.run import RunSettings, apply_sidecar
from applier.sidecar import load_sidecar
from src.input.extractor import _accept_all_paragraph_text, extract_text_from_docx
from tests.fixtures import spec_docx as fx


def entry(**overrides) -> EditEntry:
    base = dict(
        finding_id="rf-1",
        file_name="215000.docx",
        action_type="EDIT",
        existing_text=None,
        replacement_text=None,
        anchor_text=None,
        insert_position=None,
        target_element_id=None,
        evidence_element_id=None,
        edit_confidence=0.9,
        report_status="VERIFIED_SUPPORTED",
        verification_verdict="CONFIRMED",
        severity="HIGH",
        section="",
        issue="…",
        code_reference=None,
        has_per_file_original=True,
    )
    base.update(overrides)
    return EditEntry(**base)


def at(element_id: str, kind: ElementKind = ElementKind.BODY_PARAGRAPH) -> Location:
    return Location(status=LocationStatus.RESOLVED_BY_ID, element_id=element_id, kind=kind)


def one_paragraph(*inline: str):
    builder = fx.SpecDocBuilder()
    builder.add_paragraph(*inline)
    return builder.document


def body_xml(document) -> bytes:
    return document.element.body.xml.encode()


def refused(document, the_entry: EditEntry, location: Location, *, mode=TRACKED) -> str:
    """Apply, expect a refusal, and prove the document was left untouched."""
    before = body_xml(document)
    with pytest.raises(EditError) as excinfo:
        DocumentEditor(document, mode=mode).apply(the_entry, location)
    assert body_xml(document) == before, "a refused edit must not change the document"
    return str(excinfo.value)


def accept_all(document, index: int = 0) -> str:
    return _accept_all_paragraph_text(document.element.body.findall(qn("w:p"))[index])


def reject_all(p_el) -> str:
    """What survives Reject All: ``w:t`` outside ``w:ins``, plus ``w:delText``."""
    parts = []
    for node in p_el.iter():
        if node.tag == qn("w:delText"):
            parts.append(node.text or "")
        elif node.tag == qn("w:t") and not any(
            ancestor.tag == qn("w:ins") for ancestor in node.iterancestors()
        ):
            parts.append(node.text or "")
    return "".join(parts)


_CUSTOM_XML = (
    '<w:customXml w:uri="urn:fixture" w:element="material">'
    f"{fx.run('ductile iron')}</w:customXml>"
)


# ---------------------------------------------------------------------------
# The writer
# ---------------------------------------------------------------------------


class TestEachWrapperIsRefusedWithItsOwnReason:
    @pytest.mark.parametrize(
        "inline, target, reason",
        [
            (fx.inline_control(fx.run("copper"), tag="t", control_id=1), "copper",
             "inside a content control"),
            (fx.smart_tag(fx.run("Oakland")), "Oakland", "inside a smart tag"),
            (fx.simple_field(" REF bm ", "23 05 00"), "23 05 00", "field's stored result"),
            (fx.complex_field(" REF bm ", "21 13 13"), "21 13 13", "field's stored result"),
            (_CUSTOM_XML, "ductile iron", "inside a custom XML element"),
            (f'<w:dir w:val="rtl">{fx.run("A-1")}</w:dir>', "A-1", "bidirectional-text"),
        ],
        ids=["content_control", "smart_tag", "simple_field", "complex_field", "custom_xml", "dir"],
    )
    def test_the_target_inside_a_wrapper(self, inline, target, reason):
        document = one_paragraph("Use ", inline, " here.")
        message = refused(document, entry(existing_text=target, replacement_text="X"), at("p0"))
        assert reason in message and "by hand" in message

    def test_hyperlink_text_is_refused_as_a_hyperlink_not_as_a_revision(self):
        """Found by S01: the old writer searched every nested run for its
        refusal message, so link text read as "inside a tracked revision"."""
        builder = fx.SpecDocBuilder()
        builder.add_paragraph("Verify in ", builder.hyperlink(fx.HYPERLINK_URL, "the directory"), ".")
        message = refused(
            builder.document, entry(existing_text="the directory", replacement_text="x"), at("p0")
        )
        assert "inside a hyperlink" in message
        assert "tracked revision" not in message

    def test_the_outermost_container_names_the_refusal(self):
        document = fx.build_hyperlink_spec().document
        message = refused(
            document, entry(existing_text="manufacturer", replacement_text="x"), at("p1")
        )
        assert "inside a hyperlink" in message

    def test_another_authors_insertion_keeps_its_reason(self):
        document = one_paragraph("Comply with ", fx.inserted(fx.run("2025"), revision_id=1), " CBC.")
        message = refused(document, entry(existing_text="2025", replacement_text="2022"), at("p0"))
        assert "inside an existing tracked revision" in message

    def test_the_same_refusals_hold_in_direct_mode(self):
        document = one_paragraph("Use ", fx.smart_tag(fx.run("Oakland")), " here.")
        message = refused(
            document, entry(existing_text="Oakland", replacement_text="x"), at("p0"), mode=DIRECT
        )
        assert "inside a smart tag" in message


class TestTheMatchIsTheTextTheReviewSaw:
    def test_a_match_reaching_into_a_control_is_refused(self):
        document = one_paragraph(
            "Use ", fx.inline_control(fx.run("copper"), tag="t", control_id=1), " pipe."
        )
        message = refused(document, entry(existing_text="Use copper", replacement_text="x"), at("p0"))
        assert "inside a content control" in message

    def test_runs_around_a_control_are_not_matched_as_if_it_were_absent(self):
        """Before S10 the writer matched against its direct runs only, where
        this paragraph reads "Use  pipe." — so "Use pipe" was found, and the
        edit moved the control out of its sentence."""
        document = one_paragraph(
            "Use ", fx.inline_control(fx.run("copper"), tag="t", control_id=1), " pipe."
        )
        message = refused(document, entry(existing_text="Use pipe", replacement_text="x"), at("p0"))
        assert "not found in the located element" in message

    def test_a_span_across_an_empty_control_is_refused(self):
        empty_control = '<w:sdt><w:sdtPr><w:id w:val="3"/></w:sdtPr><w:sdtContent/></w:sdt>'
        document = one_paragraph("steel", empty_control, " pipe")
        message = refused(document, entry(existing_text="steel pipe", replacement_text="x"), at("p0"))
        assert "spans a field, content control" in message

    def test_a_span_across_another_authors_deletion_is_refused(self):
        """Moving the matched runs into a new deletion would put the other
        author's deletion after it, and Reject All would read "…be steelcopper"."""
        document = one_paragraph(
            "Pipe shall be ", fx.deleted(fx.deleted_run("copper "), revision_id=1), "steel."
        )
        message = refused(document, entry(existing_text="be steel", replacement_text="x"), at("p0"))
        assert "spans a field" in message

    def test_a_span_across_a_field_is_refused(self):
        no_result = (
            '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
            '<w:r><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>'
            '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
        )
        document = one_paragraph("See ", no_result, "Section 2.")
        message = refused(document, entry(existing_text="See Section", replacement_text="x"), at("p0"))
        assert "spans a field" in message

    def test_a_matched_run_carrying_a_field_character_is_refused(self):
        """A run can hold visible text and then start a field. Moving it into a
        deletion would carry the field's first character with it."""
        field_start = (
            '<w:r><w:t xml:space="preserve">See </w:t><w:fldChar w:fldCharType="begin"/></w:r>'
            '<w:r><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>'
            '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
        )
        document = one_paragraph(field_start, "Section 2.")
        assert accept_all(document) == "See Section 2."
        message = refused(document, entry(existing_text="See", replacement_text="Refer to"), at("p0"))
        assert "spans a field" in message

    def test_a_span_across_an_anchored_object_is_refused(self):
        reference = '<w:r><w:footnoteReference w:id="1"/></w:r>'
        document = one_paragraph("valves", reference, " shall close")
        message = refused(document, entry(existing_text="valves shall", replacement_text="x"), at("p0"))
        assert "spans a field" in message

    def test_bookmarks_and_proofing_marks_inside_a_span_are_fine(self):
        document = one_paragraph(
            "Install ",
            '<w:bookmarkStart w:id="0" w:name="spec"/>',
            "per ",
            '<w:proofErr w:type="spellStart"/>',
            "NFPA 13",
            '<w:proofErr w:type="spellEnd"/>',
            '<w:bookmarkEnd w:id="0"/>',
            ".",
        )
        DocumentEditor(document, mode=TRACKED).apply(
            entry(existing_text="per NFPA 13", replacement_text="per NFPA 13 (2025)"), at("p0")
        )
        assert accept_all(document) == "Install per NFPA 13 (2025)."
        assert reject_all(document.element.body[0]) == "Install per NFPA 13."

    def test_text_beside_a_control_still_edits_and_leaves_the_control_alone(self):
        document = one_paragraph(
            "Provide ", fx.inline_control(fx.run("schedule 40"), tag="t", control_id=1), " pipe for mains."
        )
        DocumentEditor(document, mode=TRACKED).apply(
            entry(existing_text="for mains", replacement_text="for risers"), at("p0")
        )
        assert accept_all(document) == "Provide schedule 40 pipe for risers."
        control = document.element.body[0].find(qn("w:sdt"))
        assert control.xpath("string(.)") == "schedule 40"


class TestTheSameTextInsideAndOutsideAControl:
    def test_in_one_paragraph_is_ambiguous(self):
        document = one_paragraph(
            "Use ", fx.inline_control(fx.run("steel"), tag="t", control_id=1), " sleeves and steel pipe."
        )
        message = refused(document, entry(existing_text="steel", replacement_text="iron"), at("p0"))
        assert "occurs 2 times" in message

    def test_in_one_row_is_ambiguous(self):
        builder = fx.SpecDocBuilder()
        control = fx.block_control(fx.paragraph("Density 0.15 gpm/sf"), tag="c", control_id=1)
        builder.add_xml(
            '<w:tbl><w:tblPr><w:tblW w:w="0" w:type="auto"/></w:tblPr>'
            '<w:tblGrid><w:gridCol w:w="2000"/><w:gridCol w:w="2000"/></w:tblGrid>'
            f"<w:tr>{fx._cell_of(control)}{fx._cell_xml('Remote area 0.15 gpm/sf')}</w:tr></w:tbl>"
        )
        message = refused(
            builder.document,
            entry(existing_text="0.15 gpm/sf", replacement_text="0.20 gpm/sf"),
            at("t0r0", ElementKind.TABLE_ROW),
        )
        assert "occurs 2 times" in message

    def test_a_row_target_only_inside_a_cell_control_is_refused(self):
        document = fx.build_table_controls_spec().document
        message = refused(
            document,
            entry(existing_text=fx.CELL_CONTROL_TEXT, replacement_text="Six inch grooved"),
            at("t0r1", ElementKind.TABLE_ROW),
        )
        assert "inside a content control" in message

    def test_a_row_target_only_inside_a_wrapped_cell_is_refused(self):
        document = fx.build_table_controls_spec().document
        message = refused(
            document,
            entry(existing_text=fx.WRAPPED_CELL_TEXT, replacement_text="Two inch grooved"),
            at("t0r2", ElementKind.TABLE_ROW),
        )
        assert "inside a content control" in message

    def test_a_plain_cell_beside_a_cell_control_still_edits(self):
        document = fx.build_table_controls_spec().document
        DocumentEditor(document, mode=TRACKED).apply(
            entry(existing_text="Riser", replacement_text="Standpipe"),
            at("t0r1", ElementKind.TABLE_ROW),
        )
        cell = document.tables[0].rows[1].cells[0]
        assert _accept_all_paragraph_text(cell.paragraphs[0]._p) == "Standpipe"
        control = document.tables[0].rows[1].cells[1]._tc.find(qn("w:sdt"))
        assert fx.CELL_CONTROL_TEXT in control.xpath("string(.)")


class TestAdditionsBesideWrappedContent:
    def test_an_anchor_inside_a_control_is_refused(self):
        document = one_paragraph(
            "Provide ", fx.inline_control(fx.run("listed valves"), tag="t", control_id=1), "."
        )
        message = refused(
            document,
            entry(action_type="ADD", replacement_text="New clause.", anchor_text="listed valves",
                  insert_position="after"),
            at("p0"),
        )
        assert "inside a content control" in message

    def test_an_anchor_in_plain_text_beside_a_control_applies(self):
        document = one_paragraph(
            "Provide ", fx.inline_control(fx.run("listed valves"), tag="t", control_id=1), " at risers."
        )
        DocumentEditor(document, mode=TRACKED).apply(
            entry(action_type="ADD", replacement_text="New clause.", anchor_text="at risers",
                  insert_position="after"),
            at("p0"),
        )
        assert accept_all(document, 1) == "New clause."

    def test_an_addition_positioned_by_a_wrapped_cell_is_refused(self):
        builder = fx.SpecDocBuilder()
        wrapped = (
            '<w:sdt><w:sdtPr><w:id w:val="2"/></w:sdtPr>'
            f"<w:sdtContent>{fx._cell_xml('Wrapped first cell')}</w:sdtContent></w:sdt>"
        )
        builder.add_xml(
            '<w:tbl><w:tblPr><w:tblW w:w="0" w:type="auto"/></w:tblPr>'
            '<w:tblGrid><w:gridCol w:w="2000"/><w:gridCol w:w="2000"/></w:tblGrid>'
            f"<w:tr>{wrapped}{fx._cell_xml('Plain cell')}</w:tr></w:tbl>"
        )
        message = refused(
            builder.document,
            entry(action_type="ADD", replacement_text="New.", insert_position="after"),
            at("t0r0", ElementKind.TABLE_ROW),
        )
        assert "inside a content control" in message


# ---------------------------------------------------------------------------
# Ids and the locator
# ---------------------------------------------------------------------------


def _spec(builder, tmp_path: Path, name: str = "spec.docx"):
    return extract_text_from_docx(fx.save_docx(builder, tmp_path, name))


class TestContentControlIds:
    def test_every_cc_id_the_extractor_mints_is_explicitly_unsupported(self, tmp_path):
        for build in (
            fx.build_block_control_spec,
            fx.build_table_controls_spec,
            fx.build_nested_controls_spec,
        ):
            spec = _spec(build(), tmp_path, f"{build.__name__}.docx")
            document = Document(spec.source_path)
            wrapped = [m.element_id for m in spec.paragraph_map if "cc" in m.element_id]
            assert wrapped, build.__name__
            for element_id in wrapped:
                assert classify_element_id(element_id) is ElementKind.CONTENT_CONTROL
                with pytest.raises(EditError, match="does not write to"):
                    DocumentEditor(document).resolve_paragraphs(element_id, ElementKind.CONTENT_CONTROL)

    def test_an_instruction_naming_a_control_id_is_reported_not_written(self, tmp_path):
        spec = _spec(fx.build_block_control_spec(), tmp_path)
        location = locate(
            entry(existing_text="listed backflow preventer", replacement_text="x",
                  target_element_id="cc1p0"),
            build_candidates(spec),
        )
        assert location.status is LocationStatus.UNSUPPORTED_ELEMENT
        assert location.kind is ElementKind.CONTENT_CONTROL
        assert "content control" in location.detail and "by hand" in location.detail


def _twin_spec(tmp_path: Path):
    """The same sentence in a body paragraph (p0) and a block control (cc1p0)."""
    builder = fx.SpecDocBuilder()
    builder.add_paragraph("Provide a listed backflow preventer at each service.")
    builder.add_xml(
        fx.block_control(
            fx.paragraph("Provide a listed backflow preventer at each service."),
            tag="twin",
            control_id=builder.next_id(),
        )
    )
    return builder, _spec(builder, tmp_path, "twin.docx")


class TestTheLocatorNeverPrefersTheWritableCopy:
    def test_with_no_id_the_match_is_ambiguous(self, tmp_path):
        _, spec = _twin_spec(tmp_path)
        location = locate(
            entry(existing_text="listed backflow preventer", replacement_text="x"),
            build_candidates(spec),
        )
        assert location.status is LocationStatus.AMBIGUOUS
        assert {c.element_id for c in location.candidates} == {"p0", "cc1p0"}
        assert "cc1p0" in location.detail

    def test_a_borrowed_id_that_misses_is_ambiguous_too(self, tmp_path):
        _, spec = _twin_spec(tmp_path)
        location = locate(
            entry(existing_text="listed backflow preventer", replacement_text="x",
                  target_element_id="p7", has_per_file_original=False),
            build_candidates(spec),
        )
        assert location.status is LocationStatus.AMBIGUOUS

    def test_an_id_that_confirms_its_text_still_resolves(self, tmp_path):
        _, spec = _twin_spec(tmp_path)
        location = locate(
            entry(existing_text="listed backflow preventer", replacement_text="x",
                  target_element_id="p0"),
            build_candidates(spec),
        )
        assert location.status is LocationStatus.RESOLVED_BY_ID
        assert location.element_id == "p0"

    def test_the_assist_tier_cannot_choose_the_control_copy(self, tmp_path):
        _, spec = _twin_spec(tmp_path)
        candidates = build_candidates(spec)
        the_entry = entry(existing_text="listed backflow preventer", replacement_text="x")
        ambiguous = locate(the_entry, candidates)

        class Choose:
            type = "tool_use"
            name = "choose_element"
            input = {"element_id": "cc1p0", "reasoning": "the control"}
            id = "tu_1"

        class Client:
            def __init__(self):
                self.messages = self

            def create(self, **_):
                return type("Message", (), {"content": [Choose()]})()

        result = assist_location(
            the_entry, candidates, ambiguous, client=Client(),
            config=AssistConfig(enabled=True, model="test-model"),
        )
        assert result.status is LocationStatus.AMBIGUOUS
        assert not result.is_applicable
        assert "does not write to — discarded" in result.detail


# ---------------------------------------------------------------------------
# End to end: sidecar in, receipt out
# ---------------------------------------------------------------------------


def _sidecar(directory: Path, file_name: str, edits: list[dict]) -> Path:
    payload = {
        "schema_version": 4,
        "generated_at": "2026-09-24T00:00:00Z",
        "report_file": "report.docx",
        "edit_count": len(edits),
        "edits": [
            {
                "finding_id": f"rf-{index:012d}",
                "fileName": file_name,
                "affected_files": [file_name],
                "has_per_file_original": True,
                "section": "",
                "severity": "HIGH",
                "issue": "Fixture edit.",
                "codeReference": None,
                "evidenceElementId": edit.get("element_id"),
                "verification_verdict": "CONFIRMED",
                "report_status": "VERIFIED_SUPPORTED",
                "edit_proposal": {
                    "action_type": "EDIT",
                    "existing_text": edit["existing"],
                    "replacement_text": edit["replacement"],
                    "anchor_text": None,
                    "insert_position": None,
                    "target_element_id": edit.get("element_id"),
                    "edit_confidence": 0.9,
                },
            }
            for index, edit in enumerate(edits)
        ],
    }
    path = directory / "report.edits.json"
    path.write_text(json.dumps(payload), encoding="utf-8")
    return path


class TestEndToEnd:
    @pytest.mark.parametrize("dry_run", [False, True])
    def test_every_wrapped_target_is_held_and_ordinary_text_still_applies(self, dry_run, tmp_path):
        source = fx.save_docx(fx.build_table_controls_spec(), tmp_path / "in", "230500.docx")
        before = source.read_bytes()
        sidecar = load_sidecar(
            _sidecar(
                tmp_path,
                "230500.docx",
                [
                    {"existing": fx.CELL_CONTROL_TEXT, "replacement": "x", "element_id": "t0r1"},
                    {"existing": fx.WRAPPED_CELL_TEXT, "replacement": "x", "element_id": "t0r2"},
                    {"existing": "Six inch welded", "replacement": "x", "element_id": "t0cc5r0"},
                    {"existing": "Pipe size", "replacement": "Pipe diameter", "element_id": "t0r0"},
                ],
            )
        )
        (result,) = apply_sidecar(
            sidecar, [source], RunSettings(output_dir=tmp_path / "out", dry_run=dry_run,
                                           allow_tracked_source=True)
        )
        assert source.read_bytes() == before
        outcomes = {o.entry.finding_id: o for o in result.outcomes}
        held = [outcomes[f"rf-{i:012d}"] for i in range(3)]
        assert [o.status for o in held] == [OutcomeStatus.UNLOCATED] * 3
        assert all("content control" in o.reason for o in held)
        assert held[2].to_dict()["element_kind"] == "content_control"
        applied = outcomes["rf-000000000003"]
        assert applied.status is (OutcomeStatus.WOULD_APPLY if dry_run else OutcomeStatus.APPLIED)
        if not dry_run:
            edited = Document(result.output_path)
            assert _accept_all_paragraph_text(
                edited.tables[0].rows[0].cells[1].paragraphs[0]._p
            ) == "Pipe diameter"

    def test_identical_text_with_no_id_is_held_as_ambiguous(self, tmp_path):
        builder, _ = _twin_spec(tmp_path)
        source = fx.save_docx(builder, tmp_path / "in", "twin.docx")
        sidecar = load_sidecar(
            _sidecar(tmp_path, "twin.docx",
                     [{"existing": "listed backflow preventer", "replacement": "x"}])
        )
        (result,) = apply_sidecar(sidecar, [source], RunSettings(output_dir=tmp_path / "out"))
        (outcome,) = result.outcomes
        assert outcome.status is OutcomeStatus.UNLOCATED
        assert outcome.location.status is LocationStatus.AMBIGUOUS
        assert not list((tmp_path / "out").glob("*.docx"))
