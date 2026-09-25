"""Compliance coverage completeness (plan WP-09).

A completed compliance response is not evidence that every controlling
requirement was assessed. These tests pin the contract that closes that gap,
in the order of the plan's acceptance list:

1. the expected set, and a prompt / schema / examples that name it;
2. normalization: omitted ids become synthetic "not assessed" rows, told
   apart from a returned "unclear"; unknown, advisory, malformed, duplicate,
   and reordered rows are handled deterministically;
3. zero controlling requirements is a valid, complete result;
4. chunk semantics: a failed, skipped, not-analyzed, or row-omitting chunk —
   or a specification excluded before the pass — cannot establish absence,
   completed findings survive, and the execution status is unchanged;
5. additions whose absence is not established are held as report-only, with
   the reason, never dropped and never emitted as edits;
6. every output path — Word, HTML, sidecar, profile export, diagnostics, the
   run log, program reports, and the headless (CLI / recovery) driver —
   carries the same completeness semantics.

Hermetic throughout: a scripted streaming client, a scripted count API, and
temporary directories.
"""
from __future__ import annotations

import dataclasses
import json
import re
import time

import pytest
from docx import Document

from src.batch.batch import BatchJob
from src.compliance import (
    CoverageCompleteness,
    ensure_coverage_completeness,
    expected_coverage_ids,
    run_chunked_compliance_check,
    run_compliance_check,
)
from src.compliance import compliance_checker as cc
from src.compliance.completeness import (
    STATE_COMPLETE,
    STATE_INCOMPLETE,
    STATE_NO_APPLICABLE_ITEMS,
    STATE_UNKNOWN,
    AssessmentUnit,
    combine,
    nothing_assessed,
    reconcile,
)
from src.core.chunked_pass import ChunkOutcome
from src.input.extractor import ExtractedSpec
from src.modules import DATACENTER_FIRE, DEFAULT_MODULE, ResearchDimension
from src.orchestration import pipeline as pl
from src.orchestration.diagnostics import DiagnosticsReport, compliance_pass_extra
from src.orchestration.pipeline import (
    BatchSubmission,
    CollectedBatchState,
    PipelineResult,
    run_batch_collection_headless,
    run_compliance_for_batch,
)
from src.output.edit_sidecar import (
    SIDECAR_SCHEMA_VERSION,
    build_edit_instructions,
    build_requirements_profile_export,
)
from src.output.html_report_exporter import render_html_report
from src.output.report_exporter import (
    _aggregate_run_diagnostics,
    _compliance_banner_row,
    _compliance_hints,
    _summarize_run_diagnostics,
    export_report,
)
from src.research import DimensionStatus, RequirementsProfile, ResearchItem
from src.review.reviewer import Finding, ReviewResult, is_held_addition
from src.review.structured_schemas import (
    COMPLIANCE_FINDINGS_SCHEMA,
    compliance_findings_tool,
)
from src.verification.verification_cache import VerificationCache
from tests.fixtures.count_api import count_response, user_words
from tests.fixtures.fake_anthropic import (
    FakeMessage,
    FakeTextBlock,
    compliance_tool_use_response,
)

A, B, C = "r-aaaaaaaaaaaa", "r-bbbbbbbbbbbb", "r-cccccccccccc"
U, P = "r-0000000000aa", "r-0000000000bb"  # unverified, process advisory


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


def _item(item_id, requirement, *, grounded=True, **overrides):
    defaults = dict(
        item_id=item_id,
        dimension_id="governing_codes",
        topic="Topic",
        category="governing_code",
        requirement=requirement,
        grounded=grounded,
        accepted_sources=["https://codes.example.gov/x"] if grounded else [],
        confidence=0.8,
    )
    defaults.update(overrides)
    return ResearchItem(**defaults)


def _profile(items=None) -> RequirementsProfile:
    if items is None:
        items = [
            _item(A, "Requirement A applies."),
            _item(B, "Requirement B applies."),
            _item(U, "Unverified requirement.", grounded=False),
            _item(P, "Permit fee schedule.", actionability="process_advisory"),
        ]
    return RequirementsProfile(
        items=list(items),
        dimension_statuses=[
            DimensionStatus(dimension_id="governing_codes", status="completed")
        ],
        research_date="2026-09-24",
        project={"city": "Ashburn", "state_or_province": "VA", "country": "US",
                 "client_name": "ExampleCo"},
    )


def _spec(content, filename="21 13 13 Wet.docx") -> ExtractedSpec:
    return ExtractedSpec(filename=filename, content=content, word_count=len(content.split()))


def _row(rid, status, evidence=None, file_name=None):
    return {"requirement_id": rid, "status": status, "evidence": evidence,
            "fileName": file_name}


def _add(rid, *, file_name="21 13 13 Wet.docx", anchor="PART 1 - GENERAL"):
    return {
        "severity": "HIGH",
        "fileName": file_name,
        "section": "1.02",
        "issue": f"Requirement {rid} is not represented in the package.",
        "actionType": "ADD",
        "existingText": None,
        "replacementText": f"C. Comply with requirement {rid}.",
        "codeReference": "Local amendment",
        "confidence": 0.8,
        "anchorText": anchor,
        "insertPosition": "after",
        "evidenceElementId": None,
    }


def _edit(file_name="21 13 13 Wet.docx"):
    return {
        "severity": "MEDIUM",
        "fileName": file_name,
        "section": "1.03",
        "issue": f"Requirement {A} is contradicted by the stated edition.",
        "actionType": "EDIT",
        "existingText": "Comply with the 2015 IBC.",
        "replacementText": "Comply with the adopted IBC.",
        "codeReference": "Local amendment",
        "confidence": 0.8,
        "anchorText": None,
        "insertPosition": None,
        "evidenceElementId": None,
    }


def _payload(coverage, findings=()):
    return {"compliance_summary": "Summary.", "coverage": list(coverage),
            "findings": list(findings)}


def _response(coverage, findings=()):
    return compliance_tool_use_response(payload=_payload(coverage, findings))


class _Stream:
    def __init__(self, message):
        self._message = message

    def __enter__(self):
        return self

    def __exit__(self, *args):
        return False

    @property
    def text_stream(self):
        return iter(())

    def get_final_message(self):
        return self._message


class _Messages:
    def __init__(self, route):
        self.route = route
        self.calls: list[dict] = []

    def stream(self, **kwargs):
        self.calls.append(kwargs)
        result = self.route(kwargs)
        if isinstance(result, Exception):
            raise result
        return _Stream(result)

    def count_tokens(self, **kwargs):
        return count_response(user_words(kwargs))


class _Client:
    def __init__(self, route):
        self.messages = _Messages(route)


@pytest.fixture(autouse=True)
def _hermetic(monkeypatch):
    monkeypatch.setattr(cc, "count_tokens", lambda text: len(text.split()))


@pytest.fixture
def client(monkeypatch):
    holder: dict = {}

    def install(route):
        fake = _Client(route)
        monkeypatch.setattr(cc, "_get_client", lambda **_: fake)
        holder["client"] = fake
        return fake

    holder["install"] = install
    return holder


def _single(response):
    return lambda kwargs: response


def _by_marker(script: dict):
    def route(kwargs):
        text = kwargs["messages"][0]["content"]
        for marker, response in script.items():
            if marker in text:
                return response
        raise AssertionError(f"no route for {text[:80]!r}")

    return route


def _cycle():
    return DEFAULT_MODULE.cycle


def _rows(result) -> dict:
    return {row["requirement_id"]: row for row in result.coverage}


def _doc_text(path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _coverage_cells(path) -> dict[str, str]:
    """``{requirement id: Coverage cell}`` from the Word report's coverage table."""
    doc = Document(str(path))
    for table in doc.tables:
        header = [cell.text.strip() for cell in table.rows[0].cells]
        if header == ["Requirement", "Coverage", "Evidence"]:
            return {
                row.cells[0].text.split("]")[0].lstrip("["): row.cells[1].text.strip()
                for row in table.rows[1:]
            }
    raise AssertionError("no Requirements Coverage table")


_HTML_COVERAGE_ROW_RE = re.compile(
    r"<tr><td>\[(r-[0-9a-f]{12})\][^<]*</td><td style=\"[^\"]*\">([^<]*)</td>"
)


def _enabled_module():
    return dataclasses.replace(
        DEFAULT_MODULE,
        project_profile_enabled=True,
        research_persona="You are a test research assistant.",
        research_dimensions=(
            ResearchDimension(
                dimension_id="governing_codes",
                title="Governing codes",
                prompt_template="Codes for {city}.",
            ),
        ),
        compliance_persona="You are a test compliance reviewer.",
        compliance_severity_definitions="- CRITICAL — permit-blocking omission.",
    )


def _pipeline_result(compliance_result, profile=None) -> PipelineResult:
    module = _enabled_module()
    return PipelineResult(
        review_result=ReviewResult(findings=[]),
        files_reviewed=["21 13 13 Wet.docx"],
        cycle_label=module.cycle.label,
        module_id=module.module_id,
        requirements_profile=(profile or _profile()).to_dict(),
        compliance_result=compliance_result,
    )


_PLAINTEXT_RE = re.compile(r'<pre id="sc-plaintext" hidden>(.*?)</pre>', re.DOTALL)
_DATA_SCRIPT_RE = re.compile(
    r'<script type="application/json" id="sc-report-data">(.*?)</script>', re.DOTALL
)


def _html_payload(html: str) -> dict:
    match = _DATA_SCRIPT_RE.search(html)
    assert match is not None
    return json.loads(match.group(1))


# Four specs in two CSI divisions: each division fits the patched ceiling,
# the whole package does not — so the chunked path runs.
_PAD = "word " * 2_000


def _chunked_specs():
    return [
        _spec(_PAD + " DIV21SPEC a", "21 13 13 Wet.docx"),
        _spec(_PAD + " DIV21SPEC b", "21 13 16 Dry.docx"),
        _spec(_PAD + " DIV22SPEC a", "22 11 13 Water.docx"),
        _spec(_PAD + " DIV22SPEC b", "22 11 16 Piping.docx"),
    ]


@pytest.fixture
def chunked(monkeypatch):
    monkeypatch.setattr(cc, "COMPLIANCE_RECOMMENDED_MAX", 5_000)


# ---------------------------------------------------------------------------
# 1. The expected set, and a prompt that names it
# ---------------------------------------------------------------------------


class TestExpectedSet:
    def test_grounded_controlling_non_process_ids_in_profile_order(self):
        profile = _profile([
            _item(B, "B."),
            _item(U, "Unverified.", grounded=False),
            _item(A, "A."),
            _item(P, "Fee.", actionability="process_advisory"),
            _item(B, "B again (a duplicate id)."),
            _item("", "An item with no id cannot be tracked."),
        ])
        assert expected_coverage_ids(profile) == (B, A)

    def test_unverified_items_are_rendered_but_never_expected(self, client):
        client["install"](_single(_response([_row(A, "represented", "q", "a"),
                                             _row(B, "represented", "q", "a")])))
        result = run_compliance_check([_spec("PART 1")], _profile(), [], cycle=_cycle())
        message = client["client"].messages.calls[0]["messages"][0]["content"]
        assert "Unverified requirement." in message  # still shown as context
        assert set(_rows(result)) == {A, B}  # but never a coverage row
        assert result.coverage_completeness.expected_ids == (A, B)

    def test_prompt_schema_tool_and_examples_name_the_controlling_set(self):
        system = cc._compliance_system_prompt(DATACENTER_FIRE.cycle)
        rules = " ".join(
            system.split("<coverage_rules>")[1].split("</coverage_rules>")[0].split()
        )
        final = cc._COMPLIANCE_FINAL_TASK_BLOCK
        schema = COMPLIANCE_FINDINGS_SCHEMA["properties"]["coverage"]["description"]
        tool = compliance_findings_tool()["description"]
        examples = cc._COMPLIANCE_EXAMPLES
        # Every surface scopes coverage to the controlling requirements ...
        for text in (rules, final, schema, tool, examples):
            assert "controlling" in text, text
            # ... and none still asks for a row per *profile* requirement id,
            # which read as including the [UNVERIFIED] ids in project context.
            assert "per profile requirement id" not in text, text
        # Unverified items are excluded wherever rows are specified.
        for text in (rules, final, schema, examples):
            assert "[UNVERIFIED]" in text, text
        # A left-out id is not silent: the prompt says what happens to it,
        # and the closing block restates only rules the system prompt holds.
        assert "none left out" in rules and "none left out" in final
        assert "reported as not assessed" in rules

    def test_examples_give_no_coverage_entry_to_the_unverified_item(self):
        rows = json.loads(
            cc._COMPLIANCE_EXAMPLES.split("represented entry backs no finding at all:")[1]
            .split("</examples>")[0]
        )
        unverified_ids = set(re.findall(r"r-[0-9a-f]{12}", cc._COMPLIANCE_EXAMPLES.split(
            "Example 3")[0].split("Example 2")[1]))
        assert unverified_ids and not unverified_ids & {r["requirement_id"] for r in rows}


# ---------------------------------------------------------------------------
# 2. Normalization — omitted, unclear, unknown, malformed, duplicate, order
# ---------------------------------------------------------------------------


class TestNormalization:
    def test_empty_coverage_with_controlling_requirements_is_incomplete(self, client):
        client["install"](_single(_response([])))
        result = run_compliance_check([_spec("PART 1")], _profile(), [], cycle=_cycle())
        assert result.cross_check_status == "completed"  # execution is separate
        completeness = result.coverage_completeness
        assert completeness.state == STATE_INCOMPLETE and not completeness.complete
        assert completeness.expected_count == 2
        assert completeness.returned_count == 0
        assert completeness.omitted_ids == (A, B) and completeness.omitted_count == 2
        rows = _rows(result)
        assert list(rows) == [A, B]
        for row in rows.values():
            assert row["status"] == "unclear"  # never represented or missing
            assert row["origin"] == "synthetic"
            assert row["assessment"] == "none"
            assert "returned no coverage row" in row["reason"]

    def test_an_omitted_row_is_distinguishable_from_a_returned_unclear(self, client):
        client["install"](_single(_response([_row(A, "unclear")])))
        result = run_compliance_check([_spec("PART 1")], _profile(), [], cycle=_cycle())
        returned, omitted = _rows(result)[A], _rows(result)[B]
        assert returned["status"] == omitted["status"] == "unclear"
        assert (returned["origin"], omitted["origin"]) == ("model", "synthetic")
        assert (returned["assessment"], omitted["assessment"]) == ("full", "none")
        assert returned["reason"] is None and omitted["reason"]
        assert result.coverage_completeness.returned_ids == (A,)
        assert result.coverage_completeness.omitted_ids == (B,)

    def test_unknown_advisory_and_malformed_rows_are_ignored_and_counted(self, client):
        client["install"](_single(_response([
            _row(A, "represented", "q", "a.docx"),
            _row("r-ffffffffffff", "missing"),  # unknown id
            _row(U, "missing"),  # unverified: advisory, never a row
            _row(P, "missing"),  # process advisory: never a row
            {"status": "missing"},  # no id
            "not an object",
            _row(B, "represented", "q2", "b.docx"),
        ])))
        result = run_compliance_check([_spec("PART 1")], _profile(), [], cycle=_cycle())
        assert list(_rows(result)) == [A, B]
        assert result.coverage_completeness.ignored_row_count == 5
        assert result.coverage_completeness.state == STATE_COMPLETE

    def test_a_coverage_value_that_is_not_a_list_is_one_ignored_row(self):
        assert cc._normalize_coverage({"requirement_id": A}, expected_ids=[A]) == ([], 1)
        assert cc._normalize_coverage(None, expected_ids=[A]) == ([], 0)

    def test_duplicate_rows_merge_by_precedence_and_keep_every_location(self, client):
        client["install"](_single(_response([
            _row(A, "represented", "found it", "a.docx"),
            _row(A, "contradicted", "wrong edition", "b.docx"),
            _row(A, "represented", "found it", "a.docx"),  # exact repeat
            _row(B, "missing"),
            _row(B, "missing"),
        ])))
        result = run_compliance_check([_spec("PART 1")], _profile(), [], cycle=_cycle())
        a = _rows(result)[A]
        assert (a["status"], a["fileName"]) == ("contradicted", "b.docx")
        assert a["also_reported"] == [
            {"status": "represented", "evidence": "found it", "fileName": "a.docx"}
        ]
        assert _rows(result)[B]["status"] == "missing"

    def test_reordered_rows_produce_identical_output(self, client):
        rows = [
            _row(B, "represented", "q2", "b.docx"),
            _row(A, "contradicted", "x", "c.docx"),
            _row(A, "represented", "q", "a.docx"),
        ]
        client["install"](_single(_response(rows)))
        first = run_compliance_check([_spec("PART 1")], _profile(), [], cycle=_cycle())
        client["install"](_single(_response(list(reversed(rows)))))
        second = run_compliance_check([_spec("PART 1")], _profile(), [], cycle=_cycle())
        assert [r["requirement_id"] for r in first.coverage] == [A, B]  # profile order
        assert first.coverage == second.coverage
        assert first.coverage_completeness == second.coverage_completeness

    def test_an_invented_status_reads_as_the_models_unclear_with_a_note(self, client):
        client["install"](_single(_response([_row(A, "PROBABLY"), _row(B, " Missing ")])))
        result = run_compliance_check([_spec("PART 1")], _profile(), [], cycle=_cycle())
        a, b = _rows(result)[A], _rows(result)[B]
        assert (a["status"], a["origin"]) == ("unclear", "model")
        assert "'PROBABLY' is not a coverage status" in a["reason"]
        assert b["status"] == "missing" and b["reason"] is None  # case/space forgiven


# ---------------------------------------------------------------------------
# 3. Zero controlling requirements
# ---------------------------------------------------------------------------


class TestNoApplicableItems:
    PROFILE_ITEMS = [
        _item(U, "Unverified.", grounded=False),
        _item(P, "Fee.", actionability="process_advisory"),
    ]

    def test_is_a_complete_result_without_an_api_call(self, client):
        client["install"](_single(_response([])))
        profile = _profile(self.PROFILE_ITEMS)
        for run in (run_compliance_check, run_chunked_compliance_check):
            result = run([_spec("PART 1")], profile, [], cycle=_cycle())
            assert result.cross_check_status == "completed"
            assert result.coverage_completeness.state == STATE_NO_APPLICABLE_ITEMS
            assert result.coverage_completeness.complete
            assert "No controlling requirements to evaluate" in result.thinking
            assert "1 process advisory is not specification content" in result.thinking
        assert client["client"].messages.calls == []

    def test_reports_say_no_applicable_requirements_not_skipped(self, client, tmp_path):
        client["install"](_single(_response([])))
        profile = _profile(self.PROFILE_ITEMS)
        result = run_compliance_check([_spec("PART 1")], profile, [], cycle=_cycle())
        pipeline = _pipeline_result(result, profile)
        out = tmp_path / "r.docx"
        export_report(pipeline, out)
        text = _doc_text(out)
        assert "no applicable requirements (none grounded)" in text
        assert "No controlling requirements to evaluate" in text
        assert "were NOT evaluated" not in text  # the old red "skipped" hint
        html = render_html_report(pipeline, include_chat=False)
        assert "no applicable requirements (none grounded)" in html
        assert "were NOT evaluated" not in html
        state = _summarize_run_diagnostics(
            findings=[], status_counts={}, edit_action_counts={},
            cross_check_result=None, pipeline_result=pipeline,
            compliance_result=result,
        )["compliance"]
        assert _compliance_banner_row(state) == (
            "no applicable requirements (none grounded)", False
        )
        assert _compliance_hints(state) == []


# ---------------------------------------------------------------------------
# 4. Chunk semantics
# ---------------------------------------------------------------------------


class TestChunkSemantics:
    def test_a_failed_chunk_prevents_global_absence_and_keeps_findings(
        self, client, chunked
    ):
        client["install"](_by_marker({
            "DIV21SPEC": _response(
                [_row(A, "missing"), _row(B, "represented", "q", "21 13 13 Wet.docx")],
                [_add(A), _edit()],
            ),
            "DIV22SPEC": RuntimeError("chunk exploded"),
        }))
        result = run_chunked_compliance_check(
            _chunked_specs(), _profile(), [], cycle=_cycle()
        )
        # Execution status and tally are unchanged by completeness.
        assert result.cross_check_status == "completed"
        assert (result.chunk_failures, result.chunk_skips) == (1, 0)
        # Missing where assessed is not missing from the package.
        a = _rows(result)[A]
        assert (a["status"], a["origin"], a["assessment"]) == (
            "unclear", "synthetic", "partial"
        )
        assert "22 11 13 Water.docx" in a["reason"]
        assert "not established" in a["reason"]
        completeness = result.coverage_completeness
        assert completeness.unassessed_specs == ("22 11 13 Water.docx", "22 11 16 Piping.docx")
        assert completeness.state == STATE_INCOMPLETE
        # The completed chunk's findings survive; the ADD is held, not dropped.
        actions = [(f.actionType, is_held_addition(f)) for f in result.findings]
        assert actions == [("REPORT_ONLY", True), ("EDIT", False)]
        assert completeness.held_addition_count == 1

    def test_a_not_analyzed_chunk_is_unassessed_scope_too(self, client, chunked):
        specs = _chunked_specs()
        specs[3] = _spec(_PAD * 3 + " DIV22SPEC huge", "22 11 16 Piping.docx")
        client["install"](_by_marker({
            "DIV21SPEC": _response([_row(A, "missing"), _row(B, "missing")], [_add(A)]),
            "DIV22SPEC": _response([_row(A, "missing"), _row(B, "missing")]),
        }))
        result = run_chunked_compliance_check(specs, _profile(), [], cycle=_cycle())
        assert result.chunk_skips == 1
        assert "22 11 16 Piping.docx" in result.coverage_completeness.unassessed_specs
        assert all(row["status"] == "unclear" for row in result.coverage)
        assert is_held_addition(result.findings[0])

    def test_a_chunk_that_omits_the_row_leaves_absence_unestablished(
        self, client, chunked
    ):
        client["install"](_by_marker({
            "DIV21SPEC": _response([_row(A, "missing"), _row(B, "missing")], [_add(A)]),
            "DIV22SPEC": _response([_row(B, "missing")]),  # A left out
        }))
        result = run_chunked_compliance_check(
            _chunked_specs(), _profile(), [], cycle=_cycle()
        )
        assert result.coverage_completeness.partially_assessed_ids == (A,)
        assert result.coverage_completeness.unassessed_specs == ()
        assert _rows(result)[A]["origin"] == "synthetic"
        assert "no coverage row was returned for it by Division 22" in _rows(result)[A]["reason"]
        assert _rows(result)[B]["status"] == "missing"  # B is established
        assert is_held_addition(result.findings[0])

    def test_an_unclear_chunk_holds_the_other_chunks_addition(self, client, chunked):
        # Master dropped this ADD silently ("not missing" read as disproven).
        client["install"](_by_marker({
            "DIV21SPEC": _response([_row(A, "missing"), _row(B, "missing")], [_add(A)]),
            "DIV22SPEC": _response([_row(A, "unclear"), _row(B, "missing")]),
        }))
        result = run_chunked_compliance_check(
            _chunked_specs(), _profile(), [], cycle=_cycle()
        )
        assert len(result.findings) == 1
        assert is_held_addition(result.findings[0])
        assert "classified it unclear" in result.findings[0].demotion_reason

    def test_unanimous_missing_across_every_chunk_keeps_the_edit(self, client, chunked):
        # Control: a silenced detector must not pass as the fix.
        client["install"](_by_marker({
            "DIV21SPEC": _response([_row(A, "missing"), _row(B, "missing")], [_add(A)]),
            "DIV22SPEC": _response(
                [_row(A, "missing"), _row(B, "missing")],
                [_add(A, file_name="22 11 13 Water.docx")],
            ),
        }))
        result = run_chunked_compliance_check(
            _chunked_specs(), _profile(), [], cycle=_cycle()
        )
        assert result.coverage_completeness.state == STATE_COMPLETE
        assert _rows(result)[A] == {
            "requirement_id": A, "status": "missing", "evidence": None,
            "fileName": None, "origin": "model", "assessment": "full",
            "reason": None, "also_reported": [],
        }
        # One executable ADD per requirement (the D-7 dedup), unchanged.
        assert [f.actionType for f in result.findings] == ["ADD"]
        assert result.findings[0].as_edit_proposal() is not None

    def test_found_elsewhere_still_drops_the_chunk_local_addition(self, client, chunked):
        client["install"](_by_marker({
            "DIV21SPEC": _response([_row(A, "missing"), _row(B, "missing")], [_add(A)]),
            "DIV22SPEC": _response(
                [_row(A, "represented", "q", "22 11 13 Water.docx"), _row(B, "missing")]
            ),
        }))
        result = run_chunked_compliance_check(
            _chunked_specs(), _profile(), [], cycle=_cycle()
        )
        assert _rows(result)[A]["status"] == "represented"
        assert result.findings == []

    def test_excluded_specs_reach_the_chunked_merge(self, client, chunked):
        client["install"](_by_marker({
            "DIV21SPEC": _response([_row(A, "missing"), _row(B, "missing")], [_add(A)]),
            "DIV22SPEC": _response([_row(A, "missing"), _row(B, "missing")]),
        }))
        result = run_chunked_compliance_check(
            _chunked_specs(), _profile(), [], cycle=_cycle(),
            excluded_specs=["23 00 00 Failed Review.docx"],
        )
        assert result.coverage_completeness.unassessed_specs == (
            "23 00 00 Failed Review.docx",
        )
        assert all(row["status"] == "unclear" for row in result.coverage)
        assert is_held_addition(result.findings[0])

    def test_every_chunk_failing_is_failed_with_nothing_assessed(self, client, chunked):
        client["install"](lambda kwargs: RuntimeError("down"))
        result = run_chunked_compliance_check(
            _chunked_specs(), _profile(), [], cycle=_cycle()
        )
        assert result.cross_check_status == "failed"
        assert result.coverage == []  # the status says it; no fabricated rows
        completeness = result.coverage_completeness
        assert completeness.omitted_ids == (A, B)
        assert len(completeness.unassessed_specs) == 4
        # The record explains itself with the pass's error, not the summary.
        assert completeness.reason == result.error and "down" in result.error

    def test_statuses_stay_the_closed_execution_set(self, client, chunked):
        seen = set()
        for route in (
            _single(_response([])),
            lambda kwargs: RuntimeError("down"),
            _by_marker({"DIV21SPEC": _response([]), "DIV22SPEC": RuntimeError("x")}),
        ):
            client["install"](route)
            seen.add(run_chunked_compliance_check(
                _chunked_specs(), _profile(), [], cycle=_cycle()
            ).cross_check_status)
            seen.add(run_compliance_check(
                [_spec("PART 1")], _profile(), [], cycle=_cycle()
            ).cross_check_status)
        assert seen <= {"completed", "failed", "skipped"}

    def test_the_finalize_hook_treats_a_chunk_result_as_rows_only(self):
        finalize = cc.coverage_finalizer(_profile())
        combined = ReviewResult(cross_check_status="completed")
        finalize(combined, [
            ChunkOutcome("div_21", "Division 21", ("a.docx",),
                         ReviewResult(cross_check_status="completed",
                                      coverage=[_row(A, "represented", "q", "a.docx")])),
            ChunkOutcome("div_22", "Division 22", ("b.docx",),
                         ReviewResult(cross_check_status="failed",
                                      coverage=[_row(B, "represented", "q", "b.docx")])),
        ])
        # A failed chunk's leftover rows are never trusted.
        assert _rows(combined)[B]["origin"] == "synthetic"
        assert combined.coverage_completeness.unassessed_specs == ("b.docx",)


# ---------------------------------------------------------------------------
# 5. Additions whose absence is not established
# ---------------------------------------------------------------------------


class TestHeldAdditions:
    def test_an_addition_for_an_omitted_requirement_is_held_with_its_text(self, client):
        client["install"](_single(_response([_row(A, "represented", "q", "a")], [_add(B)])))
        result = run_compliance_check([_spec("PART 1")], _profile(), [], cycle=_cycle())
        (finding,) = result.findings
        assert finding.actionType == "REPORT_ONLY"
        assert finding.as_edit_proposal() is None
        assert is_held_addition(finding)
        reason = finding.demotion_reason
        assert B in reason and "returned no coverage row" in reason
        # The would-be insertion stays visible, as a conditional.
        assert f'"C. Comply with requirement {B}."' in reason
        assert 'after "PART 1 - GENERAL"' in reason
        assert result.coverage_completeness.held_addition_count == 1

    def test_an_established_missing_requirement_keeps_its_edit(self, client):
        client["install"](_single(_response(
            [_row(A, "represented", "q", "a"), _row(B, "missing")], [_add(B)]
        )))
        result = run_compliance_check([_spec("PART 1")], _profile(), [], cycle=_cycle())
        assert [f.actionType for f in result.findings] == ["ADD"]
        assert result.coverage_completeness.held_addition_count == 0

    @pytest.mark.parametrize(
        ("rid", "kind"),
        [(U, "not independently verified"), (P, "a process advisory")],
    )
    def test_an_addition_resting_only_on_a_non_controlling_item_is_held(
        self, client, rid, kind
    ):
        client["install"](_single(_response(
            [_row(A, "represented", "q", "a"), _row(B, "represented", "q", "a")],
            [_add(rid)],
        )))
        result = run_compliance_check([_spec("PART 1")], _profile(), [], cycle=_cycle())
        (finding,) = result.findings
        assert is_held_addition(finding)
        assert f"{rid} ({kind})" in finding.demotion_reason

    def test_a_single_pass_self_contradiction_is_held_not_dropped(self, client):
        client["install"](_single(_response(
            [_row(A, "represented", "q", "a"), _row(B, "represented", "q", "a")],
            [_add(B)],
        )))
        result = run_compliance_check([_spec("PART 1")], _profile(), [], cycle=_cycle())
        (finding,) = result.findings
        assert is_held_addition(finding)
        assert "classified it represented" in finding.demotion_reason

    def test_non_additions_are_never_touched(self, client):
        client["install"](_single(_response([], [_edit()])))
        result = run_compliance_check([_spec("PART 1")], _profile(), [], cycle=_cycle())
        assert [f.actionType for f in result.findings] == ["EDIT"]

    def test_a_held_addition_never_reaches_the_edit_sidecar(self, client):
        client["install"](_single(_response(
            [_row(A, "represented", "q", "a")], [_add(B), _edit()]
        )))
        result = run_compliance_check([_spec("PART 1")], _profile(), [], cycle=_cycle())
        sidecar = build_edit_instructions(_pipeline_result(result))
        actions = [e["edit_proposal"]["action_type"] for e in sidecar["edits"]]
        assert actions == ["EDIT"]

    def test_the_reports_render_a_hold_as_conditional(self, client, tmp_path):
        client["install"](_single(_response([_row(A, "represented", "q", "a")], [_add(B)])))
        result = run_compliance_check([_spec("PART 1")], _profile(), [], cycle=_cycle())
        pipeline = _pipeline_result(result)
        out = tmp_path / "r.docx"
        export_report(pipeline, out)
        text = _doc_text(out)
        assert "Addition held as REPORT_ONLY and not emitted as an edit." in text
        assert "demoted to REPORT_ONLY at parse time: Absence not established" not in text
        html = render_html_report(pipeline, include_chat=False)
        assert "Addition held as REPORT_ONLY and not emitted as an edit." in html


# ---------------------------------------------------------------------------
# 6. Every output path carries the same semantics
# ---------------------------------------------------------------------------


def _incomplete_result(client):
    client["install"](_single(_response([_row(A, "represented", "q", "a.docx")], [_add(B)])))
    return run_compliance_check([_spec("PART 1")], _profile(), [], cycle=_cycle())


class TestReportSurfaces:
    def test_word_report_shows_the_notice_matrix_and_banner(self, client, tmp_path):
        pipeline = _pipeline_result(_incomplete_result(client))
        out = tmp_path / "r.docx"
        export_report(pipeline, out)
        text = _doc_text(out)
        # Banner row (the one finding is the held addition), highlighted, and
        # the notice (banner + coverage section).
        assert "1 finding — 0 missing / 0 contradicted — coverage incomplete" in text
        assert text.count("⚠ Compliance coverage is incomplete: 1 of 2 controlling "
                          "requirements was not assessed") == 2
        assert "1 addition that depended on a requirement's absence is shown as report-only" in text
        # The matrix names the synthetic row for what it is — in the table
        # cell itself, not only in the notice's prose.
        assert _coverage_cells(out) == {A: "Represented", B: "NOT ASSESSED"}
        # And the section never calls an incomplete pass clean.
        assert "no missing or contradicted requirements found.\n" not in text + "\n"
        assert "Coverage is incomplete" in text

    def test_html_report_matches_the_word_report(self, client):
        pipeline = _pipeline_result(_incomplete_result(client))
        html = render_html_report(pipeline, include_chat=True)
        notice = "⚠ Compliance coverage is incomplete: 1 of 2 controlling "
        digest = _PLAINTEXT_RE.search(html).group(1)
        rendered = _PLAINTEXT_RE.sub("", html)
        # Banner + coverage section, as in Word — and the chat's plain-text
        # digest carries both, so Ask AI sees the report's own caveat.
        assert rendered.count(notice) == 2
        assert digest.count(notice) == 2
        assert "coverage incomplete" in rendered
        assert dict(_HTML_COVERAGE_ROW_RE.findall(rendered)) == {
            A: "Represented", B: "NOT ASSESSED"
        }
        completeness = _html_payload(html)["compliance"]["completeness"]
        assert completeness["state"] == "incomplete"
        assert completeness["omitted_ids"] == [B]
        assert completeness["complete"] is False

    def test_a_result_without_a_record_is_never_read_as_complete(self, tmp_path):
        legacy = ReviewResult(
            cross_check_status="completed",
            coverage=[_row(A, "represented", "q", "a.docx"), _row(B, "represented")],
        )
        pipeline = _pipeline_result(legacy)
        state = _summarize_run_diagnostics(
            findings=[], status_counts={}, edit_action_counts={},
            cross_check_result=None, pipeline_result=pipeline, compliance_result=legacy,
        )["compliance"]
        value, highlight = _compliance_banner_row(state)
        assert value.endswith("coverage completeness not recorded") and highlight
        out = tmp_path / "r.docx"
        export_report(pipeline, out)
        assert "carries no coverage completeness record" in _doc_text(out)
        assert build_edit_instructions(pipeline)["requirements_coverage_completeness"] is None

    def test_a_skipped_pass_shows_its_skip_hint_not_a_second_notice(self):
        skipped = ReviewResult(
            cross_check_status="skipped", thinking="too large",
            coverage_completeness=nothing_assessed(
                (A, B), unassessed_specs=("a.docx",), reason="too large"
            ),
        )
        state = _summarize_run_diagnostics(
            findings=[], status_counts={}, edit_action_counts={},
            cross_check_result=None, pipeline_result=_pipeline_result(skipped),
            compliance_result=skipped,
        )["compliance"]
        assert _compliance_banner_row(state) == ("skipped", True)
        hints = _compliance_hints(state)
        assert len(hints) == 1 and "were NOT evaluated" in hints[0][0]

    def test_the_evidence_cell_keeps_every_reported_location(self):
        from src.output.report_exporter import _coverage_evidence_parts

        parts = _coverage_evidence_parts({
            "requirement_id": A, "status": "contradicted", "evidence": "wrong",
            "fileName": "b.docx", "origin": "model", "reason": None,
            "also_reported": [
                {"status": "represented", "evidence": "right", "fileName": "a.docx"}
            ],
        })
        assert parts == ["wrong — b.docx", "Also reported represented: right — a.docx"]
        synthetic = _coverage_evidence_parts({
            "requirement_id": A, "status": "unclear", "evidence": "ignored",
            "fileName": "ignored.docx", "origin": "synthetic", "reason": "Why.",
            "also_reported": [],
        })
        assert synthetic == ["Why."]

    def test_a_complete_pass_renders_exactly_as_before(self, client, tmp_path):
        client["install"](_single(_response(
            [_row(A, "represented", "q", "a.docx"), _row(B, "missing")], [_add(B)]
        )))
        result = run_compliance_check([_spec("PART 1")], _profile(), [], cycle=_cycle())
        pipeline = _pipeline_result(result)
        out = tmp_path / "r.docx"
        export_report(pipeline, out)
        text = _doc_text(out)
        assert "coverage incomplete" not in text
        assert "NOT ASSESSED" not in text
        assert "1 finding — 1 missing / 0 contradicted" in text
        assert "q — a.docx" in text  # evidence cell unchanged for a model row


class TestJsonAndDiagnostics:
    def test_sidecar_and_profile_export_carry_the_record(self, client):
        pipeline = _pipeline_result(_incomplete_result(client))
        sidecar = build_edit_instructions(pipeline)
        # The record is additive (plan WP-09): it moved no schema number.
        assert sidecar["schema_version"] == SIDECAR_SCHEMA_VERSION
        record = sidecar["requirements_coverage_completeness"]
        assert record["state"] == "incomplete" and record["omitted_ids"] == [B]
        assert {r["requirement_id"]: r["origin"] for r in sidecar["requirements_coverage"]} == {
            A: "model", B: "synthetic"
        }
        exported = build_requirements_profile_export(pipeline)
        assert exported["requirements_coverage_completeness"] == record

    def test_diagnostics_extra_and_run_log_name_the_gap(self, client):
        result = _incomplete_result(client)
        extra = compliance_pass_extra(result)
        assert extra["coverage_completeness"]["state"] == "incomplete"
        assert extra["coverage_count"] == 2
        lines: list[tuple[str, str]] = []
        pl._log_compliance_status(lambda msg, level="info", **_: lines.append((level, msg)), result)
        warnings = [msg for level, msg in lines if level == "warning"]
        assert len(warnings) == 1
        assert warnings[0].startswith("Compliance coverage INCOMPLETE — 1 of 2 controlling")

    def test_a_complete_pass_logs_no_warning(self, client):
        client["install"](_single(_response(
            [_row(A, "represented", "q", "a"), _row(B, "represented", "q", "b")]
        )))
        result = run_compliance_check([_spec("PART 1")], _profile(), [], cycle=_cycle())
        lines: list[str] = []
        pl._log_compliance_status(lambda msg, level="info", **_: lines.append(level), result)
        assert lines == ["success"]


# ---------------------------------------------------------------------------
# Pipeline stage + headless driver (the CLI / recovery path)
# ---------------------------------------------------------------------------


def _submission(module, specs, *, requirements_profile) -> BatchSubmission:
    request_map = {
        f"review__{i}": {"filename": spec.filename, "index": i, "type": "review"}
        for i, spec in enumerate(specs)
    }
    return BatchSubmission(
        job=BatchJob(batch_id="msgbatch_WP09", job_type="review",
                     request_map=request_map, created_at=1_700_000_000.0),
        files_reviewed=[spec.filename for spec in specs],
        review_request_ids=list(request_map),
        model="claude-opus-4-8",
        project_context="Hyperscale data-center program.",
        prepared_specs=list(specs),
        cycle_label=module.cycle.label,
        module_id=module.module_id,
        project_profile={"city": "Ashburn", "state_or_province": "VA",
                         "country": "US", "client_name": "ExampleCo"},
        requirements_profile=requirements_profile.to_dict()
        if requirements_profile is not None else None,
    )


class TestPipelineStage:
    def _state(self, specs, *, profile=_profile(), truncated=()):
        return CollectedBatchState(
            submission=_submission(DATACENTER_FIRE, specs, requirements_profile=profile),
            review_result=ReviewResult(findings=[]),
            truncated_specs=list(truncated),
        )

    def test_a_missing_profile_is_unknown_not_no_applicable_items(self):
        state = self._state([_spec("PART 1")], profile=None)
        run_compliance_for_batch(state)
        completeness = state.compliance_result.coverage_completeness
        assert state.compliance_result.cross_check_status == "skipped"
        assert completeness.state == STATE_UNKNOWN and not completeness.complete

    def test_unavailable_content_names_the_unassessed_specs(self):
        # Codex review on PR #381: a recovery with a saved profile but no
        # re-extracted content must still say which specs went unassessed.
        state = self._state(
            [_spec("x", "21 13 13 Wet.docx"), _spec("y", "21 13 16 Dry.docx")]
        )
        state.submission.prepared_specs = []
        run_compliance_for_batch(state)
        result = state.compliance_result
        assert result.cross_check_status == "skipped"
        completeness = result.coverage_completeness
        assert completeness.omitted_ids == (A, B)
        assert completeness.unassessed_specs == ("21 13 13 Wet.docx", "21 13 16 Dry.docx")
        record = build_edit_instructions(PipelineResult(
            review_result=ReviewResult(findings=[]), compliance_result=result,
        ))["requirements_coverage_completeness"]
        assert record["unassessed_specs"] == ["21 13 13 Wet.docx", "21 13 16 Dry.docx"]

    def test_a_missing_profile_names_the_unassessed_specs_too(self):
        state = self._state([_spec("x", "21 13 13 Wet.docx")], profile=None)
        run_compliance_for_batch(state)
        completeness = state.compliance_result.coverage_completeness
        assert completeness.state == STATE_UNKNOWN
        assert completeness.unassessed_specs == ("21 13 13 Wet.docx",)

    def test_a_spec_whose_review_failed_is_unassessed_scope(self, client):
        client["install"](_single(_response(
            [_row(A, "missing"), _row(B, "represented", "q", "21 13 13 Wet.docx")],
            [_add(A)],
        )))
        state = self._state(
            [_spec("PART 1 - GENERAL", "21 13 13 Wet.docx"),
             _spec("PART 1 - GENERAL", "21 13 16 Dry.docx")],
            truncated=["21 13 16 Dry.docx"],
        )
        run_compliance_for_batch(state)
        result = state.compliance_result
        assert result.coverage_completeness.unassessed_specs == ("21 13 16 Dry.docx",)
        assert _rows(result)[A]["status"] == "unclear"
        assert "21 13 16 Dry.docx" in _rows(result)[A]["reason"]
        assert is_held_addition(result.findings[0])
        # The excluded spec never reached the request.
        message = client["client"].messages.calls[0]["messages"][0]["content"]
        assert "21 13 16 Dry.docx" not in message

    def test_every_spec_failing_review_names_them(self):
        state = self._state([_spec("x", "21 13 13 Wet.docx")], truncated=["21 13 13 Wet.docx"])
        run_compliance_for_batch(state)
        completeness = state.compliance_result.coverage_completeness
        assert completeness.unassessed_specs == ("21 13 13 Wet.docx",)
        assert completeness.omitted_ids == (A, B)

    def test_a_stand_in_result_is_stamped_before_it_is_stored(self, monkeypatch):
        import src.compliance as compliance_pkg

        stand_in = ReviewResult(
            cross_check_status="completed",
            coverage=[_row(A, "missing")],
            findings=[Finding(
                severity="HIGH", fileName="21 13 13 Wet.docx", section="1.02",
                issue=f"Requirement {A} is missing.", actionType="ADD",
                existingText=None, replacementText="C. Add A.", codeReference=None,
                anchorText="PART 1", insertPosition="after",
            )],
        )
        monkeypatch.setattr(
            compliance_pkg, "run_chunked_compliance_check", lambda *a, **k: stand_in
        )
        state = self._state([_spec("PART 1", "21 13 13 Wet.docx")])
        run_compliance_for_batch(state)
        completeness = state.compliance_result.coverage_completeness
        assert completeness is not None and completeness.omitted_ids == (B,)
        assert state.compliance_result.findings[0].actionType == "ADD"  # A is missing

    def test_ensure_is_a_no_op_when_a_record_exists(self):
        record = CoverageCompleteness(expected_ids=(A,), returned_ids=(A,))
        result = ReviewResult(cross_check_status="completed", coverage_completeness=record)
        assert ensure_coverage_completeness(result, _profile()).coverage_completeness is record


class TestIncompleteAnalysisScenario:
    """Plan §28 scenario B, through the headless driver the CLI and recovery use.

    Mocked request counts force chunking. One chunk fails, another omits a
    controlling coverage row, and another returns useful findings. The final
    report keeps the useful findings, shows incomplete coverage, and emits no
    unsupported global-absence addition.
    """

    SPECS = [
        _spec(_PAD + " DIV21SPEC a\n\nPART 1 - GENERAL", "21 13 13 Wet.docx"),
        _spec(_PAD + " DIV21SPEC b", "21 13 16 Dry.docx"),
        _spec(_PAD + " DIV28SPEC a", "28 31 00 Alarm.docx"),
        _spec(_PAD + " DIV28SPEC b", "28 46 00 Detection.docx"),
        _spec(_PAD + " DIV22SPEC a\n\nComply with the 2015 IBC.", "22 11 13 Water.docx"),
        _spec(_PAD + " DIV22SPEC b", "22 11 16 Piping.docx"),
    ]

    def _run(self, monkeypatch, client):
        monkeypatch.setattr(cc, "COMPLIANCE_RECOMMENDED_MAX", 5_000)
        client["install"](_by_marker({
            "DIV28SPEC": RuntimeError("chunk exploded"),
            # Division 21 leaves B out and wants A added everywhere.
            "DIV21SPEC": _response([_row(A, "missing")], [_add(A)]),
            # Division 22 returns every row and a useful edit.
            "DIV22SPEC": _response(
                [_row(A, "missing"), _row(B, "represented", "q", "22 11 13 Water.docx")],
                [_edit(file_name="22 11 13 Water.docx")],
            ),
        }))
        monkeypatch.setattr(pl, "retrieve_review_results", lambda job, *, model: {
            cid: ReviewResult(findings=[], parse_status="ok") for cid in job.request_map
        })
        monkeypatch.setattr(pl, "start_batch_verification", lambda findings, **kw: None)
        diagnostics = DiagnosticsReport()
        submission = _submission(
            DATACENTER_FIRE, self.SPECS, requirements_profile=_profile()
        )
        self.log: list[tuple[str, str]] = []
        result = run_batch_collection_headless(
            submission, cache=VerificationCache(),
            log=lambda msg, level="info", **_: self.log.append((level, msg)),
            diagnostics=diagnostics,
        )
        return result, diagnostics

    def test_useful_findings_survive_and_no_absence_edit_is_emitted(
        self, monkeypatch, client, tmp_path
    ):
        result, diagnostics = self._run(monkeypatch, client)
        comp = result.compliance_result
        assert comp.cross_check_status == "completed" and comp.chunk_failures == 1
        by_action = {f.actionType: f for f in comp.findings}
        assert set(by_action) == {"EDIT", "REPORT_ONLY"}
        assert is_held_addition(by_action["REPORT_ONLY"])
        completeness = comp.coverage_completeness
        assert completeness.state == STATE_INCOMPLETE
        assert completeness.partially_assessed_ids == (B,)
        assert completeness.unassessed_specs == ("28 31 00 Alarm.docx", "28 46 00 Detection.docx")

        sidecar = build_edit_instructions(result)
        assert [e["edit_proposal"]["action_type"] for e in sidecar["edits"]] == ["EDIT"]
        assert sidecar["requirements_coverage_completeness"]["state"] == "incomplete"

        out = tmp_path / "r.docx"
        export_report(result, out)
        text = _doc_text(out)
        assert "coverage incomplete" in text
        assert "2 specifications were not assessed at all: 28 31 00 Alarm.docx" in text
        assert _coverage_cells(out) == {A: "NOT FULLY ASSESSED", B: "Represented"}

        # The diagnostics event and the run log say it too.
        records = [
            event.data["coverage_completeness"]
            for event in diagnostics.events
            if event.phase == "compliance" and "coverage_completeness" in (event.data or {})
        ]
        assert [record["state"] for record in records] == ["incomplete"]
        assert any(
            level == "warning" and msg.startswith("Compliance coverage INCOMPLETE")
            for level, msg in self.log
        )


class TestProgramReports:
    def test_program_rollup_and_report_carry_the_notice(self, client, tmp_path):
        from src.orchestration.program_pipeline import ProgramPipelineResult
        from src.programs.assignments import SpecAssignment
        from src.programs.models import (
            RoutingEvidence,
            RoutingEvidenceSource,
            RoutingState,
            SpecRoutingDecision,
        )

        def assignment(spec_id, module_id):
            return SpecAssignment(
                source_path=f"/specs/{spec_id}",
                decision=SpecRoutingDecision(
                    spec_id=spec_id, program_id="hyperscale_datacenter",
                    automatic_state=RoutingState.SUPPORTED,
                    automatic_module_ids=(module_id,), confidence=0.9,
                    evidence=(RoutingEvidence(
                        source=RoutingEvidenceSource.CSI_SECTION, signal="csi",
                        detail="prefix", module_id=module_id, weight=0.9,
                    ),),
                ),
            )

        incomplete = _incomplete_result(client)
        client["install"](_single(_response(
            [_row(A, "represented", "q", "a"), _row(B, "represented", "q", "b")]
        )))
        complete = run_compliance_check([_spec("PART 1")], _profile(), [], cycle=_cycle())

        def child(module_id, spec, compliance):
            from src.modules import require_module

            return PipelineResult(
                review_result=ReviewResult(findings=[]),
                files_reviewed=[spec],
                cycle_label=require_module(module_id).cycle.label,
                module_id=module_id,
                requirements_profile=_profile().to_dict(),
                compliance_result=compliance,
            )

        program = ProgramPipelineResult(
            program_id="hyperscale_datacenter",
            assignments=(assignment("21 13 13 Wet.docx", "datacenter_fire"),
                         assignment("26 05 00 Power.docx", "datacenter_electrical")),
            module_results={
                "datacenter_fire": child("datacenter_fire", "21 13 13 Wet.docx", incomplete),
                "datacenter_electrical": child(
                    "datacenter_electrical", "26 05 00 Power.docx", complete
                ),
            },
        )
        merged = program.compliance_result.coverage_completeness
        assert merged.state == STATE_INCOMPLETE and merged.omitted_ids == (B,)
        assert merged.expected_count == 4

        out = tmp_path / "program.docx"
        export_report(program, out)
        text = _doc_text(out)
        assert "coverage incomplete" in text
        assert "1 of 4 controlling requirements was not assessed" in text
        html = render_html_report(program, include_chat=False)
        assert "1 of 4 controlling requirements was not assessed" in html
        sidecar = build_edit_instructions(program)
        by_module = sidecar["requirements_coverage_completeness_by_module"]
        assert by_module["datacenter_fire"]["state"] == "incomplete"
        assert by_module["datacenter_electrical"]["state"] == "complete"

    def test_combine_never_reads_missing_metadata_as_complete(self):
        known = CoverageCompleteness(expected_ids=(A,), returned_ids=(A,))
        assert combine([("fire", known), ("elec", None)]).state == STATE_UNKNOWN
        assert combine([("fire", known)]).state == STATE_COMPLETE
        gap = nothing_assessed((B,), unassessed_specs=("x.docx",))
        merged = combine([("fire", known), ("elec", gap)])
        assert merged.unassessed_specs == ("elec: x.docx",)
        assert combine([]) is None

    def test_rollup_sums_only_completed_modules_gaps(self):
        completed_gap = {
            "status": "completed", "coverage_state": "incomplete",
            "coverage_expected": 3, "coverage_not_assessed": 1,
            "coverage_partial": 0, "unassessed_specs": ["a.docx"], "held_additions": 1,
        }
        skipped = {
            "status": "skipped", "reason": "profile unavailable",
            "coverage_state": "unknown", "coverage_expected": 0,
            "coverage_not_assessed": 0, "coverage_partial": 0,
            "unassessed_specs": [], "held_additions": 0,
        }
        base = {"edit_suggested": 0, "report_only": 0, "failed_review_specs": []}
        agg = _aggregate_run_diagnostics([
            ("Fire", {**base, "compliance": completed_gap}),
            ("Electrical", {**base, "compliance": skipped}),
        ])["compliance"]
        assert agg["status"] == "skipped"
        assert agg["coverage_state"] == "unknown"
        assert agg["unassessed_specs"] == ["Fire: a.docx"]
        tones = [tone for _text, tone in _compliance_hints(agg)]
        texts = " ".join(text for text, _tone in _compliance_hints(agg))
        # The skipped module's hint AND the completed module's gap both show.
        assert tones[:2] == ["red", "red"]
        assert "were NOT evaluated" in texts and "Fire: a.docx" in texts


# ---------------------------------------------------------------------------
# The pure reconcile, directly
# ---------------------------------------------------------------------------


class TestReconcile:
    def _unit(self, label, rows, *, completed=True, files=("a.docx",)):
        return AssessmentUnit(
            label=label, filenames=tuple(files), completed=completed,
            rows=tuple(dict(_row(*row), origin="model", reason=None) for row in rows),
        )

    def test_no_completed_unit_yields_no_rows(self):
        rows, completeness = reconcile(
            [self._unit("A", [], completed=False, files=("x.docx",))],
            expected_ids=[A, B],
        )
        assert rows == []
        assert completeness.omitted_ids == (A, B)
        assert completeness.unassessed_specs == ("x.docx",)

    def test_represented_anywhere_stands_even_with_unassessed_scope(self):
        rows, completeness = reconcile(
            [self._unit("A", [(A, "represented", "q", "a.docx")]),
             self._unit("B", [], completed=False, files=("b.docx",))],
            expected_ids=[A],
        )
        assert (rows[0]["status"], rows[0]["origin"], rows[0]["assessment"]) == (
            "represented", "model", "partial"
        )
        assert completeness.state == STATE_INCOMPLETE

    def test_excluded_specs_leave_missing_unestablished(self):
        rows, completeness = reconcile(
            [self._unit("", [(A, "missing", None, None)])],
            expected_ids=[A],
            excluded_specs=["b.docx", "b.docx", ""],
        )
        assert rows[0]["status"] == "unclear" and rows[0]["origin"] == "synthetic"
        assert completeness.unassessed_specs == ("b.docx",)

    def test_to_dict_is_json_ready(self):
        record = nothing_assessed((A,), unassessed_specs=("x.docx",), reason="failed")
        data = json.loads(json.dumps(record.to_dict()))
        assert data == {
            "state": "incomplete", "complete": False, "expected_count": 1,
            "returned_count": 0, "omitted_count": 1, "expected_ids": [A],
            "omitted_ids": [A], "partially_assessed_ids": [],
            "unassessed_specs": ["x.docx"], "ignored_row_count": 0,
            "held_addition_count": 0, "reason": "failed",
        }
        assert nothing_assessed(None).to_dict()["expected_count"] is None
