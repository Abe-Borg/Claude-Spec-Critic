"""Quick status check for the implementation plan's reproducible defects.

Runs the review's reproductions (mostly the plan's own Appendix A examples)
against the current code and prints one line per check:

    OPEN   the defect is still present
    FIXED  the code now behaves as the plan requires
    ERROR  the check itself broke (a function it uses was renamed or
           reshaped by a fix) -- update or delete that check

It covers the packages whose defects can be reproduced in a few lines. The
others (WP-01, WP-03, WP-08, WP-09, WP-14, WP-15, WP-16) are tracked only in
plans/PROGRESS.md, which is the authority on what is done. This script is a
starting-state snapshot taken on 2026-09-23 (every check OPEN); chunk S01
turns these checks into strict-xfail tests and deletes this file.

Hermetic: no API key, no network. Writes only to temporary directories.

Usage (from anywhere):  python plans/check_plan_status.py
"""
from __future__ import annotations

import os
import sys
import tempfile
import traceback
from pathlib import Path
from types import SimpleNamespace

REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT))
os.environ.setdefault("ANTHROPIC_API_KEY", "sk-placeholder-not-a-real-key")
os.environ["SPEC_CRITIC_VERIFICATION_CACHE_PERSIST"] = "0"
os.environ["SPEC_CRITIC_TRACE"] = "0"

RESULTS: list[tuple[str, str, str]] = []


def check(label: str):
    """Register a check; the function returns (fixed: bool, detail: str)."""

    def decorator(fn):
        try:
            fixed, detail = fn()
            RESULTS.append(("FIXED" if fixed else "OPEN", label, detail))
        except Exception as exc:  # the check needs updating, not the code
            last = traceback.extract_tb(exc.__traceback__)[-1]
            RESULTS.append(
                ("ERROR", label, f"{type(exc).__name__}: {exc} (line {last.lineno})")
            )
        return fn

    return decorator


def _finding(issue: str = "Wrong valve type.", **overrides):
    from src.review.reviewer import Finding

    fields = dict(
        severity="HIGH",
        fileName="210500.docx",
        section="2.01",
        issue=issue,
        actionType="REPORT_ONLY",
        existingText=None,
        replacementText=None,
        codeReference=None,
    )
    fields.update(overrides)
    return Finding(**fields)


CLEAN_THREE_PART = "\n\n".join(
    [
        "PART 1 GENERAL",
        "1.01 SUMMARY",
        "A. Provide the specified piping system.",
        "1.02 SUBMITTALS",
        "A. Submit product data before fabrication.",
        "PART 2 PRODUCTS",
        "2.01 MATERIALS",
        "A. Provide materials meeting the scheduled requirements.",
        "PART 3 EXECUTION",
        "3.01 INSTALLATION",
        "A. Install in accordance with the approved product instructions.",
    ]
)


# ----------------------------------------------------------------- WP-04 ---
@check("WP-04A clean 3-PART spec has no empty-section alerts")
def _():
    from src.input.preprocessor import detect_empty_sections

    alerts = detect_empty_sections(CLEAN_THREE_PART, "clean.docx")
    return not alerts, f"{len(alerts)} alert(s): {[a['match'] for a in alerts]}"


@check("WP-04A quantity lines are not headings")
def _():
    from src.input.preprocessor import detect_duplicate_headings, detect_empty_sections

    text = "\n\n".join(
        ["3.01 PAINTING", "2 coats of primer shall be applied.", "3.02 CLEANING", "A. Clean surfaces."]
    )
    empty = [a["match"] for a in detect_empty_sections(text, "q.docx")]
    dupes = detect_duplicate_headings(
        "\n\n".join(["1.01 SUMMARY", "2 coats of primer.", "A. Text.", "2 coats of primer.", "B. Text."]),
        "q.docx",
    )
    bad = [m for m in empty if m.startswith("2 coats")] + [a.get("match") for a in dupes]
    return not bad, f"misread as headings: {bad}"


def _stale(sentence: str) -> int:
    from src.core.code_cycles import CALIFORNIA_2025
    from src.input.preprocessor import detect_stale_code_cycle_references

    return len(detect_stale_code_cycle_references(sentence, "s.docx", CALIFORNIA_2025))


for _sentence in (
    "Submit shop drawings prior to fabrication in accordance with 2022 CBC Section 1704.",
    "Coordinate with the historical society and comply with 2022 CBC.",
    "Contractor may not deviate from 2022 CBC Chapter 17.",
):

    @check(f"WP-04B flags: {_sentence[:52]}...")
    def _(sentence=_sentence):
        n = _stale(sentence)
        return n > 0, f"{n} alert(s)"


for _token in ("ASCE/SEI 7-16", "ASCE 7–16", "ASCE 7-2016"):

    @check(f"WP-04C recognizes {_token!r}")
    def _(token=_token):
        n = _stale(f"Design loads per {token}.")
        return n > 0, f"{n} alert(s)"


@check("WP-04D bare TBD is a placeholder")
def _():
    from src.input.preprocessor import detect_placeholders

    n = len(detect_placeholders("Pipe size: TBD by engineer.", "p.docx"))
    return n > 0, f"{n} alert(s)"


@check("WP-04D [EDITION]/[SELECTED] are not EDIT/SELECT placeholders")
def _():
    from src.input.preprocessor import detect_placeholders

    alerts = detect_placeholders("See [EDITION 2024] and [SELECTED ITEMS].", "p.docx")
    return not alerts, f"false alerts: {[a.get('type') for a in alerts]}"


@check("WP-04E a mix of naming styles is reported")
def _():
    from src.input.preprocessor import detect_inconsistent_file_naming

    alerts = detect_inconsistent_file_naming(
        ["21 05 00.docx", "211313.docx", "SECTION 21 13 16.DOCX"]
    )
    return bool(alerts), f"{len(alerts)} alert(s)"


# ----------------------------------------------------------------- WP-05 ---
for _name, _heading in (
    ("210500.docx", "SECTION 21 05 00\n\nCOMMON WORK RESULTS FOR FIRE SUPPRESSION"),
    ("211313.docx", "SECTION 21 13 13\n\nWET-PIPE SPRINKLER SYSTEMS"),
):

    @check(f"WP-05 {_name} + its SECTION heading routes to fire suppression")
    def _(name=_name, heading=_heading):
        from src.programs.catalog import get_program
        from src.programs.models import SpecRoutingInput
        from src.programs.routing import route_spec

        decision = route_spec(
            SpecRoutingInput(spec_id=name, section_title=name, content=heading + "\n\nPART 1 GENERAL"),
            program=get_program("hyperscale_datacenter"),
        )
        state = getattr(decision.automatic_state, "value", decision.automatic_state)
        modules = decision.automatic_module_ids
        return (
            state == "supported" and modules == ("datacenter_fire",),
            f"state={state}, modules={modules}",
        )


# ---------------------------------------------------------------- WP-06A ---
@check("WP-06A copper and PVC findings stay distinct")
def _():
    from src.orchestration import pipeline

    copper = _finding("Section 21 05 00 requires copper pipe in 210500.docx.")
    pvc = _finding("Section 21 05 00 requires PVC pipe in 210500.docx.")
    survivors = pipeline._deduplicate_findings([copper, pvc])
    return len(survivors) == 2, f"{len(survivors)} finding(s) survive dedup"


# ---------------------------------------------------------------- WP-06B ---
@check("WP-06B same edit at p4 and p8 gives two sidecar entries")
def _():
    from src.orchestration import pipeline
    from src.output import edit_sidecar
    from src.review.reviewer import ReviewResult

    edit = dict(actionType="EDIT", existingText="gate valve", replacementText="ball valve")
    merged = pipeline._deduplicate_findings(
        [_finding(evidenceElementId="p4", **edit), _finding(evidenceElementId="p8", **edit)]
    )
    payload = edit_sidecar.build_edit_instructions(
        SimpleNamespace(review_result=ReviewResult(findings=merged), module_id="datacenter_fire")
    )
    entries = list(payload.get("edits") or [])
    targets = [
        e.get("evidenceElementId") or (e.get("edit_proposal") or {}).get("target_element_id")
        for e in entries
    ]
    return len(entries) == 2, f"{len(entries)} entr(ies), targets={targets}"


# ----------------------------------------------------------------- WP-07 ---
@check("WP-07 two different spec.docx inputs are refused as ambiguous")
def _():
    from applier import run as applier_run

    with tempfile.TemporaryDirectory() as tmp:
        first = Path(tmp, "projA", "spec.docx")
        second = Path(tmp, "projB", "spec.docx")
        try:
            forward = applier_run._index_specs([first, second])
            backward = applier_run._index_specs([second, first])
        except Exception as exc:  # a fixed index may refuse outright
            return True, f"refused: {type(exc).__name__}"
    bound = forward.get("spec.docx")
    if isinstance(bound, Path):
        return False, (
            f"first wins: {bound.parent.name}; reversed input binds "
            f"{backward['spec.docx'].parent.name}"
        )
    return bound is not None and len(list(bound)) == 2, f"index value: {bound!r}"


# ----------------------------------------------------------------- WP-10 ---
@check("WP-10 a grounded UNVERIFIED is not replayed from the cache")
def _():
    from src.core.code_cycles import CALIFORNIA_2025
    from src.verification.verification_cache import VerificationCache
    from src.verification.verifier import VerificationResult

    cache = VerificationCache()
    finding = _finding("ASME B31.9 requires X.", codeReference="ASME B31.9")
    cache.put(
        finding,
        cycle=CALIFORNIA_2025,
        result=VerificationResult(
            verdict="UNVERIFIED",
            explanation="could not settle",
            grounded=True,
            searched_sources=["https://example.org/a"],
            successful_source_count=1,
        ),
    )
    hit = cache.get(finding, cycle=CALIFORNIA_2025)
    return hit is None, "cache returned a hit" if hit is not None else "no hit"


@check("WP-10 a blank source does not count as a citation")
def _():
    from src.output.report_status import ReportStatus, classify_status
    from src.verification.verifier import VerificationResult

    finding = _finding()
    finding.verification = VerificationResult(
        verdict="DISPUTED", grounded=True, sources=[""], accepted_sources=[""]
    )
    status = classify_status(finding)
    return status != ReportStatus.DISPUTED, f"classified as {status}"


# ----------------------------------------------------------------- WP-02 ---
def _extract_docx_with_wrappers() -> str:
    from docx import Document
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    from src.input.extractor import extract_text_from_docx

    doc = Document()
    doc.add_paragraph("Before the control.")
    doc.element.body.insert(
        1,
        parse_xml(
            f"<w:sdt {nsdecls('w')}><w:sdtContent><w:p><w:r>"
            "<w:t>Text inside a block content control.</w:t>"
            "</w:r></w:p></w:sdtContent></w:sdt>"
        ),
    )
    paragraph = doc.add_paragraph("Section ")
    paragraph._p.append(
        parse_xml(
            f'<w:fldSimple {nsdecls("w")} w:instr="REF sec \\h">'
            "<w:r><w:t>23 05 00</w:t></w:r></w:fldSimple>"
        )
    )
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp, "wrappers.docx")
        doc.save(path)
        spec = extract_text_from_docx(path)
    return getattr(spec, "content", str(spec))


@check("WP-02 text inside a block content control is extracted")
def _():
    present = "inside a block content control" in _extract_docx_with_wrappers()
    return present, "present" if present else "missing from extraction"


@check("WP-02 a simple field's stored result is extracted")
def _():
    present = "23 05 00" in _extract_docx_with_wrappers()
    return present, "present" if present else "missing from extraction"


# ------------------------------------------------------ source-level hints ---
def _read(rel: str) -> str:
    return (REPO_ROOT / rel).read_text(encoding="utf-8")


@check("WP-11 Retry-After is honored (source hint)")
def _():
    text = "\n".join(p.read_text(encoding="utf-8") for p in (REPO_ROOT / "src").rglob("*.py"))
    found = "retry-after" in text.lower()
    return found, "Retry-After handled somewhere in src/" if found else "no Retry-After handling in src/"


@check("WP-12 the chat never writes the API key to web storage (source hint)")
def _():
    stored = 'sessionStorage.setItem("sc_api_key"' in _read("src/output/html_report_exporter.py")
    return not stored, "key written to sessionStorage" if stored else "no stored key"


@check("WP-13 the GUI never copies the key into os.environ (source hint)")
def _():
    sites = [
        f"{p.relative_to(REPO_ROOT).as_posix()}"
        for p in (REPO_ROOT / "src" / "gui").glob("*.py")
        if 'os.environ["ANTHROPIC_API_KEY"] =' in p.read_text(encoding="utf-8")
    ]
    return not sites, f"sites: {sites}" if sites else "none"


@check("WP-17 Haiku cache minimum is not described as 2048 (source hint)")
def _():
    text = _read("src/core/api_config.py")
    stale = "2048 tokens for Haiku" in text or "2048-token Haiku" in text
    return not stale, "api_config.py still says 2048" if stale else "corrected"


@check("WP-17 the banner counts a hand-built no-op EDIT as a demotion")
def _():
    from src.orchestration import pipeline
    from src.output.report_exporter import _summarize_run_diagnostics
    from src.output.report_status import summarize_edit_actions

    # Built directly (not parsed), then sent through dedup, the review path's
    # normalization step. The parser stamps demotion_reason; this path doesn't.
    noop = _finding(actionType="EDIT", existingText="same", replacementText="same")
    findings = pipeline._deduplicate_findings([noop])
    summary = _summarize_run_diagnostics(
        findings=findings,
        status_counts={},
        edit_action_counts=summarize_edit_actions(findings),
        cross_check_result=None,
    )
    count = summary.get("demotion_count")
    return count == 1, f"banner demotion_count={count} (expected 1)"


def main() -> int:
    width = max(len(label) for _, label, _ in RESULTS)
    for state, label, detail in RESULTS:
        print(f"{state:<6} {label:<{width}}  {detail}")
    counts = {s: sum(1 for r in RESULTS if r[0] == s) for s in ("OPEN", "FIXED", "ERROR")}
    print()
    print(f"OPEN: {counts['OPEN']}   FIXED: {counts['FIXED']}   ERROR: {counts['ERROR']}")
    print("plans/PROGRESS.md is the authority on what is done; this is a quick check.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
